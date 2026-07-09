from __future__ import annotations

import math
import re
from collections import OrderedDict
from pathlib import Path
from typing import Any
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

WORD_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
SUPPLEMENTAL_METRICS_AFTER = {
    "\\u4e0a\\u4e0b\\u6e38\\u5c65\\u7ea6\\u8d44\\u6599\\u95ed\\u73af\\u7387": "\\u975e\\u8bc9\\u6536\\u6b3e",
}


def zh(value: str) -> str:
    return value.encode("utf-8").decode("unicode_escape")


def find_desktop_workbook(pattern: str) -> Path:
    desktop_name = zh("\\u684c\\u9762")
    desktop = next((path for path in Path("E:/").iterdir() if path.is_dir() and path.name == desktop_name), None)
    if desktop is None:
        raise FileNotFoundError("E:/ desktop folder not found")
    candidates = [
        path
        for path in desktop.glob(pattern)
        if path.is_file() and not path.name.startswith("~$")
    ]
    if not candidates:
        raise FileNotFoundError(f"no workbook matched {pattern}")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def find_adjusted_template_docx() -> Path:
    desktop_name = zh("\\u684c\\u9762")
    template_name = zh("\\u5404\\u56e2\\u961f\\u7ee9\\u6548\\u6c47\\u62a5\\u6a21\\u677f\\uff08\\u8c03\\u6574\\u7248\\uff09.docx")
    desktop = next((path for path in Path("E:/").iterdir() if path.is_dir() and path.name == desktop_name), None)
    if desktop is None:
        raise FileNotFoundError("E:/ desktop folder not found")
    template = desktop / template_name
    if not template.exists():
        raise FileNotFoundError(template)
    return template


def find_defendant_rate_workbook() -> Path:
    desktop_name = zh("\\u684c\\u9762")
    workbook_name = zh("\\u6cd5\\u52a1\\u90e8\\u95e8\\u6848\\u4ef6\\u4e0b\\u964d\\u7387\\u7edf\\u8ba1\\u8868_\\u5df2\\u586b.xlsx")
    desktop = next((path for path in Path("E:/").iterdir() if path.is_dir() and path.name == desktop_name), None)
    if desktop is None:
        raise FileNotFoundError("E:/ desktop folder not found")
    workbook = desktop / workbook_name
    if not workbook.exists():
        raise FileNotFoundError(workbook)
    return workbook


def find_main_performance_workbook() -> Path:
    desktop_name = zh("\\u684c\\u9762")
    desktop = next((path for path in Path("E:/").iterdir() if path.is_dir() and path.name == desktop_name), None)
    if desktop is None:
        raise FileNotFoundError("E:/ desktop folder not found")
    preferred_names = [
        zh("\\u6cd5\\u52a1\\u90e8\\u95e86\\u6708\\u7ee9\\u6548\\u5b8c\\u6210\\u60c5\\u51b5.xlsx"),
        zh("\\u5404\\u56e2\\u961f\\u6307\\u6807\\u5b8c\\u6210\\u60c5\\u51b57.1.xlsx"),
    ]
    for name in preferred_names:
        candidate = desktop / name
        if candidate.exists() and not candidate.name.startswith("~$"):
            return candidate
    return find_desktop_workbook("*7.1.xlsx")


def cell(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def clean(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str):
        return re.sub(r"\s+", "", value.strip())
    return value


def is_num(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def fmt_num(value: Any, unit: str = "") -> str:
    if value is None or value == "":
        return zh("\\u672a\\u586b")
    if isinstance(value, str):
        return value.strip()
    if is_num(value):
        if unit == "%":
            number = value * 100 if abs(value) <= 1 else value
            return f"{number:.2f}%".replace(".00%", "%")
        if abs(value - round(value)) < 1e-9:
            return str(int(round(value)))
        return f"{value:.2f}".rstrip("0").rstrip(".")
    return str(value)


def fmt_value(value: Any, unit: str) -> str:
    text = fmt_num(value, "%" if unit == "%" else "")
    if text in ("", zh("\\u672a\\u586b")):
        return text
    if unit == "%":
        return text
    return f"{text}{unit}" if unit else text


def calc_rate(actual: Any, target: Any) -> float | None:
    if is_num(target) and abs(target) > 1e-12 and is_num(actual):
        return actual / target
    return None


def display_rate(cached: Any, actual: Any, target: Any) -> str:
    if is_num(target) and abs(target) <= 1e-12:
        return zh("\\u65e0")
    calculated = calc_rate(actual, target)
    if calculated is not None:
        return f"{calculated * 100:.2f}%".replace(".00%", "%")
    if cached not in (None, ""):
        return fmt_num(cached, "%")
    return zh("\\u672a\\u586b")


def read_rows(source: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(source, data_only=True)
    sheet = workbook.active
    rows: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        rows.append(
            {
                "seq": clean(row[0]),
                "team": clean(cell(row, 1)),
                "leader": clean(cell(row, 2)),
                "metric": clean(cell(row, 3)),
                "unit": clean(cell(row, 4)),
                "year_target": clean(cell(row, 5)),
                "year_actual": clean(cell(row, 6)),
                "year_rate_cached": clean(cell(row, 7)),
                "yoy": clean(cell(row, 8)),
                "month_target": clean(cell(row, 9)),
                "month_actual": clean(cell(row, 10)),
                "month_rate_cached": clean(cell(row, 11)),
                "mom": clean(cell(row, 12)),
                "half_plan": clean(cell(row, 13)),
                "half_forecast": clean(cell(row, 14)),
            }
        )
    return rows


def read_defendant_rate_rows(source: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(source, data_only=True)
    sheet = workbook.active
    result: dict[str, dict[str, Any]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        team = clean(row[1])
        metric = clean(row[3])
        if not is_defendant_inventory_metric(metric):
            continue
        result[str(team)] = {
            "seq": clean(row[0]),
            "team": team,
            "leader": clean(cell(row, 2)),
            "metric": metric,
            "unit": clean(cell(row, 4)),
            "year_target": clean(cell(row, 5)),
            "year_actual": clean(cell(row, 6)),
            "year_rate_cached": clean(cell(row, 7)),
            "yoy": clean(cell(row, 8)),
            "month_target": clean(cell(row, 9)),
            "month_actual": clean(cell(row, 10)),
            "month_rate_cached": clean(cell(row, 11)),
            "mom": clean(cell(row, 12)),
            "half_plan": "",
            "half_forecast": "",
        }
    return result


def merge_defendant_rate_rows(rows: list[dict[str, Any]], defendant_rows: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    merged: list[dict[str, Any]] = []
    for row in rows:
        if is_defendant_inventory_metric(row["metric"]) and row["team"] in defendant_rows:
            merged.append({**row, **defendant_rows[row["team"]]})
        else:
            merged.append(row)
    return merged


def read_template_metric_order(template: Path) -> dict[str, list[str]]:
    with ZipFile(template) as archive:
        document_xml = archive.read("word/document.xml")
    root = ET.fromstring(document_xml)
    paragraphs = []
    for paragraph in root.findall(".//w:p", WORD_NS):
        text = "".join((node.text or "") for node in paragraph.findall(".//w:t", WORD_NS)).strip()
        if text:
            paragraphs.append(text)

    title_suffix = zh("X\\u6708\\u5ea6\\u7ee9\\u6548\\u5de5\\u4f5c\\u6c47\\u62a5")
    known_departments = {
        zh("\\u6cd5\\u52a1\\u4e00\\u90e8"),
        zh("\\u6cd5\\u52a1\\u4e8c\\u90e8"),
        zh("\\u6cd5\\u52a1\\u4e09\\u90e8"),
        zh("\\u6cd5\\u52a1\\u56db\\u90e8"),
        zh("\\u6cd5\\u52a1\\u4e94\\u90e8"),
        zh("\\u6cd5\\u52a1\\u516d\\u90e8"),
        zh("\\u7efc\\u5408\\u7ba1\\u7406\\u90e8"),
    }
    sections: dict[str, list[str]] = {}
    current = ""
    for text in paragraphs:
        compact = "".join(text.split())
        if compact.endswith(title_suffix):
            department = compact[: -len(title_suffix)]
            current = department if department in known_departments else ""
            if current:
                sections.setdefault(current, [])
            continue
        match = re.match(r"^(\d+)[、](.+)$", text)
        if current and match:
            sections[current].append(clean(match.group(2)))
    return sections


def group_by_team(rows: list[dict[str, Any]]) -> OrderedDict[str, dict[str, Any]]:
    teams: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for row in rows:
        teams.setdefault(row["team"], {"leader": row["leader"], "rows": []})["rows"].append(row)
    return teams


def is_overall_team(team: str) -> bool:
    return "整体" in str(team)


def is_defendant_inventory_metric(metric_name: str) -> bool:
    return metric_name == zh("\\u88ab\\u544a\\u5b58\\u91cf/\\u65b0\\u589e\\u6848\\u4ef6\\u6570\\u91cf\\u4e0b\\u964d\\u7387")


def display_unit_for_metric(row: dict[str, Any]) -> str:
    unit = str(row.get("unit") or "").strip()
    if unit:
        return unit
    metric_name = str(row.get("metric") or "")
    amount_markers = ("收款", "收入", "结算增加额", "索赔", "非诉")
    if any(marker in metric_name for marker in amount_markers):
        return zh("\\u4e07\\u5143")
    return ""


def display_rows_for_team(team: str, rows: list[dict[str, Any]], template_order: dict[str, list[str]]) -> list[dict[str, Any]]:
    if is_overall_team(team):
        return rows
    order = template_order.get(team)
    if not order:
        return rows
    by_metric = {row["metric"]: row for row in rows}
    result: list[dict[str, Any]] = []
    added = set()
    supplemental_after = {zh(name): zh(after) for name, after in SUPPLEMENTAL_METRICS_AFTER.items()}
    for name in order:
        if name in by_metric:
            result.append(by_metric[name])
            added.add(name)
        for supplemental_name, anchor in supplemental_after.items():
            if name == anchor and supplemental_name in by_metric and supplemental_name not in added:
                result.append(by_metric[supplemental_name])
                added.add(supplemental_name)
    return result


def template_extra_rows(teams: OrderedDict[str, dict[str, Any]], template_order: dict[str, list[str]]) -> list[tuple[str, str]]:
    extras: list[tuple[str, str]] = []
    supplemental_names = {zh(name) for name in SUPPLEMENTAL_METRICS_AFTER}
    for team, item in teams.items():
        if is_overall_team(team):
            continue
        allowed = set(template_order.get(team, [])) | supplemental_names
        for row in item["rows"]:
            if allowed and row["metric"] not in allowed:
                extras.append((team, row["metric"]))
    return extras


def fmt_compound_text(value: Any) -> str:
    if value is None or value == "":
        return zh("\\u672a\\u586b")
    if is_num(value):
        return fmt_num(value, "%")
    text = str(value).strip()
    text = re.sub(r"\s+", "", text)
    new_case_label = zh("\\u3010\\u65b0\\u589e\\u3011")
    if new_case_label in text:
        text = text.replace(new_case_label, zh("\\uff0c\\u3010\\u65b0\\u589e\\u3011"))
    else:
        text = re.sub(zh(r"(?<![\\uff1b\\u3001\\uff0c])\\u65b0\\u589e"), zh("\\uff0c\\u65b0\\u589e"), text, count=1)
    text = text.replace("【存量】", "存量").replace("【新增】", "新增")
    return text


def fmt_compound_comparison(value: Any) -> str:
    text = fmt_compound_text(value)
    replacements = {
        "比目标低": "低于目标",
        "同比低去年": "同比下降",
        "同比高去年": "同比上升",
        "环比低": "环比下降",
        "环比高": "环比上升",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def fmt_defendant_decline_target(value: Any) -> str:
    if is_num(value):
        target = fmt_num(value, "%")
        stock_label = zh("\\u5b58\\u91cf\\u4e0b\\u964d")
        new_label = zh("\\uff0c\\u65b0\\u589e\\u4e0b\\u964d")
        return f"{stock_label}{target}{new_label}{target}"
    text = fmt_compound_text(value)
    parts = []
    for part in re.split(r"[，,；;]+", text):
        item = part.strip()
        if not item:
            continue
        if any(word in item for word in ("下降", "增长", "持平")):
            parts.append(item)
        elif item.startswith("存量"):
            parts.append(item.replace("存量", "存量下降", 1))
        elif item.startswith("新增"):
            parts.append(item.replace("新增", "新增下降", 1))
        else:
            parts.append(item)
    return "，".join(parts) if parts else zh("\\u672a\\u586b")


def metric_display_lines(row: dict[str, Any]) -> list[str]:
    unit = display_unit_for_metric(row)
    month_target_label = zh("\\u6708\\u5ea6\\u76ee\\u6807")
    month_actual_label = zh("\\u5b9e\\u9645\\u5b8c\\u6210")
    completion_rate_label = zh("\\u5b8c\\u6210\\u7387")
    month_mom_label = zh("\\u73af\\u6bd4")
    month_mom_full_label = zh("\\u73af\\u6bd4\\u4e0a\\u5347/\\u4e0b\\u964d")
    year_target_label = zh("\\u5e74\\u5ea6\\u76ee\\u6807")
    year_actual_label = zh("\\u7d2f\\u8ba1\\u5b9e\\u9645\\u5b8c\\u6210")
    year_rate_label = zh("\\u7d2f\\u8ba1\\u5b8c\\u6210\\u7387")
    yoy_label = zh("\\u540c\\u6bd4")
    yoy_full_label = zh("\\u540c\\u6bd4\\u4e0a\\u5347/\\u4e0b\\u964d")
    unfilled = zh("\\u672a\\u586b")
    if is_defendant_inventory_metric(row["metric"]):
        target = fmt_value(row["year_target"], unit)
        stock_target = zh("\\u5b58\\u91cf\\u4e0b\\u964d")
        new_target = zh("\\uff0c\\u65b0\\u589e\\u4e0b\\u964d")
        return [
            f"   {month_target_label}：{fmt_defendant_decline_target(row['month_target'])}；"
            f"{month_actual_label}：{fmt_compound_text(row['month_actual'])}；"
            f"{completion_rate_label}：{fmt_compound_comparison(row['month_rate_cached'])}；"
            f"{month_mom_full_label}：{fmt_compound_comparison(row['mom'])}；",
            f"   {year_target_label}：{stock_target}{target}{new_target}{target}；"
            f"{year_actual_label}：{fmt_compound_text(row['year_actual'])}；"
            f"{year_rate_label}：{fmt_compound_comparison(row['year_rate_cached'])}；"
            f"{yoy_full_label}：{fmt_compound_comparison(row['yoy'])}；",
        ]
    return [
        f"   {month_target_label}：{fmt_value(row['month_target'], unit)}，"
        f"{month_actual_label}：{fmt_value(row['month_actual'], unit)}，"
        f"{completion_rate_label}：{display_rate(row['month_rate_cached'], row['month_actual'], row['month_target'])}，"
        f"{month_mom_label}：{fmt_num(row['mom'], '%') if row['mom'] != '' else unfilled}；",
        f"   {year_target_label}：{fmt_value(row['year_target'], unit)}，"
        f"{year_actual_label}：{fmt_value(row['year_actual'], unit)}，"
        f"{year_rate_label}：{display_rate(row['year_rate_cached'], row['year_actual'], row['year_target'])}，"
        f"{yoy_label}：{fmt_num(row['yoy'], '%') if row['yoy'] != '' else unfilled}；",
    ]


def quality_counts(rows: list[dict[str, Any]]) -> tuple[int, int]:
    mismatches = 0
    zero_targets = 0
    for row in rows:
        for target_key, actual_key, rate_key in (
            ("year_target", "year_actual", "year_rate_cached"),
            ("month_target", "month_actual", "month_rate_cached"),
        ):
            target = row[target_key]
            actual = row[actual_key]
            cached = row[rate_key]
            if is_num(target) and abs(target) <= 1e-12:
                zero_targets += 1
            calculated = calc_rate(actual, target)
            if calculated is not None and is_num(cached) and abs(calculated - cached) > 0.005:
                mismatches += 1
    return mismatches, zero_targets


def build_preview(rows: list[dict[str, Any]], source: Path, template: Path, defendant_source: Path | None = None) -> str:
    teams = group_by_team(rows)
    template_order = read_template_metric_order(template)
    extras = template_extra_rows(teams, template_order)
    mismatches, zero_targets = quality_counts(rows)
    lines: list[str] = []
    lines.append(zh("# \\u5404\\u56e2\\u961f\\u6307\\u6807\\u5b8c\\u6210\\u60c5\\u51b57.1 \\u6570\\u636e\\u9884\\u89c8"))
    lines.append("")
    lines.append(zh("\\u6570\\u636e\\u6e90\\uff1a") + str(source))
    if defendant_source is not None:
        lines.append(zh("\\u88ab\\u544a\\u6307\\u6807\\u6570\\u636e\\u6e90\\uff1a") + str(defendant_source))
    lines.append(zh("\\u987a\\u5e8f\\u6e90\\uff1a") + str(template))
    lines.append(zh("\\u8fd9\\u4efd\\u9884\\u89c8\\u4ec5\\u7528\\u4e8e\\u786e\\u8ba4\\u6307\\u6807\\u89e3\\u6790\\u548c\\u56e2\\u961f\\u586b\\u62a5\\u53e3\\u5f84\\uff0c\\u5c1a\\u672a\\u6b63\\u5f0f\\u53d1\\u9001\\u3002"))
    lines.append("")
    lines.append(zh("## \\u4e00\\u3001\\u56e2\\u961f\\u4e0e\\u6307\\u6807\\u6570\\u91cf"))
    metric_count_label = zh("\\u4e2a\\u586b\\u62a5\\u6307\\u6807")
    leader_label = zh("\\u8d1f\\u8d23\\u4eba")
    for team, item in teams.items():
        display_rows = display_rows_for_team(team, item["rows"], template_order)
        note = zh("\\uff0c\\u6574\\u4f53\\u53c2\\u8003\\u6570\\u636e\\uff0c\\u4e0d\\u751f\\u6210\\u8d1f\\u8d23\\u4eba\\u586b\\u62a5\\u6bb5\\u843d") if is_overall_team(team) else ""
        lines.append(f"- {team}：{len(display_rows)} {metric_count_label}，{leader_label}：{item['leader']}{note}")
    lines.append("")
    lines.append(zh("## \\u4e8c\\u3001\\u6570\\u636e\\u8d28\\u91cf\\u63d0\\u793a"))
    row_count_label = zh("\\u6307\\u6807\\u884c\\u6570")
    mismatch_label = zh("\\u5b8c\\u6210\\u7387\\u7f13\\u5b58\\u503c\\u4e0e\\u76ee\\u6807/\\u5b9e\\u9645\\u91cd\\u7b97\\u4e0d\\u4e00\\u81f4")
    mismatch_suffix = zh("\\u5904\\uff0c\\u540e\\u7eed\\u4ee5\\u91cd\\u7b97\\u503c\\u4e3a\\u51c6")
    zero_target_label = zh("\\u76ee\\u6807\\u4e3a0\\u7684\\u5b8c\\u6210\\u7387")
    zero_target_suffix = zh("\\u5904\\uff0c\\u7edf\\u4e00\\u5c55\\u793a\\u4e3a\\u201c\\u65e0\\u201d")
    lines.append(f"- {row_count_label}：{len(rows)}")
    lines.append(f"- {mismatch_label}：{mismatches} {mismatch_suffix}")
    lines.append(f"- {zero_target_label}：{zero_targets} {zero_target_suffix}")
    lines.append(zh("- \\u5de5\\u4f5c\\u7c3f\\u672a\\u63d0\\u4f9b\\u5468\\u76ee\\u6807/\\u5468\\u5b9e\\u9645/\\u5468\\u5b8c\\u6210\\u7387\\uff0c\\u4e0d\\u5efa\\u8bae\\u5728\\u6b64\\u6b21\\u53d1\\u9001\\u5185\\u5bb9\\u91cc\\u9020\\u5468\\u7ef4\\u5ea6\\u6570\\u636e\\u3002"))
    lines.append(zh("- \\u201c\\u88ab\\u544a\\u5b58\\u91cf/\\u65b0\\u589e\\u6848\\u4ef6\\u6570\\u91cf\\u4e0b\\u964d\\u7387\\u201d\\u4e3a\\u590d\\u5408\\u6bd4\\u7387\\u6307\\u6807\\uff0c\\u4e0d\\u505a\\u7b80\\u5355\\u6c47\\u603b\\uff0c\\u6309\\u90e8\\u95e8\\u539f\\u6587\\u5c55\\u793a\\u3002"))
    if extras:
        extra_text = zh("\\u3001").join(f"{team}-{metric}" for team, metric in extras)
        extras_label = zh("\\u5de5\\u4f5c\\u7c3f\\u5b58\\u5728\\u6a21\\u677f\\u5916\\u6307\\u6807")
        extras_suffix = zh("\\u672a\\u8fdb\\u5165\\u8d1f\\u8d23\\u4eba\\u586b\\u62a5\\u6bb5\\u843d\\u3002")
        lines.append(f"- {extras_label}：{extra_text}；{extras_suffix}")
    lines.append("")

    for team, item in teams.items():
        display_rows = display_rows_for_team(team, item["rows"], template_order)
        lines.append(zh("## \\u4e09\\u3001") + f"{team} - " + zh("\\u6307\\u6807\\u5b8c\\u6210\\u60c5\\u51b5\\u6982\\u89c8"))
        lines.append(f"{leader_label}：{item['leader']}")
        lines.append("")
        for index, row in enumerate(display_rows, start=1):
            lines.append(f"{index}. **{row['metric']}**")
            lines.extend(metric_display_lines(row))
            lines.append("")
        if is_overall_team(team):
            lines.append(zh("### \\u8bf4\\u660e"))
            lines.append(zh("\\u8be5\\u90e8\\u5206\\u4ec5\\u4f5c\\u4e3a\\u6700\\u7ec8\\u90e8\\u95e8\\u6708\\u62a5\\u6574\\u4f53\\u53c2\\u8003\\uff0c\\u4e0d\\u5411\\u8d1f\\u8d23\\u4eba\\u6536\\u96c6\\u56de\\u590d\\u3002"))
            lines.append("")
            continue
        lines.append(zh("### \\u7ed9\\u8d1f\\u8d23\\u4eba\\u7684\\u586b\\u62a5\\u6bb5\\u843d\\uff08\\u9884\\u89c8\\uff09"))
        lines.append(zh("\\u3010\\u8bf7\\u56de\\u590d\\u3011\\u672c\\u6b21\\u9700\\u8981\\u586b\\u5199\\u7684\\u6307\\u6807\\uff1a"))
        for index, row in enumerate(display_rows, start=1):
            unit = display_unit_for_metric(row)
            target_label = zh("\\u4e0b\\u6708\\u76ee\\u6807")
            if unit:
                target_label += f"（{unit}）"
            lines.append(f"{index}. 【{row['metric']}】")
            lines.append(zh("\\u672a\\u5b8c\\u6210\\u539f\\u56e0/\\u5b58\\u5728\\u95ee\\u9898\\uff1a"))
            lines.append(f"{target_label}：")
            lines.append(zh("\\u884c\\u52a8\\u65b9\\u6848\\uff1a"))
            lines.append("")
        lines.append("")
    return "\n".join(lines)


def main() -> None:
    source = find_main_performance_workbook()
    template = find_adjusted_template_docx()
    defendant_source = find_defendant_rate_workbook()
    rows = read_rows(source)
    rows = merge_defendant_rate_rows(rows, read_defendant_rate_rows(defendant_source))
    output = Path("outputs/monthly_report_data_preview_20260701.md")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(build_preview(rows, source, template, defendant_source), encoding="utf-8")
    print(output.resolve())
    print(f"rows={len(rows)}")


if __name__ == "__main__":
    main()
