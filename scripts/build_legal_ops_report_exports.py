from __future__ import annotations

from pathlib import Path
import json
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.legal_ops.seed import build_phase0_seed


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "artifacts" / "legal-ops-exports"
NAVY = "101B2D"
GOLD = "C79B50"
LIGHT = "F2F4F7"
RED = "9B1C1C"


def tenant_rows(seed: dict, collection: str) -> list[dict]:
    return [row for row in seed[collection] if row["tenant_id"] == "sandbox-alpha"]


def format_metric_value(value: float, unit: str) -> str:
    return f"{value:.0%}" if unit == "%" else f"{value:,.0f}"


def set_font(run, *, name: str = "Microsoft YaHei", size: float = 11, bold: bool = False, color: str = "172033") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shd = properties.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        properties.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    tcw = properties.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        properties.append(tcw)
    tcw.set(qn("w:w"), str(width_dxa))
    tcw.set(qn("w:type"), "dxa")


def configure_table(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), "9360")
    width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    properties.append(indent)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for cell, value in zip(row.cells, widths):
            set_cell_width(cell, value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                cell._tc.get_or_add_tcPr().append(margins)
            for edge, amount in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                node = margins.find(qn(f"w:{edge}"))
                if node is None:
                    node = OxmlElement(f"w:{edge}")
                    margins.append(node)
                node.set(qn("w:w"), str(amount))
                node.set(qn("w:type"), "dxa")


def add_title_block(document: Document, report: dict) -> None:
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(70)
    kicker.paragraph_format.space_after = Pt(14)
    set_font(kicker.add_run("法务中心绩效报告"), size=10, bold=True, color=GOLD)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    set_font(title.add_run(report["title"]), size=28, bold=True, color=NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    set_font(subtitle.add_run(f"统计周期：{report['period_start']} 至 {report['period_end']}"), size=12, color="566274")
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(70)
    set_font(meta.add_run(f"状态：{report['status']}  |  生成时间：{report['generated_at']}"), size=10, color="6D7688")


def build_docx(report: dict, metrics: list[dict], output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 6), ("Heading 3", 12, 8, 4)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("2E5D86")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    header = section.header.paragraphs[0]
    set_font(header.add_run("法务运营作战中心"), size=9, bold=True, color="6D7688")
    header.add_run("                                      ")
    set_font(header.add_run("Sandbox 固定格式报告"), size=9, color="6D7688")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("仅用于产品演示 | 非生产业务事实"), size=8, color="7A8494")
    add_title_block(document, report)
    document.add_page_break()
    document.add_heading("一、法务中心总体情况", level=1)
    document.add_paragraph(report["overall_summary"])
    document.add_heading("二、团队指标", level=1)
    table = document.add_table(rows=1, cols=7)
    headers = ("团队", "指标", "目标", "实际", "完成率", "来源", "确认状态")
    for index, text in enumerate(headers):
        shade(table.rows[0].cells[index], NAVY)
        paragraph = table.rows[0].cells[index].paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(paragraph.add_run(text), size=9, bold=True, color="FFFFFF")
    for metric in metrics:
        row = table.add_row().cells
        values = (
            metric["team_name"],
            metric["metric_name"],
            format_metric_value(metric["target_value"], metric["unit"]),
            format_metric_value(metric["actual_value"], metric["unit"]),
            f"{metric['completion_rate']:.0%}",
            metric["data_source_name"],
            metric["confirmation_status"],
        )
        for index, text in enumerate(values):
            paragraph = row[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in (0, 2, 3, 4, 6) else WD_ALIGN_PARAGRAPH.LEFT
            set_font(paragraph.add_run(str(text)), size=8.5, color="172033")
        if metric["confirmation_status"] != "已确认":
            shade(row[-1], "FCE9E8")
    configure_table(table, [1250, 2050, 950, 950, 850, 2050, 1260])
    document.add_heading("三、重点问题与风险", level=1)
    risk = document.add_paragraph()
    risk.paragraph_format.left_indent = Inches(0.16)
    risk.paragraph_format.space_before = Pt(4)
    risk.paragraph_format.space_after = Pt(8)
    set_font(risk.add_run(report["risk_summary"]), size=11, bold=True, color=RED)
    document.add_heading("四、需协调事项", level=1)
    document.add_paragraph(report["coordination_needed"])
    document.add_heading("五、下阶段目标", level=1)
    document.add_paragraph(report["next_period_goal"])
    source = document.add_paragraph()
    source.paragraph_format.space_before = Pt(12)
    source.paragraph_format.space_after = Pt(4)
    set_font(source.add_run("数据来源：WorkBuddy Skill、人工维护、团队负责人机器人结构化回复。信息中心接口尚未接入。"), size=8.5, color="6D7688")
    document.core_properties.title = report["title"]
    document.core_properties.subject = "法务绩效固定格式报告"
    document.core_properties.author = "Legal Operations Sandbox"
    document.save(output)


def build_pdf(report: dict, metrics: list[dict], output: Path) -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ChineseTitle", parent=styles["Title"], fontName="STSong-Light", fontSize=24, leading=30, textColor=colors.HexColor(f"#{NAVY}"), alignment=TA_CENTER, spaceAfter=14)
    heading_style = ParagraphStyle("ChineseHeading", parent=styles["Heading1"], fontName="STSong-Light", fontSize=14, leading=18, textColor=colors.HexColor("#2E5D86"), spaceBefore=12, spaceAfter=8)
    body_style = ParagraphStyle("ChineseBody", parent=styles["BodyText"], fontName="STSong-Light", fontSize=10, leading=16, textColor=colors.HexColor("#172033"), alignment=TA_LEFT)
    small_style = ParagraphStyle("ChineseSmall", parent=body_style, fontSize=8, leading=11, textColor=colors.HexColor("#6D7688"))
    doc = SimpleDocTemplate(str(output), pagesize=letter, rightMargin=0.65*inch, leftMargin=0.65*inch, topMargin=0.7*inch, bottomMargin=0.65*inch, title=report["title"], author="Legal Operations Sandbox")
    story = [Spacer(1, 0.65*inch), Paragraph(report["title"], title_style), Paragraph(f"统计周期：{report['period_start']} 至 {report['period_end']}", body_style), Spacer(1, 0.55*inch), Paragraph(report["overall_summary"], body_style), PageBreak(), Paragraph("团队指标", heading_style)]
    data = [["团队", "指标", "目标", "实际", "完成率", "来源", "状态"]]
    for metric in metrics:
        data.append([metric["team_name"], metric["metric_name"], format_metric_value(metric["target_value"], metric["unit"]), format_metric_value(metric["actual_value"], metric["unit"]), f"{metric['completion_rate']:.0%}", metric["data_source_name"], metric["confirmation_status"]])
    table = Table(data, colWidths=[0.8*inch, 1.32*inch, 0.65*inch, 0.65*inch, 0.55*inch, 1.25*inch, 0.7*inch], repeatRows=1)
    table.setStyle(TableStyle([("FONTNAME", (0,0), (-1,-1), "STSong-Light"), ("FONTSIZE", (0,0), (-1,-1), 7.5), ("BACKGROUND", (0,0), (-1,0), colors.HexColor(f"#{NAVY}")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("ALIGN", (0,0), (-1,-1), "CENTER"), ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#D9DEE6")), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F7F8FA")]), ("TOPPADDING", (0,0), (-1,-1), 6), ("BOTTOMPADDING", (0,0), (-1,-1), 6)]))
    story.extend([table, Paragraph("重点问题与风险", heading_style), Paragraph(report["risk_summary"], body_style), Paragraph("需协调事项", heading_style), Paragraph(report["coordination_needed"], body_style), Paragraph("下阶段目标", heading_style), Paragraph(report["next_period_goal"], body_style), Spacer(1, 12), Paragraph("数据来源：WorkBuddy Skill、人工维护、团队负责人机器人结构化回复。信息中心接口尚未接入。", small_style)])
    doc.build(story)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    seed = build_phase0_seed()
    reports = tenant_rows(seed, "report_runs")
    all_metrics = tenant_rows(seed, "performance_metrics")
    (OUTPUT / "report-data.json").write_text(
        json.dumps({"reports": reports, "metrics": all_metrics}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    for report in reports:
        period = report["period_type"]
        metrics = [row for row in all_metrics if row["period_type"] == period]
        build_docx(report, metrics, OUTPUT / f"legal-ops-{period}-report.docx")
        build_pdf(report, metrics, OUTPUT / f"legal-ops-{period}-report.pdf")
    print(OUTPUT)


if __name__ == "__main__":
    main()
