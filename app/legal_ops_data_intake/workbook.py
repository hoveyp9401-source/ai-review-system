from __future__ import annotations

import csv
import io
import re
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import PurePath
from typing import Any, Literal

from openpyxl import Workbook, load_workbook


class FileValidationError(ValueError):
    pass


ColumnType = Literal["text", "date", "decimal", "percentage", "integer"]


@dataclass(frozen=True)
class TableColumn:
    key: str
    name: str
    data_type: ColumnType
    required: bool = False
    unique: bool = False


@dataclass(frozen=True)
class TableSchema:
    key: str
    name: str
    columns: tuple[TableColumn, ...]
    allow_extra_columns: bool = False


@dataclass(frozen=True)
class RowError:
    code: str
    field: str
    message: str
    critical: bool = True


@dataclass
class ParsedRow:
    row_number: int
    raw: dict[str, Any]
    normalized: dict[str, Any]
    errors: list[RowError]


@dataclass(frozen=True)
class FormulaCell:
    formula: str


@dataclass(frozen=True)
class ParsedTable:
    filename: str
    headers: tuple[str, ...]
    rows: tuple[ParsedRow, ...]
    sheet_name: str = ""

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def error_count(self) -> int:
        return sum(bool(row.errors) for row in self.rows)


_MAX_FILE_BYTES = 25 * 1024 * 1024
_MAX_ROWS = 100_000
_MAX_XLSX_ENTRIES = 2_000
_MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
_MAX_XLSX_COLUMNS = 512
_MAX_XLSX_CELLS = 5_000_000


def parse_tabular_file(
    content: bytes,
    filename: str,
    schema: TableSchema,
    *,
    max_bytes: int = _MAX_FILE_BYTES,
    max_rows: int = _MAX_ROWS,
    preferred_sheet_name: str = "",
    header_mapping: Mapping[str, str] | None = None,
) -> ParsedTable:
    safe_name, suffix = _validate_name(filename)
    if max_bytes < 1 or max_rows < 1:
        raise ValueError("文件大小和行数限制必须大于零")
    if len(content) > max_bytes:
        raise FileValidationError("文件超过允许大小")
    source_header_by_key = {
        column.key: str((header_mapping or {}).get(column.key) or column.name).strip()
        for column in schema.columns
    }
    expected_source_headers = {
        value for value in source_header_by_key.values() if value
    }
    if len(expected_source_headers) != len(schema.columns):
        raise FileValidationError("字段对应关系存在重复或空值")
    if suffix == ".csv":
        raw_rows = _read_csv(content, max_rows=max_rows)
        sheet_name = ""
    elif suffix == ".xlsx":
        sheet_name, raw_rows = _read_xlsx(
            content,
            max_rows=max_rows,
            required_headers=expected_source_headers,
            preferred_sheet_name=preferred_sheet_name,
        )
    else:
        raise FileValidationError("只支持 xlsx 和 csv 文件")
    if not raw_rows:
        raise FileValidationError("文件没有表头")
    headers = tuple(str(value or "").strip() for value in raw_rows[0])
    if any(not value for value in headers) or len(headers) != len(set(headers)):
        raise FileValidationError("表头为空或重复")
    missing_headers = [
        column.name
        for column in schema.columns
        if source_header_by_key[column.key] not in headers
    ]
    if missing_headers:
        raise FileValidationError(f"缺少表头：{'、'.join(missing_headers)}")
    unknown_headers = [
        header for header in headers if header not in expected_source_headers
    ]
    if unknown_headers and not schema.allow_extra_columns:
        raise FileValidationError(f"存在未定义表头：{'、'.join(unknown_headers)}")

    column_by_name = {
        source_header_by_key[column.key]: column for column in schema.columns
    }
    rows: list[ParsedRow] = []
    for row_number, values in enumerate(raw_rows[1:], start=2):
        padded = list(values) + [None] * max(0, len(headers) - len(values))
        if all(_is_blank(value) for value in padded[: len(headers)]):
            continue
        raw = {header: padded[index] for index, header in enumerate(headers)}
        normalized: dict[str, Any] = {}
        errors: list[RowError] = []
        for header, value in raw.items():
            if header not in column_by_name:
                continue
            column = column_by_name[header]
            if _is_blank(value):
                normalized[column.key] = None
                if column.required:
                    errors.append(
                        RowError("required", column.key, f"{column.name}不能为空")
                    )
                continue
            try:
                normalized[column.key] = _normalize(value, column)
            except FileValidationError as exc:
                errors.append(
                    RowError(_error_code(column.data_type), column.key, str(exc))
                )
        rows.append(ParsedRow(row_number, raw, normalized, errors))

    for column in schema.columns:
        if not column.unique:
            continue
        seen: dict[str, ParsedRow] = {}
        for row in rows:
            value = row.normalized.get(column.key)
            if value is None:
                continue
            marker = str(value)
            if marker in seen:
                message = f"{column.name}在文件内重复"
                if not any(
                    error.code == "duplicate" and error.field == column.key
                    for error in seen[marker].errors
                ):
                    seen[marker].errors.append(
                        RowError("duplicate", column.key, message)
                    )
                row.errors.append(RowError("duplicate", column.key, message))
            else:
                seen[marker] = row
    return ParsedTable(safe_name, headers, tuple(rows), sheet_name)


def _validate_name(filename: str) -> tuple[str, str]:
    value = str(filename or "").strip()
    if (
        not value
        or "/" in value
        or "\\" in value
        or ":" in value
        or value in {".", ".."}
        or PurePath(value).name != value
    ):
        raise FileValidationError("文件名不安全")
    suffix = PurePath(value).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise FileValidationError("只支持 xlsx 和 csv 文件，不支持宏文件")
    return value, suffix


def _read_csv(content: bytes, *, max_rows: int) -> list[list[Any]]:
    text: str | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise FileValidationError("CSV 编码无法识别，请使用 UTF-8")
    try:
        rows: list[list[Any]] = []
        for index, row in enumerate(csv.reader(io.StringIO(text))):
            if index > max_rows:
                raise FileValidationError(f"文件超过单次 {max_rows} 行限制")
            rows.append(list(row))
        return rows
    except csv.Error as exc:
        raise FileValidationError(f"CSV 文件损坏：{exc}") from exc


def read_tabular_sheets(
    content: bytes,
    filename: str,
    *,
    max_bytes: int = _MAX_FILE_BYTES,
    max_rows: int = _MAX_ROWS,
) -> tuple[str, dict[str, list[list[Any]]]]:
    """Read a CSV or every worksheet from an XLSX after the common safety checks.

    The caller is responsible for selecting a business data sheet. This is used by
    controlled source profiles whose official workbooks also contain dashboards.
    """

    safe_name, suffix = _validate_name(filename)
    if max_bytes < 1 or max_rows < 1:
        raise ValueError("文件大小和行数限制必须大于零")
    if len(content) > max_bytes:
        raise FileValidationError("文件超过允许大小")
    if suffix == ".csv":
        return safe_name, {"": _read_csv(content, max_rows=max_rows)}
    return safe_name, _read_xlsx_sheets(content, max_rows=max_rows)


def inspect_tabular_source(
    content: bytes,
    filename: str,
    *,
    max_bytes: int = _MAX_FILE_BYTES,
    max_rows: int = _MAX_ROWS,
) -> dict[str, Any]:
    """Describe an existing business workbook without requiring a rule schema.

    Only sheet names, headers, row counts and inferred column types are returned.
    Cell values are deliberately omitted so this profile is safe to show in the
    rule-review workspace and useful as context for rule understanding.
    """

    safe_name, sheets = read_tabular_sheets(
        content,
        filename,
        max_bytes=max_bytes,
        max_rows=max_rows,
    )
    profiles: list[dict[str, Any]] = []
    warnings: list[str] = []
    total_rows = 0
    ignored_sheet_names = {"仪表板", "填写说明", "说明", "使用说明"}
    for sheet_name, raw_rows in sheets.items():
        first_non_blank = next(
            (
                index
                for index, row in enumerate(raw_rows)
                if any(not _is_blank(value) for value in row)
            ),
            None,
        )
        display_name = sheet_name or "CSV"
        if first_non_blank is None:
            warnings.append(f"工作表“{display_name}”为空，未作为底表数据")
            continue
        header_row = raw_rows[first_non_blank]
        last_header_index = max(
            (
                index
                for index, value in enumerate(header_row)
                if not _is_blank(value)
            ),
            default=-1,
        )
        headers = [
            "" if _is_blank(value) else str(value).strip()
            for value in header_row[: last_header_index + 1]
        ]
        data_rows = [
            list(row[: len(headers)])
            for row in raw_rows[first_non_blank + 1 :]
            if any(not _is_blank(value) for value in row[: len(headers)])
        ]
        non_blank_headers = [header for header in headers if header]
        looks_like_display_sheet = (
            display_name.strip() in ignored_sheet_names
            or len(non_blank_headers) < 2
        )
        if looks_like_display_sheet:
            reason = (
                "是说明或展示页"
                if display_name.strip() in ignored_sheet_names
                else "没有形成至少两列的有效表头"
            )
            if data_rows:
                warnings.append(
                    f"工作表“{display_name}”{reason}，未作为底表数据"
                )
            elif non_blank_headers:
                warnings.append(
                    f"工作表“{display_name}”只有表头、没有数据行"
                )
            else:
                warnings.append(f"工作表“{display_name}”为空")
            continue
        if not data_rows:
            warnings.append(f"工作表“{display_name}”只有表头、没有数据行")
            continue
        duplicate_headers = sorted(
            {
                header
                for header in non_blank_headers
                if non_blank_headers.count(header) > 1
            }
        )
        if duplicate_headers:
            warnings.append(
                f"工作表“{display_name}”存在重复表头："
                f"{'、'.join(duplicate_headers[:10])}"
            )
        blank_positions = [
            str(index + 1) for index, header in enumerate(headers) if not header
        ]
        if blank_positions:
            warnings.append(
                f"工作表“{display_name}”第"
                f"{'、'.join(blank_positions[:10])}列没有表头"
            )
        formula_count = sum(
            isinstance(value, FormulaCell)
            for row in data_rows
            for value in row
        )
        if formula_count:
            warnings.append(
                f"工作表“{display_name}”包含{formula_count}个公式单元格；"
                "正式校验时必须使用计算后的值"
            )
        columns: list[dict[str, Any]] = []
        for column_index, header in enumerate(headers):
            if not header:
                continue
            samples: list[Any] = []
            for row in data_rows:
                if (
                    column_index < len(row)
                    and not _is_blank(row[column_index])
                    and not isinstance(row[column_index], FormulaCell)
                ):
                    samples.append(row[column_index])
                    if len(samples) == 50:
                        break
            columns.append(
                {
                    "position": column_index + 1,
                    "header": header,
                    "inferred_type": _infer_source_column_type(samples),
                    "non_blank_sample_count": len(samples),
                }
            )
        profiles.append(
            {
                "sheet_name": display_name,
                "header_row": first_non_blank + 1,
                "row_count": len(data_rows),
                "column_count": len(columns),
                "columns": columns,
                "formula_cell_count": formula_count,
            }
        )
        total_rows += len(data_rows)

    if not profiles and not warnings:
        warnings.append("文件没有可识别的数据工作表")
    return {
        "file_name": safe_name,
        "sheet_count": len(sheets),
        "usable_sheet_count": len(profiles),
        "total_data_rows": total_rows,
        "sheets": profiles,
        "warnings": warnings,
    }


def _infer_source_column_type(values: list[Any]) -> str:
    if not values:
        return "unknown"
    if all(isinstance(value, (date, datetime)) for value in values):
        return "date"
    if all(isinstance(value, bool) for value in values):
        return "boolean"
    if all(
        isinstance(value, int) and not isinstance(value, bool)
        for value in values
    ):
        return "integer"
    if all(
        isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)
        for value in values
    ):
        return "decimal"
    texts = [str(value).strip() for value in values]
    if texts and all(
        re.fullmatch(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}", value)
        for value in texts
    ):
        return "date"
    if texts and all(re.fullmatch(r"[-+]?\d+", value) for value in texts):
        return "integer"
    if texts and all(
        re.fullmatch(r"[-+]?(?:\d+(?:\.\d+)?|\.\d+)%?", value.replace(",", ""))
        for value in texts
    ):
        return "percentage" if all(value.endswith("%") for value in texts) else "decimal"
    return "text"


def _read_xlsx(
    content: bytes,
    *,
    max_rows: int,
    required_headers: set[str],
    preferred_sheet_name: str = "",
) -> tuple[str, list[list[Any]]]:
    sheets = _read_xlsx_sheets(content, max_rows=max_rows)
    ignored_sheet_names = {"填写说明", "仪表板", "说明", "使用说明"}
    data_sheets: dict[str, list[list[Any]]] = {}
    matching_sheets: dict[str, list[list[Any]]] = {}
    for name, rows in sheets.items():
        if name.strip() in ignored_sheet_names:
            continue
        first_non_blank = next(
            (
                index
                for index, row in enumerate(rows)
                if any(not _is_blank(value) for value in row)
            ),
            None,
        )
        if first_non_blank is None:
            continue
        normalized_rows = rows[first_non_blank:]
        data_sheets[name] = normalized_rows
        headers = {
            str(value or "").strip()
            for value in normalized_rows[0]
            if str(value or "").strip()
        }
        if required_headers <= headers:
            matching_sheets[name] = normalized_rows
    preferred = str(preferred_sheet_name or "").strip()
    if preferred and preferred in matching_sheets:
        selected = (preferred, matching_sheets[preferred])
    elif len(matching_sheets) == 1:
        selected = next(iter(matching_sheets.items()))
    elif len(matching_sheets) > 1:
        raise FileValidationError(
            "多个工作表都符合规则表头，请在规则中明确工作表名称"
        )
    elif len(data_sheets) == 1:
        selected = next(iter(data_sheets.items()))
    else:
        raise FileValidationError(
            "没有找到唯一符合规则表头的数据工作表"
        )
    sheet_name, rows = selected
    if any(isinstance(value, FormulaCell) for row in rows for value in row):
        raise FileValidationError(
            f"数据工作表“{sheet_name}”包含公式，请上传公式计算后的值"
        )
    return sheet_name, rows


def _read_xlsx_sheets(
    content: bytes,
    *,
    max_rows: int,
) -> dict[str, list[list[Any]]]:
    if not content.startswith(b"PK") or not zipfile.is_zipfile(io.BytesIO(content)):
        raise FileValidationError("文件内容与扩展名不一致")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        if len(members) > _MAX_XLSX_ENTRIES:
            raise FileValidationError("Excel 工作簿内部文件数量超过安全限制")
        if sum(member.file_size for member in members) > _MAX_XLSX_UNCOMPRESSED_BYTES:
            raise FileValidationError("Excel 工作簿解压后超过安全限制")
        if any(member.flag_bits & 0x1 for member in members):
            raise FileValidationError("不支持加密的 Excel 工作簿")
        lower_names = {member.filename.lower() for member in members}
        if any(name.endswith("vbaproject.bin") for name in lower_names):
            raise FileValidationError("不支持包含宏的工作簿")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise FileValidationError("Excel 工作簿损坏或无法读取") from exc
    try:
        output_sheets: dict[str, list[list[Any]]] = {}
        total_rows = 0
        total_cells = 0
        for name in workbook.sheetnames:
            sheet = workbook[name]
            # Some ERP exports incorrectly declare the worksheet dimension as A1
            # even though the XML contains thousands of rows and columns.
            if sheet.calculate_dimension() == "A1:A1":
                sheet.reset_dimensions()
            rows: list[list[Any]] = []
            for values in sheet.iter_rows(values_only=False):
                total_rows += 1
                if total_rows > max_rows + 1:
                    raise FileValidationError(f"文件超过单次 {max_rows} 行限制")
                if len(values) > _MAX_XLSX_COLUMNS:
                    raise FileValidationError(
                        f"文件超过单表 {_MAX_XLSX_COLUMNS} 列安全限制"
                    )
                total_cells += len(values)
                if total_cells > _MAX_XLSX_CELLS:
                    raise FileValidationError("Excel 工作簿总单元格数量超过安全限制")
                output: list[Any] = []
                for cell in values:
                    if cell.data_type == "f":
                        output.append(FormulaCell(str(cell.value or "")))
                    else:
                        output.append(cell.value)
                rows.append(output)
            output_sheets[name] = rows
        return output_sheets
    finally:
        workbook.close()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def normalize_cell_value(value: Any, column: TableColumn) -> Any:
    """Normalize one uploaded cell with the same rules as standard templates."""

    return _normalize(value, column)


def _normalize(value: Any, column: TableColumn) -> Any:
    if column.data_type == "text":
        return re.sub(r"\s+", " ", str(value).strip())
    if column.data_type == "date":
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()
        for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日"):
            try:
                # The parsed value is a calendar date, so a timezone is intentionally absent.
                return datetime.strptime(text, pattern).date()  # noqa: DTZ007
            except ValueError:
                pass
        raise FileValidationError(f"{column.name}日期格式不正确")
    if column.data_type == "integer":
        try:
            decimal = Decimal(str(value).strip())
        except InvalidOperation:
            raise FileValidationError(f"{column.name}必须是整数") from None
        if not decimal.is_finite():
            raise FileValidationError(f"{column.name}必须是有限整数")
        if decimal != decimal.to_integral_value():
            raise FileValidationError(f"{column.name}必须是整数")
        return int(decimal)
    if column.data_type in {"decimal", "percentage"}:
        text = str(value).strip().replace(",", "")
        percent = text.endswith("%")
        if percent:
            text = text[:-1].strip()
        try:
            result = Decimal(text)
        except InvalidOperation:
            raise FileValidationError(
                f"{column.name}{'比例' if column.data_type == 'percentage' else '金额'}格式不正确"
            ) from None
        if not result.is_finite():
            raise FileValidationError(f"{column.name}必须是有限数值")
        if column.data_type == "percentage":
            if percent:
                result /= Decimal(100)
            if result < 0 or result > 1:
                raise FileValidationError(f"{column.name}比例必须在0%到100%之间")
        return result
    raise FileValidationError(f"{column.name}的数据类型不受支持")


def _error_code(data_type: ColumnType) -> str:
    return {
        "date": "invalid_date",
        "decimal": "invalid_decimal",
        "percentage": "invalid_percentage",
        "integer": "invalid_integer",
        "text": "invalid_text",
    }[data_type]


def build_error_workbook(parsed: ParsedTable, schema: TableSchema) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "错误数据"
    sheet.append(
        [safe_excel_cell(column.name) for column in schema.columns]
        + ["原始行号", "错误原因"]
    )
    for row in parsed.rows:
        if not row.errors:
            continue
        sheet.append(
            [safe_excel_cell(row.raw.get(column.name)) for column in schema.columns]
            + [
                row.row_number,
                safe_excel_cell("；".join(error.message for error in row.errors)),
            ]
        )
    sheet.freeze_panes = "A2"
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def safe_excel_cell(value: object) -> object:
    """Keep uploaded text as text when a downloaded workbook is opened."""

    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value
