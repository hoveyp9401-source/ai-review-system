from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
from dataclasses import asdict
from datetime import date, datetime, time, timezone
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from typing import Any, Literal

from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import select, text
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.ext.asyncio import AsyncSession

from app.agent2.business.models import (
    Agent2Case,
    Agent2IdentityBinding,
    CaseLifecycleState,
    CaseProgress,
)
from app.config import Settings
from app.legal_ops_data_intake.calculator import (
    CalculationBlocked,
    DeterministicCalculator,
    aggregate_row_relevant_to_any_metric,
    aggregate_table_row_key,
    validate_rule_spec,
)
from app.legal_ops_data_intake.case_progress_profiles import (
    CaseProgressSourceProfileError,
    parse_case_progress_source_file,
)
from app.legal_ops_data_intake.case_source_profiles import (
    CaseSourceProfileError,
    parse_case_master_source_file,
)
from app.legal_ops_data_intake.cases import (
    CaseMasterIndex,
    CaseProgressIndex,
    ImportErrorDetail,
    ImportPreview,
    PreviewRow,
    preview_case_master,
    preview_case_progress,
)
from app.legal_ops_data_intake.column_mapping import (
    mapping_payload,
    propose_column_mapping,
    validate_column_mapping,
)
from app.legal_ops_data_intake.performance_reporting import (
    DefendantPerformanceReport,
    PerformanceReportError,
    build_defendant_performance_report,
    export_defendant_performance_docx,
    export_defendant_performance_xlsx,
)
from app.legal_ops_data_intake.rule_literals import extract_safe_rule_literals
from app.legal_ops_data_intake.rule_package import (
    RulePackageError,
    inspect_rule_package,
)
from app.legal_ops_data_intake.rule_understanding import RuleDraftInterpreter
from app.legal_ops_data_intake.schemas import (
    CASE_MASTER_SCHEMA,
    CASE_PROGRESS_SCHEMA,
    performance_table_schema,
    schema_to_dict,
)
from app.legal_ops_data_intake.storage import (
    FileStorageError,
    IntakeFileStore,
    StoredFile,
)
from app.legal_ops_data_intake.workbook import (
    FileValidationError,
    ParsedRow,
    ParsedTable,
    RowError,
    inspect_tabular_source,
    parse_tabular_file,
    safe_excel_cell,
)
from app.models import Team


class DataIntakeError(HTTPException):
    def __init__(
        self, message: str, *, code: str = "intake_error", status_code: int = 400
    ):
        super().__init__(
            status_code=status_code,
            detail={"code": code, "message": message},
        )
        self.message = message
        self.code = code

    def __str__(self) -> str:
        return self.message


BUSINESS_LABELS = {
    "performance_rule_package": "绩效规则包",
    "performance_source_table": "绩效源表",
    "case_master": "案件主表",
    "case_progress": "案件进展",
}
STATUS_LABELS = {
    "uploaded": "已上传",
    "validating": "校验中",
    "ready": "待发布",
    "validation_failed": "存在错误",
    "pending_publish": "待发布",
    "published": "已发布",
    "abandoned": "已放弃",
    "failed": "处理失败",
}
ACTION_LABELS = {
    "新增": "新增",
    "更新": "更新",
    "无变化": "无变化",
    "重复跳过": "重复跳过",
    "冲突": "冲突",
    "失败": "失败",
    "待核验": "待核验",
}
PERFORMANCE_VALIDATION_SEMANTICS_VERSION = "2"


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _batch_no(prefix: str) -> str:
    stamp = _utcnow().strftime("%Y%m%d%H%M%S")
    return f"{prefix}-{stamp}-{uuid.uuid4().hex[:8].upper()}"


def _stable_uuid(*parts: str) -> uuid.UUID:
    return uuid.uuid5(uuid.NAMESPACE_URL, "|".join(parts))


def _jsonable(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, uuid.UUID):
        return str(value)
    if isinstance(value, dict):
        return {str(key): _jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_jsonable(item) for item in value]
    return value


def _json(value: Any) -> str:
    return json.dumps(_jsonable(value), ensure_ascii=False, sort_keys=True)


def _hash_json(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(
            _jsonable(value),
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()


def _header_mapping_hash(mapping: dict[str, str] | None) -> str:
    return _hash_json(
        {
            str(key): str(value)
            for key, value in sorted((mapping or {}).items())
            if str(key).strip() and str(value).strip()
        }
    )


def _batch_matches_header_mapping(
    metadata: dict[str, Any] | None,
    mapping_hash: str,
) -> bool:
    """Legacy batches had no mapping hash and must be revalidated once."""

    stored_hash = str((metadata or {}).get("header_mapping_hash") or "")
    return bool(stored_hash) and stored_hash == mapping_hash


def _batch_matches_validation_semantics(
    metadata: dict[str, Any] | None,
) -> bool:
    return (
        str((metadata or {}).get("validation_semantics_version") or "")
        == PERFORMANCE_VALIDATION_SEMANTICS_VERSION
    )


def _aggregate_protected_missing_fields(
    rule_spec: dict[str, Any],
    table: dict[str, Any],
) -> set[str]:
    table_key = str(table.get("key") or "")
    protected = {aggregate_table_row_key(rule_spec, table)}
    subject = rule_spec.get("subject") or {}
    direct_field = str((subject.get("fields") or {}).get(table_key) or "")
    if direct_field.startswith(f"{table_key}."):
        protected.add(direct_field.split(".", 1)[1])
    lookup = (subject.get("lookups") or {}).get(table_key) or {}
    protected.add(str(lookup.get("source_field") or ""))
    return {field for field in protected if field}


def _downgrade_irrelevant_aggregate_missing_errors(
    parsed: ParsedTable,
    *,
    rule_spec: dict[str, Any],
    table: dict[str, Any],
    period_start: date | str | None,
    period_end: date | str | None,
) -> ParsedTable:
    """Keep out-of-period data-quality gaps visible without blocking calculation."""

    table_key = str(table.get("key") or "")
    protected = _aggregate_protected_missing_fields(rule_spec, table)
    column_labels = {
        str(column.get("key") or ""): str(
            column.get("name") or column.get("key") or "字段"
        )
        for column in table.get("columns") or []
        if isinstance(column, dict)
    }
    for row in parsed.rows:
        if not row.errors:
            continue
        updated: list[RowError] = []
        for error in row.errors:
            if error.code != "required" or error.field in protected:
                updated.append(error)
                continue
            try:
                blocks_metric = aggregate_row_relevant_to_any_metric(
                    rule_spec,
                    table_key=table_key,
                    field_key=error.field,
                    row=row.normalized,
                    period_start=period_start,
                    period_end=period_end,
                )
            except CalculationBlocked:
                blocks_metric = True
            updated.append(
                error
                if blocks_metric
                else RowError(
                    "missing_out_of_scope",
                    error.field,
                    (
                        f"{column_labels.get(error.field, error.field)}为空；"
                        "依照 Skill 计算条件，该行不会进入依赖此字段的"
                        "本周期指标，已保留为数据提醒"
                    ),
                    critical=False,
                )
            )
        row.errors = updated
    return parsed


def _rule_draft_section_items(draft: dict[str, Any]) -> dict[str, list[str]]:
    def objects(key: str) -> list[dict[str, Any]]:
        values = draft.get(key) or []
        return [dict(item) for item in values if isinstance(item, dict)]

    source_tables = [
        str(item.get("name") or item.get("key") or "未命名源表")
        for item in objects("source_tables")
    ]
    scopes = [
        str(
            item.get("scope_name")
            or item.get("team_name")
            or item.get("name")
            or "适用范围待确认"
        )
        for item in objects("applicability_scopes")
    ]
    metrics = [
        str(item.get("metric_name") or item.get("name") or "指标名称待确认")
        for item in objects("metric_catalog")
    ]
    targets: list[str] = []
    for item in objects("target_versions"):
        name = str(item.get("metric_name") or item.get("name") or "指标名称待确认")
        scope = str(item.get("scope_name") or item.get("team_name") or "").strip()
        value = str(item.get("target_value") or item.get("value") or "待确认")
        unit = str(item.get("unit") or "")
        comparison = {
            "at_least": "不低于",
            "at_most": "不高于",
            "equal": "等于",
        }.get(str(item.get("comparison") or ""), "")
        period = str(item.get("effective_period") or item.get("period") or "").strip()
        label = f"{name}：{comparison}{value}{unit}"
        if scope:
            label = f"{scope} · {label}"
        if period:
            label = f"{label}（{period}）"
        targets.append(label)
    rules = [
        f"{item['name']}：{item['explanation']}"
        for item in _rule_explanation_items(draft.get("calculation_rules") or [])
    ]
    stable_keys = [
        str(item).strip()
        for item in draft.get("stable_person_keys") or []
        if str(item).strip()
    ]
    unresolved = [
        str(item).strip() for item in draft.get("unresolved") or [] if str(item).strip()
    ]
    executable = (
        _rule_spec_review_items(draft["rule_spec"])
        if isinstance(draft.get("rule_spec"), dict)
        else ["尚未形成"]
    )
    return {
        "源表": source_tables,
        "稳定关联字段": stable_keys,
        "适用板块": scopes,
        "指标目录": metrics,
        "目标值": targets,
        "计算口径": rules,
        "待确认事项": unresolved,
        "固定计算结构": executable,
    }


def _rule_draft_changes(
    before: dict[str, Any],
    after: dict[str, Any],
) -> list[dict[str, Any]]:
    before_sections = _rule_draft_section_items(before)
    after_sections = _rule_draft_section_items(after)
    changes: list[dict[str, Any]] = []
    for section in before_sections:
        before_items = before_sections[section]
        after_items = after_sections[section]
        if before_items == after_items:
            continue
        if section == "固定计算结构":
            before_items, after_items = _rule_spec_change_sides(
                before.get("rule_spec"),
                after.get("rule_spec"),
            )
        changes.append(
            {
                "section": section,
                "before": before_items[:100],
                "after": after_items[:100],
            }
        )
    return changes


def _rule_spec_change_sides(
    before_spec: Any,
    after_spec: Any,
) -> tuple[list[str], list[str]]:
    if not isinstance(before_spec, dict) or not isinstance(after_spec, dict):
        return (
            _rule_spec_review_items(before_spec)
            if isinstance(before_spec, dict)
            else ["尚未形成"],
            _rule_spec_review_items(after_spec)
            if isinstance(after_spec, dict)
            else ["尚未形成"],
        )
    before_items: list[str] = []
    after_items: list[str] = []
    before_lookups = _lookup_step_maps(before_spec)
    after_lookups = _lookup_step_maps(after_spec)
    for identity in sorted(set(before_lookups).union(after_lookups)):
        label = identity[1]
        before_map = before_lookups.get(identity)
        after_map = after_lookups.get(identity)
        if before_map is None:
            before_items.append(f"精确映射 {label}：尚未配置")
            after_items.append(f"精确映射 {label}：新增{len(after_map or {})}条")
            continue
        if after_map is None:
            before_items.append(f"精确映射 {label}：原有{len(before_map)}条")
            after_items.append(f"精确映射 {label}：已移除")
            continue
        changed_keys = [
            key
            for key in sorted(set(before_map).union(after_map))
            if before_map.get(key) != after_map.get(key)
        ]
        for key in changed_keys[:100]:
            before_items.append(f"{label}：{key} → {before_map.get(key, '未配置')}")
            after_items.append(f"{label}：{key} → {after_map.get(key, '已移除')}")
        if len(changed_keys) > 100:
            remaining = len(changed_keys) - 100
            before_items.append(f"{label}：另有{remaining}条变化")
            after_items.append(f"{label}：另有{remaining}条变化")
    for collection_label, before_values, after_values in _rule_collection_changes(
        before_spec,
        after_spec,
    ):
        before_items.append(f"{collection_label}：{'、'.join(before_values) or '无'}")
        after_items.append(f"{collection_label}：{'、'.join(after_values) or '无'}")
    for item_type, before_definitions, after_definitions in (
        (
            "基础指标",
            _keyed_rule_items(before_spec.get("metrics")),
            _keyed_rule_items(after_spec.get("metrics")),
        ),
        (
            "结果公式",
            _keyed_rule_items(before_spec.get("outputs")),
            _keyed_rule_items(after_spec.get("outputs")),
        ),
    ):
        for key in sorted(set(before_definitions).union(after_definitions)):
            before_value = before_definitions.get(key)
            after_value = after_definitions.get(key)
            if _hash_json(before_value) == _hash_json(after_value):
                continue
            before_items.append(
                f"{item_type} {key}："
                + _rule_item_review(before_spec, item_type, before_value)
            )
            after_items.append(
                f"{item_type} {key}："
                + _rule_item_review(after_spec, item_type, after_value)
            )
    if not before_items and _hash_json(before_spec) != _hash_json(after_spec):
        return _rule_spec_review_items(before_spec), _rule_spec_review_items(after_spec)
    return before_items or ["无变化"], after_items or ["无变化"]


def _lookup_step_maps(
    spec: dict[str, Any],
) -> dict[tuple[str, str], dict[str, str]]:
    output: dict[tuple[str, str], dict[str, str]] = {}
    subject = spec.get("subject") or {}
    for table_key, lookup in (subject.get("lookups") or {}).items():
        if not isinstance(lookup, dict):
            continue
        for step in lookup.get("steps") or []:
            if not isinstance(step, dict):
                continue
            name = str(step.get("name") or "未命名映射")
            output[(str(table_key), name)] = {
                str(key): str(value)
                for key, value in (step.get("mapping") or {}).items()
            }
    return output


def _rule_collection_changes(
    before_spec: dict[str, Any],
    after_spec: dict[str, Any],
) -> list[tuple[str, list[str], list[str]]]:
    output = []
    before_subject = before_spec.get("subject") or {}
    after_subject = after_spec.get("subject") or {}
    before_lookups = before_subject.get("lookups") or {}
    after_lookups = after_subject.get("lookups") or {}
    for table_key in sorted(set(before_lookups).union(after_lookups)):
        before_values = sorted(
            str(value)
            for value in (
                (before_lookups.get(table_key) or {}).get(
                    "exclude_source_values",
                    [],
                )
            )
        )
        after_values = sorted(
            str(value)
            for value in (
                (after_lookups.get(table_key) or {}).get(
                    "exclude_source_values",
                    [],
                )
            )
        )
        if before_values != after_values:
            output.append((f"{table_key}明确排除清单", before_values, after_values))
    return output


def _keyed_rule_items(values: Any) -> dict[str, dict[str, Any]]:
    return {
        str(item.get("name") or item.get("key") or "未命名"): dict(item)
        for item in (values or [])
        if isinstance(item, dict)
    }


def _performance_source_state_hash(
    versions: dict[str, Any],
    *,
    rule_version: str,
) -> str:
    """Fingerprint the exact published source set used by an upload preview."""

    return _hash_json(
        {
            "rule_version": rule_version,
            "sources": [
                {
                    "table_key": table_key,
                    "source_version_ref": str(row["source_version_id"]),
                    "version": int(row["version"]),
                    "file_hash": str(row["file_hash"]),
                }
                for table_key, row in sorted(versions.items())
            ],
        }
    )


def _safe_sheet_title(value: object) -> str:
    title = re.sub(r"[\\/*?:\[\]]", " ", str(value or "")).strip().strip("'")
    return (title or "数据模板")[:31]


def _media_type(filename: str) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1]
    return {
        "csv": "text/csv",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "zip": "application/zip",
        "md": "text/markdown",
    }.get(suffix, "application/octet-stream")


class DataIntakeService:
    """Single application seam for the three supported data intake workflows."""

    def __init__(
        self,
        session: AsyncSession,
        settings: Settings,
        *,
        tenant_id: str,
        actor_user_id: str,
        llm_client: Any | None = None,
        code_version: str = "",
    ):
        self.session = session
        self.settings = settings
        self.tenant_id = tenant_id
        self.actor_user_id = actor_user_id
        self.llm_client = llm_client
        self.code_version = code_version or "workspace"
        self.files = IntakeFileStore(
            settings.legal_ops_data_intake_storage_path,
            max_bytes=settings.legal_ops_data_intake_max_file_mb * 1024 * 1024,
        )

    def _store_file(self, content: bytes, filename: str) -> StoredFile:
        try:
            return self.files.store(content, filename)
        except FileStorageError as exc:
            raise DataIntakeError(
                str(exc),
                code="unsafe_file",
                status_code=422,
            ) from exc

    async def _register_original_file(self, stored: StoredFile) -> uuid.UUID:
        """Make every retained raw file traceable before later processing."""

        await self.session.rollback()
        async with self.session.begin():
            return await self._register_file(stored)

    async def dashboard(self) -> dict[str, Any]:
        periods = await self.list_periods()
        batches = await self.list_batches(limit=12)
        errors = await self.list_errors(limit=12)
        active_rule = await self._latest_active_rule()
        return {
            "title": "数据接入中心",
            "scope_notice": "只处理绩效数据包、案件主表和案件进展；不会写回ERP，也不会触发钉钉消息。",
            "rule_state": self._rule_card(active_rule),
            "periods": periods,
            "recent_batches": batches,
            "open_errors": errors,
            "empty_state": not periods and not batches,
        }

    async def create_period(
        self,
        *,
        label: str,
        period_type: str,
        starts_on: date,
        ends_on: date,
    ) -> dict[str, Any]:
        if not label.strip() or len(label.strip()) > 128:
            raise DataIntakeError("考核周期名称不能为空且不能超过128字")
        if period_type not in {"month", "quarter", "half_year", "year", "custom"}:
            raise DataIntakeError("考核周期类型无效")
        if ends_on < starts_on:
            raise DataIntakeError("结束日期不能早于开始日期")
        period_id = uuid.uuid4()
        try:
            async with self.session.begin():
                await self.session.execute(
                    text(
                        """
                        INSERT INTO legal_ops_performance_periods (
                            period_id, tenant_id, label, period_type, starts_on, ends_on,
                            status, created_by
                        ) VALUES (
                            :period_id, :tenant_id, :label, :period_type, :starts_on, :ends_on,
                            'preparing', :actor
                        )
                        """
                    ),
                    {
                        "period_id": period_id,
                        "tenant_id": self.tenant_id,
                        "label": label.strip(),
                        "period_type": period_type,
                        "starts_on": starts_on,
                        "ends_on": ends_on,
                        "actor": self.actor_user_id,
                    },
                )
                await self._audit(
                    "create_period",
                    "performance_period",
                    str(period_id),
                    f"新建考核周期：{label.strip()}",
                    {"starts_on": starts_on, "ends_on": ends_on},
                )
        except Exception as exc:
            if "legal_ops_performance_period_key" in str(exc):
                raise DataIntakeError(
                    "同名考核周期已经存在", code="period_exists"
                ) from exc
            raise
        return {
            "period_ref": str(period_id),
            "label": label.strip(),
            "period_type": period_type,
            "period_type_label": _period_type_label(period_type),
            "starts_on": starts_on.isoformat(),
            "ends_on": ends_on.isoformat(),
            "status": "preparing",
            "status_label": "准备中",
        }

    async def list_periods(self) -> list[dict[str, Any]]:
        result = await self.session.execute(
            text(
                """
                SELECT period_id, label, period_type, starts_on, ends_on, status,
                       created_by, created_at
                FROM legal_ops_performance_periods
                WHERE tenant_id = :tenant_id
                ORDER BY starts_on DESC, created_at DESC
                """
            ),
            {"tenant_id": self.tenant_id},
        )
        items = []
        for row in result.mappings().all():
            active_rules = await self._rules_for_period(row["period_id"])
            item = {
                "period_ref": str(row["period_id"]),
                "label": row["label"],
                "period_type": row["period_type"],
                "period_type_label": _period_type_label(row["period_type"]),
                "starts_on": row["starts_on"].isoformat(),
                "ends_on": row["ends_on"].isoformat(),
                "status": row["status"],
                "status_label": _period_status_label(row["status"]),
                "created_by": row["created_by"],
                "created_at": row["created_at"].isoformat(),
            }
            item["packages"] = [
                await self._period_package(row["period_id"], active_rule)
                for active_rule in active_rules
            ]
            item["package"] = (
                item["packages"][0]
                if len(item["packages"]) == 1
                else _empty_period_package(
                    "请选择一个已启用的负责板块；各板块的规则和底表相互隔离"
                    if item["packages"]
                    else "规则文件尚未配置，暂不可上传绩效源表或试算"
                )
            )
            items.append(item)
        return items

    async def upload_rule_package(
        self,
        *,
        content: bytes,
        filename: str,
        period_ref: str | None,
        business_scope_name: str = "",
    ) -> dict[str, Any]:
        try:
            inspected = inspect_rule_package(content, filename)
        except RulePackageError as exc:
            raise DataIntakeError(
                str(exc),
                code="invalid_rule_package",
                status_code=422,
            ) from exc
        stored = self._store_file(content, filename)
        await self._register_original_file(stored)
        period_id = self._uuid(period_ref, "考核周期") if period_ref else None
        business_scope_name = (
            str(business_scope_name or "").strip() or inspected.skill_name
        )
        scope_key = _business_scope_key(business_scope_name)
        async with self.session.begin():
            period = None
            if period_id:
                period = await self._require_period(period_id)
            file_id = await self._register_file(stored)
            duplicate = await self._find_duplicate_batch(
                business_type="performance_rule_package",
                data_source="Workbuddy技能包",
                file_hash=stored.file_hash,
                period_id=period_id,
                table_key=scope_key,
                import_mode="replace",
            )
            if duplicate:
                return {**self._batch_payload(duplicate), "duplicate_upload": True}
            batch_id = uuid.uuid4()
            batch_no = _batch_no("RULE")
            await self._insert_batch(
                batch_id=batch_id,
                batch_no=batch_no,
                business_type="performance_rule_package",
                data_source="Workbuddy技能包",
                file_id=file_id,
                stored=stored,
                period_id=period_id,
                table_key=scope_key,
                import_mode="replace",
                status="ready",
                original_rows=0,
                counts={},
                metadata={
                    "skill_name": inspected.skill_name,
                    "skill_version": inspected.skill_version,
                    "skill_path": inspected.skill_path,
                    "references": list(inspected.references),
                    "reference_documents": [
                        {
                            "file_name": PurePosixPath(item.path).name,
                            "path": item.path,
                            "file_hash": item.file_hash,
                            "paragraph_count": item.paragraph_count,
                        }
                        for item in inspected.reference_documents
                    ],
                    "blocking_reasons": list(inspected.blocking_reasons),
                    "business_scope_name": business_scope_name,
                },
            )
            rule_version_id = uuid.uuid4()
            rule_version = (
                f"{inspected.skill_version}-{inspected.package_hash[:8]}"
                if inspected.skill_version != "未声明"
                else f"skill-{inspected.package_hash[:12]}"
            )
            if scope_key:
                rule_version = f"{rule_version}-{scope_key.removeprefix('scope-')[:8]}"
            rule_spec_hash = (
                _hash_json(inspected.rule_spec) if inspected.rule_spec else ""
            )
            await self.session.execute(
                text(
                    """
                    INSERT INTO legal_ops_performance_rule_versions (
                        rule_version_id, tenant_id, period_id, source_batch_id,
                        rule_version, skill_name, skill_version, source_skill_path,
                        source_file_hash, skill_markdown, package_inventory_json,
                        rule_spec_json, rule_spec_hash, rule_summary,
                        applicable_period, business_scope_name, business_scope_key,
                        status, source_kind,
                        code_version, created_by
                    ) VALUES (
                        :rule_id, :tenant_id, :period_id, :batch_id,
                        :rule_version, :skill_name, :skill_version, :skill_path,
                        :file_hash, :skill_markdown, CAST(:inventory AS jsonb),
                        CAST(:rule_spec AS jsonb), :rule_spec_hash, :rule_summary,
                        :applicable_period, :business_scope_name, :business_scope_key,
                        'uploaded',
                        'workbuddy_skill', :code_version, :actor
                    )
                    """
                ),
                {
                    "rule_id": rule_version_id,
                    "tenant_id": self.tenant_id,
                    "period_id": period_id,
                    "batch_id": batch_id,
                    "rule_version": rule_version,
                    "skill_name": inspected.skill_name,
                    "skill_version": inspected.skill_version,
                    "skill_path": inspected.skill_path,
                    "file_hash": inspected.package_hash,
                    "skill_markdown": inspected.skill_markdown,
                    "inventory": _json(
                        {
                            "files": list(inspected.references),
                            "reference_documents": [
                                {
                                    "file_name": PurePosixPath(item.path).name,
                                    "path": item.path,
                                    "file_hash": item.file_hash,
                                    "paragraph_count": item.paragraph_count,
                                }
                                for item in inspected.reference_documents
                            ],
                        }
                    ),
                    "rule_spec": _json(inspected.rule_spec)
                    if inspected.rule_spec
                    else None,
                    "rule_spec_hash": rule_spec_hash,
                    "rule_summary": "Workbuddy 技能包已上传，等待规则理解与人工确认",
                    "applicable_period": (
                        str(period["label"]) if period is not None else "通用规则"
                    ),
                    "business_scope_name": business_scope_name,
                    "business_scope_key": scope_key,
                    "code_version": self.code_version,
                    "actor": self.actor_user_id,
                },
            )
            await self._audit(
                "upload_rule_package",
                "performance_rule",
                rule_version,
                f"上传绩效规则包：{stored.safe_file_name}",
                {
                    "batch_no": batch_no,
                    "file_hash": stored.file_hash,
                    "activation_ready": False,
                    "business_scope_name": business_scope_name,
                },
                batch_id=batch_id,
            )
        return {
            "batch_no": batch_no,
            "rule_ref": str(rule_version_id),
            "rule_version": rule_version,
            "skill_name": inspected.skill_name,
            "skill_version": inspected.skill_version,
            "file_hash": inspected.package_hash,
            "status": "uploaded",
            "status_label": "已上传，待理解",
            "activation_ready": False,
            "business_scope_name": business_scope_name,
            "blocking_reasons": list(inspected.blocking_reasons),
            "reference_documents": [
                {
                    "file_name": PurePosixPath(item.path).name,
                    "file_hash": item.file_hash,
                    "paragraph_count": item.paragraph_count,
                }
                for item in inspected.reference_documents
            ],
            "duplicate_upload": False,
        }

    async def interpret_rule(self, rule_ref: str) -> dict[str, Any]:
        rule_id = self._uuid(rule_ref, "规则")
        row = await self._rule_row(rule_id)
        self._require_rule_interpretable(row)
        inspected = await self._inspect_stored_rule_package(row)
        interpreter = RuleDraftInterpreter(
            self.llm_client,
            model=self.settings.legal_ops_rule_understanding_model,
            timeout_seconds=self.settings.legal_ops_rule_understanding_timeout_seconds,
            enabled=self.settings.legal_ops_rule_understanding_enabled,
            fallback_model=self.settings.llm_extract_model,
            long_document_timeout_seconds=(
                self.settings.legal_ops_rule_understanding_long_document_timeout_seconds
            ),
        )
        draft = await interpreter.interpret(
            inspected,
            data_profile=await self._rule_source_file_context(rule_id),
        )
        status_value = (
            "awaiting_confirmation" if draft.activation_ready else "understanding_draft"
        )
        await self.session.rollback()
        async with self.session.begin():
            locked_row = await self._rule_row(rule_id, for_update=True)
            self._require_rule_interpretable(locked_row)
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_performance_rule_versions
                    SET understanding_json = CAST(:understanding AS jsonb),
                        understanding_hash = :understanding_hash,
                        rule_spec_json = CAST(:rule_spec AS jsonb),
                        rule_spec_hash = :rule_spec_hash,
                        rule_summary = :rule_summary,
                        status = :status,
                        interpreted_by = :actor,
                        interpreted_at = now()
                    WHERE tenant_id = :tenant_id AND rule_version_id = :rule_id
                    """
                ),
                {
                    "understanding": _json(draft.as_dict()),
                    "understanding_hash": draft.draft_hash,
                    "rule_spec": _json(draft.rule_spec) if draft.rule_spec else None,
                    "rule_spec_hash": _hash_json(draft.rule_spec)
                    if draft.rule_spec
                    else "",
                    "rule_summary": (
                        f"识别出{len(draft.source_tables)}张源表、"
                        f"{len(draft.outputs)}项结果，"
                        f"仍有{len(draft.unresolved)}项待确认"
                    ),
                    "status": status_value,
                    "actor": self.actor_user_id,
                    "tenant_id": self.tenant_id,
                    "rule_id": rule_id,
                },
            )
            await self._audit(
                "interpret_rule",
                "performance_rule",
                str(rule_id),
                "生成规则理解草稿，等待人工核对",
                {
                    "draft_hash": draft.draft_hash,
                    "activation_ready": draft.activation_ready,
                    "unresolved_count": len(draft.unresolved),
                },
                batch_id=locked_row["source_batch_id"],
            )
        return {
            "rule_ref": str(rule_id),
            "rule_version": row["rule_version"],
            "status": status_value,
            "status_label": "待人工确认" if draft.activation_ready else "仍有待确认项",
            "understanding_hash": draft.draft_hash,
            "draft": draft.as_dict(),
            "activation_ready": draft.activation_ready,
            "notice": "这只是规则理解草稿；尚未启用，也没有计算任何绩效结果。",
        }

    async def list_rule_amendments(
        self,
        rule_ref: str,
    ) -> list[dict[str, Any]]:
        rule_id = self._uuid(rule_ref, "规则")
        await self._rule_row(rule_id)
        result = await self.session.execute(
            text(
                """
                SELECT *
                FROM legal_ops_performance_rule_amendments
                WHERE tenant_id = :tenant_id AND rule_version_id = :rule_id
                ORDER BY amendment_no DESC
                """
            ),
            {"tenant_id": self.tenant_id, "rule_id": rule_id},
        )
        return [self._rule_amendment_payload(row) for row in result.mappings().all()]

    async def create_rule_working_copy(
        self,
        rule_ref: str,
    ) -> dict[str, Any]:
        """Create or reuse an editable successor without mutating the active rule."""

        rule_id = self._uuid(rule_ref, "规则")
        created = False
        async with self.session.begin():
            active = await self._rule_row(rule_id, for_update=True)
            if not (bool(active["is_current"]) and str(active["status"]) == "active"):
                if not bool(active["is_current"]) and str(active["status"]) in {
                    "understanding_draft",
                    "awaiting_confirmation",
                }:
                    return {
                        "rule_ref": str(active["rule_version_id"]),
                        "rule_version": str(active["rule_version"]),
                        "status": str(active["status"]),
                        "status_label": _rule_status_label(str(active["status"])),
                        "understanding_hash": str(active["understanding_hash"] or ""),
                        "working_copy_created": False,
                    }
                raise DataIntakeError(
                    "只有当前生效规则可以建立修改草稿",
                    code="rule_working_copy_unavailable",
                    status_code=409,
                )

            lock_key = f"performance-rule-working-copy|{self.tenant_id}|{rule_id}"
            await self.session.execute(
                text("SELECT pg_advisory_xact_lock(hashtextextended(:lock_key, 0))"),
                {"lock_key": lock_key},
            )
            existing = (
                (
                    await self.session.execute(
                        text(
                            """
                            SELECT *
                            FROM legal_ops_performance_rule_versions
                            WHERE tenant_id = :tenant_id
                              AND source_kind = 'active_working_copy'
                              AND source_batch_id = :source_batch_id
                              AND period_id IS NOT DISTINCT FROM :period_id
                              AND business_scope_key = :business_scope_key
                              AND status IN (
                                  'understanding_draft',
                                  'awaiting_confirmation'
                              )
                              AND NOT is_current
                            ORDER BY created_at DESC
                            LIMIT 1
                            FOR UPDATE
                            """
                        ),
                        {
                            "tenant_id": self.tenant_id,
                            "source_batch_id": active["source_batch_id"],
                            "period_id": active["period_id"],
                            "business_scope_key": active["business_scope_key"],
                        },
                    )
                )
                .mappings()
                .first()
            )
            if existing:
                working = existing
            else:
                created = True
                working_id = uuid.uuid4()
                stamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
                suffix = f"-修改稿-{stamp}-{working_id.hex[:6]}"
                base_version = str(active["rule_version"] or "规则")
                working_version = f"{base_version[: 128 - len(suffix)]}{suffix}"
                understanding = dict(active["understanding_json"] or {})
                status_value = (
                    "awaiting_confirmation"
                    if bool(understanding.get("activation_ready", False))
                    else "understanding_draft"
                )
                await self.session.execute(
                    text(
                        """
                        INSERT INTO legal_ops_performance_rule_versions (
                            rule_version_id, tenant_id, period_id,
                            source_batch_id, rule_version, skill_name,
                            skill_version, source_skill_path,
                            source_file_hash, skill_markdown,
                            package_inventory_json, understanding_json,
                            understanding_hash, rule_spec_json,
                            rule_spec_hash, rule_summary, applicable_period,
                            business_scope_name, business_scope_key,
                            draft_revision, status, source_kind,
                            code_version, created_by, interpreted_by,
                            interpreted_at, is_current
                        ) VALUES (
                            :working_id, :tenant_id, :period_id,
                            :source_batch_id, :rule_version, :skill_name,
                            :skill_version, :source_skill_path,
                            :source_file_hash, :skill_markdown,
                            CAST(:package_inventory AS jsonb),
                            CAST(:understanding AS jsonb),
                            :understanding_hash, CAST(:rule_spec AS jsonb),
                            :rule_spec_hash, :rule_summary,
                            :applicable_period, :business_scope_name,
                            :business_scope_key, :draft_revision, :status,
                            'active_working_copy', :code_version, :actor,
                            :actor, now(), false
                        )
                        """
                    ),
                    {
                        "working_id": working_id,
                        "tenant_id": self.tenant_id,
                        "period_id": active["period_id"],
                        "source_batch_id": active["source_batch_id"],
                        "rule_version": working_version,
                        "skill_name": active["skill_name"],
                        "skill_version": active["skill_version"],
                        "source_skill_path": active["source_skill_path"],
                        "source_file_hash": active["source_file_hash"],
                        "skill_markdown": active["skill_markdown"],
                        "package_inventory": _json(
                            active["package_inventory_json"] or {}
                        ),
                        "understanding": _json(understanding),
                        "understanding_hash": active["understanding_hash"],
                        "rule_spec": (
                            _json(active["rule_spec_json"])
                            if active["rule_spec_json"]
                            else None
                        ),
                        "rule_spec_hash": active["rule_spec_hash"],
                        "rule_summary": active["rule_summary"],
                        "applicable_period": active["applicable_period"],
                        "business_scope_name": active["business_scope_name"],
                        "business_scope_key": active["business_scope_key"],
                        "draft_revision": int(active["draft_revision"] or 1),
                        "status": status_value,
                        "code_version": self.code_version,
                        "actor": self.actor_user_id,
                    },
                )
                await self.session.execute(
                    text(
                        """
                        INSERT INTO legal_ops_performance_rule_source_files (
                            source_file_id, tenant_id, rule_version_id,
                            file_id, source_file_no, business_label,
                            inspection_json, status, created_by
                        )
                        SELECT
                            gen_random_uuid(), tenant_id, :working_id,
                            file_id, source_file_no, business_label,
                            inspection_json, 'active', :actor
                        FROM legal_ops_performance_rule_source_files
                        WHERE tenant_id = :tenant_id
                          AND rule_version_id = :active_id
                          AND status = 'active'
                        ORDER BY source_file_no
                        """
                    ),
                    {
                        "working_id": working_id,
                        "tenant_id": self.tenant_id,
                        "active_id": rule_id,
                        "actor": self.actor_user_id,
                    },
                )
                await self._audit(
                    "create_rule_working_copy",
                    "performance_rule",
                    str(working_id),
                    "基于当前生效规则建立修改草稿",
                    {
                        "base_rule_ref": str(rule_id),
                        "base_rule_version": str(active["rule_version"]),
                        "working_rule_version": working_version,
                        "active_rule_unchanged": True,
                    },
                    batch_id=active["source_batch_id"],
                )
                working = await self._rule_row(working_id)

        return {
            "rule_ref": str(working["rule_version_id"]),
            "rule_version": str(working["rule_version"]),
            "status": str(working["status"]),
            "status_label": _rule_status_label(str(working["status"])),
            "understanding_hash": str(working["understanding_hash"] or ""),
            "working_copy_created": created,
            "notice": (
                "已建立新的修改草稿；当前生效规则没有变化"
                if created
                else "继续使用尚未确认的修改草稿；当前生效规则没有变化"
            ),
        }

    async def save_rule_draft(
        self,
        rule_ref: str,
        *,
        expected_understanding_hash: str,
        business_scope_name: str,
        rule_summary: str,
        applicable_period: str,
        applicability_scopes: list[dict[str, Any]],
        target_versions: list[dict[str, Any]],
        issue_resolutions: list[dict[str, str]],
    ) -> dict[str, Any]:
        rule_id = self._uuid(rule_ref, "规则")
        async with self.session.begin():
            row = await self._rule_row(rule_id, for_update=True)
            self._require_rule_draft_mutable(row)
            if str(row["understanding_hash"] or "") != expected_understanding_hash:
                raise DataIntakeError(
                    "规则草稿已经变化，请刷新后再保存",
                    code="stale_rule_draft",
                    status_code=409,
                )
            before = dict(row["understanding_json"] or {})
            after = json.loads(json.dumps(before, ensure_ascii=False))
            normalized_scopes = [
                {
                    key: value
                    for key, value in {
                        "scope_name": str(item.get("scope_name") or "").strip(),
                        "scope_key": str(item.get("scope_key") or "").strip(),
                        "description": str(item.get("description") or "").strip(),
                    }.items()
                    if value
                }
                for item in applicability_scopes
                if str(item.get("scope_name") or "").strip()
            ]
            normalized_targets = [
                {
                    key: value
                    for key, value in {
                        "metric_name": str(item.get("metric_name") or "").strip(),
                        "scope_name": str(item.get("scope_name") or "").strip(),
                        "target_value": str(item.get("target_value") or "").strip(),
                        "unit": str(item.get("unit") or "").strip(),
                        "comparison": str(item.get("comparison") or "").strip(),
                        "effective_period": str(
                            item.get("effective_period") or ""
                        ).strip(),
                    }.items()
                    if value
                }
                for item in target_versions
                if str(item.get("metric_name") or "").strip()
            ]
            normalized_resolutions = [
                {
                    "issue": str(item.get("issue") or "").strip(),
                    "response": str(item.get("response") or "").strip(),
                }
                for item in issue_resolutions
                if str(item.get("issue") or "").strip()
                and str(item.get("response") or "").strip()
            ]
            # Targets are deliberately maintained outside the fixed arithmetic
            # program. Changing a target only changes completion comparison, so it
            # is safe to edit directly. Scope changes may alter which records
            # participate and must still be reconciled into the fixed rule.
            structures_changed = normalized_scopes != list(
                before.get("applicability_scopes") or []
            )
            after["applicability_scopes"] = normalized_scopes
            after["target_versions"] = normalized_targets
            after["issue_resolutions"] = normalized_resolutions
            if structures_changed:
                unresolved = [
                    str(item)
                    for item in after.get("unresolved") or []
                    if str(item).strip()
                ]
                synchronization_notice = (
                    "页面直接编辑内容尚未同步到固定计算结构；请生成修改建议并核对差异"
                )
                if synchronization_notice not in unresolved:
                    unresolved.append(synchronization_notice)
                after["unresolved"] = unresolved
                after["activation_ready"] = False
            new_summary = str(rule_summary or "").strip() or str(
                row["rule_summary"] or ""
            )
            new_period = str(applicable_period or "").strip() or str(
                row["applicable_period"] or ""
            )
            new_scope = str(business_scope_name or "").strip()
            changes = _rule_draft_changes(before, after)
            for section, previous, current in (
                (
                    "负责板块",
                    str(row["business_scope_name"] or ""),
                    new_scope,
                ),
                ("规则摘要", str(row["rule_summary"] or ""), new_summary),
                ("适用周期", str(row["applicable_period"] or ""), new_period),
            ):
                if previous != current:
                    changes.append(
                        {
                            "section": section,
                            "before": [previous] if previous else [],
                            "after": [current] if current else [],
                        }
                    )
            new_hash = _hash_json(after)
            if not changes and new_hash == expected_understanding_hash:
                return {
                    "rule_ref": str(rule_id),
                    "understanding_hash": new_hash,
                    "saved": False,
                    "message": "草稿内容没有变化",
                }
            amendment_id = uuid.uuid4()
            amendment_no = await self._next_rule_amendment_no(rule_id)
            await self.session.execute(
                text(
                    """
                    INSERT INTO legal_ops_performance_rule_amendments (
                        amendment_id, tenant_id, rule_version_id, amendment_no,
                        change_kind, instruction, issue_resolutions_json,
                        base_understanding_hash, proposed_understanding_json,
                        proposed_understanding_hash, proposed_rule_spec_json,
                        proposed_rule_spec_hash, changes_json, validation_json,
                        status, created_by, applied_by, applied_at
                    ) VALUES (
                        :amendment_id, :tenant_id, :rule_id, :amendment_no,
                        'manual_edit', :instruction, CAST(:resolutions AS jsonb),
                        :base_hash, CAST(:understanding AS jsonb),
                        :proposed_hash, CAST(:rule_spec AS jsonb),
                        :rule_spec_hash, CAST(:changes AS jsonb),
                        CAST(:validation AS jsonb), 'applied', :actor, :actor, now()
                    )
                    """
                ),
                {
                    "amendment_id": amendment_id,
                    "tenant_id": self.tenant_id,
                    "rule_id": rule_id,
                    "amendment_no": amendment_no,
                    "instruction": "网页表单直接编辑",
                    "resolutions": _json(normalized_resolutions),
                    "base_hash": expected_understanding_hash,
                    "understanding": _json(after),
                    "proposed_hash": new_hash,
                    "rule_spec": _json(row["rule_spec_json"])
                    if row["rule_spec_json"]
                    else None,
                    "rule_spec_hash": str(row["rule_spec_hash"] or ""),
                    "changes": _json(changes),
                    "validation": _json(
                        {
                            "activation_ready": bool(
                                after.get("activation_ready", False)
                            ),
                            "remaining_questions": len(after.get("unresolved") or []),
                            "fixed_program_ready": bool(row["rule_spec_json"]),
                        }
                    ),
                    "actor": self.actor_user_id,
                },
            )
            status_value = (
                "awaiting_confirmation"
                if bool(after.get("activation_ready", False))
                else "understanding_draft"
            )
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_performance_rule_versions
                    SET understanding_json = CAST(:understanding AS jsonb),
                        understanding_hash = :understanding_hash,
                        rule_summary = :rule_summary,
                        applicable_period = :applicable_period,
                        business_scope_name = :business_scope_name,
                        status = :status,
                        draft_revision = draft_revision + 1,
                        interpreted_by = :actor,
                        interpreted_at = now()
                    WHERE tenant_id = :tenant_id AND rule_version_id = :rule_id
                    """
                ),
                {
                    "understanding": _json(after),
                    "understanding_hash": new_hash,
                    "rule_summary": new_summary,
                    "applicable_period": new_period,
                    "business_scope_name": new_scope,
                    "status": status_value,
                    "actor": self.actor_user_id,
                    "tenant_id": self.tenant_id,
                    "rule_id": rule_id,
                },
            )
            await self._audit(
                "edit_rule_draft",
                "performance_rule",
                str(rule_id),
                "保存绩效规则网页草稿",
                {
                    "amendment_ref": str(amendment_id),
                    "base_hash": expected_understanding_hash,
                    "new_hash": new_hash,
                    "changed_sections": [str(item["section"]) for item in changes],
                },
                batch_id=row["source_batch_id"],
            )
        return {
            "rule_ref": str(rule_id),
            "amendment_ref": str(amendment_id),
            "understanding_hash": new_hash,
            "status": status_value,
            "status_label": _rule_status_label(status_value),
            "saved": True,
            "changes": changes,
            "message": "草稿已保存；涉及计算的改动仍需生成修改建议并核对",
        }

    async def propose_rule_amendment(
        self,
        rule_ref: str,
        *,
        expected_understanding_hash: str,
        instruction: str,
        issue_resolutions: list[dict[str, str]],
    ) -> dict[str, Any]:
        rule_id = self._uuid(rule_ref, "规则")
        row = await self._rule_row(rule_id)
        self._require_rule_draft_mutable(row)
        if str(row["understanding_hash"] or "") != expected_understanding_hash:
            raise DataIntakeError(
                "规则草稿已经变化，请刷新后重新生成修改建议",
                code="stale_rule_draft",
                status_code=409,
            )
        before = dict(row["understanding_json"] or {})
        normalized_resolutions = [
            {
                "issue": str(item.get("issue") or "").strip(),
                "response": str(item.get("response") or "").strip(),
            }
            for item in issue_resolutions
            if str(item.get("issue") or "").strip()
            and str(item.get("response") or "").strip()
        ]
        guidance_parts = [str(instruction or "").strip()]
        if normalized_resolutions:
            guidance_parts.append("对待确认事项的逐项答复：")
            guidance_parts.extend(
                f"- 问题：{item['issue']}\n  答复：{item['response']}"
                for item in normalized_resolutions
            )
        guidance = "\n".join(part for part in guidance_parts if part).strip()[:20_000]
        if not guidance:
            raise DataIntakeError("请填写需要修改或补充的业务口径")
        inspected = await self._inspect_stored_rule_package(row)
        interpreter = RuleDraftInterpreter(
            self.llm_client,
            model=self.settings.legal_ops_rule_understanding_model,
            timeout_seconds=self.settings.legal_ops_rule_understanding_timeout_seconds,
            enabled=self.settings.legal_ops_rule_understanding_enabled,
            fallback_model=self.settings.llm_extract_model,
            long_document_timeout_seconds=(
                self.settings.legal_ops_rule_understanding_long_document_timeout_seconds
            ),
        )
        draft = await interpreter.interpret(
            inspected,
            guidance=guidance,
            current_draft=before,
            data_profile=await self._rule_source_file_context(rule_id),
        )
        proposed = draft.as_dict()
        proposed["issue_resolutions"] = normalized_resolutions
        proposed_hash = _hash_json(proposed)
        changes = _rule_draft_changes(before, proposed)
        validation = {
            "activation_ready": draft.activation_ready,
            "remaining_questions": len(draft.unresolved),
            "fixed_program_ready": bool(draft.rule_spec),
            "evidence_count": len(draft.evidence),
            "can_apply": bool(changes),
            "notice": (
                "修改建议已形成；必须查看差异并人工应用后才会改变草稿"
                if changes
                else "没有识别到可应用的变化，请补充更明确的修改说明"
            ),
        }
        await self.session.rollback()
        async with self.session.begin():
            locked_row = await self._rule_row(rule_id, for_update=True)
            self._require_rule_draft_mutable(locked_row)
            if (
                str(locked_row["understanding_hash"] or "")
                != expected_understanding_hash
            ):
                raise DataIntakeError(
                    "规则草稿已经变化，请刷新后重新生成修改建议",
                    code="stale_rule_draft",
                    status_code=409,
                )
            duplicate = (
                (
                    await self.session.execute(
                        text(
                            """
                            SELECT *
                            FROM legal_ops_performance_rule_amendments
                            WHERE tenant_id = :tenant_id
                              AND rule_version_id = :rule_id
                              AND base_understanding_hash = :base_hash
                              AND proposed_understanding_hash = :proposed_hash
                              AND status = 'proposed'
                            LIMIT 1
                            """
                        ),
                        {
                            "tenant_id": self.tenant_id,
                            "rule_id": rule_id,
                            "base_hash": expected_understanding_hash,
                            "proposed_hash": proposed_hash,
                        },
                    )
                )
                .mappings()
                .first()
            )
            if duplicate:
                return {
                    **self._rule_amendment_payload(duplicate),
                    "duplicate_proposal": True,
                }
            amendment_id = uuid.uuid4()
            amendment_no = await self._next_rule_amendment_no(rule_id)
            await self.session.execute(
                text(
                    """
                    INSERT INTO legal_ops_performance_rule_amendments (
                        amendment_id, tenant_id, rule_version_id, amendment_no,
                        change_kind, instruction, issue_resolutions_json,
                        base_understanding_hash, proposed_understanding_json,
                        proposed_understanding_hash, proposed_rule_spec_json,
                        proposed_rule_spec_hash, changes_json, validation_json,
                        status, created_by
                    ) VALUES (
                        :amendment_id, :tenant_id, :rule_id, :amendment_no,
                        'conversation', :instruction, CAST(:resolutions AS jsonb),
                        :base_hash, CAST(:understanding AS jsonb),
                        :proposed_hash, CAST(:rule_spec AS jsonb),
                        :rule_spec_hash, CAST(:changes AS jsonb),
                        CAST(:validation AS jsonb), 'proposed', :actor
                    )
                    """
                ),
                {
                    "amendment_id": amendment_id,
                    "tenant_id": self.tenant_id,
                    "rule_id": rule_id,
                    "amendment_no": amendment_no,
                    "instruction": str(instruction or "").strip(),
                    "resolutions": _json(normalized_resolutions),
                    "base_hash": expected_understanding_hash,
                    "understanding": _json(proposed),
                    "proposed_hash": proposed_hash,
                    "rule_spec": _json(draft.rule_spec) if draft.rule_spec else None,
                    "rule_spec_hash": _hash_json(draft.rule_spec)
                    if draft.rule_spec
                    else "",
                    "changes": _json(changes),
                    "validation": _json(validation),
                    "actor": self.actor_user_id,
                },
            )
            await self._audit(
                "propose_rule_amendment",
                "performance_rule",
                str(rule_id),
                "根据业务人员说明生成规则修改建议",
                {
                    "amendment_ref": str(amendment_id),
                    "base_hash": expected_understanding_hash,
                    "proposed_hash": proposed_hash,
                    "changed_sections": [str(item["section"]) for item in changes],
                    "activation_ready": draft.activation_ready,
                },
                batch_id=locked_row["source_batch_id"],
            )
            amendment_row = (
                (
                    await self.session.execute(
                        text(
                            """
                            SELECT *
                            FROM legal_ops_performance_rule_amendments
                            WHERE tenant_id = :tenant_id
                              AND amendment_id = :amendment_id
                            """
                        ),
                        {
                            "tenant_id": self.tenant_id,
                            "amendment_id": amendment_id,
                        },
                    )
                )
                .mappings()
                .one()
            )
        return {
            **self._rule_amendment_payload(amendment_row),
            "duplicate_proposal": False,
        }

    async def apply_rule_amendment(
        self,
        rule_ref: str,
        amendment_ref: str,
        *,
        proposal_hash: str,
        confirmed: bool,
    ) -> dict[str, Any]:
        if not confirmed:
            raise DataIntakeError("必须明确确认已经查看修改前后差异")
        rule_id = self._uuid(rule_ref, "规则")
        amendment_id = self._uuid(amendment_ref, "修改建议")
        stale = False
        async with self.session.begin():
            row = await self._rule_row(rule_id, for_update=True)
            self._require_rule_draft_mutable(row)
            amendment = (
                (
                    await self.session.execute(
                        text(
                            """
                            SELECT *
                            FROM legal_ops_performance_rule_amendments
                            WHERE tenant_id = :tenant_id
                              AND rule_version_id = :rule_id
                              AND amendment_id = :amendment_id
                            FOR UPDATE
                            """
                        ),
                        {
                            "tenant_id": self.tenant_id,
                            "rule_id": rule_id,
                            "amendment_id": amendment_id,
                        },
                    )
                )
                .mappings()
                .first()
            )
            if not amendment:
                raise DataIntakeError(
                    "修改建议不存在",
                    code="not_found",
                    status_code=404,
                )
            if amendment["status"] != "proposed":
                raise DataIntakeError(
                    "该修改建议已经处理",
                    code="amendment_not_pending",
                    status_code=409,
                )
            if str(row["understanding_hash"] or "") != str(
                amendment["base_understanding_hash"]
            ):
                await self.session.execute(
                    text(
                        """
                        UPDATE legal_ops_performance_rule_amendments
                        SET status = 'stale'
                        WHERE tenant_id = :tenant_id
                          AND amendment_id = :amendment_id
                        """
                    ),
                    {
                        "tenant_id": self.tenant_id,
                        "amendment_id": amendment_id,
                    },
                )
                stale = True
            else:
                if proposal_hash != str(
                    amendment["proposed_understanding_hash"]
                ) or proposal_hash != _hash_json(
                    amendment["proposed_understanding_json"]
                ):
                    raise DataIntakeError(
                        "修改建议内容已经变化，请刷新后重新核对",
                        code="stale_amendment",
                        status_code=409,
                    )
                proposed = dict(amendment["proposed_understanding_json"] or {})
                proposed_spec = amendment["proposed_rule_spec_json"]
                if proposed_spec is not None:
                    try:
                        proposed_spec = validate_rule_spec(dict(proposed_spec))
                    except CalculationBlocked as exc:
                        raise DataIntakeError(
                            f"修改后的固定计算结构未通过校验：{exc}",
                            code="invalid_rule_spec",
                            status_code=409,
                        ) from exc
                derived_ready = bool(
                    proposed_spec
                    and not (proposed.get("unresolved") or [])
                    and (proposed.get("evidence") or [])
                )
                if bool(proposed.get("activation_ready", False)) != derived_ready:
                    raise DataIntakeError(
                        "修改建议的校验状态不一致，请重新生成",
                        code="invalid_amendment",
                        status_code=409,
                    )
                status_value = (
                    "awaiting_confirmation" if derived_ready else "understanding_draft"
                )
                await self.session.execute(
                    text(
                        """
                        UPDATE legal_ops_performance_rule_versions
                        SET understanding_json = CAST(:understanding AS jsonb),
                            understanding_hash = :understanding_hash,
                            rule_spec_json = CAST(:rule_spec AS jsonb),
                            rule_spec_hash = :rule_spec_hash,
                            status = :status,
                            draft_revision = draft_revision + 1,
                            interpreted_by = :actor,
                            interpreted_at = now()
                        WHERE tenant_id = :tenant_id
                          AND rule_version_id = :rule_id
                        """
                    ),
                    {
                        "understanding": _json(proposed),
                        "understanding_hash": proposal_hash,
                        "rule_spec": _json(proposed_spec) if proposed_spec else None,
                        "rule_spec_hash": _hash_json(proposed_spec)
                        if proposed_spec
                        else "",
                        "status": status_value,
                        "actor": self.actor_user_id,
                        "tenant_id": self.tenant_id,
                        "rule_id": rule_id,
                    },
                )
                await self.session.execute(
                    text(
                        """
                        UPDATE legal_ops_performance_rule_amendments
                        SET status = 'applied', applied_by = :actor,
                            applied_at = now()
                        WHERE tenant_id = :tenant_id
                          AND amendment_id = :amendment_id
                        """
                    ),
                    {
                        "actor": self.actor_user_id,
                        "tenant_id": self.tenant_id,
                        "amendment_id": amendment_id,
                    },
                )
                await self.session.execute(
                    text(
                        """
                        UPDATE legal_ops_performance_rule_amendments
                        SET status = 'stale'
                        WHERE tenant_id = :tenant_id
                          AND rule_version_id = :rule_id
                          AND amendment_id <> :amendment_id
                          AND status = 'proposed'
                        """
                    ),
                    {
                        "tenant_id": self.tenant_id,
                        "rule_id": rule_id,
                        "amendment_id": amendment_id,
                    },
                )
                await self._audit(
                    "apply_rule_amendment",
                    "performance_rule",
                    str(rule_id),
                    "人工核对差异后应用规则修改建议",
                    {
                        "amendment_ref": str(amendment_id),
                        "base_hash": str(amendment["base_understanding_hash"]),
                        "applied_hash": proposal_hash,
                        "activation_ready": derived_ready,
                    },
                    batch_id=row["source_batch_id"],
                )
        if stale:
            raise DataIntakeError(
                "规则草稿已经变化，这条修改建议已失效，请重新生成",
                code="stale_amendment",
                status_code=409,
            )
        return {
            "rule_ref": str(rule_id),
            "amendment_ref": str(amendment_id),
            "understanding_hash": proposal_hash,
            "status": status_value,
            "status_label": _rule_status_label(status_value),
            "activation_ready": derived_ready,
            "message": (
                "修改已应用，等待发布管理员最终确认"
                if derived_ready
                else "修改已应用，仍可继续编辑或沟通完善"
            ),
        }

    async def confirm_rule(
        self,
        rule_ref: str,
        *,
        understanding_hash: str,
        confirmed: bool,
        assignment_confirmed: bool = False,
    ) -> dict[str, Any]:
        if not confirmed:
            raise DataIntakeError("必须明确勾选已核对原文、源表和公式")
        rule_id = self._uuid(rule_ref, "规则")
        async with self.session.begin():
            row_hint = await self._rule_row(rule_id)
            if row_hint["period_id"]:
                await self._lock_performance_period(row_hint["period_id"])
            else:
                period_ids = await self.session.execute(
                    text(
                        """
                        SELECT period_id
                        FROM legal_ops_performance_periods
                        WHERE tenant_id = :tenant_id
                        ORDER BY period_id
                        """
                    ),
                    {"tenant_id": self.tenant_id},
                )
                for period_id in period_ids.scalars():
                    await self._lock_performance_period(period_id)
            row = await self._rule_row(rule_id, for_update=True)
            if row["status"] != "awaiting_confirmation":
                raise DataIntakeError(
                    "该规则尚未形成可确认的结构化草稿",
                    code="rule_not_ready",
                    status_code=409,
                )
            if not row["rule_spec_json"]:
                raise DataIntakeError(
                    "规则文件尚未配置完整，暂不可启用",
                    code="rule_spec_missing",
                    status_code=409,
                )
            rule_spec = dict(row["rule_spec_json"])
            confirmed_assignment_additions: list[dict[str, str | int]] = []
            confirmation_rows = await self.session.execute(
                text(
                    """
                    SELECT validation_json
                    FROM legal_ops_performance_rule_amendments
                    WHERE tenant_id = :tenant_id
                      AND rule_version_id = :rule_id
                      AND status = 'applied'
                      AND applied_by IS NOT NULL
                      AND applied_at IS NOT NULL
                    ORDER BY amendment_no
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "rule_id": rule_id,
                },
            )
            for confirmation_row in confirmation_rows.mappings().all():
                validation = dict(confirmation_row["validation_json"] or {})
                if validation.get("user_confirmed") is not True:
                    continue
                for addition in validation.get("assignment_additions") or []:
                    if isinstance(addition, dict):
                        confirmed_assignment_additions.append(dict(addition))
            _require_structural_assignment_chain(
                rule_spec,
                dict(row["understanding_json"] or {}),
                str(row["skill_markdown"] or ""),
                confirmed_additions=confirmed_assignment_additions,
            )
            subject = rule_spec.get("subject") or {}
            assignment_confirmation_required = str(
                rule_spec.get("schema_version") or ""
            ) == "2" and bool(
                (subject.get("fields") or {}) or (subject.get("lookups") or {})
            )
            if assignment_confirmation_required and not assignment_confirmed:
                raise DataIntakeError(
                    "请先确认已经核对团队或人员归属链，再启用规则",
                    code="assignment_not_confirmed",
                    status_code=409,
                )
            if (
                not understanding_hash
                or understanding_hash != row["understanding_hash"]
            ):
                raise DataIntakeError(
                    "规则草稿已经变化，请刷新后重新核对",
                    code="stale_rule_draft",
                    status_code=409,
                )
            replaced_rule_ids = list(
                (
                    await self.session.execute(
                        text(
                            """
                            SELECT rule_version_id
                            FROM legal_ops_performance_rule_versions
                            WHERE tenant_id = :tenant_id
                              AND is_current
                              AND period_id IS NOT DISTINCT FROM :period_id
                              AND business_scope_key = :business_scope_key
                            FOR UPDATE
                            """
                        ),
                        {
                            "tenant_id": self.tenant_id,
                            "period_id": row["period_id"],
                            "business_scope_key": row["business_scope_key"],
                        },
                    )
                ).scalars()
            )
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_performance_rule_versions
                    SET is_current = false,
                        status = CASE WHEN status = 'active' THEN 'superseded' ELSE status END
                    WHERE tenant_id = :tenant_id AND is_current
                      AND period_id IS NOT DISTINCT FROM :period_id
                      AND business_scope_key = :business_scope_key
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "period_id": row["period_id"],
                    "business_scope_key": row["business_scope_key"],
                },
            )
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_performance_rule_versions
                    SET status = 'active', is_current = true,
                        confirmed_by = :actor, confirmed_at = now(), activated_at = now()
                    WHERE tenant_id = :tenant_id AND rule_version_id = :rule_id
                    """
                ),
                {
                    "actor": self.actor_user_id,
                    "tenant_id": self.tenant_id,
                    "rule_id": rule_id,
                },
            )
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_intake_batches
                    SET status = 'published', published_by = :actor,
                        published_at = now(), updated_at = now()
                    WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                    """
                ),
                {
                    "actor": self.actor_user_id,
                    "tenant_id": self.tenant_id,
                    "batch_id": row["source_batch_id"],
                },
            )
            if replaced_rule_ids:
                await self._invalidate_calculations(
                    period_id=row["period_id"],
                    reason="绩效规则版本发生变化，需重新试算",
                    rule_version_ids=tuple(replaced_rule_ids),
                )
            if row["period_id"]:
                await self._refresh_period_status(row["period_id"])
            else:
                period_ids = await self.session.execute(
                    text(
                        """
                        SELECT period_id
                        FROM legal_ops_performance_periods
                        WHERE tenant_id = :tenant_id
                        """
                    ),
                    {"tenant_id": self.tenant_id},
                )
                for period_id in period_ids.scalars():
                    await self._refresh_period_status(period_id)
            await self._audit(
                "activate_rule",
                "performance_rule",
                str(rule_id),
                f"确认并启用规则版本：{row['rule_version']}",
                {
                    "understanding_hash": understanding_hash,
                    "rule_spec_hash": row["rule_spec_hash"],
                    "assignment_confirmed": assignment_confirmed,
                },
                batch_id=row["source_batch_id"],
            )
        return {
            "rule_ref": str(rule_id),
            "rule_version": row["rule_version"],
            "status": "active",
            "status_label": "已确认并启用",
            "confirmed_by": self.actor_user_id,
        }

    async def list_rules(self) -> list[dict[str, Any]]:
        result = await self.session.execute(
            text(
                """
                SELECT rule_version_id, rule_version, skill_name, skill_version,
                       source_file_hash, status, understanding_json,
                       understanding_hash, rule_spec_hash, created_by, created_at,
                       interpreted_by, interpreted_at, confirmed_by, confirmed_at,
                       is_current, package_inventory_json,
                       business_scope_name, business_scope_key, draft_revision
                FROM legal_ops_performance_rule_versions
                WHERE tenant_id = :tenant_id
                ORDER BY created_at DESC
                """
            ),
            {"tenant_id": self.tenant_id},
        )
        return [self._rule_card(row) for row in result.mappings().all()]

    async def rule_detail(self, rule_ref: str) -> dict[str, Any]:
        rule_id = self._uuid(rule_ref, "规则")
        row = await self._rule_row(rule_id)
        understanding = dict(row["understanding_json"] or {})
        rule_spec = dict(row["rule_spec_json"] or {})
        field_labels = {
            f"{table['key']}.{column['key']}": (
                f"{table.get('name', table['key'])}·{column.get('name', column['key'])}"
            )
            for table in rule_spec.get("source_tables") or []
            for column in table.get("columns") or []
        }
        metric_labels = {
            str(metric.get("key") or ""): str(
                metric.get("name") or metric.get("key") or "基础指标"
            )
            for metric in [
                *(rule_spec.get("metrics") or []),
                *(rule_spec.get("outputs") or []),
            ]
            if isinstance(metric, dict)
        }
        formulas = [
            {
                "result_name": metric_labels.get(
                    str(metric.get("key") or ""),
                    str(metric.get("key") or "基础指标"),
                ),
                "explanation": _describe_aggregate_metric(
                    metric,
                    field_labels,
                ),
                "rounding": "按底表精确汇总",
                "decimal_places": None,
                "zero_policy": "不适用",
                "subject_scope": (
                    "仅整体"
                    if metric.get("subject_scope") == "total_only"
                    else "团队与整体"
                ),
                "unit": str(metric.get("unit") or ""),
                "source_reference": "；".join(
                    part
                    for part in (
                        (
                            f"代码变量 {metric['source_symbol']}"
                            if metric.get("source_symbol")
                            else ""
                        ),
                        (
                            f"筛选变量 {metric['filter_source_symbol']}"
                            if metric.get("filter_source_symbol")
                            else ""
                        ),
                    )
                    if part
                ),
            }
            for metric in rule_spec.get("metrics") or []
            if isinstance(metric, dict)
        ]
        for output in rule_spec.get("outputs") or []:
            formulas.append(
                {
                    "result_name": output.get("name", output.get("key", "结果")),
                    "explanation": _describe_expression(
                        output.get("expression") or {},
                        field_labels,
                        metric_labels,
                    ),
                    "rounding": _rounding_label(str(output.get("rounding", "half_up"))),
                    "decimal_places": int(output.get("decimal_places", 2)),
                    "zero_policy": _divide_by_zero_label(
                        str(output.get("on_divide_by_zero", "error"))
                    ),
                    "subject_scope": (
                        "仅整体"
                        if output.get("subject_scope") == "total_only"
                        else "团队与整体"
                    ),
                    "unit": str(output.get("unit") or ""),
                    "source_reference": (
                        f"代码变量 {output['source_symbol']}"
                        if output.get("source_symbol")
                        else ""
                    ),
                }
            )
        return {
            **self._rule_card(row),
            "source": "Workbuddy 技能包",
            "skill_path": row["source_skill_path"],
            "skill_file_name": (
                PurePosixPath(str(row["source_skill_path"] or "SKILL.md")).name
                or "SKILL.md"
            ),
            "skill_markdown": str(row["skill_markdown"] or ""),
            "applicable_period": row["applicable_period"] or "未声明",
            "business_scope_name": row["business_scope_name"] or "尚未填写",
            "draft_revision": int(row["draft_revision"] or 1),
            "rule_summary": row["rule_summary"],
            "source_tables": understanding.get(
                "source_tables",
                rule_spec.get("source_tables", []),
            ),
            "stable_person_keys": understanding.get("stable_person_keys", []),
            "fixed_lookup_catalog": understanding.get(
                "fixed_lookup_catalog",
                [],
            ),
            "fixed_lookup_details": _rule_lookup_details(rule_spec),
            "assignment_chains": _rule_assignment_chains(rule_spec),
            "requires_assignment_confirmation": (
                str(rule_spec.get("schema_version") or "") == "2"
                and bool(
                    ((rule_spec.get("subject") or {}).get("fields") or {})
                    or ((rule_spec.get("subject") or {}).get("lookups") or {})
                )
            ),
            "applicability_scopes": understanding.get("applicability_scopes", []),
            "metric_catalog": understanding.get("metric_catalog", []),
            "target_versions": understanding.get("target_versions", []),
            "formulas": formulas,
            "calculation_rules": understanding.get("calculation_rules", []),
            "rule_explanations": _rule_explanation_items(
                understanding.get("calculation_rules", [])
            ),
            "caps_and_rounding": understanding.get("caps_and_rounding", []),
            "exception_handling": understanding.get("exception_handling", []),
            "outputs": understanding.get("outputs", []),
            "formula_audits": understanding.get("formula_audits", []),
            "unresolved": understanding.get("unresolved", []),
            "issue_resolutions": understanding.get("issue_resolutions", []),
            "evidence": understanding.get("evidence", []),
            "reference_documents": _rule_reference_documents(
                row.get("package_inventory_json")
            ),
            "activation_notice": (
                "规则已确认，可用于确定性计算"
                if row["status"] == "active"
                else (
                    "已经形成固定计算草稿，可以上传底表验证指标；"
                    "待确认事项解决前不能启用，也不会生成正式绩效"
                    if row["rule_spec_json"]
                    else "当前仍是理解草稿，尚未形成可验证的固定计算结构"
                )
            ),
        }

    async def list_rule_source_files(
        self,
        rule_ref: str,
        *,
        period_ref: str | None = None,
    ) -> list[dict[str, Any]]:
        rule_id = self._uuid(rule_ref, "规则")
        await self._rule_row(rule_id)
        result = await self.session.execute(
            text(
                """
                SELECT s.*, f.safe_file_name, f.file_hash, f.size_bytes,
                       f.uploaded_by, f.uploaded_at
                FROM legal_ops_performance_rule_source_files s
                JOIN legal_ops_intake_files f ON f.file_id = s.file_id
                WHERE s.tenant_id = :tenant_id
                  AND s.rule_version_id = :rule_id
                ORDER BY s.source_file_no, s.created_at
                """
            ),
            {"tenant_id": self.tenant_id, "rule_id": rule_id},
        )
        items = [self._rule_source_file_payload(row) for row in result.mappings().all()]
        if not period_ref:
            return items
        period_id = self._uuid(period_ref, "考核周期")
        await self._require_period(period_id)
        validation_result = await self.session.execute(
            text(
                """
                SELECT DISTINCT ON (b.file_hash, s.table_key)
                       b.file_hash, b.batch_no, b.status AS batch_status,
                       b.uploaded_at AS validated_at, b.error_summary,
                       b.warning_rows,
                       s.table_key, s.table_name, s.version, s.row_count,
                       s.error_count, s.status AS source_status
                FROM legal_ops_performance_source_versions s
                JOIN legal_ops_intake_batches b ON b.batch_id = s.batch_id
                WHERE s.tenant_id = :tenant_id
                  AND s.period_id = :period_id
                  AND s.rule_version_id = :rule_id
                  AND b.metadata_json ->> 'rule_ref' = :rule_ref
                  AND b.status NOT IN ('abandoned', 'failed')
                ORDER BY b.file_hash, s.table_key, s.version DESC
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "period_id": period_id,
                "rule_id": rule_id,
                "rule_ref": str(rule_id),
            },
        )
        validations_by_hash: dict[str, list[dict[str, Any]]] = {}
        for row in validation_result.mappings().all():
            detail = await self.batch_detail(
                str(row["batch_no"]),
                offset=0,
                limit=6,
            )
            status_value = str(row["batch_status"])
            validations_by_hash.setdefault(str(row["file_hash"]), []).append(
                {
                    "table_key": str(row["table_key"]),
                    "table_name": str(row["table_name"]),
                    "version": int(row["version"]),
                    "row_count": int(row["row_count"]),
                    "error_count": int(row["error_count"]),
                    "warning_count": int(row["warning_rows"] or 0),
                    "status": status_value,
                    "status_label": (
                        (
                            f"校验通过，有{int(row['warning_rows'])}行提醒"
                            if int(row["warning_rows"] or 0)
                            else "校验通过，可查看指标"
                        )
                        if status_value in {"ready", "published"}
                        else f"发现{int(row['error_count'])}行问题"
                    ),
                    "validated_at": row["validated_at"].isoformat(),
                    "error_summary": str(row["error_summary"] or ""),
                    "batch_no": str(row["batch_no"]),
                    "preview_rows": list(detail["rows"]),
                    "total_preview_rows": int(
                        detail.get("row_page", {}).get("total", 0)
                    ),
                }
            )
        for item in items:
            item["validations"] = validations_by_hash.get(
                str(item["file_hash"]),
                [],
            )
        return items

    async def upload_rule_source_file(
        self,
        rule_ref: str,
        *,
        content: bytes,
        filename: str,
        business_label: str = "",
    ) -> dict[str, Any]:
        rule_id = self._uuid(rule_ref, "规则")
        try:
            inspection = inspect_tabular_source(
                content,
                filename,
                max_bytes=self.files.max_bytes,
                max_rows=self.settings.legal_ops_data_intake_max_rows,
            )
        except FileValidationError as exc:
            raise DataIntakeError(
                str(exc),
                code="invalid_rule_source_file",
                status_code=422,
            ) from exc
        stored = self._store_file(content, filename)
        file_id = await self._register_original_file(stored)
        label = (
            str(business_label or "").strip()
            or PurePosixPath(stored.safe_file_name).stem
        )
        label = label[:256]
        async with self.session.begin():
            rule = await self._rule_row(rule_id, for_update=True)
            self._require_rule_source_files_mutable(rule)
            lock_key = (
                f"performance-rule-source|{self.tenant_id}|{rule_id}|{stored.file_hash}"
            )
            await self.session.execute(
                text("SELECT pg_advisory_xact_lock(hashtextextended(:lock_key, 0))"),
                {"lock_key": lock_key},
            )
            existing = (
                (
                    await self.session.execute(
                        text(
                            """
                            SELECT s.*, f.safe_file_name, f.file_hash,
                                   f.size_bytes, f.uploaded_by, f.uploaded_at
                            FROM legal_ops_performance_rule_source_files s
                            JOIN legal_ops_intake_files f ON f.file_id = s.file_id
                            WHERE s.tenant_id = :tenant_id
                              AND s.rule_version_id = :rule_id
                              AND s.file_id = :file_id
                            FOR UPDATE
                            """
                        ),
                        {
                            "tenant_id": self.tenant_id,
                            "rule_id": rule_id,
                            "file_id": file_id,
                        },
                    )
                )
                .mappings()
                .first()
            )
            if existing and str(existing["status"]) == "active":
                return {
                    **self._rule_source_file_payload(existing),
                    "duplicate_upload": True,
                }
            if existing:
                await self.session.execute(
                    text(
                        """
                        UPDATE legal_ops_performance_rule_source_files
                        SET business_label = :business_label,
                            inspection_json = CAST(:inspection AS jsonb),
                            status = 'active', abandoned_by = '',
                            abandoned_at = NULL
                        WHERE tenant_id = :tenant_id
                          AND source_file_id = :source_file_id
                        """
                    ),
                    {
                        "business_label": label,
                        "inspection": _json(inspection),
                        "tenant_id": self.tenant_id,
                        "source_file_id": existing["source_file_id"],
                    },
                )
                source_file_id = existing["source_file_id"]
                source_file_no = int(existing["source_file_no"])
                action = "restore_rule_source_file"
                summary = f"恢复规则核对底表：{stored.safe_file_name}"
            else:
                source_file_no = int(
                    await self.session.scalar(
                        text(
                            """
                            SELECT COALESCE(MAX(source_file_no), 0) + 1
                            FROM legal_ops_performance_rule_source_files
                            WHERE tenant_id = :tenant_id
                              AND rule_version_id = :rule_id
                            """
                        ),
                        {"tenant_id": self.tenant_id, "rule_id": rule_id},
                    )
                    or 1
                )
                source_file_id = uuid.uuid4()
                await self.session.execute(
                    text(
                        """
                        INSERT INTO legal_ops_performance_rule_source_files (
                            source_file_id, tenant_id, rule_version_id, file_id,
                            source_file_no, business_label, inspection_json,
                            status, created_by
                        ) VALUES (
                            :source_file_id, :tenant_id, :rule_id, :file_id,
                            :source_file_no, :business_label,
                            CAST(:inspection AS jsonb), 'active', :actor
                        )
                        """
                    ),
                    {
                        "source_file_id": source_file_id,
                        "tenant_id": self.tenant_id,
                        "rule_id": rule_id,
                        "file_id": file_id,
                        "source_file_no": source_file_no,
                        "business_label": label,
                        "inspection": _json(inspection),
                        "actor": self.actor_user_id,
                    },
                )
                action = "upload_rule_source_file"
                summary = f"上传规则核对底表：{stored.safe_file_name}"
            await self._audit(
                action,
                "performance_rule_source_file",
                str(source_file_id),
                summary,
                {
                    "rule_ref": str(rule_id),
                    "file_hash": stored.file_hash,
                    "usable_sheet_count": inspection["usable_sheet_count"],
                    "total_data_rows": inspection["total_data_rows"],
                    "warnings": inspection["warnings"],
                },
                batch_id=rule["source_batch_id"],
            )
        source_row = await self._rule_source_file_row(source_file_id)
        return {
            **self._rule_source_file_payload(source_row),
            "duplicate_upload": False,
        }

    async def abandon_rule_source_file(
        self,
        rule_ref: str,
        source_file_ref: str,
    ) -> dict[str, Any]:
        rule_id = self._uuid(rule_ref, "规则")
        source_file_id = self._uuid(source_file_ref, "底表")
        async with self.session.begin():
            rule = await self._rule_row(rule_id, for_update=True)
            self._require_rule_source_files_mutable(rule)
            row = (
                (
                    await self.session.execute(
                        text(
                            """
                            SELECT *
                            FROM legal_ops_performance_rule_source_files
                            WHERE tenant_id = :tenant_id
                              AND rule_version_id = :rule_id
                              AND source_file_id = :source_file_id
                            FOR UPDATE
                            """
                        ),
                        {
                            "tenant_id": self.tenant_id,
                            "rule_id": rule_id,
                            "source_file_id": source_file_id,
                        },
                    )
                )
                .mappings()
                .first()
            )
            if not row:
                raise DataIntakeError(
                    "规则核对底表不存在",
                    code="not_found",
                    status_code=404,
                )
            if str(row["status"]) == "abandoned":
                return {
                    "source_file_ref": str(source_file_id),
                    "status": "abandoned",
                    "status_label": "已移出当前规则核对",
                    "duplicate_action": True,
                }
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_performance_rule_source_files
                    SET status = 'abandoned', abandoned_by = :actor,
                        abandoned_at = now()
                    WHERE tenant_id = :tenant_id
                      AND source_file_id = :source_file_id
                    """
                ),
                {
                    "actor": self.actor_user_id,
                    "tenant_id": self.tenant_id,
                    "source_file_id": source_file_id,
                },
            )
            await self._audit(
                "abandon_rule_source_file",
                "performance_rule_source_file",
                str(source_file_id),
                "将底表移出当前规则核对；原始文件与审计记录保留",
                {"rule_ref": str(rule_id)},
                batch_id=rule["source_batch_id"],
            )
        return {
            "source_file_ref": str(source_file_id),
            "status": "abandoned",
            "status_label": "已移出当前规则核对",
            "duplicate_action": False,
        }

    async def validate_rule_source_file(
        self,
        rule_ref: str,
        source_file_ref: str,
        *,
        period_ref: str,
        table_key: str,
        column_mapping: dict[str, str] | None = None,
        confirm_mapping: bool = False,
    ) -> dict[str, Any]:
        rule_id = self._uuid(rule_ref, "规则")
        source_file_id = self._uuid(source_file_ref, "底表")
        rule = await self._rule_row(rule_id)
        source = await self._rule_source_file_row(source_file_id)
        if source["rule_version_id"] != rule_id or str(source["status"]) != "active":
            raise DataIntakeError(
                "这张底表不属于当前规则或已被移出",
                code="invalid_rule_source_file",
                status_code=409,
            )
        try:
            content = self.files.read(str(source["storage_key"]))
        except FileStorageError as exc:
            raise DataIntakeError(
                "原始底表不存在或校验失败，请重新上传",
                code="source_file_missing",
                status_code=409,
            ) from exc
        table = next(
            (
                item
                for item in (rule["rule_spec_json"] or {}).get("source_tables") or []
                if str(item.get("key") or "") == str(table_key or "").strip()
            ),
            None,
        )
        if table is None:
            raise DataIntakeError(
                "所选表格不属于当前规则版本",
                code="unknown_source_table",
            )
        inspection = dict(source["inspection_json"] or {})
        sheets = [
            dict(item)
            for item in inspection.get("sheets") or []
            if isinstance(item, dict) and (item.get("columns") or [])
        ]
        if not sheets:
            raise DataIntakeError(
                "没有识别到可用的数据工作表，请检查原文件",
                code="source_sheet_missing",
                status_code=422,
            )
        saved_record = (inspection.get("field_mappings") or {}).get(
            str(table_key)
        ) or {}
        saved_mapping = (
            saved_record.get("mapping") or {} if isinstance(saved_record, dict) else {}
        )
        if not saved_mapping:
            saved_mapping = await self._saved_rule_source_mapping(
                rule_id=rule_id,
                source_file_id=source_file_id,
                table_key=str(table_key),
            )
        candidates: list[tuple[int, dict[str, Any], Any, list[str]]] = []
        for sheet in sheets:
            workbook_headers = [
                str(column.get("header") or "").strip()
                for column in sheet.get("columns") or []
                if isinstance(column, dict) and str(column.get("header") or "").strip()
            ]
            proposal = (
                validate_column_mapping(
                    table,
                    workbook_headers,
                    column_mapping or {},
                )
                if confirm_mapping
                else propose_column_mapping(
                    table,
                    workbook_headers,
                    skill_markdown=str(rule["skill_markdown"] or ""),
                    saved_mapping=saved_mapping,
                )
            )
            candidates.append(
                (
                    len(proposal.mapping),
                    sheet,
                    proposal,
                    workbook_headers,
                )
            )
        candidates.sort(key=lambda item: item[0], reverse=True)
        _, selected_sheet, proposal, workbook_headers = candidates[0]
        if not proposal.complete:
            return {
                "status": "mapping_required",
                "status_label": "请确认原表字段",
                "source_file_ref": str(source_file_id),
                "table_key": str(table_key),
                "table_name": str(table.get("name") or table_key),
                "sheet_name": str(selected_sheet.get("sheet_name") or ""),
                "workbook_headers": workbook_headers,
                "mapping": mapping_payload(proposal),
                "message": (
                    "原表不需要修改。请只确认一次字段对应关系，以后同类文件会自动沿用。"
                ),
            }
        await self._persist_rule_source_mapping(
            rule_id=rule_id,
            source_file_id=source_file_id,
            table_key=str(table_key),
            table_name=str(table.get("name") or table_key),
            sheet_name=str(selected_sheet.get("sheet_name") or ""),
            proposal=proposal,
            confirmed=confirm_mapping,
        )
        batch = await self.upload_performance_table(
            period_ref=period_ref,
            table_key=str(table_key or "").strip(),
            content=content,
            filename=str(source["safe_file_name"]),
            rule_ref=str(rule_id),
            header_mapping=proposal.mapping,
        )
        return {
            **batch,
            "mapping_saved": True,
            "field_mapping": mapping_payload(proposal),
            "sheet_name": str(selected_sheet.get("sheet_name") or ""),
        }

    async def _persist_rule_source_mapping(
        self,
        *,
        rule_id: uuid.UUID,
        source_file_id: uuid.UUID,
        table_key: str,
        table_name: str,
        sheet_name: str,
        proposal: Any,
        confirmed: bool,
    ) -> None:
        await self.session.rollback()
        async with self.session.begin():
            row = (
                (
                    await self.session.execute(
                        text(
                            """
                            SELECT inspection_json
                            FROM legal_ops_performance_rule_source_files
                            WHERE tenant_id = :tenant_id
                              AND rule_version_id = :rule_id
                              AND source_file_id = :source_file_id
                              AND status = 'active'
                            FOR UPDATE
                            """
                        ),
                        {
                            "tenant_id": self.tenant_id,
                            "rule_id": rule_id,
                            "source_file_id": source_file_id,
                        },
                    )
                )
                .mappings()
                .first()
            )
            if not row:
                raise DataIntakeError(
                    "底表已经变化，请重新打开后再确认字段",
                    code="source_file_changed",
                    status_code=409,
                )
            inspection = dict(row["inspection_json"] or {})
            field_mappings = dict(inspection.get("field_mappings") or {})
            payload = mapping_payload(proposal)
            field_mappings[table_key] = {
                "table_name": table_name,
                "sheet_name": sheet_name,
                "mapping": dict(proposal.mapping),
                "items": payload["items"],
                "confirmed": bool(confirmed),
                "confirmed_by": self.actor_user_id,
                "confirmed_at": datetime.now(timezone.utc).isoformat(),
                "mapping_hash": _hash_json(dict(proposal.mapping)),
            }
            inspection["field_mappings"] = field_mappings
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_performance_rule_source_files
                    SET inspection_json = CAST(:inspection AS jsonb)
                    WHERE tenant_id = :tenant_id
                      AND rule_version_id = :rule_id
                      AND source_file_id = :source_file_id
                    """
                ),
                {
                    "inspection": _json(inspection),
                    "tenant_id": self.tenant_id,
                    "rule_id": rule_id,
                    "source_file_id": source_file_id,
                },
            )
            await self._audit(
                (
                    "confirm_rule_source_mapping"
                    if confirmed
                    else "auto_map_rule_source_headers"
                ),
                "performance_rule_source_file",
                str(source_file_id),
                (
                    f"确认“{table_name}”原表字段对应关系"
                    if confirmed
                    else f"按 Skill 明确口径识别“{table_name}”原表字段"
                ),
                {
                    "rule_ref": str(rule_id),
                    "table_key": table_key,
                    "sheet_name": sheet_name,
                    "mapping_hash": _hash_json(dict(proposal.mapping)),
                    "field_count": len(proposal.mapping),
                },
            )

    async def _saved_rule_source_mapping(
        self,
        *,
        rule_id: uuid.UUID,
        source_file_id: uuid.UUID,
        table_key: str,
    ) -> dict[str, str]:
        result = await self.session.execute(
            text(
                """
                SELECT inspection_json
                FROM legal_ops_performance_rule_source_files
                WHERE tenant_id = :tenant_id
                  AND rule_version_id = :rule_id
                  AND source_file_id <> :source_file_id
                  AND status = 'active'
                ORDER BY created_at DESC
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "rule_id": rule_id,
                "source_file_id": source_file_id,
            },
        )
        for row in result.mappings().all():
            inspection = dict(row["inspection_json"] or {})
            record = (inspection.get("field_mappings") or {}).get(table_key) or {}
            if not isinstance(record, dict):
                continue
            mapping = record.get("mapping") or {}
            if isinstance(mapping, dict) and mapping:
                return {
                    str(key): str(value)
                    for key, value in mapping.items()
                    if str(key).strip() and str(value).strip()
                }
        return {}

    async def upload_performance_table(
        self,
        *,
        period_ref: str,
        table_key: str,
        content: bytes,
        filename: str,
        rule_ref: str | None = None,
        header_mapping: dict[str, str] | None = None,
    ) -> dict[str, Any]:
        period_id = self._uuid(period_ref, "考核周期")
        period = await self._require_period(period_id)
        validation_rule = await self._rule_for_validation(
            period_id,
            rule_ref=rule_ref,
        )
        if not validation_rule or not validation_rule.get("rule_spec_json"):
            raise DataIntakeError(
                "规则草稿尚未形成固定计算结构，无法校验底表",
                code="rule_missing",
                status_code=409,
            )
        table = next(
            (
                item
                for item in validation_rule["rule_spec_json"].get("source_tables") or []
                if str(item.get("key")) == table_key
            ),
            None,
        )
        if table is None:
            raise DataIntakeError(
                "所选表格不属于当前规则版本", code="unknown_source_table"
            )
        validation_spec = dict(validation_rule["rule_spec_json"])
        aggregate_rule = str(validation_spec.get("schema_version") or "") == "2"
        schema = performance_table_schema(
            table,
            row_key=aggregate_table_row_key(validation_spec, table)
            if aggregate_rule
            else "",
        )
        stored = self._store_file(content, filename)
        await self._register_original_file(stored)
        try:
            parsed = parse_tabular_file(
                content,
                filename,
                schema,
                max_bytes=self.files.max_bytes,
                max_rows=self.settings.legal_ops_data_intake_max_rows,
                preferred_sheet_name=str(
                    table.get("sheet_name") or table.get("name") or ""
                ),
                header_mapping=header_mapping,
            )
        except FileValidationError as exc:
            parsed = ParsedTable(
                stored.safe_file_name,
                (),
                (
                    ParsedRow(
                        1,
                        {},
                        {},
                        [RowError("file_validation", "file", str(exc))],
                    ),
                ),
            )
        if aggregate_rule:
            parsed = _downgrade_irrelevant_aggregate_missing_errors(
                parsed,
                rule_spec=validation_spec,
                table=table,
                period_start=period["starts_on"],
                period_end=period["ends_on"],
            )
        header_mapping_hash = _header_mapping_hash(header_mapping)
        await self.session.rollback()
        async with self.session.begin():
            await self._require_period(period_id)
            await self._lock_performance_period(period_id)
            locked_rule = await self._rule_for_validation(
                period_id,
                rule_ref=rule_ref,
                for_update=bool(rule_ref),
            )
            if (
                not locked_rule
                or str(locked_rule["rule_version"])
                != str(validation_rule["rule_version"])
                or str(locked_rule["rule_spec_hash"])
                != str(validation_rule["rule_spec_hash"])
            ):
                raise DataIntakeError(
                    "规则版本在文件校验期间发生变化，请重新上传",
                    code="rule_changed",
                    status_code=409,
                )
            if not aggregate_rule:
                await self._apply_person_and_cross_table_checks(
                    parsed,
                    period_id=period_id,
                    table_key=table_key,
                    person_key=str(locked_rule["rule_spec_json"]["person_key"]),
                    rule_version_id=locked_rule["rule_version_id"],
                    rule_version=str(locked_rule["rule_version"]),
                )
            error_rows = sum(
                any(error.critical for error in row.errors) for row in parsed.rows
            )
            warning_rows = sum(
                bool(row.errors) and not any(error.critical for error in row.errors)
                for row in parsed.rows
            )
            status_value = "validation_failed" if error_rows else "ready"
            current_versions = await self._current_source_versions(
                period_id,
                rule_version_id=locked_rule["rule_version_id"],
                rule_version=str(locked_rule["rule_version"]),
            )
            preview_source_state_hash = _performance_source_state_hash(
                current_versions,
                rule_version=str(locked_rule["rule_version"]),
            )
            file_id = await self._register_file(stored)
            current_same_table = current_versions.get(table_key)
            if (
                current_same_table is not None
                and str(current_same_table["file_hash"]) == stored.file_hash
                and _batch_matches_header_mapping(
                    dict(current_same_table["metadata_json"] or {}),
                    header_mapping_hash,
                )
                and _batch_matches_validation_semantics(
                    dict(current_same_table["metadata_json"] or {})
                )
            ):
                current_batch = await self._batch_by_id(current_same_table["batch_id"])
                return {
                    **self._batch_payload(current_batch),
                    "duplicate_upload": True,
                }
            duplicate = await self._find_duplicate_batch(
                business_type="performance_source_table",
                data_source="绩效数据包",
                file_hash=stored.file_hash,
                period_id=period_id,
                table_key=table_key,
                import_mode="replace",
                rule_version=str(locked_rule["rule_version"]),
                rule_spec_hash=str(locked_rule["rule_spec_hash"]),
            )
            if duplicate:
                duplicate_metadata = dict(duplicate["metadata_json"] or {})
                if (
                    duplicate["status"] != "published"
                    and _batch_matches_header_mapping(
                        duplicate_metadata,
                        header_mapping_hash,
                    )
                    and _batch_matches_validation_semantics(duplicate_metadata)
                    and str(duplicate_metadata.get("preview_source_state_hash") or "")
                    == preview_source_state_hash
                ):
                    return {
                        **self._batch_payload(duplicate),
                        "duplicate_upload": True,
                    }
                if duplicate["status"] != "published":
                    await self._mark_stale_performance_batch_in_transaction(
                        duplicate["batch_id"],
                        (
                            "原表字段对应关系已更新，旧校验结果已失效"
                            if not _batch_matches_header_mapping(
                                duplicate_metadata,
                                header_mapping_hash,
                            )
                            else (
                                "底表校验口径已升级，旧校验结果已失效"
                                if not _batch_matches_validation_semantics(
                                    duplicate_metadata
                                )
                                else "预览依据已变化，请使用新批次"
                            )
                        ),
                    )
            preview_current_source_version_id = await self.session.scalar(
                text(
                    """
                    SELECT source_version_id
                    FROM legal_ops_performance_source_versions
                    WHERE tenant_id = :tenant_id
                      AND period_id = :period_id
                      AND rule_version_id = :rule_id
                      AND table_key = :table_key
                      AND is_current
                    FOR SHARE
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "period_id": period_id,
                    "rule_id": locked_rule["rule_version_id"],
                    "table_key": table_key,
                },
            )
            batch_id = uuid.uuid4()
            batch_no = _batch_no("PERF")
            await self._insert_batch(
                batch_id=batch_id,
                batch_no=batch_no,
                business_type="performance_source_table",
                data_source="绩效数据包",
                file_id=file_id,
                stored=stored,
                period_id=period_id,
                table_key=table_key,
                import_mode="replace",
                status=status_value,
                original_rows=parsed.row_count,
                counts={"failed": error_rows, "warning": warning_rows},
                metadata={
                    "table_name": schema.name,
                    "purpose": str(table.get("purpose") or ""),
                    "rule_version": locked_rule["rule_version"],
                    "rule_ref": str(locked_rule["rule_version_id"]),
                    "rule_spec_hash": str(locked_rule["rule_spec_hash"]),
                    "understanding_hash": str(locked_rule["understanding_hash"] or ""),
                    "draft_revision": int(locked_rule["draft_revision"] or 1),
                    "draft_validation": str(locked_rule["status"]) != "active",
                    "preview_source_state_hash": preview_source_state_hash,
                    "header_mapping_hash": header_mapping_hash,
                    "validation_semantics_version": (
                        PERFORMANCE_VALIDATION_SEMANTICS_VERSION
                    ),
                    "preview_current_source_version_ref": str(
                        preview_current_source_version_id or ""
                    ),
                    "field_mapping": [
                        {
                            "business_field": str(
                                column.get("name") or column.get("key") or ""
                            ),
                            "source_header": str(
                                (header_mapping or {}).get(
                                    str(column.get("key") or ""),
                                    column.get("name") or column.get("key") or "",
                                )
                            ),
                        }
                        for column in table.get("columns") or []
                        if isinstance(column, dict)
                    ],
                },
            )
            await self._insert_parsed_rows(batch_id, parsed)
            version = int(
                await self.session.scalar(
                    text(
                        """
                        SELECT COALESCE(MAX(version), 0) + 1
                        FROM legal_ops_performance_source_versions
                        WHERE tenant_id = :tenant_id AND period_id = :period_id
                          AND rule_version_id = :rule_id
                          AND table_key = :table_key
                        """
                    ),
                    {
                        "tenant_id": self.tenant_id,
                        "period_id": period_id,
                        "rule_id": locked_rule["rule_version_id"],
                        "table_key": table_key,
                    },
                )
                or 1
            )
            source_version_id = uuid.uuid4()
            await self.session.execute(
                text(
                    """
                    INSERT INTO legal_ops_performance_source_versions (
                        source_version_id, tenant_id, period_id, table_key, table_name,
                        purpose, version, batch_id, schema_json, row_count, error_count,
                        status, is_current, created_by, rule_version_id
                    ) VALUES (
                        :source_version_id, :tenant_id, :period_id, :table_key, :table_name,
                        :purpose, :version, :batch_id, CAST(:schema AS jsonb), :row_count,
                        :error_count, :status, false, :actor, :rule_id
                    )
                    """
                ),
                {
                    "source_version_id": source_version_id,
                    "tenant_id": self.tenant_id,
                    "period_id": period_id,
                    "rule_id": locked_rule["rule_version_id"],
                    "table_key": table_key,
                    "table_name": schema.name,
                    "purpose": str(table.get("purpose") or ""),
                    "version": version,
                    "batch_id": batch_id,
                    "schema": _json(schema_to_dict(schema)),
                    "row_count": parsed.row_count,
                    "error_count": error_rows,
                    "status": "validation_failed" if error_rows else "staged",
                    "actor": self.actor_user_id,
                },
            )
            await self._audit(
                "upload_performance_table",
                "performance_source_version",
                str(source_version_id),
                f"上传{schema.name}第{version}版",
                {
                    "batch_no": batch_no,
                    "error_rows": error_rows,
                    "warning_rows": warning_rows,
                },
                batch_id=batch_id,
            )
        return {
            "batch_no": batch_no,
            "source_version_ref": str(source_version_id),
            "table_key": table_key,
            "table_name": schema.name,
            "version": version,
            "row_count": parsed.row_count,
            "error_count": error_rows,
            "warning_count": warning_rows,
            "status": status_value,
            "status_label": STATUS_LABELS[status_value],
            "can_publish": not error_rows,
            "duplicate_upload": False,
        }

    async def publish_performance_table(self, batch_no: str) -> dict[str, Any]:
        try:
            return await self._publish_performance_table_once(batch_no)
        except DataIntakeError as exc:
            if exc.code in {
                "performance_source_changed",
                "performance_sources_changed",
                "cross_table_changed",
                "rule_changed",
            }:
                try:
                    await self._mark_stale_performance_batch(batch_no, exc.message)
                except SQLAlchemyError:
                    await self.session.rollback()
            raise

    async def _publish_performance_table_once(self, batch_no: str) -> dict[str, Any]:
        async with self.session.begin():
            batch_hint = await self._batch_by_no(batch_no)
            self._require_business(batch_hint, "performance_source_table")
            await self._lock_performance_period(batch_hint["period_id"])
            batch = await self._batch_by_no(batch_no, for_update=True)
            self._require_batch_ready(batch)
            source = (
                (
                    await self.session.execute(
                        text(
                            """
                        SELECT *
                        FROM legal_ops_performance_source_versions
                        WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                        FOR UPDATE
                        """
                        ),
                        {"tenant_id": self.tenant_id, "batch_id": batch["batch_id"]},
                    )
                )
                .mappings()
                .one()
            )
            if source["error_count"]:
                raise DataIntakeError(
                    "该表仍有错误，不能发布", code="critical_errors", status_code=409
                )
            metadata = dict(batch["metadata_json"] or {})
            active_rule = await self._active_rule_for_period_ref(
                source["period_id"],
                str(metadata.get("rule_ref") or ""),
            )
            if (
                not active_rule
                or source["rule_version_id"] != active_rule["rule_version_id"]
                or str(active_rule["rule_version"])
                != str(metadata.get("rule_version") or "")
                or str(active_rule["rule_spec_hash"])
                != str(metadata.get("rule_spec_hash") or "")
            ):
                raise DataIntakeError(
                    "规则版本或计算结构在预览后已经变化，请按当前规则重新上传",
                    code="rule_changed",
                    status_code=409,
                )
            current_versions = await self._current_source_versions(
                source["period_id"],
                rule_version_id=active_rule["rule_version_id"],
                rule_version=str(active_rule["rule_version"]),
            )
            expected_source_state_hash = str(
                (batch["metadata_json"] or {}).get("preview_source_state_hash") or ""
            )
            current_source_state_hash = _performance_source_state_hash(
                current_versions,
                rule_version=str(active_rule["rule_version"]),
            )
            if current_source_state_hash != expected_source_state_hash:
                raise DataIntakeError(
                    "本周期的已发布源表在预览后发生变化，请重新上传并核对",
                    code="performance_sources_changed",
                    status_code=409,
                )
            current_source_version_id = await self.session.scalar(
                text(
                    """
                    SELECT source_version_id
                    FROM legal_ops_performance_source_versions
                    WHERE tenant_id = :tenant_id
                      AND period_id = :period_id
                      AND rule_version_id = :rule_id
                      AND table_key = :table_key
                      AND is_current
                    FOR UPDATE
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "period_id": source["period_id"],
                    "rule_id": active_rule["rule_version_id"],
                    "table_key": source["table_key"],
                },
            )
            expected_source_version_ref = str(
                (batch["metadata_json"] or {}).get("preview_current_source_version_ref")
                or ""
            )
            if str(current_source_version_id or "") != expected_source_version_ref:
                raise DataIntakeError(
                    "该表的当前有效版本在预览后已经变化，请重新上传并核对",
                    code="performance_source_changed",
                    status_code=409,
                )
            active_spec = dict(active_rule["rule_spec_json"])
            aggregate_rule = str(active_spec.get("schema_version") or "") == "2"
            active_table = next(
                (
                    item
                    for item in active_spec.get("source_tables") or []
                    if str(item.get("key")) == str(source["table_key"])
                ),
                {},
            )
            record_key = (
                aggregate_table_row_key(active_spec, active_table)
                if aggregate_rule
                else str(active_spec.get("person_key") or "")
            )
            rows = await self._batch_rows(batch["batch_id"])
            publish_check = ParsedTable(
                str(batch["file_name"]),
                (),
                tuple(
                    ParsedRow(
                        int(row["source_row_number"]),
                        dict(row["raw_snapshot_json"] or {}),
                        dict(row["normalized_json"] or {}),
                        [],
                    )
                    for row in rows
                    if row["validation_status"] != "error"
                ),
            )
            if not aggregate_rule:
                await self._apply_person_and_cross_table_checks(
                    publish_check,
                    period_id=source["period_id"],
                    table_key=str(source["table_key"]),
                    person_key=record_key,
                    rule_version_id=active_rule["rule_version_id"],
                    rule_version=str(active_rule["rule_version"]),
                )
            if publish_check.error_count:
                raise DataIntakeError(
                    "其他已发布源表在预览后发生变化，跨表人员记录无法一一关联，请重新上传",
                    code="cross_table_changed",
                    status_code=409,
                )
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_performance_source_versions
                    SET is_current = false,
                        status = CASE WHEN status = 'published' THEN 'replaced' ELSE status END
                    WHERE tenant_id = :tenant_id AND period_id = :period_id
                      AND rule_version_id = :rule_id
                      AND table_key = :table_key AND is_current
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "period_id": source["period_id"],
                    "rule_id": active_rule["rule_version_id"],
                    "table_key": source["table_key"],
                },
            )
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_performance_source_versions
                    SET is_current = true, status = 'published',
                        published_by = :actor, published_at = now()
                    WHERE source_version_id = :source_version_id
                    """
                ),
                {
                    "actor": self.actor_user_id,
                    "source_version_id": source["source_version_id"],
                },
            )
            for row in rows:
                if row["validation_status"] == "error":
                    continue
                values = dict(row["normalized_json"] or {})
                record_id = str(values.get(record_key) or "").strip()
                await self.session.execute(
                    text(
                        """
                        INSERT INTO legal_ops_performance_source_rows (
                            tenant_id, source_version_id, source_row_number,
                            stable_person_id, values_json
                        ) VALUES (
                            :tenant_id, :source_version_id, :row_number,
                            :person_id, CAST(:values AS jsonb)
                        )
                        ON CONFLICT (tenant_id, source_version_id, stable_person_id)
                        DO NOTHING
                        """
                    ),
                    {
                        "tenant_id": self.tenant_id,
                        "source_version_id": source["source_version_id"],
                        "row_number": row["source_row_number"],
                        "person_id": record_id,
                        "values": _json(values),
                    },
                )
            await self._mark_batch_published(batch["batch_id"])
            await self._invalidate_calculations(
                period_id=source["period_id"],
                reason=f"{source['table_name']}源数据已替换，需重新试算",
                rule_version_ids=(active_rule["rule_version_id"],),
            )
            await self._refresh_period_status(source["period_id"])
            await self._audit(
                "publish_performance_table",
                "performance_source_version",
                str(source["source_version_id"]),
                f"发布{source['table_name']}第{source['version']}版",
                {"batch_no": batch_no},
                batch_id=batch["batch_id"],
            )
        return {
            "batch_no": batch_no,
            "table_name": source["table_name"],
            "version": source["version"],
            "status": "published",
            "status_label": "已发布为当前有效版本",
            "published_by": self.actor_user_id,
        }

    async def preview_rule_calculation(
        self,
        rule_ref: str,
        period_ref: str,
    ) -> dict[str, Any]:
        """Run a draft-only verification without publishing source data or results."""

        rule_id = self._uuid(rule_ref, "规则")
        period_id = self._uuid(period_ref, "考核周期")
        period = await self._require_period(period_id)
        rule = await self._rule_for_validation(period_id, rule_ref=rule_ref)
        if not rule:
            raise DataIntakeError(
                "规则草稿不存在",
                code="rule_missing",
                status_code=409,
            )
        rule_spec = dict(rule["rule_spec_json"] or {})
        required = {
            str(item["key"])
            for item in rule_spec.get("source_tables") or []
            if bool(item.get("required", True))
        }
        version_result = await self.session.execute(
            text(
                """
                SELECT DISTINCT ON (s.table_key)
                       s.*, b.file_hash, b.file_name, b.status AS batch_status,
                       b.metadata_json
                FROM legal_ops_performance_source_versions s
                JOIN legal_ops_intake_batches b ON b.batch_id = s.batch_id
                WHERE s.tenant_id = :tenant_id
                  AND s.period_id = :period_id
                  AND s.rule_version_id = :rule_id
                  AND s.error_count = 0
                  AND s.status IN ('staged','published')
                  AND b.status IN ('ready','published')
                  AND b.metadata_json ->> 'rule_ref' = :rule_ref
                  AND b.metadata_json ->> 'rule_spec_hash' = :rule_spec_hash
                ORDER BY s.table_key, s.version DESC
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "period_id": period_id,
                "rule_id": rule_id,
                "rule_ref": str(rule_id),
                "rule_spec_hash": str(rule["rule_spec_hash"]),
            },
        )
        versions = {
            str(item["table_key"]): item for item in version_result.mappings().all()
        }
        missing = sorted(required.difference(versions))
        if missing:
            table_labels = {
                str(item.get("key")): str(item.get("name") or item.get("key") or "源表")
                for item in rule_spec.get("source_tables") or []
            }
            raise DataIntakeError(
                "还缺少通过校验的底表："
                + "、".join(table_labels.get(key, key) for key in missing),
                code="required_tables_missing",
                status_code=409,
            )

        sources: dict[str, list[dict[str, Any]]] = {}
        source_versions: list[dict[str, Any]] = []
        for table_key, version in sorted(versions.items()):
            row_result = await self.session.execute(
                text(
                    """
                    SELECT source_row_number, normalized_json
                    FROM legal_ops_intake_rows
                    WHERE tenant_id = :tenant_id
                      AND batch_id = :batch_id
                      AND validation_status IN ('valid','warning')
                    ORDER BY source_row_number
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "batch_id": version["batch_id"],
                },
            )
            values: list[dict[str, Any]] = []
            for source_row in row_result.mappings().all():
                payload = dict(source_row["normalized_json"] or {})
                payload["__row_number__"] = int(source_row["source_row_number"])
                values.append(payload)
            sources[table_key] = values
            source_versions.append(
                {
                    "table_key": table_key,
                    "table_name": str(version["table_name"]),
                    "version": int(version["version"]),
                    "source_version_ref": str(version["source_version_id"]),
                    "file_name": str(version["file_name"]),
                    "file_hash": str(version["file_hash"]),
                    "published": str(version["status"]) == "published",
                }
            )
        preview_rule_version = (
            f"{rule['rule_version']} · 草稿第{int(rule['draft_revision'] or 1)}版"
        )
        try:
            calculated = DeterministicCalculator(
                rule_spec,
                rule_version=preview_rule_version,
            ).calculate(
                sources,
                period_start=period["starts_on"],
                period_end=period["ends_on"],
            )
        except CalculationBlocked as exc:
            raise DataIntakeError(
                str(exc),
                code="calculation_blocked",
                status_code=409,
            ) from exc

        latest_rule = await self._rule_row(rule_id)
        if str(latest_rule["understanding_hash"] or "") != str(
            rule["understanding_hash"] or ""
        ) or str(latest_rule["rule_spec_hash"] or "") != str(
            rule["rule_spec_hash"] or ""
        ):
            raise DataIntakeError(
                "规则草稿在验证期间发生变化，请重新运行验证",
                code="calculation_stale",
                status_code=409,
            )
        names = await self._person_display_names()
        output_labels = {
            str(output.get("key") or ""): str(
                output.get("name") or output.get("key") or "计算结果"
            )
            for output in rule_spec.get("outputs") or []
            if isinstance(output, dict)
        }
        output_definitions = {
            str(output.get("key") or ""): dict(output)
            for output in rule_spec.get("outputs") or []
            if isinstance(output, dict) and str(output.get("key") or "").strip()
        }
        understanding = dict(rule["understanding_json"] or {})
        targets = [
            dict(item)
            for item in understanding.get("target_versions") or []
            if isinstance(item, dict)
        ]
        preview_results = []
        achievement_counts = {
            "achieved": 0,
            "not_achieved": 0,
            "unknown": 0,
        }
        for item in calculated.results:
            value_items = []
            for key, value in item.values.items():
                name = output_labels.get(key, key)
                target = _matching_metric_target(
                    targets,
                    key=key,
                    name=name,
                    subject_key=item.person_id,
                    subject_name=item.display_name,
                )
                achievement = _metric_achievement(value, target)
                achievement_counts[achievement["status"]] += 1
                unit = str(
                    output_definitions.get(key, {}).get("unit")
                    or achievement.get("target_unit")
                    or ""
                ).strip()
                value_items.append(
                    {
                        "key": key,
                        "name": name,
                        "value": str(value),
                        "unit": unit,
                        **_metric_presentation(
                            key=key,
                            name=name,
                            value=value,
                            unit=unit,
                        ),
                        **achievement,
                    }
                )
            preview_results.append(
                {
                    "person_key": item.person_id,
                    "person_name": item.display_name
                    or names.get(item.person_id, "")
                    or item.person_id
                    or "未匹配名称",
                    "value_items": value_items,
                    "lineage_summary": _lineage_summary(
                        item.lineage,
                        rule_spec,
                        preview_rule_version,
                    ),
                    "errors": [],
                }
            )
        for error in calculated.errors:
            preview_results.append(
                {
                    "person_key": str(error.get("person_id") or ""),
                    "person_name": str(error.get("subject_name") or "")
                    or names.get(str(error.get("person_id") or ""), "")
                    or str(error.get("person_id") or "未匹配名称"),
                    "value_items": [],
                    "lineage_summary": _lineage_summary(
                        {"source_rows": error.get("source_rows") or []},
                        rule_spec,
                        preview_rule_version,
                    ),
                    "errors": [dict(error)],
                }
            )

        preview_no = f"VERIFY-{calculated.input_hash[:12].upper()}"
        await self.session.rollback()
        async with self.session.begin():
            await self._audit(
                "preview_rule_calculation",
                "performance_rule",
                str(rule_id),
                f"运行规则底表验证：{preview_no}",
                {
                    "period_ref": str(period_id),
                    "rule_spec_hash": str(rule["rule_spec_hash"]),
                    "input_hash": calculated.input_hash,
                    "output_hash": calculated.output_hash,
                    "source_versions": source_versions,
                    "participant_count": len(preview_results),
                    "error_count": len(calculated.errors),
                },
                batch_id=rule["source_batch_id"],
            )
        return {
            "preview": True,
            "preview_no": preview_no,
            "status": ("validation_failed" if calculated.errors else "preview"),
            "status_label": (
                f"验证完成，有{len(calculated.errors)}项待处理（未发布）"
                if calculated.errors
                else "底表验证完成（未发布）"
            ),
            "rule_ref": str(rule_id),
            "rule_version": str(rule["rule_version"]),
            "draft_revision": int(rule["draft_revision"] or 1),
            "rule_spec_hash": str(rule["rule_spec_hash"]),
            "period_ref": str(period_id),
            "participant_count": len(preview_results),
            "success_count": len(calculated.results),
            "error_count": len(calculated.errors),
            "achievement_counts": achievement_counts,
            "input_hash": calculated.input_hash,
            "output_hash": calculated.output_hash,
            "source_versions": source_versions,
            "results": preview_results,
            "result_subject_name": str(
                (rule_spec.get("subject") or {}).get("name") or "人员"
            ),
            "notice": ("这是规则核对结果，不会写入正式绩效，也不会自动发布底表。"),
        }

    async def performance_report(
        self,
        rule_ref: str,
        period_ref: str,
        *,
        view: Literal["week", "month"],
        anchor_date: date,
        scope_key: str = "",
        published_only: bool = False,
    ) -> dict[str, Any]:
        """Build the weekly/monthly report from validated source rows, read-only."""

        report, context = await self._load_performance_report(
            rule_ref,
            period_ref,
            view=view,
            anchor_date=anchor_date,
            published_only=published_only,
        )
        payload = report.as_dict()
        selected_scope = self._select_performance_report_scope(
            payload,
            scope_key,
        )
        return {
            **payload,
            **context,
            "selected_scope": selected_scope,
            "scope_options": [
                {
                    "scope_key": str(item.get("scope_key") or ""),
                    "scope_name": str(item.get("scope_name") or ""),
                }
                for item in payload.get("scopes") or []
            ],
            "notice": (
                "本页依据当前已校验的底表实时汇总，只读展示，"
                "不会创建试算、发布结果或修改底表。"
            ),
        }

    async def performance_report_preview(
        self,
        rule_ref: str,
        period_ref: str,
        *,
        view: Literal["week", "month"],
        anchor_date: date,
        scope_key: str = "",
    ) -> dict[str, Any]:
        """Compatibility name used by the report preview HTTP endpoint."""

        return await self.performance_report(
            rule_ref,
            period_ref,
            view=view,
            anchor_date=anchor_date,
            scope_key=scope_key,
        )

    async def current_performance_report_catalog(
        self,
        anchor_date: date,
    ) -> dict[str, Any] | None:
        """Return current weekly/monthly defendant reports for bot reuse."""

        result = await self.session.execute(
            text(
                """
                SELECT r.rule_version_id, r.rule_version, r.rule_spec_json,
                       r.activated_at, p.period_id, p.label,
                       p.starts_on, p.ends_on
                FROM legal_ops_performance_rule_versions r
                JOIN legal_ops_performance_periods p
                  ON p.tenant_id = r.tenant_id
                 AND (r.period_id IS NULL OR r.period_id = p.period_id)
                WHERE r.tenant_id = :tenant_id
                  AND r.is_current
                  AND r.status = 'active'
                  AND r.confirmed_at IS NOT NULL
                  AND r.rule_spec_json IS NOT NULL
                  AND r.business_scope_name = :business_scope_name
                  AND p.starts_on <= :anchor_date
                ORDER BY (p.ends_on >= :anchor_date) DESC,
                         (r.period_id = p.period_id) DESC,
                         p.ends_on DESC, p.starts_on DESC,
                         r.activated_at DESC
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "anchor_date": anchor_date,
                "business_scope_name": "被告案件",
            },
        )
        selected: Any | None = None
        for row in result.mappings().all():
            try:
                build_defendant_performance_report(
                    dict(row["rule_spec_json"] or {}),
                    {},
                    view="month",
                    anchor_date=anchor_date,
                )
            except PerformanceReportError:
                continue
            selected = row
            break
        if selected is None:
            return None
        rule_ref = str(selected["rule_version_id"])
        period_ref = str(selected["period_id"])
        effective_anchor_date = min(anchor_date, selected["ends_on"])
        week = await self.performance_report(
            rule_ref,
            period_ref,
            view="week",
            anchor_date=effective_anchor_date,
            published_only=True,
        )
        month = await self.performance_report(
            rule_ref,
            period_ref,
            view="month",
            anchor_date=effective_anchor_date,
            published_only=True,
        )
        return {
            "rule_ref": rule_ref,
            "period_ref": period_ref,
            "rule_version": str(week["rule"]["rule_version"]),
            "data_status_label": str(week["data_status_label"]),
            "week": week,
            "month": month,
            "requested_anchor_date": anchor_date.isoformat(),
            "effective_anchor_date": effective_anchor_date.isoformat(),
            "read_only": True,
        }

    async def export_performance_report(
        self,
        rule_ref: str,
        period_ref: str,
        *,
        view: Literal["week", "month"],
        anchor_date: date,
        scope_key: str,
        file_type: Literal["xlsx", "docx"],
    ) -> tuple[bytes, str]:
        """Export the same read-only snapshot shown by the report preview."""

        report, context = await self._load_performance_report(
            rule_ref,
            period_ref,
            view=view,
            anchor_date=anchor_date,
        )
        selected_scope = self._select_performance_report_scope(
            report.as_dict(),
            scope_key,
        )
        resolved_scope_key = str(selected_scope.get("scope_key") or "")
        if file_type == "xlsx":
            content = export_defendant_performance_xlsx(
                report,
                scope_key=resolved_scope_key,
            )
        elif file_type == "docx":
            content = export_defendant_performance_docx(
                report,
                scope_key=resolved_scope_key,
            )
        else:
            raise DataIntakeError(
                "只支持导出 Excel 或 Word",
                code="unsupported_report_format",
                status_code=422,
            )
        period_label = str((report.as_dict().get("period") or {}).get("label") or "")
        scope_label = str(selected_scope.get("scope_name") or "整体")
        period_label = re.sub(r'[\\/:*?"<>|\r\n]+', "-", period_label).strip()
        scope_label = re.sub(r'[\\/:*?"<>|\r\n]+', "-", scope_label).strip()
        report_hash = str(context.get("report_hash") or "")[:8].upper()
        filename = (
            f"{scope_label}被告案件{period_label}绩效报告-{report_hash}.{file_type}"
        )
        return content, filename

    async def _load_performance_report(
        self,
        rule_ref: str,
        period_ref: str,
        *,
        view: Literal["week", "month"],
        anchor_date: date,
        published_only: bool = False,
    ) -> tuple[DefendantPerformanceReport, dict[str, Any]]:
        """Read one immutable report input snapshot without any audit writes."""

        rule_id = self._uuid(rule_ref, "规则")
        period_id = self._uuid(period_ref, "考核周期")
        period = await self._require_period(period_id)
        if not period["starts_on"] <= anchor_date <= period["ends_on"]:
            raise DataIntakeError(
                "统计截止日期必须位于所选考核周期内",
                code="report_date_outside_period",
                status_code=422,
            )
        rule = await self._active_rule_for_period_ref(period_id, rule_id)
        if not rule or not rule.get("rule_spec_json"):
            raise DataIntakeError(
                "所选板块尚无已确认并启用的计算规则",
                code="rule_missing",
                status_code=409,
            )
        rule_spec = dict(rule["rule_spec_json"] or {})
        required = {
            str(item.get("key") or "")
            for item in rule_spec.get("source_tables") or []
            if isinstance(item, dict) and bool(item.get("required", True))
        }
        version_result = await self.session.execute(
            text(
                """
                SELECT DISTINCT ON (s.table_key)
                       s.*, b.file_hash, b.file_name, b.batch_no,
                       b.status AS batch_status, b.warning_rows,
                       b.uploaded_at, b.uploaded_by
                FROM legal_ops_performance_source_versions s
                JOIN legal_ops_intake_batches b ON b.batch_id = s.batch_id
                WHERE s.tenant_id = :tenant_id
                  AND s.period_id = :period_id
                  AND s.rule_version_id = :rule_id
                  AND s.error_count = 0
                  AND s.status IN ('staged','published')
                  AND b.status IN ('ready','published')
                  AND (
                      NOT :published_only
                      OR (
                          s.status = 'published'
                          AND b.status = 'published'
                      )
                  )
                  AND b.metadata_json ->> 'rule_ref' = :rule_ref
                  AND b.metadata_json ->> 'rule_spec_hash' = :rule_spec_hash
                ORDER BY s.table_key, s.version DESC
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "period_id": period_id,
                "rule_id": rule_id,
                "rule_ref": str(rule_id),
                "rule_spec_hash": str(rule["rule_spec_hash"]),
                "published_only": published_only,
            },
        )
        versions = {
            str(row["table_key"]): row for row in version_result.mappings().all()
        }
        missing = sorted(required.difference(versions))
        if missing:
            labels = {
                str(item.get("key") or ""): str(
                    item.get("name") or item.get("key") or "底表"
                )
                for item in rule_spec.get("source_tables") or []
                if isinstance(item, dict)
            }
            raise DataIntakeError(
                "还缺少通过校验的底表："
                + "、".join(labels.get(key, key) for key in missing),
                code="required_tables_missing",
                status_code=409,
            )

        sources: dict[str, list[dict[str, Any]]] = {}
        source_versions: list[dict[str, Any]] = []
        for table_key, version in sorted(versions.items()):
            row_result = await self.session.execute(
                text(
                    """
                    SELECT source_row_number, normalized_json, raw_snapshot_json
                    FROM legal_ops_intake_rows
                    WHERE tenant_id = :tenant_id
                      AND batch_id = :batch_id
                      AND validation_status IN ('valid','warning')
                    ORDER BY source_row_number
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "batch_id": version["batch_id"],
                },
            )
            values: list[dict[str, Any]] = []
            for source_row in row_result.mappings().all():
                payload = dict(source_row["normalized_json"] or {})
                payload["__raw__"] = dict(source_row["raw_snapshot_json"] or {})
                payload["__row_number__"] = int(source_row["source_row_number"])
                values.append(payload)
            sources[table_key] = values
            source_versions.append(
                {
                    "table_key": table_key,
                    "table_name": str(version["table_name"]),
                    "version": int(version["version"]),
                    "source_version_ref": str(version["source_version_id"]),
                    "batch_no": str(version["batch_no"]),
                    "file_name": str(version["file_name"]),
                    "file_hash": str(version["file_hash"]),
                    "row_count": int(version["row_count"]),
                    "error_count": int(version["error_count"]),
                    "warning_count": int(version["warning_rows"] or 0),
                    "source_status": str(version["status"]),
                    "batch_status": str(version["batch_status"]),
                    "uploaded_by": str(version["uploaded_by"] or ""),
                    "uploaded_at": (
                        version["uploaded_at"].isoformat()
                        if version["uploaded_at"]
                        else ""
                    ),
                }
            )
        targets = [
            dict(item)
            for item in (
                (rule.get("understanding_json") or {}).get(
                    "target_versions",
                    [],
                )
            )
            if isinstance(item, dict)
        ]
        try:
            report = build_defendant_performance_report(
                rule_spec,
                sources,
                view=view,
                anchor_date=anchor_date,
                targets=targets,
            )
        except PerformanceReportError as exc:
            raise DataIntakeError(
                str(exc),
                code="performance_report_blocked",
                status_code=409,
            ) from exc
        payload = report.as_dict()
        source_state_hash = _performance_source_state_hash(
            versions,
            rule_version=str(rule["rule_version"]),
        )
        report_hash = _hash_json(
            {
                "rule_ref": str(rule_id),
                "rule_spec_hash": str(rule["rule_spec_hash"]),
                "source_state_hash": source_state_hash,
                "view": view,
                "anchor_date": anchor_date.isoformat(),
                "report": payload,
            }
        )
        staged = any(item["source_status"] == "staged" for item in source_versions)
        assignment_error_count = len(payload.get("assignment_errors") or [])
        data_status = (
            "validated_with_assignment_errors"
            if assignment_error_count
            else "validated_staged"
            if staged
            else "published"
        )
        data_status_label = (
            f"已校验，有{assignment_error_count}条归属待确认"
            if assignment_error_count
            else "已校验，尚未发布"
            if staged
            else "已发布"
        )
        return report, {
            "rule": {
                "rule_ref": str(rule_id),
                "rule_version": str(rule["rule_version"]),
                "rule_spec_hash": str(rule["rule_spec_hash"]),
                "source_file_hash": str(rule["source_file_hash"] or ""),
                "status": str(rule["status"]),
                "confirmed_by": str(rule["confirmed_by"] or ""),
                "confirmed_at": (
                    rule["confirmed_at"].isoformat() if rule["confirmed_at"] else ""
                ),
            },
            "performance_period": {
                "period_ref": str(period_id),
                "label": str(period["label"]),
                "starts_on": period["starts_on"].isoformat(),
                "ends_on": period["ends_on"].isoformat(),
            },
            "source_versions": source_versions,
            "source_state_hash": source_state_hash,
            "report_hash": report_hash,
            "data_status": data_status,
            "data_status_label": data_status_label,
            "read_only": True,
        }

    @staticmethod
    def _select_performance_report_scope(
        payload: dict[str, Any],
        scope_key: str,
    ) -> dict[str, Any]:
        scopes = [
            dict(item) for item in payload.get("scopes") or [] if isinstance(item, dict)
        ]
        requested = str(scope_key or "").strip()
        if not requested and scopes:
            return scopes[0]
        for item in scopes:
            if requested in {
                str(item.get("scope_key") or ""),
                str(item.get("scope_name") or ""),
            }:
                return item
        raise DataIntakeError(
            "未找到所选团队的绩效报告",
            code="report_scope_not_found",
            status_code=404,
        )

    async def trial_calculate(
        self,
        period_ref: str,
        *,
        rule_ref: str = "",
    ) -> dict[str, Any]:
        period_id = self._uuid(period_ref, "考核周期")
        period = await self._require_period(period_id)
        active_rule = await self._active_rule_for_period_ref(period_id, rule_ref)
        if not active_rule or not active_rule.get("rule_spec_json"):
            raise DataIntakeError(
                "规则文件尚未配置，暂不可计算", code="rule_missing", status_code=409
            )
        rule_spec = active_rule["rule_spec_json"]
        versions = await self._current_source_versions(
            period_id,
            rule_version_id=active_rule["rule_version_id"],
            rule_version=str(active_rule["rule_version"]),
        )
        required = {
            str(item["key"])
            for item in rule_spec.get("source_tables") or []
            if bool(item.get("required", True))
        }
        missing = sorted(required.difference(versions))
        if missing:
            raise DataIntakeError(
                f"缺少必需表格：{'、'.join(missing)}，暂不可试算",
                code="required_tables_missing",
                status_code=409,
            )
        sources: dict[str, list[dict[str, Any]]] = {}
        source_versions: list[dict[str, Any]] = []
        for table_key, version in versions.items():
            result = await self.session.execute(
                text(
                    """
                    SELECT source_row_number, stable_person_id, values_json
                    FROM legal_ops_performance_source_rows
                    WHERE tenant_id = :tenant_id AND source_version_id = :version_id
                    ORDER BY source_row_number
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "version_id": version["source_version_id"],
                },
            )
            values = []
            for row in result.mappings().all():
                payload = dict(row["values_json"])
                payload["__row_number__"] = row["source_row_number"]
                values.append(payload)
            sources[table_key] = values
            source_versions.append(
                {
                    "table_key": table_key,
                    "table_name": version["table_name"],
                    "version": version["version"],
                    "source_version_ref": str(version["source_version_id"]),
                    "file_hash": version["file_hash"],
                    "rule_spec_hash": active_rule["rule_spec_hash"],
                }
            )
        try:
            calculated = DeterministicCalculator(
                rule_spec,
                rule_version=active_rule["rule_version"],
            ).calculate(
                sources,
                period_start=period["starts_on"],
                period_end=period["ends_on"],
            )
        except CalculationBlocked as exc:
            raise DataIntakeError(
                str(exc), code="calculation_blocked", status_code=409
            ) from exc

        names = await self._person_display_names()
        await self.session.rollback()
        async with self.session.begin():
            await self._require_period(period_id)
            await self._lock_performance_period(period_id)
            try:
                locked_rule = await self._active_rule_for_period_ref(
                    period_id,
                    active_rule["rule_version_id"],
                )
            except DataIntakeError as exc:
                if exc.code != "rule_not_active_for_period":
                    raise
                locked_rule = None
            if (
                not locked_rule
                or locked_rule["rule_version_id"] != active_rule["rule_version_id"]
                or str(locked_rule["rule_spec_hash"])
                != str(active_rule["rule_spec_hash"])
            ):
                raise DataIntakeError(
                    "规则版本在试算期间发生变化，请重新试算",
                    code="calculation_stale",
                    status_code=409,
                )
            locked_versions = await self._current_source_versions(
                period_id,
                rule_version_id=locked_rule["rule_version_id"],
                rule_version=str(locked_rule["rule_version"]),
            )
            captured_refs = {
                str(item["source_version_ref"]) for item in source_versions
            }
            locked_refs = {
                str(item["source_version_id"]) for item in locked_versions.values()
            }
            if captured_refs != locked_refs:
                raise DataIntakeError(
                    "源表版本在试算期间发生变化，请重新试算",
                    code="calculation_stale",
                    status_code=409,
                )
            calculation_identity = "|".join(
                (
                    self.tenant_id,
                    str(period_id),
                    str(active_rule["rule_version_id"]),
                    calculated.input_hash,
                    calculated.output_hash,
                )
            )
            await self.session.execute(
                text(
                    "SELECT pg_advisory_xact_lock("
                    "hashtextextended(:calculation_identity, 0))"
                ),
                {"calculation_identity": calculation_identity},
            )
            concurrent_existing = await self.session.scalar(
                text(
                    """
                    SELECT calculation_id
                    FROM legal_ops_performance_calculation_runs
                    WHERE tenant_id = :tenant_id
                      AND period_id = :period_id
                      AND rule_version_id = :rule_id
                      AND input_hash = :input_hash
                      AND output_hash = :output_hash
                      AND status IN ('trial','published')
                    ORDER BY calculated_at DESC
                    LIMIT 1
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "period_id": period_id,
                    "rule_id": active_rule["rule_version_id"],
                    "input_hash": calculated.input_hash,
                    "output_hash": calculated.output_hash,
                },
            )
            if concurrent_existing is not None:
                payload = await self.calculation_detail(str(concurrent_existing))
                payload["duplicate_calculation"] = True
                return payload
            calculation_id = uuid.uuid4()
            calculation_no = _batch_no("CALC")
            await self.session.execute(
                text(
                    """
                    INSERT INTO legal_ops_performance_calculation_runs (
                        calculation_id, calculation_no, tenant_id, period_id,
                        rule_version_id, source_versions_json, input_hash, output_hash,
                        status, participant_count, success_count, error_count,
                        code_version, calculated_by
                    ) VALUES (
                        :calculation_id, :calculation_no, :tenant_id, :period_id,
                        :rule_id, CAST(:source_versions AS jsonb), :input_hash, :output_hash,
                        'trial', :participants, :success, :errors, :code_version, :actor
                    )
                    """
                ),
                {
                    "calculation_id": calculation_id,
                    "calculation_no": calculation_no,
                    "tenant_id": self.tenant_id,
                    "period_id": period_id,
                    "rule_id": active_rule["rule_version_id"],
                    "source_versions": _json(source_versions),
                    "input_hash": calculated.input_hash,
                    "output_hash": calculated.output_hash,
                    "participants": len(calculated.results) + len(calculated.errors),
                    "success": len(calculated.results),
                    "errors": len(calculated.errors),
                    "code_version": self.code_version,
                    "actor": self.actor_user_id,
                },
            )
            for result in calculated.results:
                await self.session.execute(
                    text(
                        """
                        INSERT INTO legal_ops_performance_results (
                            tenant_id, calculation_id, stable_person_id, display_name,
                            result_json, lineage_json, errors_json
                        ) VALUES (
                            :tenant_id, :calculation_id, :person_id, :display_name,
                            CAST(:result AS jsonb), CAST(:lineage AS jsonb), '[]'::jsonb
                        )
                        """
                    ),
                    {
                        "tenant_id": self.tenant_id,
                        "calculation_id": calculation_id,
                        "person_id": result.person_id,
                        "display_name": result.display_name
                        or names.get(result.person_id, ""),
                        "result": _json(result.values),
                        "lineage": _json(result.lineage),
                    },
                )
            for error in calculated.errors:
                await self.session.execute(
                    text(
                        """
                        INSERT INTO legal_ops_performance_results (
                            tenant_id, calculation_id, stable_person_id, display_name,
                            result_json, lineage_json, errors_json
                        ) VALUES (
                            :tenant_id, :calculation_id, :person_id, :display_name,
                            '{}'::jsonb, CAST(:lineage AS jsonb), CAST(:errors AS jsonb)
                        )
                        """
                    ),
                    {
                        "tenant_id": self.tenant_id,
                        "calculation_id": calculation_id,
                        "person_id": error["person_id"],
                        "display_name": str(error.get("subject_name") or "")
                        or names.get(str(error["person_id"]), ""),
                        "lineage": _json({"source_rows": error.get("source_rows", [])}),
                        "errors": _json([error]),
                    },
                )
            await self._audit(
                "trial_calculation",
                "performance_calculation",
                str(calculation_id),
                f"完成绩效试算：{calculation_no}",
                {
                    "input_hash": calculated.input_hash,
                    "output_hash": calculated.output_hash,
                    "rule_version": active_rule["rule_version"],
                },
            )
        return await self.calculation_detail(str(calculation_id))

    async def calculation_detail(self, calculation_ref: str) -> dict[str, Any]:
        calculation_id = self._uuid(calculation_ref, "计算批次")
        run = (
            (
                await self.session.execute(
                    text(
                        """
                    SELECT c.*, p.label AS period_label, r.rule_version,
                           r.rule_spec_hash, r.rule_spec_json,
                           r.understanding_json
                    FROM legal_ops_performance_calculation_runs c
                    JOIN legal_ops_performance_periods p ON p.period_id = c.period_id
                    JOIN legal_ops_performance_rule_versions r
                      ON r.rule_version_id = c.rule_version_id
                    WHERE c.tenant_id = :tenant_id AND c.calculation_id = :calculation_id
                    """
                    ),
                    {"tenant_id": self.tenant_id, "calculation_id": calculation_id},
                )
            )
            .mappings()
            .first()
        )
        if not run:
            raise DataIntakeError("计算批次不存在", code="not_found", status_code=404)
        output_labels = {
            str(output.get("key") or ""): str(
                output.get("name") or output.get("key") or "计算结果"
            )
            for output in (run["rule_spec_json"] or {}).get("outputs") or []
            if isinstance(output, dict)
        }
        output_definitions = {
            str(output.get("key") or ""): dict(output)
            for output in (run["rule_spec_json"] or {}).get("outputs") or []
            if isinstance(output, dict) and str(output.get("key") or "").strip()
        }
        result_subject_name = str(
            ((run["rule_spec_json"] or {}).get("subject") or {}).get("name") or "人员"
        )
        results = (
            (
                await self.session.execute(
                    text(
                        """
                    SELECT stable_person_id, display_name, result_json, lineage_json,
                           errors_json
                    FROM legal_ops_performance_results
                    WHERE tenant_id = :tenant_id AND calculation_id = :calculation_id
                    ORDER BY display_name, stable_person_id
                    """
                    ),
                    {"tenant_id": self.tenant_id, "calculation_id": calculation_id},
                )
            )
            .mappings()
            .all()
        )
        total_key = str(
            ((run["rule_spec_json"] or {}).get("subject") or {}).get(
                "total_key",
                "__total__",
            )
        )
        results = sorted(
            results,
            key=lambda item: (
                0 if str(item["stable_person_id"]) == total_key else 1,
                str(item["display_name"] or ""),
                str(item["stable_person_id"]),
            ),
        )
        targets = [
            dict(item)
            for item in (run["understanding_json"] or {}).get(
                "target_versions",
                [],
            )
            if isinstance(item, dict)
        ]
        achievement_counts = {
            "achieved": 0,
            "not_achieved": 0,
            "unknown": 0,
        }
        result_payloads = []
        for item in results:
            display_name = item["display_name"] or f"未匹配{result_subject_name}"
            value_items = []
            for key, value in item["result_json"].items():
                name = output_labels.get(key, key)
                achievement = _metric_achievement(
                    Decimal(str(value)),
                    _matching_metric_target(
                        targets,
                        key=key,
                        name=name,
                        subject_key=str(item["stable_person_id"]),
                        subject_name=str(display_name),
                    ),
                )
                achievement_counts[achievement["status"]] += 1
                unit = str(
                    output_definitions.get(key, {}).get("unit")
                    or achievement.get("target_unit")
                    or ""
                ).strip()
                value_items.append(
                    {
                        "key": key,
                        "name": name,
                        "value": value,
                        "unit": unit,
                        **_metric_presentation(
                            key=key,
                            name=name,
                            value=Decimal(str(value)),
                            unit=unit,
                        ),
                        **achievement,
                    }
                )
            result_payloads.append(
                {
                    "person_key": item["stable_person_id"],
                    "person_name": display_name,
                    "values": item["result_json"],
                    "value_items": value_items,
                    "lineage": item["lineage_json"],
                    "lineage_summary": _lineage_summary(
                        item["lineage_json"],
                        run["rule_spec_json"] or {},
                        run["rule_version"],
                    ),
                    "errors": item["errors_json"],
                }
            )
        return {
            "calculation_ref": str(calculation_id),
            "calculation_no": run["calculation_no"],
            "period_label": run["period_label"],
            "status": run["status"],
            "status_label": _calculation_status_label(run["status"]),
            "participant_count": run["participant_count"],
            "success_count": run["success_count"],
            "error_count": run["error_count"],
            "rule_version": run["rule_version"],
            "rule_spec_hash": run["rule_spec_hash"],
            "source_versions": run["source_versions_json"],
            "input_hash": run["input_hash"],
            "output_hash": run["output_hash"],
            "calculated_by": run["calculated_by"],
            "calculated_at": run["calculated_at"].isoformat(),
            "result_subject_name": result_subject_name,
            "achievement_counts": achievement_counts,
            "results": result_payloads,
            "duplicate_calculation": False,
        }

    async def publish_calculation(self, calculation_ref: str) -> dict[str, Any]:
        calculation_id = self._uuid(calculation_ref, "计算批次")
        async with self.session.begin():
            run_hint = (
                (
                    await self.session.execute(
                        text(
                            """
                        SELECT period_id
                        FROM legal_ops_performance_calculation_runs
                        WHERE tenant_id = :tenant_id
                          AND calculation_id = :calculation_id
                        """
                        ),
                        {
                            "tenant_id": self.tenant_id,
                            "calculation_id": calculation_id,
                        },
                    )
                )
                .mappings()
                .first()
            )
            if not run_hint:
                raise DataIntakeError(
                    "计算批次不存在", code="not_found", status_code=404
                )
            await self._lock_performance_period(run_hint["period_id"])
            run = (
                (
                    await self.session.execute(
                        text(
                            """
                        SELECT * FROM legal_ops_performance_calculation_runs
                        WHERE tenant_id = :tenant_id AND calculation_id = :calculation_id
                        FOR UPDATE
                        """
                        ),
                        {"tenant_id": self.tenant_id, "calculation_id": calculation_id},
                    )
                )
                .mappings()
                .first()
            )
            if not run:
                raise DataIntakeError(
                    "计算批次不存在", code="not_found", status_code=404
                )
            if run["status"] == "published":
                return {
                    "calculation_ref": str(calculation_id),
                    "status": "published",
                    "status_label": "已发布",
                    "duplicate_publish": True,
                }
            if run["status"] != "trial":
                raise DataIntakeError(
                    "只有有效试算可以发布",
                    code="calculation_not_trial",
                    status_code=409,
                )
            if int(run["error_count"] or 0) > 0:
                raise DataIntakeError(
                    "试算仍有异常人员，请先修正源数据并重新试算",
                    code="calculation_has_errors",
                    status_code=409,
                )
            try:
                current_rule = await self._active_rule_for_period_ref(
                    run["period_id"],
                    run["rule_version_id"],
                )
            except DataIntakeError as exc:
                if exc.code != "rule_not_active_for_period":
                    raise
                current_rule = None
            if (
                not current_rule
                or current_rule["rule_version_id"] != run["rule_version_id"]
            ):
                raise DataIntakeError(
                    "规则版本已经变化，必须重新试算",
                    code="calculation_stale",
                    status_code=409,
                )
            run_rule_hashes = {
                str(item.get("rule_spec_hash") or "")
                for item in run["source_versions_json"]
            }
            if run_rule_hashes != {str(current_rule["rule_spec_hash"])}:
                raise DataIntakeError(
                    "规则内容已经变化，必须重新试算",
                    code="calculation_stale",
                    status_code=409,
                )
            current_versions = await self._current_source_versions(
                run["period_id"],
                rule_version_id=current_rule["rule_version_id"],
                rule_version=str(current_rule["rule_version"]),
            )
            current_refs = {
                str(value["source_version_id"]) for value in current_versions.values()
            }
            run_refs = {
                str(item["source_version_ref"]) for item in run["source_versions_json"]
            }
            if current_refs != run_refs:
                raise DataIntakeError(
                    "源表版本已经变化，必须重新试算",
                    code="calculation_stale",
                    status_code=409,
                )
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_performance_calculation_runs
                    SET status = 'published', published_by = :actor, published_at = now()
                    WHERE calculation_id = :calculation_id
                    """
                ),
                {"actor": self.actor_user_id, "calculation_id": calculation_id},
            )
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_performance_periods
                    SET status = 'published', updated_at = now()
                    WHERE tenant_id = :tenant_id AND period_id = :period_id
                    """
                ),
                {"tenant_id": self.tenant_id, "period_id": run["period_id"]},
            )
            await self._audit(
                "publish_calculation",
                "performance_calculation",
                str(calculation_id),
                f"发布正式绩效结果：{run['calculation_no']}",
                {
                    "input_hash": run["input_hash"],
                    "output_hash": run["output_hash"],
                },
            )
        return {
            "calculation_ref": str(calculation_id),
            "calculation_no": run["calculation_no"],
            "status": "published",
            "status_label": "正式结果已发布",
            "published_by": self.actor_user_id,
            "duplicate_publish": False,
        }

    async def abandon_calculation(self, calculation_ref: str) -> dict[str, Any]:
        calculation_id = self._uuid(calculation_ref, "计算批次")
        async with self.session.begin():
            result = await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_performance_calculation_runs
                    SET status = 'abandoned'
                    WHERE tenant_id = :tenant_id AND calculation_id = :calculation_id
                      AND status = 'trial'
                    RETURNING calculation_no
                    """
                ),
                {"tenant_id": self.tenant_id, "calculation_id": calculation_id},
            )
            calculation_no = result.scalar_one_or_none()
            if not calculation_no:
                raise DataIntakeError(
                    "只有未发布的试算可以放弃", code="cannot_abandon", status_code=409
                )
            await self._audit(
                "abandon_calculation",
                "performance_calculation",
                str(calculation_id),
                f"放弃绩效试算：{calculation_no}",
                {},
            )
        return {
            "calculation_ref": str(calculation_id),
            "status": "abandoned",
            "status_label": "已放弃",
        }

    async def export_calculation(self, calculation_ref: str) -> bytes:
        detail = await self.calculation_detail(calculation_ref)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "绩效结果"
        output_names: list[str] = []
        for result in detail["results"]:
            for item in result["value_items"]:
                if item["name"] not in output_names:
                    output_names.append(item["name"])
        subject_name = str(detail.get("result_subject_name") or "人员")
        sheet.append(
            [f"{subject_name}稳定标识", subject_name]
            + [safe_excel_cell(name) for name in output_names]
            + ["计算批次", "规则版本"]
        )
        for result in detail["results"]:
            values = {item["name"]: item["value"] for item in result["value_items"]}
            sheet.append(
                [
                    safe_excel_cell(result["person_key"]),
                    safe_excel_cell(result["person_name"]),
                ]
                + [safe_excel_cell(values.get(name, "")) for name in output_names]
                + [
                    safe_excel_cell(detail["calculation_no"]),
                    safe_excel_cell(detail["rule_version"]),
                ]
            )
        lineage = workbook.create_sheet("数据血缘")
        lineage.append(
            [f"{subject_name}稳定标识", "结果项", "规则版本", "使用字段", "来源记录"]
        )
        for result in detail["results"]:
            for item in result["lineage_summary"]:
                lineage.append(
                    [
                        safe_excel_cell(result["person_key"]),
                        safe_excel_cell(item["result_name"]),
                        safe_excel_cell(item["rule_version"]),
                        safe_excel_cell("、".join(item["fields"]) or "固定规则值"),
                        safe_excel_cell(_lineage_source_text(item)),
                    ]
                )
        errors = workbook.create_sheet("异常明细")
        errors.append([f"{subject_name}稳定标识", subject_name, "异常原因", "来源记录"])
        for result in detail["results"]:
            for error in result["errors"]:
                errors.append(
                    [
                        safe_excel_cell(result["person_key"]),
                        safe_excel_cell(result["person_name"]),
                        safe_excel_cell(error.get("message", "计算异常")),
                        safe_excel_cell(
                            "、".join(
                                f"{row.get('table', '源表')}第{row.get('row_number') or '未记录'}行"
                                for row in error.get("source_rows") or []
                            )
                        ),
                    ]
                )
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    async def export_calculation_errors(self, calculation_ref: str) -> bytes:
        detail = await self.calculation_detail(calculation_ref)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "异常明细"
        subject_name = str(detail.get("result_subject_name") or "人员")
        sheet.append([f"{subject_name}稳定标识", subject_name, "异常原因", "来源记录"])
        for result in detail["results"]:
            source_rows = [
                row for item in result["lineage_summary"] for row in item["source_rows"]
            ]
            source_text = "；".join(
                _lineage_source_text(item) for item in result["lineage_summary"]
            )
            if not source_text:
                source_text = "、".join(
                    f"{row['table_name']}第{row['row_number'] or '未记录'}行"
                    for row in source_rows
                )
            for error in result["errors"]:
                sheet.append(
                    [
                        safe_excel_cell(result["person_key"]),
                        safe_excel_cell(result["person_name"]),
                        safe_excel_cell(error.get("message", "计算异常")),
                        safe_excel_cell(source_text),
                    ]
                )
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    async def upload_case_master(
        self,
        *,
        content: bytes,
        filename: str,
        source_system: str,
        import_mode: Literal["full", "incremental"],
        profile_key: str = "auto",
    ) -> dict[str, Any]:
        stored = self._store_file(content, filename)
        await self._register_original_file(stored)
        try:
            detected = parse_case_master_source_file(
                content,
                filename,
                profile_key=profile_key,
                source_system=source_system,
                max_bytes=self.files.max_bytes,
                max_rows=self.settings.legal_ops_data_intake_max_rows,
            )
            parsed = detected.parsed
            source_system = _source_system(detected.source_system)
            source_metadata = detected.metadata()
        except (FileValidationError, CaseSourceProfileError) as exc:
            fallback_source = _source_system(source_system)
            return await self._save_case_preview(
                business_type="case_master",
                source_system=fallback_source,
                import_mode=import_mode,
                content=content,
                filename=filename,
                preview=_file_error_preview(str(exc)),
                metadata={
                    "source_profile_key": profile_key,
                    "source_profile_label": "尚未识别",
                },
            )
        await self._canonicalize_person_column(
            parsed,
            "owner_user_id",
            preserve_unlinked=detected.profile.preserve_unlinked_owner,
            source_system=source_system,
        )
        await self._canonicalize_team_column(
            parsed,
            "team_id",
            preserve_unlinked=detected.profile.preserve_unlinked_team,
            source_system=source_system,
        )
        normalized = []
        for row in parsed.rows:
            payload = _jsonable(row.normalized)
            payload["__row_number__"] = row.row_number
            payload["__raw_snapshot__"] = _jsonable(row.raw)
            payload["source_system"] = source_system
            payload["source_profile_key"] = detected.profile.key
            payload["source_profile_version"] = detected.profile.version
            payload["source_sheet"] = parsed.sheet_name
            payload["source_header_hash"] = detected.header_hash
            normalized.append(payload)
        case_index = await self._case_index()
        people, teams = await self._known_people_and_teams()
        people.update(
            str(row.normalized.get("owner_user_id") or "").strip()
            for row in parsed.rows
            if str(row.normalized.get("owner_user_id") or "").strip()
        )
        teams.update(
            str(row.normalized.get("team_id") or "").strip()
            for row in parsed.rows
            if str(row.normalized.get("team_id") or "").strip()
        )
        preview = preview_case_master(
            normalized,
            existing=case_index,
            import_mode=import_mode,
            known_people=people,
            known_teams=teams,
        )
        preview = _merge_parse_errors(preview, parsed)
        await self.session.rollback()
        return await self._save_case_preview(
            business_type="case_master",
            source_system=source_system,
            import_mode=import_mode,
            content=content,
            filename=filename,
            preview=preview,
            metadata=source_metadata,
        )

    async def upload_case_progress(
        self,
        *,
        content: bytes,
        filename: str,
        source_system: str,
        profile_key: str = "auto",
        snapshot_date: str = "",
        reporter_id: str = "",
    ) -> dict[str, Any]:
        stored = self._store_file(content, filename)
        await self._register_original_file(stored)
        if profile_key == "auto":
            try:
                parse_tabular_file(
                    content,
                    filename,
                    CASE_PROGRESS_SCHEMA,
                    max_bytes=self.files.max_bytes,
                    max_rows=self.settings.legal_ops_data_intake_max_rows,
                )
            except FileValidationError:
                pass
            else:
                profile_key = "legal_ops_standard_case_progress_v1"
        try:
            if profile_key == "legal_ops_standard_case_progress_v1":
                source_system = _source_system(source_system)
                parsed = parse_tabular_file(
                    content,
                    filename,
                    CASE_PROGRESS_SCHEMA,
                    max_bytes=self.files.max_bytes,
                    max_rows=self.settings.legal_ops_data_intake_max_rows,
                )
                source_metadata = {
                    "source_profile_key": profile_key,
                    "source_profile_version": "1",
                    "source_profile_label": "标准案件进展接口表",
                    "detected_sheet": parsed.sheet_name,
                    "detected_columns": len(parsed.headers),
                }
            else:
                detected = parse_case_progress_source_file(
                    content,
                    filename,
                    profile_key=profile_key,
                    snapshot_date=snapshot_date,
                    reporter_id=reporter_id,
                    max_bytes=self.files.max_bytes,
                    max_rows=self.settings.legal_ops_data_intake_max_rows,
                )
                parsed = detected.parsed
                source_system = _source_system(detected.source_system)
                source_metadata = detected.metadata()
        except (FileValidationError, CaseProgressSourceProfileError) as exc:
            fallback_source = _source_system(source_system)
            return await self._save_case_preview(
                business_type="case_progress",
                source_system=fallback_source,
                import_mode="append_or_update",
                content=content,
                filename=filename,
                preview=_file_error_preview(str(exc)),
                metadata={
                    "source_profile_key": profile_key,
                    "source_profile_label": "尚未识别",
                },
                idempotency_scope=f"{profile_key}|{snapshot_date}|{reporter_id}",
            )
        await self._canonicalize_person_column(parsed, "reporter_id")
        normalized = []
        for row in parsed.rows:
            payload = _jsonable(row.normalized)
            payload["__row_number__"] = row.row_number
            payload["__raw_snapshot__"] = _jsonable(row.raw)
            payload["source_system"] = source_system
            normalized.append(payload)
        case_index = await self._case_index()
        progress_index = await self._progress_index()
        people, _ = await self._known_people_and_teams()
        nodes = await self._known_nodes()
        preview = preview_case_progress(
            normalized,
            cases=case_index,
            existing_progress=progress_index,
            known_nodes=nodes,
            known_people=people,
        )
        preview = _merge_parse_errors(preview, parsed)
        await self.session.rollback()
        return await self._save_case_preview(
            business_type="case_progress",
            source_system=source_system,
            import_mode="append_or_update",
            content=content,
            filename=filename,
            preview=preview,
            metadata=source_metadata,
            idempotency_scope=f"{profile_key}|{snapshot_date}|{reporter_id}",
        )

    async def publish_case_master(self, batch_no: str) -> dict[str, Any]:
        async with self.session.begin():
            batch = await self._batch_by_no(batch_no, for_update=True)
            self._require_business(batch, "case_master")
            self._require_batch_ready(batch)
            source_lock_key = (
                f"case-master-source|{self.tenant_id}|{batch['data_source']}"
            )
            await self.session.execute(
                text(
                    "SELECT pg_advisory_xact_lock("
                    "hashtextextended(:source_lock_key, 0))"
                ),
                {"source_lock_key": source_lock_key},
            )
            if batch["import_mode"] == "full":
                current_source_hash = await self._current_case_source_hash(
                    str(batch["data_source"]),
                    lock_rows=True,
                )
                preview_source_hash = str(
                    (batch["metadata_json"] or {}).get("preview_source_snapshot_hash")
                    or ""
                )
                if current_source_hash != preview_source_hash:
                    raise DataIntakeError(
                        "ERP案件来源集合在预览后已经变化，请重新上传全量文件并核对",
                        code="case_source_snapshot_changed",
                        status_code=409,
                    )
            rows = await self._batch_rows(batch["batch_id"])
            for row in rows:
                action = row["proposed_action"]
                before = row["before_json"] or {}
                after = row["after_json"] or {}
                if action == "新增":
                    await self._require_new_case_identity(after)
                    case_id = _stable_uuid(
                        "legal-ops-case",
                        self.tenant_id,
                        str(after["source_system"]),
                        str(after["source_case_id"]),
                    )
                    case = Agent2Case(
                        case_id=case_id,
                        tenant_id=self.tenant_id,
                        company_id=str(after.get("company_id") or ""),
                        department_id=str(after.get("department_id") or ""),
                        team_id=str(after.get("team_id") or ""),
                        external_case_id=str(after["source_case_id"]),
                        case_number=str(after.get("case_number") or ""),
                        case_name=str(after.get("case_name") or ""),
                        case_type=str(after.get("case_type") or ""),
                        status=str(after.get("status") or "open"),
                        owner_user_id=str(after.get("owner_user_id") or ""),
                        source_type=str(after["source_system"]),
                        source_id=f"{batch_no}:{row['source_row_number']}",
                        source_json=dict(after.get("source_json") or {}),
                        version=1,
                    )
                    self.session.add(case)
                elif action == "更新":
                    case = await self._lock_case_preview(before)
                    case_id = case.case_id
                    case.company_id = str(after.get("company_id") or case.company_id)
                    case.department_id = str(
                        after.get("department_id") or case.department_id
                    )
                    case.team_id = str(after.get("team_id") or "")
                    case.case_number = str(after.get("case_number") or "")
                    case.case_name = str(after.get("case_name") or "")
                    case.case_type = str(after.get("case_type") or "")
                    case.status = str(after.get("status") or "")
                    case.owner_user_id = str(after.get("owner_user_id") or "")
                    case.source_type = str(after["source_system"])
                    case.source_id = f"{batch_no}:{row['source_row_number']}"
                    case.source_json = dict(after.get("source_json") or {})
                    case.version += 1
                elif action == "无变化":
                    case = await self._lock_case_preview(before)
                    case_id = case.case_id
                elif action == "待核验":
                    case = await self._lock_case_preview(before)
                    case_id = case.case_id
                    await self._upsert_case_source(
                        batch=batch,
                        row=row,
                        case_id=case_id,
                        snapshot=before,
                        missing=True,
                    )
                    continue
                else:
                    continue
                await self.session.flush()
                await self._upsert_case_source(
                    batch=batch,
                    row=row,
                    case_id=case_id,
                    snapshot=after,
                    missing=False,
                )
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_intake_rows SET published = true
                    WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                      AND validation_status <> 'error'
                    """
                ),
                {"tenant_id": self.tenant_id, "batch_id": batch["batch_id"]},
            )
            await self._mark_batch_published(batch["batch_id"])
            await self._audit(
                "publish_case_master",
                "case_master_batch",
                batch_no,
                f"原子发布案件主表批次：{batch_no}",
                {
                    "added": batch["added_rows"],
                    "updated": batch["updated_rows"],
                    "skipped": batch["skipped_rows"],
                    "warnings": batch["warning_rows"],
                },
                batch_id=batch["batch_id"],
            )
        return {
            "batch_no": batch_no,
            "status": "published",
            "status_label": "案件主表已原子发布",
            "published_by": self.actor_user_id,
            "counts": self._counts_from_batch(batch),
        }

    async def publish_case_progress(self, batch_no: str) -> dict[str, Any]:
        async with self.session.begin():
            batch = await self._batch_by_no(batch_no, for_update=True)
            self._require_business(batch, "case_progress")
            self._require_batch_ready(batch)
            rows = await self._batch_rows(batch["batch_id"])
            for row in rows:
                if row["proposed_action"] not in {"新增", "更新"}:
                    continue
                after = dict(row["after_json"] or {})
                before = dict(row["before_json"] or {})
                occurred_at = datetime.combine(
                    date.fromisoformat(str(after["occurred_at"])),
                    time.min,
                    tzinfo=timezone.utc,
                )
                fingerprint = str(after["fingerprint"])
                external_id = str(after.get("external_progress_id") or "")
                idempotency_key = (
                    f"file-progress:{self.tenant_id}:{after['source_system']}:"
                    f"{external_id or fingerprint}"
                )
                details = _progress_details(after)
                if row["proposed_action"] == "新增":
                    if await self._progress_source_row(after):
                        raise DataIntakeError(
                            "预览后已出现同一来源进展，请重新上传生成新预览",
                            code="progress_changed",
                            status_code=409,
                        )
                    progress_id = _stable_uuid("legal-ops-progress", idempotency_key)
                    progress = CaseProgress(
                        progress_id=progress_id,
                        tenant_id=self.tenant_id,
                        case_id=self._uuid(str(after["case_id"]), "案件"),
                        occurred_at=occurred_at,
                        recorded_at=_utcnow(),
                        reporter_id=str(after["reporter_id"]),
                        progress_type=str(after["progress_type"]),
                        summary=str(after["summary"]),
                        details=details,
                        source_message_id=f"{batch_no}:{row['source_row_number']}",
                        source_channel="file_import",
                        content_origin="imported_record",
                        related_party_ids=[],
                        related_document_ids=[],
                        related_travel_intent_ids=[],
                        confidence=Decimal(1),
                        confirmation_status="confirmed_by_import_publisher",
                        version=1,
                        idempotency_key=idempotency_key,
                    )
                    self.session.add(progress)
                else:
                    progress_id = self._uuid(str(before.get("progress_id")), "案件进展")
                    progress = await self.session.scalar(
                        select(CaseProgress)
                        .where(
                            CaseProgress.tenant_id == self.tenant_id,
                            CaseProgress.progress_id == progress_id,
                            CaseProgress.content_origin == "imported_record",
                        )
                        .with_for_update()
                    )
                    if progress is None:
                        raise DataIntakeError(
                            "文件来源进展已变化，发布已中止",
                            code="progress_changed",
                            status_code=409,
                        )
                    source_map = await self._progress_source_row(before)
                    self._require_progress_preview_current(
                        progress=progress,
                        source_map=source_map,
                        before=before,
                    )
                    await self.session.execute(
                        text(
                            """
                            INSERT INTO legal_ops_case_progress_source_history (
                                tenant_id, source_map_id, source_version,
                                snapshot_json, changed_by, source_batch_id
                            ) VALUES (
                                :tenant_id, :source_map_id, :source_version,
                                CAST(:snapshot AS jsonb), :actor, :batch_id
                            )
                            """
                        ),
                        {
                            "tenant_id": self.tenant_id,
                            "source_map_id": source_map["source_map_id"],
                            "source_version": source_map["source_version"],
                            "snapshot": _json(source_map["current_snapshot_json"]),
                            "actor": self.actor_user_id,
                            "batch_id": batch["batch_id"],
                        },
                    )
                    progress.occurred_at = occurred_at
                    progress.reporter_id = str(after["reporter_id"])
                    progress.progress_type = str(after["progress_type"])
                    progress.summary = str(after["summary"])
                    progress.details = details
                    progress.source_message_id = (
                        f"{batch_no}:{row['source_row_number']}"
                    )
                    progress.version += 1
                await self.session.flush()
                await self._upsert_progress_source(
                    batch=batch,
                    row=row,
                    progress_id=progress_id,
                    snapshot=after,
                )
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_intake_rows SET published = true
                    WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                      AND validation_status <> 'error'
                    """
                ),
                {"tenant_id": self.tenant_id, "batch_id": batch["batch_id"]},
            )
            await self._mark_batch_published(batch["batch_id"])
            await self._audit(
                "publish_case_progress",
                "case_progress_batch",
                batch_no,
                f"原子发布案件进展批次：{batch_no}",
                self._counts_from_batch(batch),
                batch_id=batch["batch_id"],
            )
        return {
            "batch_no": batch_no,
            "status": "published",
            "status_label": "案件进展已原子发布",
            "published_by": self.actor_user_id,
            "counts": self._counts_from_batch(batch),
        }

    async def list_batches(
        self,
        *,
        business_type: str | None = None,
        limit: int = 100,
    ) -> list[dict[str, Any]]:
        condition = "AND business_type = :business_type" if business_type else ""
        result = await self.session.execute(
            text(
                f"""
                SELECT *
                FROM legal_ops_intake_batches
                WHERE tenant_id = :tenant_id {condition}
                ORDER BY uploaded_at DESC
                LIMIT :limit
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "business_type": business_type,
                "limit": limit,
            },
        )
        return [self._batch_payload(row) for row in result.mappings().all()]

    async def batch_detail(
        self,
        batch_no: str,
        *,
        offset: int = 0,
        limit: int = 100,
        row_status: str | None = None,
    ) -> dict[str, Any]:
        batch = await self._batch_by_no(batch_no)
        rows, total_rows = await self._batch_rows_page(
            batch["batch_id"],
            offset=offset,
            limit=limit,
            row_status=row_status,
        )
        return {
            **self._batch_payload(batch),
            "row_page": {
                "offset": offset,
                "limit": limit,
                "total": total_rows,
                "has_more": offset + len(rows) < total_rows,
                "row_status": row_status or "",
            },
            "rows": [
                {
                    "row_ref": str(row["row_id"]),
                    "source_row_number": row["source_row_number"],
                    "status": row["validation_status"],
                    "status_label": _validation_status_label(row["validation_status"]),
                    "action": row["proposed_action"],
                    "action_label": ACTION_LABELS.get(
                        row["proposed_action"], row["proposed_action"]
                    ),
                    "data": _display_row(
                        row["raw_snapshot_json"] or row["normalized_json"]
                    ),
                    "errors": [
                        {
                            "code": item.get("code", ""),
                            "field": item.get("field", ""),
                            "message": item.get("message", ""),
                            "severity": "错误"
                            if item.get("critical", True)
                            else "提醒",
                        }
                        for item in (row["errors_json"] or [])
                    ],
                    "before": _display_row(row["before_json"])
                    if row["before_json"]
                    else None,
                    "after": _display_row(row["after_json"])
                    if row["after_json"]
                    else None,
                    "published": row["published"],
                }
                for row in rows
            ],
        }

    async def list_errors(
        self,
        *,
        error_type: str | None = None,
        limit: int = 100,
    ) -> list[dict[str, Any]]:
        result = await self.session.execute(
            text(
                """
                SELECT r.row_id, b.batch_no, b.business_type, b.file_name, b.uploaded_at,
                       r.source_row_number, r.errors_json, r.normalized_json,
                       r.raw_snapshot_json
                FROM legal_ops_intake_rows r
                JOIN legal_ops_intake_batches b ON b.batch_id = r.batch_id
                WHERE r.tenant_id = :tenant_id AND r.validation_status = 'error'
                ORDER BY b.uploaded_at DESC, r.source_row_number
                LIMIT :limit
                """
            ),
            {"tenant_id": self.tenant_id, "limit": limit * 5},
        )
        items: list[dict[str, Any]] = []
        for row in result.mappings().all():
            errors = [
                item
                for item in row["errors_json"] or []
                if not error_type or item.get("code") == error_type
            ]
            if not errors:
                continue
            items.append(
                {
                    "row_ref": str(row["row_id"]),
                    "batch_no": row["batch_no"],
                    "business_type": row["business_type"],
                    "business_type_label": BUSINESS_LABELS[row["business_type"]],
                    "file_name": row["file_name"],
                    "source_row_number": row["source_row_number"],
                    "uploaded_at": row["uploaded_at"].isoformat(),
                    "data": _display_row(
                        row["raw_snapshot_json"] or row["normalized_json"]
                    ),
                    "errors": [
                        {
                            "code": item.get("code", ""),
                            "message": item.get("message", ""),
                            "field": item.get("field", ""),
                        }
                        for item in errors
                    ],
                    "can_resolve": row["business_type"] == "case_progress",
                }
            )
            if len(items) >= limit:
                break
        return items

    async def error_resolution_options(self, row_ref: str) -> dict[str, Any]:
        row_id = self._uuid(row_ref, "错误记录")
        row = await self._error_row(row_id)
        if row["business_type"] != "case_progress":
            raise DataIntakeError(
                "当前只支持在网页中处理案件进展匹配错误",
                code="error_not_editable",
                status_code=409,
            )
        cases = (await self._case_index()).cases
        people_result = await self.session.execute(
            select(
                Agent2IdentityBinding.user_id,
                Agent2IdentityBinding.display_name,
                Agent2IdentityBinding.permission_scope_json,
            ).where(
                Agent2IdentityBinding.tenant_id == self.tenant_id,
                Agent2IdentityBinding.active.is_(True),
            )
        )
        people = []
        for user_id, display_name, scope in people_result.all():
            employee = str(
                (scope or {}).get("employee_no")
                or (scope or {}).get("employee_id")
                or ""
            ).strip()
            people.append(
                {
                    "value": str(user_id),
                    "label": (
                        f"{display_name}（员工编号 {employee}）"
                        if employee
                        else str(display_name or "未命名人员")
                    ),
                }
            )
        return {
            "row_ref": str(row_id),
            "batch_no": str(row["batch_no"]),
            "source_row_number": int(row["source_row_number"]),
            "current": _display_row(dict(row["normalized_json"] or {})),
            "errors": list(row["errors_json"] or []),
            "cases": [
                {
                    "value": str(item.get("case_id") or ""),
                    "label": " · ".join(
                        value
                        for value in (
                            str(item.get("case_number") or "").strip(),
                            str(item.get("case_name") or "").strip(),
                            str(item.get("status") or "").strip(),
                        )
                        if value
                    )
                    or "未命名案件",
                }
                for item in sorted(
                    cases,
                    key=lambda value: (
                        str(value.get("case_number") or ""),
                        str(value.get("case_name") or ""),
                    ),
                )
            ],
            "procedure_nodes": sorted(await self._known_nodes()),
            "people": sorted(people, key=lambda item: item["label"]),
            "notice": (
                "系统不会根据模糊案件名替你选择；只有你明确选定后，该行才会重新校验。"
            ),
        }

    async def resolve_case_progress_error(
        self,
        row_ref: str,
        *,
        case_ref: str = "",
        procedure_node: str = "",
        reporter_ref: str = "",
        progress_date: str = "",
        content: str = "",
        next_plan: str = "",
        apply_same_node: bool = True,
        apply_same_reporter: bool = True,
    ) -> dict[str, Any]:
        row_id = self._uuid(row_ref, "错误记录")
        async with self.session.begin():
            target = await self._error_row(row_id, for_update=True)
            if target["business_type"] != "case_progress":
                raise DataIntakeError(
                    "当前只支持在网页中处理案件进展匹配错误",
                    code="error_not_editable",
                    status_code=409,
                )
            if target["batch_status"] not in {"validation_failed", "ready"}:
                raise DataIntakeError(
                    "该批次已经发布、放弃或失效，不能再修改",
                    code="batch_not_editable",
                    status_code=409,
                )
            normalized = dict(target["normalized_json"] or {})
            original_node = str(normalized.get("procedure_node") or "")
            original_reporter = str(normalized.get("reporter_id") or "")
            if case_ref:
                normalized["manual_case_id"] = str(self._uuid(case_ref, "人工选择案件"))
            if procedure_node:
                normalized["procedure_node"] = procedure_node.strip()
            if reporter_ref:
                normalized["reporter_id"] = reporter_ref.strip()
            if progress_date:
                normalized["progress_date"] = progress_date.strip()
            if content:
                normalized["content"] = content.strip()
            if next_plan or "next_plan" in normalized:
                normalized["next_plan"] = next_plan.strip()
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_intake_rows
                    SET normalized_json = CAST(:normalized AS jsonb)
                    WHERE tenant_id = :tenant_id AND row_id = :row_id
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "row_id": row_id,
                    "normalized": _json(normalized),
                },
            )
            all_rows = await self._batch_rows(target["batch_id"])
            incoming = []
            for row in all_rows:
                payload = dict(
                    normalized
                    if row["row_id"] == row_id
                    else row["normalized_json"] or {}
                )
                if (
                    procedure_node
                    and apply_same_node
                    and str(payload.get("procedure_node") or "") == original_node
                ):
                    payload["procedure_node"] = procedure_node.strip()
                if (
                    reporter_ref
                    and apply_same_reporter
                    and str(payload.get("reporter_id") or "") == original_reporter
                ):
                    payload["reporter_id"] = reporter_ref.strip()
                payload["__row_number__"] = int(row["source_row_number"])
                incoming.append(payload)
            preview = preview_case_progress(
                incoming,
                cases=await self._case_index(),
                existing_progress=await self._progress_index(),
                known_nodes=await self._known_nodes(),
                known_people=(await self._known_people_and_teams())[0],
            )
            preview_by_row = {item.row_number: item for item in preview.rows}
            for row in all_rows:
                item = preview_by_row[int(row["source_row_number"])]
                status_value = (
                    "error"
                    if any(error.critical for error in item.errors)
                    else "warning"
                    if item.errors
                    else "skipped"
                    if item.action == "重复跳过"
                    else "valid"
                )
                await self.session.execute(
                    text(
                        """
                        UPDATE legal_ops_intake_rows
                        SET normalized_json = CAST(:normalized AS jsonb),
                            validation_status = :status,
                            proposed_action = :action,
                            matched_object_type = :object_type,
                            matched_object_id = :object_id,
                            before_json = CAST(:before AS jsonb),
                            after_json = CAST(:after AS jsonb),
                            errors_json = CAST(:errors AS jsonb),
                            fingerprint = :fingerprint
                        WHERE tenant_id = :tenant_id AND row_id = :row_id
                        """
                    ),
                    {
                        "tenant_id": self.tenant_id,
                        "row_id": row["row_id"],
                        "normalized": _json(item.normalized),
                        "status": status_value,
                        "action": item.action,
                        "object_type": (
                            "案件"
                            if (item.before or item.after or {}).get("case_id")
                            else ""
                        ),
                        "object_id": str(
                            (item.before or item.after or {}).get("case_id") or ""
                        ),
                        "before": (
                            _json(item.before) if item.before is not None else None
                        ),
                        "after": _json(item.after) if item.after is not None else None,
                        "errors": _json([asdict(error) for error in item.errors]),
                        "fingerprint": str(
                            (item.after or item.normalized).get("fingerprint") or ""
                        )
                        or None,
                    },
                )
            counts = dict(preview.counts)
            critical_errors = sum(
                any(error.critical for error in item.errors) for item in preview.rows
            )
            warning_rows = sum(
                bool(item.errors) and not any(error.critical for error in item.errors)
                for item in preview.rows
            )
            batch_status = "validation_failed" if critical_errors else "ready"
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_intake_batches
                    SET added_rows = :added, updated_rows = :updated,
                        skipped_rows = :skipped, conflict_rows = :conflict,
                        failed_rows = :failed, warning_rows = :warnings,
                        status = :status, error_summary = :summary,
                        updated_at = now()
                    WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "batch_id": target["batch_id"],
                    "added": counts.get("新增", 0),
                    "updated": counts.get("更新", 0),
                    "skipped": counts.get("重复跳过", 0),
                    "conflict": counts.get("冲突", 0),
                    "failed": counts.get("失败", 0),
                    "warnings": warning_rows,
                    "status": batch_status,
                    "summary": (
                        f"{critical_errors}条错误仍待处理" if critical_errors else ""
                    ),
                },
            )
            await self._audit(
                "resolve_case_progress_error",
                "case_progress_import_row",
                str(row_id),
                f"人工处理案件进展导入第{target['source_row_number']}行",
                {
                    "batch_no": target["batch_no"],
                    "remaining_errors": critical_errors,
                },
                batch_id=target["batch_id"],
            )
        return {
            "row_ref": str(row_id),
            "batch_no": target["batch_no"],
            "status": batch_status,
            "status_label": STATUS_LABELS[batch_status],
            "remaining_error_rows": critical_errors,
            "can_publish": not critical_errors,
        }

    async def _error_row(
        self,
        row_id: uuid.UUID,
        *,
        for_update: bool = False,
    ) -> Any:
        suffix = " FOR UPDATE OF r, b" if for_update else ""
        result = await self.session.execute(
            text(
                f"""
                SELECT r.*, b.batch_no, b.business_type,
                       b.status AS batch_status
                FROM legal_ops_intake_rows r
                JOIN legal_ops_intake_batches b ON b.batch_id = r.batch_id
                WHERE r.tenant_id = :tenant_id AND r.row_id = :row_id
                {suffix}
                """
            ),
            {"tenant_id": self.tenant_id, "row_id": row_id},
        )
        row = result.mappings().first()
        if not row:
            raise DataIntakeError(
                "错误记录不存在",
                code="not_found",
                status_code=404,
            )
        return row

    async def abandon_batch(self, batch_no: str) -> dict[str, Any]:
        # The API may have read the batch first to determine its business
        # permission. End that read transaction before opening the atomic write.
        await self.session.rollback()
        async with self.session.begin():
            batch = await self._batch_by_no(batch_no, for_update=True)
            if batch["status"] == "published":
                raise DataIntakeError(
                    "已发布批次不可放弃", code="published_batch", status_code=409
                )
            if batch["status"] == "abandoned":
                return {**self._batch_payload(batch), "duplicate_abandon": True}
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_intake_batches
                    SET status = 'abandoned', abandoned_by = :actor,
                        abandoned_at = now(), updated_at = now()
                    WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                    """
                ),
                {
                    "actor": self.actor_user_id,
                    "tenant_id": self.tenant_id,
                    "batch_id": batch["batch_id"],
                },
            )
            if batch["business_type"] == "performance_source_table":
                await self.session.execute(
                    text(
                        """
                        UPDATE legal_ops_performance_source_versions
                        SET status = 'abandoned'
                        WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                          AND status IN ('staged','validation_failed')
                        """
                    ),
                    {"tenant_id": self.tenant_id, "batch_id": batch["batch_id"]},
                )
            await self._audit(
                "abandon_batch",
                "intake_batch",
                batch_no,
                f"放弃未发布批次：{batch_no}",
                {},
                batch_id=batch["batch_id"],
            )
        return {
            "batch_no": batch_no,
            "status": "abandoned",
            "status_label": "已放弃",
            "duplicate_abandon": False,
        }

    async def download_batch_errors(self, batch_no: str) -> bytes:
        batch = await self._batch_by_no(batch_no)
        rows = await self._batch_rows(batch["batch_id"])
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "错误数据"
        raw_headers: list[str] = []
        for row in rows:
            for header in row["raw_snapshot_json"] or {}:
                if header not in raw_headers:
                    raw_headers.append(str(header))
        sheet.append(
            [
                "原始行号",
                *(safe_excel_cell(header) for header in raw_headers),
                "错误原因",
            ]
        )
        for row in rows:
            if row["validation_status"] != "error":
                continue
            raw = row["raw_snapshot_json"] or {}
            sheet.append(
                [
                    row["source_row_number"],
                    *(safe_excel_cell(raw.get(header, "")) for header in raw_headers),
                    safe_excel_cell(
                        "；".join(
                            str(error.get("message") or "校验失败")
                            for error in row["errors_json"] or []
                        )
                    ),
                ]
            )
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    async def template(
        self,
        template_type: str,
        *,
        table_key: str = "",
        period_ref: str = "",
        rule_ref: str = "",
    ) -> bytes:
        if template_type == "case_master":
            schema = CASE_MASTER_SCHEMA
        elif template_type == "case_progress":
            schema = CASE_PROGRESS_SCHEMA
        elif template_type == "performance":
            period_id = self._uuid(period_ref, "考核周期")
            await self._require_period(period_id)
            active = (
                await self._active_rule_for_period_ref(period_id, rule_ref)
                if rule_ref
                else await self._rule_for_period(period_id)
            )
            if not active or not active.get("rule_spec_json"):
                raise DataIntakeError(
                    "规则文件尚未配置，暂时没有绩效源表模板",
                    code="rule_missing",
                    status_code=409,
                )
            table = next(
                (
                    item
                    for item in active["rule_spec_json"].get("source_tables") or []
                    if str(item.get("key")) == table_key
                ),
                None,
            )
            if not table:
                raise DataIntakeError(
                    "未找到该绩效源表模板", code="template_not_found", status_code=404
                )
            schema = performance_table_schema(
                table,
                row_key=aggregate_table_row_key(
                    dict(active["rule_spec_json"]),
                    table,
                )
                if str(active["rule_spec_json"].get("schema_version") or "") == "2"
                else "",
            )
        else:
            raise DataIntakeError("模板类型无效")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = _safe_sheet_title(schema.name)
        sheet.append([safe_excel_cell(column.name) for column in schema.columns])
        sheet.freeze_panes = "A2"
        note = workbook.create_sheet("填写说明")
        note.append(["字段", "是否必填", "数据格式", "是否唯一"])
        if schema.allow_extra_columns:
            note.append(
                [
                    "现有业务表其他列",
                    "可保留",
                    "系统只读取规则所需列，不要求删减ERP原表",
                    "否",
                ]
            )
        for column in schema.columns:
            note.append(
                [
                    safe_excel_cell(column.name),
                    "是" if column.required else "否",
                    _column_type_label(column.data_type),
                    "是" if column.unique else "否",
                ]
            )
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    # ---- internal persistence helpers -------------------------------------------------

    async def _save_case_preview(
        self,
        *,
        business_type: str,
        source_system: str,
        import_mode: str,
        content: bytes,
        filename: str,
        preview: ImportPreview,
        metadata: dict[str, Any] | None = None,
        idempotency_scope: str = "",
    ) -> dict[str, Any]:
        stored = self._store_file(content, filename)
        await self._register_original_file(stored)
        critical_errors = sum(
            any(error.critical for error in row.errors) for row in preview.rows
        )
        warning_rows = sum(
            bool(row.errors) and not any(error.critical for error in row.errors)
            for row in preview.rows
        )
        status_value = "validation_failed" if critical_errors else "ready"
        prefix = "CASE" if business_type == "case_master" else "PROG"
        async with self.session.begin():
            file_id = await self._register_file(stored)
            duplicate = await self._find_duplicate_batch(
                business_type=business_type,
                data_source=source_system,
                file_hash=stored.file_hash,
                period_id=None,
                table_key=_batch_scope_key(idempotency_scope),
                import_mode=import_mode,
            )
            if duplicate:
                return {**self._batch_payload(duplicate), "duplicate_upload": True}
            batch_id = uuid.uuid4()
            batch_no = _batch_no(prefix)
            action_counts = {key: value for key, value in preview.counts.items()}
            await self._insert_batch(
                batch_id=batch_id,
                batch_no=batch_no,
                business_type=business_type,
                data_source=source_system,
                file_id=file_id,
                stored=stored,
                period_id=None,
                table_key=_batch_scope_key(idempotency_scope),
                import_mode=import_mode,
                status=status_value,
                original_rows=sum(row.row_number > 0 for row in preview.rows),
                counts={
                    "added": action_counts.get("新增", 0),
                    "updated": action_counts.get("更新", 0),
                    "skipped": action_counts.get("无变化", 0)
                    + action_counts.get("重复跳过", 0),
                    "conflict": action_counts.get("冲突", 0),
                    "failed": action_counts.get("失败", 0),
                    "warning": action_counts.get("待核验", 0) + warning_rows,
                },
                metadata={
                    **(metadata or {}),
                    "source_system": source_system,
                    "preview_counts": action_counts,
                    "preview_source_snapshot_hash": (
                        _preview_case_source_hash(
                            preview,
                            source_system=source_system,
                        )
                        if business_type == "case_master" and import_mode == "full"
                        else ""
                    ),
                    "write_boundary": (
                        "仅ERP权威字段"
                        if business_type == "case_master"
                        else "仅文件导入来源进展"
                    ),
                },
            )
            await self._insert_preview_rows(batch_id, preview)
            await self._audit(
                f"upload_{business_type}",
                f"{business_type}_batch",
                batch_no,
                f"上传并完成{BUSINESS_LABELS[business_type]}预览：{stored.safe_file_name}",
                {
                    "counts": action_counts,
                    "publishable": preview.publishable,
                },
                batch_id=batch_id,
            )
        return {
            "batch_no": batch_no,
            "business_type": business_type,
            "business_type_label": BUSINESS_LABELS[business_type],
            "status": status_value,
            "status_label": STATUS_LABELS[status_value],
            "file_name": stored.safe_file_name,
            "file_hash": stored.file_hash,
            "row_count": sum(row.row_number > 0 for row in preview.rows),
            "counts": preview.counts,
            "error_count": critical_errors,
            "warning_count": warning_rows + preview.counts.get("待核验", 0),
            "can_publish": preview.publishable,
            "duplicate_upload": False,
            "source_profile": (metadata or {}).get("source_profile_label", ""),
            "source_sheet": (metadata or {}).get("detected_sheet", ""),
            "detected_columns": (metadata or {}).get("detected_columns", 0),
        }

    async def _register_file(self, stored: StoredFile) -> uuid.UUID:
        result = await self.session.execute(
            text(
                """
                INSERT INTO legal_ops_intake_files (
                    file_id, tenant_id, file_hash, safe_file_name, media_type,
                    size_bytes, storage_key, uploaded_by
                ) VALUES (
                    :file_id, :tenant_id, :file_hash, :file_name, :media_type,
                    :size_bytes, :storage_key, :actor
                )
                ON CONFLICT (tenant_id, file_hash)
                DO UPDATE SET safe_file_name = legal_ops_intake_files.safe_file_name
                RETURNING file_id
                """
            ),
            {
                "file_id": uuid.uuid4(),
                "tenant_id": self.tenant_id,
                "file_hash": stored.file_hash,
                "file_name": stored.safe_file_name,
                "media_type": _media_type(stored.safe_file_name),
                "size_bytes": stored.size_bytes,
                "storage_key": stored.storage_key,
                "actor": self.actor_user_id,
            },
        )
        return result.scalar_one()

    async def _find_duplicate_batch(
        self,
        *,
        business_type: str,
        data_source: str,
        file_hash: str,
        period_id: uuid.UUID | None,
        table_key: str,
        import_mode: str,
        rule_version: str = "",
        rule_spec_hash: str = "",
    ) -> Any | None:
        lock_key = "|".join(
            (
                self.tenant_id,
                business_type,
                data_source,
                file_hash,
                str(period_id or ""),
                table_key,
                import_mode,
                rule_version,
                rule_spec_hash,
            )
        )
        await self.session.execute(
            text("SELECT pg_advisory_xact_lock(hashtextextended(:lock_key, 0))"),
            {"lock_key": lock_key},
        )
        result = await self.session.execute(
            text(
                """
                SELECT *
                FROM legal_ops_intake_batches
                WHERE tenant_id = :tenant_id
                  AND business_type = :business_type
                  AND data_source = :data_source
                  AND file_hash = :file_hash
                  AND period_id IS NOT DISTINCT FROM :period_id
                  AND table_key = :table_key
                  AND import_mode = :import_mode
                  AND (
                    :rule_version = ''
                    OR metadata_json ->> 'rule_version' = :rule_version
                  )
                  AND (
                    :rule_spec_hash = ''
                    OR metadata_json ->> 'rule_spec_hash' = :rule_spec_hash
                  )
                  AND status NOT IN ('abandoned','failed')
                ORDER BY uploaded_at DESC
                LIMIT 1
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "business_type": business_type,
                "data_source": data_source,
                "file_hash": file_hash,
                "period_id": period_id,
                "table_key": table_key,
                "import_mode": import_mode,
                "rule_version": rule_version,
                "rule_spec_hash": rule_spec_hash,
            },
        )
        return result.mappings().first()

    async def _insert_batch(
        self,
        *,
        batch_id: uuid.UUID,
        batch_no: str,
        business_type: str,
        data_source: str,
        file_id: uuid.UUID,
        stored: StoredFile,
        period_id: uuid.UUID | None,
        table_key: str,
        import_mode: str,
        status: str,
        original_rows: int,
        counts: dict[str, int],
        metadata: dict[str, Any],
    ) -> None:
        await self.session.execute(
            text(
                """
                INSERT INTO legal_ops_intake_batches (
                    batch_id, batch_no, tenant_id, business_type, data_source,
                    file_id, file_name, file_hash, uploaded_by, period_id,
                    table_key, import_mode, original_rows, added_rows, updated_rows,
                    skipped_rows, conflict_rows, failed_rows, warning_rows, status,
                    code_version, error_summary, metadata_json
                ) VALUES (
                    :batch_id, :batch_no, :tenant_id, :business_type, :data_source,
                    :file_id, :file_name, :file_hash, :actor, :period_id,
                    :table_key, :import_mode, :original_rows, :added, :updated,
                    :skipped, :conflict, :failed, :warning, :status,
                    :code_version, :error_summary, CAST(:metadata AS jsonb)
                )
                """
            ),
            {
                "batch_id": batch_id,
                "batch_no": batch_no,
                "tenant_id": self.tenant_id,
                "business_type": business_type,
                "data_source": data_source,
                "file_id": file_id,
                "file_name": stored.safe_file_name,
                "file_hash": stored.file_hash,
                "actor": self.actor_user_id,
                "period_id": period_id,
                "table_key": table_key,
                "import_mode": import_mode,
                "original_rows": original_rows,
                "added": counts.get("added", 0),
                "updated": counts.get("updated", 0),
                "skipped": counts.get("skipped", 0),
                "conflict": counts.get("conflict", 0),
                "failed": counts.get("failed", 0),
                "warning": counts.get("warning", 0),
                "status": status,
                "code_version": self.code_version,
                "error_summary": (
                    f"{counts.get('failed', 0)}条错误，{counts.get('conflict', 0)}条冲突"
                    if counts.get("failed", 0) or counts.get("conflict", 0)
                    else ""
                ),
                "metadata": _json(metadata),
            },
        )

    async def _insert_parsed_rows(
        self, batch_id: uuid.UUID, parsed: ParsedTable
    ) -> None:
        for row in parsed.rows:
            await self.session.execute(
                text(
                    """
                    INSERT INTO legal_ops_intake_rows (
                        tenant_id, batch_id, source_row_number, raw_snapshot_json,
                        normalized_json, validation_status, errors_json
                    ) VALUES (
                        :tenant_id, :batch_id, :row_number, CAST(:raw AS jsonb),
                        CAST(:normalized AS jsonb), :validation_status, CAST(:errors AS jsonb)
                    )
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "batch_id": batch_id,
                    "row_number": row.row_number,
                    "raw": _json(row.raw),
                    "normalized": _json(row.normalized),
                    "validation_status": (
                        "error"
                        if any(error.critical for error in row.errors)
                        else "warning"
                        if row.errors
                        else "valid"
                    ),
                    "errors": _json([asdict(error) for error in row.errors]),
                },
            )

    async def _insert_preview_rows(
        self, batch_id: uuid.UUID, preview: ImportPreview
    ) -> None:
        for row in preview.rows:
            status_value = (
                "error"
                if any(error.critical for error in row.errors)
                else "warning"
                if row.errors
                else "skipped"
                if row.action in {"无变化", "重复跳过"}
                else "valid"
            )
            fingerprint = (
                str((row.after or row.normalized).get("fingerprint") or "") or None
            )
            await self.session.execute(
                text(
                    """
                    INSERT INTO legal_ops_intake_rows (
                        tenant_id, batch_id, source_row_number, raw_snapshot_json,
                        normalized_json, validation_status, proposed_action,
                        matched_object_type, matched_object_id, before_json, after_json,
                        errors_json, fingerprint
                    ) VALUES (
                        :tenant_id, :batch_id, :row_number, CAST(:raw AS jsonb),
                        CAST(:normalized AS jsonb), :status, :action,
                        :object_type, :object_id, CAST(:before AS jsonb), CAST(:after AS jsonb),
                        CAST(:errors AS jsonb), :fingerprint
                    )
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "batch_id": batch_id,
                    "row_number": row.row_number,
                    "raw": _json(row.raw),
                    "normalized": _json(row.normalized),
                    "status": status_value,
                    "action": row.action,
                    "object_type": (
                        "案件" if (row.before or row.after or {}).get("case_id") else ""
                    ),
                    "object_id": str(
                        (row.before or row.after or {}).get("case_id") or ""
                    ),
                    "before": _json(row.before) if row.before is not None else None,
                    "after": _json(row.after) if row.after is not None else None,
                    "errors": _json([asdict(error) for error in row.errors]),
                    "fingerprint": fingerprint,
                },
            )

    async def _batch_by_no(self, batch_no: str, *, for_update: bool = False) -> Any:
        suffix = " FOR UPDATE" if for_update else ""
        result = await self.session.execute(
            text(
                f"""
                SELECT * FROM legal_ops_intake_batches
                WHERE tenant_id = :tenant_id AND batch_no = :batch_no
                {suffix}
                """
            ),
            {"tenant_id": self.tenant_id, "batch_no": batch_no},
        )
        row = result.mappings().first()
        if not row:
            raise DataIntakeError("导入批次不存在", code="not_found", status_code=404)
        return row

    async def _batch_by_id(self, batch_id: uuid.UUID) -> Any:
        result = await self.session.execute(
            text(
                """
                SELECT * FROM legal_ops_intake_batches
                WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                """
            ),
            {"tenant_id": self.tenant_id, "batch_id": batch_id},
        )
        row = result.mappings().first()
        if not row:
            raise DataIntakeError("导入批次不存在", code="not_found", status_code=404)
        return row

    async def _lock_performance_period(self, period_id: uuid.UUID) -> None:
        lock_key = f"performance-source-period|{self.tenant_id}|{period_id}"
        await self.session.execute(
            text("SELECT pg_advisory_xact_lock(hashtextextended(:lock_key, 0))"),
            {"lock_key": lock_key},
        )

    async def _mark_stale_performance_batch(
        self,
        batch_no: str,
        reason: str,
    ) -> None:
        await self.session.rollback()
        async with self.session.begin():
            batch = await self._batch_by_no(batch_no)
            self._require_business(batch, "performance_source_table")
            await self._lock_performance_period(batch["period_id"])
            await self._mark_stale_performance_batch_in_transaction(
                batch["batch_id"],
                reason,
            )

    async def _mark_stale_performance_batch_in_transaction(
        self,
        batch_id: uuid.UUID,
        reason: str,
    ) -> None:
        result = await self.session.execute(
            text(
                """
                UPDATE legal_ops_intake_batches
                SET status = 'failed', error_summary = :reason, updated_at = now()
                WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                  AND status NOT IN ('published','abandoned','failed')
                RETURNING batch_no
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "batch_id": batch_id,
                "reason": reason,
            },
        )
        batch_no = result.scalar_one_or_none()
        if batch_no is None:
            return
        await self.session.execute(
            text(
                """
                UPDATE legal_ops_performance_source_versions
                SET status = 'abandoned'
                WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                  AND status IN ('staged','validation_failed')
                """
            ),
            {"tenant_id": self.tenant_id, "batch_id": batch_id},
        )
        await self._audit(
            "invalidate_performance_source_batch",
            "performance_source_batch",
            str(batch_no),
            reason,
            {},
            batch_id=batch_id,
        )

    async def _batch_rows(self, batch_id: uuid.UUID) -> list[Any]:
        result = await self.session.execute(
            text(
                """
                SELECT * FROM legal_ops_intake_rows
                WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                ORDER BY source_row_number, created_at
                """
            ),
            {"tenant_id": self.tenant_id, "batch_id": batch_id},
        )
        return list(result.mappings().all())

    async def _batch_rows_page(
        self,
        batch_id: uuid.UUID,
        *,
        offset: int,
        limit: int,
        row_status: str | None,
    ) -> tuple[list[Any], int]:
        condition = (
            "AND validation_status = :row_status"
            if row_status in {"valid", "warning", "error", "skipped"}
            else ""
        )
        parameters = {
            "tenant_id": self.tenant_id,
            "batch_id": batch_id,
            "row_status": row_status,
            "offset": offset,
            "limit": limit,
        }
        total = await self.session.scalar(
            text(
                f"""
                SELECT count(*)
                FROM legal_ops_intake_rows
                WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                {condition}
                """
            ),
            parameters,
        )
        result = await self.session.execute(
            text(
                f"""
                SELECT * FROM legal_ops_intake_rows
                WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                {condition}
                ORDER BY source_row_number, created_at
                LIMIT :limit OFFSET :offset
                """
            ),
            parameters,
        )
        return list(result.mappings().all()), int(total or 0)

    async def _inspect_stored_rule_package(self, row: Any):
        file_row = await self.session.execute(
            text(
                """
                SELECT f.storage_key, f.safe_file_name
                FROM legal_ops_intake_batches b
                JOIN legal_ops_intake_files f ON f.file_id = b.file_id
                WHERE b.tenant_id = :tenant_id AND b.batch_id = :batch_id
                """
            ),
            {"tenant_id": self.tenant_id, "batch_id": row["source_batch_id"]},
        )
        stored_rule = file_row.mappings().first()
        if not stored_rule:
            raise DataIntakeError(
                "规则原始文件不存在",
                code="source_file_missing",
                status_code=409,
            )
        try:
            content = self.files.read(str(stored_rule["storage_key"]))
        except FileStorageError as exc:
            raise DataIntakeError(
                "规则原始文件不存在或校验失败",
                code="source_file_missing",
                status_code=409,
            ) from exc
        try:
            return inspect_rule_package(
                content,
                str(stored_rule["safe_file_name"]),
            )
        except RulePackageError as exc:
            raise DataIntakeError(
                f"规则原始文件无法再次验证：{exc}",
                code="invalid_rule_package",
                status_code=409,
            ) from exc

    async def _rule_source_file_row(self, source_file_id: uuid.UUID) -> Any:
        result = await self.session.execute(
            text(
                """
                SELECT s.*, f.safe_file_name, f.file_hash, f.size_bytes,
                       f.storage_key, f.uploaded_by, f.uploaded_at
                FROM legal_ops_performance_rule_source_files s
                JOIN legal_ops_intake_files f ON f.file_id = s.file_id
                WHERE s.tenant_id = :tenant_id
                  AND s.source_file_id = :source_file_id
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "source_file_id": source_file_id,
            },
        )
        row = result.mappings().first()
        if not row:
            raise DataIntakeError(
                "规则核对底表不存在",
                code="not_found",
                status_code=404,
            )
        return row

    async def _rule_source_file_context(self, rule_id: uuid.UUID) -> str:
        result = await self.session.execute(
            text(
                """
                SELECT s.business_label, s.inspection_json,
                       f.safe_file_name, f.file_hash
                FROM legal_ops_performance_rule_source_files s
                JOIN legal_ops_intake_files f ON f.file_id = s.file_id
                WHERE s.tenant_id = :tenant_id
                  AND s.rule_version_id = :rule_id
                  AND s.status = 'active'
                ORDER BY s.source_file_no
                """
            ),
            {"tenant_id": self.tenant_id, "rule_id": rule_id},
        )
        lines: list[str] = []
        for row in result.mappings().all():
            inspection = dict(row["inspection_json"] or {})
            lines.append(
                f"底表：{row['business_label'] or row['safe_file_name']}；"
                f"文件：{row['safe_file_name']}；"
                f"文件哈希：{str(row['file_hash'])[:12]}…；"
                f"数据行：{inspection.get('total_data_rows', 0)}"
            )
            for sheet in inspection.get("sheets") or []:
                columns = "；".join(
                    f"第{column.get('position')}列“{column.get('header')}”"
                    f"（{column.get('inferred_type', 'unknown')}）"
                    for column in (sheet.get("columns") or [])
                )
                lines.append(
                    f"- 工作表“{sheet.get('sheet_name', '')}”："
                    f"{sheet.get('row_count', 0)}行、"
                    f"{sheet.get('column_count', 0)}列。表头：{columns}"
                )
            for warning in inspection.get("warnings") or []:
                lines.append(f"- 底表检查提醒：{warning}")
        return "\n".join(lines)[:80_000]

    async def _next_rule_amendment_no(self, rule_id: uuid.UUID) -> int:
        value = await self.session.scalar(
            text(
                """
                SELECT COALESCE(MAX(amendment_no), 0) + 1
                FROM legal_ops_performance_rule_amendments
                WHERE tenant_id = :tenant_id AND rule_version_id = :rule_id
                """
            ),
            {"tenant_id": self.tenant_id, "rule_id": rule_id},
        )
        return int(value or 1)

    async def _rule_row(self, rule_id: uuid.UUID, *, for_update: bool = False) -> Any:
        suffix = " FOR UPDATE" if for_update else ""
        result = await self.session.execute(
            text(
                f"""
                SELECT * FROM legal_ops_performance_rule_versions
                WHERE tenant_id = :tenant_id AND rule_version_id = :rule_id
                {suffix}
                """
            ),
            {"tenant_id": self.tenant_id, "rule_id": rule_id},
        )
        row = result.mappings().first()
        if not row:
            raise DataIntakeError("规则版本不存在", code="not_found", status_code=404)
        return row

    async def _latest_active_rule(self) -> Any | None:
        result = await self.session.execute(
            text(
                """
                SELECT * FROM legal_ops_performance_rule_versions
                WHERE tenant_id = :tenant_id AND is_current AND status = 'active'
                ORDER BY activated_at DESC
                LIMIT 1
                """
            ),
            {"tenant_id": self.tenant_id},
        )
        return result.mappings().first()

    async def _rule_for_period(self, period_id: uuid.UUID) -> Any | None:
        rules = await self._rules_for_period(period_id)
        if len(rules) > 1:
            raise DataIntakeError(
                "本周期有多个负责板块，请明确选择要使用的规则",
                code="rule_selection_required",
                status_code=409,
            )
        return rules[0] if rules else None

    async def _rules_for_period(self, period_id: uuid.UUID) -> list[Any]:
        result = await self.session.execute(
            text(
                """
                SELECT DISTINCT ON (business_scope_key) *
                FROM legal_ops_performance_rule_versions
                WHERE tenant_id = :tenant_id
                  AND is_current
                  AND status = 'active'
                  AND (period_id = :period_id OR period_id IS NULL)
                ORDER BY business_scope_key,
                         (period_id = :period_id) DESC,
                         activated_at DESC
                """
            ),
            {"tenant_id": self.tenant_id, "period_id": period_id},
        )
        return list(result.mappings().all())

    async def _rule_for_validation(
        self,
        period_id: uuid.UUID,
        *,
        rule_ref: str | None,
        for_update: bool = False,
    ) -> Any | None:
        if not rule_ref:
            rules = await self._rules_for_period(period_id)
            if len(rules) > 1:
                raise DataIntakeError(
                    "本周期有多个负责板块，请明确选择要使用的规则",
                    code="rule_selection_required",
                    status_code=409,
                )
            return rules[0] if rules else None
        rule_id = self._uuid(rule_ref, "规则")
        row = await self._rule_row(rule_id, for_update=for_update)
        if row["period_id"] and row["period_id"] != period_id:
            raise DataIntakeError(
                "该规则草稿不适用于所选考核周期",
                code="rule_period_mismatch",
                status_code=409,
            )
        if str(row["status"]) not in {
            "understanding_draft",
            "awaiting_confirmation",
            "active",
        }:
            raise DataIntakeError(
                "请先读取并理解规则，再上传验证底表",
                code="rule_not_interpreted",
                status_code=409,
            )
        if not row["rule_spec_json"]:
            raise DataIntakeError(
                "规则草稿尚未形成固定计算结构；请先完善计算口径",
                code="rule_spec_missing",
                status_code=409,
            )
        try:
            validate_rule_spec(dict(row["rule_spec_json"]))
        except CalculationBlocked as exc:
            raise DataIntakeError(
                f"固定计算结构未通过校验：{exc}",
                code="invalid_rule_spec",
                status_code=409,
            ) from exc
        return row

    async def _active_rule_for_period_ref(
        self,
        period_id: uuid.UUID,
        rule_ref: str | uuid.UUID | None,
    ) -> Any | None:
        rules = await self._rules_for_period(period_id)
        if not rule_ref:
            if len(rules) > 1:
                raise DataIntakeError(
                    "本周期有多个负责板块，请先选择要试算或维护的板块",
                    code="rule_selection_required",
                    status_code=409,
                )
            return rules[0] if rules else None
        requested = self._uuid(str(rule_ref), "规则")
        for row in rules:
            if row["rule_version_id"] == requested:
                return row
        raise DataIntakeError(
            "所选规则不是本周期当前有效版本",
            code="rule_not_active_for_period",
            status_code=409,
        )

    async def _require_period(self, period_id: uuid.UUID) -> Any:
        result = await self.session.execute(
            text(
                """
                SELECT * FROM legal_ops_performance_periods
                WHERE tenant_id = :tenant_id AND period_id = :period_id
                """
            ),
            {"tenant_id": self.tenant_id, "period_id": period_id},
        )
        row = result.mappings().first()
        if not row:
            raise DataIntakeError("考核周期不存在", code="not_found", status_code=404)
        return row

    async def _current_source_versions(
        self,
        period_id: uuid.UUID,
        *,
        rule_version_id: uuid.UUID | None = None,
        rule_version: str = "",
    ) -> dict[str, Any]:
        result = await self.session.execute(
            text(
                """
                SELECT s.*, b.file_hash, b.metadata_json
                FROM legal_ops_performance_source_versions s
                JOIN legal_ops_intake_batches b ON b.batch_id = s.batch_id
                WHERE s.tenant_id = :tenant_id AND s.period_id = :period_id
                  AND s.is_current AND s.status = 'published'
                  AND (
                    CAST(:rule_version_id AS uuid) IS NULL
                    OR s.rule_version_id = CAST(:rule_version_id AS uuid)
                  )
                  AND (
                    :rule_version = ''
                    OR b.metadata_json ->> 'rule_version' = :rule_version
                  )
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "period_id": period_id,
                "rule_version_id": rule_version_id,
                "rule_version": rule_version,
            },
        )
        return {str(row["table_key"]): row for row in result.mappings().all()}

    async def _period_package(
        self, period_id: uuid.UUID, active_rule: Any | None
    ) -> dict[str, Any]:
        if not active_rule or not active_rule.get("rule_spec_json"):
            return _empty_period_package("规则文件尚未配置，暂不可上传绩效源表或试算")
        spec = active_rule["rule_spec_json"]
        current = await self._current_source_versions(
            period_id,
            rule_version_id=active_rule["rule_version_id"],
            rule_version=str(active_rule["rule_version"]),
        )
        staged_result = await self.session.execute(
            text(
                """
                SELECT DISTINCT ON (s.table_key)
                       s.table_key, s.table_name, s.purpose, s.version, s.row_count,
                       s.error_count, s.status, s.created_by, s.created_at,
                       b.batch_no, b.file_hash
                FROM legal_ops_performance_source_versions s
                JOIN legal_ops_intake_batches b ON b.batch_id = s.batch_id
                WHERE s.tenant_id = :tenant_id AND s.period_id = :period_id
                  AND s.rule_version_id = :rule_id
                  AND b.metadata_json ->> 'rule_version' = :rule_version
                  AND s.status IN ('staged','validation_failed','published')
                  AND b.status IN ('ready','validation_failed','published')
                ORDER BY s.table_key, s.version DESC
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "period_id": period_id,
                "rule_id": active_rule["rule_version_id"],
                "rule_version": active_rule["rule_version"],
            },
        )
        latest = {str(row["table_key"]): row for row in staged_result.mappings().all()}
        tables = []
        for table in spec.get("source_tables") or []:
            key = str(table["key"])
            version = latest.get(key)
            current_version = current.get(key)
            if not version:
                status_value = "missing"
                status_label = "未上传"
            elif version["error_count"]:
                status_value = "error"
                status_label = f"存在{version['error_count']}条错误"
                if current_version:
                    status_label += f"；当前正式版为第{current_version['version']}版"
            elif version["status"] == "staged":
                status_value = "ready"
                status_label = f"第{version['version']}版校验通过，待发布"
                if current_version:
                    status_label += f"；当前正式版为第{current_version['version']}版"
            elif current_version:
                status_value = "published"
                status_label = f"第{current_version['version']}版已发布"
            else:
                status_value = "ready"
                status_label = "校验通过，待发布"
            tables.append(
                {
                    "table_key": key,
                    "table_name": table.get("name", key),
                    "purpose": table.get("purpose", ""),
                    "required": bool(table.get("required", True)),
                    "status": status_value,
                    "status_label": status_label,
                    "version": int(version["version"]) if version else None,
                    "current_version": (
                        int(current_version["version"]) if current_version else None
                    ),
                    "row_count": int(version["row_count"]) if version else 0,
                    "error_count": int(version["error_count"]) if version else 0,
                    "uploaded_by": version["created_by"] if version else "",
                    "uploaded_at": version["created_at"].isoformat() if version else "",
                    "batch_no": version["batch_no"] if version else "",
                    "source_file_hash": str(version["file_hash"]) if version else "",
                }
            )
        required = [item for item in tables if item["required"]]
        return {
            "rule_configured": True,
            "rule_ref": str(active_rule["rule_version_id"]),
            "rule_version": active_rule["rule_version"],
            "business_scope_name": (
                active_rule.get("business_scope_name") or active_rule["skill_name"]
            ),
            "business_scope_key": active_rule.get("business_scope_key") or "",
            "required_table_count": len(required),
            "uploaded_count": sum(item["status"] != "missing" for item in required),
            "passed_count": sum(
                item["status"] in {"ready", "published"} for item in required
            ),
            "error_count": sum(item["status"] == "error" for item in required),
            "missing_count": sum(item["status"] == "missing" for item in required),
            "can_calculate": bool(required)
            and all(item["status"] == "published" for item in required),
            "tables": tables,
        }

    async def _apply_person_and_cross_table_checks(
        self,
        parsed: ParsedTable,
        *,
        period_id: uuid.UUID,
        table_key: str,
        person_key: str,
        rule_version_id: uuid.UUID,
        rule_version: str,
    ) -> None:
        identity_map = await self._identity_value_map()
        rows_by_user: dict[str, list[tuple[ParsedRow, str]]] = {}
        for row in parsed.rows:
            value = str(row.normalized.get(person_key) or "").strip()
            if not value:
                continue
            matches = identity_map.get(value, set())
            if not matches:
                row.errors.append(
                    RowError("person_not_found", person_key, "人员稳定标识无法匹配")
                )
            elif len(matches) > 1:
                row.errors.append(
                    RowError("ambiguous_person", person_key, "一个人员标识匹配到多个人")
                )
            else:
                user_id = next(iter(matches))
                rows_by_user.setdefault(user_id, []).append((row, value))
        for matched_rows in rows_by_user.values():
            identifiers = {value for _, value in matched_rows}
            if len(identifiers) <= 1:
                continue
            for row, _ in matched_rows:
                row.errors.append(
                    RowError(
                        "multiple_person_identifiers",
                        person_key,
                        "同一人员在本表出现多个稳定标识，必须人工核对",
                    )
                )
        current = await self._current_source_versions(
            period_id,
            rule_version_id=rule_version_id,
            rule_version=rule_version,
        )
        incoming = {
            str(row.normalized.get(person_key) or "").strip()
            for row in parsed.rows
            if row.normalized.get(person_key)
        }
        for other_key, version in current.items():
            if other_key == table_key:
                continue
            result = await self.session.execute(
                text(
                    """
                    SELECT stable_person_id
                    FROM legal_ops_performance_source_rows
                    WHERE tenant_id = :tenant_id AND source_version_id = :version_id
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "version_id": version["source_version_id"],
                },
            )
            other = {str(value) for value in result.scalars().all()}
            for row in parsed.rows:
                value = str(row.normalized.get(person_key) or "").strip()
                if value and value not in other:
                    row.errors.append(
                        RowError(
                            "cross_table_missing",
                            person_key,
                            f"该人员在已发布的{version['table_name']}中不存在",
                        )
                    )
            missing_in_incoming = other.difference(incoming)
            if missing_in_incoming and parsed.rows:
                parsed.rows[0].errors.append(
                    RowError(
                        "cross_table_missing",
                        person_key,
                        f"本表缺少{len(missing_in_incoming)}名已在{version['table_name']}中的人员",
                    )
                )

    async def _identity_value_map(self) -> dict[str, set[str]]:
        result = await self.session.execute(
            select(
                Agent2IdentityBinding.user_id,
                Agent2IdentityBinding.dingtalk_user_id,
                Agent2IdentityBinding.permission_scope_json,
            ).where(
                Agent2IdentityBinding.tenant_id == self.tenant_id,
                Agent2IdentityBinding.active.is_(True),
            )
        )
        values: dict[str, set[str]] = {}
        for user_id, dingtalk_id, scope in result.all():
            stable_values = {
                str(user_id or "").strip(),
                str(dingtalk_id or "").strip(),
            }
            for key in ("employee_id", "erp_person_id", "system_user_id"):
                stable_values.add(str((scope or {}).get(key) or "").strip())
            for value in stable_values:
                if value:
                    values.setdefault(value, set()).add(str(user_id))
        return values

    async def _canonicalize_person_column(
        self,
        parsed: ParsedTable,
        column_key: str,
        *,
        preserve_unlinked: bool = False,
        source_system: str = "",
    ) -> None:
        identity_map, name_map = await self._case_person_maps()
        for row in parsed.rows:
            value = str(row.normalized.get(column_key) or "").strip()
            if not value:
                continue
            stable_candidate = _prefixed_employee_identifier(value)
            matches = identity_map.get(
                (stable_candidate or value).casefold(),
                set(),
            )
            if len(matches) == 1:
                row.normalized[column_key] = next(iter(matches))
                if column_key == "owner_user_id":
                    row.normalized["owner_source_value"] = value
                    row.normalized["owner_link_status"] = "已按稳定编号关联"
                    row.normalized["owner_display_name"] = (
                        _display_name_without_identifier(value)
                    )
                continue
            if len(matches) > 1:
                row.errors.append(
                    RowError(
                        "ambiguous_person",
                        column_key,
                        "该人员标识匹配到多个人，必须人工处理",
                    )
                )
                continue
            display_candidate = _display_name_without_identifier(value)
            name_matches = name_map.get(display_candidate.casefold(), set())
            if len(name_matches) == 1:
                row.normalized[column_key] = next(iter(name_matches))
                if column_key == "owner_user_id":
                    row.normalized["owner_source_value"] = value
                    row.normalized["owner_link_status"] = (
                        "仅按唯一姓名关联，待发布人核对"
                    )
                    row.normalized["owner_display_name"] = display_candidate
                row.errors.append(
                    RowError(
                        "matched_by_unique_name",
                        column_key,
                        "原文件未提供可匹配的人员稳定编号，系统仅按唯一姓名匹配；"
                        "发布前必须人工核对",
                        critical=False,
                    )
                )
            elif not name_matches and preserve_unlinked:
                row.normalized[column_key] = _unlinked_person_reference(
                    source_system,
                    stable_candidate or value,
                )
                row.normalized["owner_source_value"] = value
                row.normalized["owner_source_stable_id"] = stable_candidate
                row.normalized["owner_display_name"] = display_candidate
                row.normalized["owner_link_status"] = (
                    "来源人员编号尚未关联系统账号"
                    if stable_candidate
                    else "来源仅有姓名，尚未关联系统账号"
                )
                row.errors.append(
                    RowError(
                        "owner_unlinked",
                        column_key,
                        (
                            "来源负责人尚未关联现有登录账号；系统保留原始负责人，"
                            "不会把案件错误分配给同名人员。发布后个人案件权限暂不生效"
                        ),
                        critical=False,
                    )
                )
            elif not name_matches:
                row.errors.append(
                    RowError(
                        "person_not_found",
                        column_key,
                        "承办法务无法匹配到现有人员，请补充人员关系后重试",
                    )
                )
            else:
                row.errors.append(
                    RowError(
                        "ambiguous_person",
                        column_key,
                        "存在同名人员，不能仅按姓名自动关联",
                    )
                )

    async def _case_person_maps(
        self,
    ) -> tuple[dict[str, set[str]], dict[str, set[str]]]:
        result = await self.session.execute(
            select(
                Agent2IdentityBinding.user_id,
                Agent2IdentityBinding.dingtalk_user_id,
                Agent2IdentityBinding.display_name,
                Agent2IdentityBinding.permission_scope_json,
            ).where(
                Agent2IdentityBinding.tenant_id == self.tenant_id,
                Agent2IdentityBinding.active.is_(True),
            )
        )
        stable: dict[str, set[str]] = {}
        names: dict[str, set[str]] = {}
        for user_id, dingtalk_id, display_name, scope in result.all():
            target = str(user_id)
            for value in (
                user_id,
                dingtalk_id,
                (scope or {}).get("employee_id"),
                (scope or {}).get("employee_no"),
                (scope or {}).get("erp_person_id"),
                (scope or {}).get("system_user_id"),
            ):
                text_value = str(value or "").strip()
                if text_value:
                    stable.setdefault(text_value.casefold(), set()).add(target)
            name = str(display_name or "").strip()
            if name:
                names.setdefault(name.casefold(), set()).add(target)
        return stable, names

    async def _canonicalize_team_column(
        self,
        parsed: ParsedTable,
        column_key: str,
        *,
        preserve_unlinked: bool = False,
        source_system: str = "",
    ) -> None:
        team_map = await self._case_team_map()
        for row in parsed.rows:
            value = str(row.normalized.get(column_key) or "").strip()
            if not value:
                continue
            matches = team_map.get(value.casefold(), set())
            if len(matches) == 1:
                row.normalized[column_key] = next(iter(matches))
                row.normalized["team_source_value"] = value
                row.normalized["team_link_status"] = "已关联现有团队"
            elif not matches and preserve_unlinked:
                row.normalized[column_key] = _unlinked_team_reference(
                    source_system,
                    value,
                )
                row.normalized["team_source_value"] = value
                row.normalized["team_link_status"] = "历史来源团队尚未关联现有团队"
                row.errors.append(
                    RowError(
                        "team_unlinked",
                        column_key,
                        (
                            "来源团队未在当前法务团队目录中找到；系统保留原始团队，"
                            "不会错误归入其他团队。发布后该案件暂不进入现任团队范围"
                        ),
                        critical=False,
                    )
                )
            elif not matches:
                row.errors.append(
                    RowError(
                        "team_not_found",
                        column_key,
                        "所属团队无法匹配，请先维护团队关系",
                    )
                )
            else:
                row.errors.append(
                    RowError(
                        "ambiguous_team",
                        column_key,
                        "该团队名称对应多个团队，必须人工处理",
                    )
                )

    async def _case_team_map(self) -> dict[str, set[str]]:
        bindings_result = await self.session.execute(
            select(
                Agent2IdentityBinding.team_id,
                Agent2IdentityBinding.permission_scope_json,
            ).where(
                Agent2IdentityBinding.tenant_id == self.tenant_id,
                Agent2IdentityBinding.active.is_(True),
            )
        )
        values: dict[str, set[str]] = {}
        for team_id, scope in bindings_result.all():
            target = str(team_id or "").strip()
            if not target:
                continue
            for alias in (
                target,
                str((scope or {}).get("team_id") or "").strip(),
                str((scope or {}).get("team_code") or "").strip(),
                str((scope or {}).get("team_name") or "").strip(),
                str((scope or {}).get("department_name") or "").strip(),
            ):
                if alias:
                    values.setdefault(alias.casefold(), set()).add(target)
        teams_result = await self.session.execute(
            select(Team.id, Team.code, Team.name).where(
                Team.active.is_(True),
                Team.department_name == "法务合约中心",
            )
        )
        for team_id, code, name in teams_result.all():
            target = str(code or team_id)
            for alias in (
                str(team_id),
                str(code or "").strip(),
                str(name or "").strip(),
            ):
                if alias:
                    values.setdefault(alias.casefold(), set()).add(target)
        return values

    async def _person_display_names(self) -> dict[str, str]:
        result = await self.session.execute(
            select(
                Agent2IdentityBinding.user_id,
                Agent2IdentityBinding.dingtalk_user_id,
                Agent2IdentityBinding.display_name,
                Agent2IdentityBinding.permission_scope_json,
            ).where(
                Agent2IdentityBinding.tenant_id == self.tenant_id,
                Agent2IdentityBinding.active.is_(True),
            )
        )
        names: dict[str, str] = {}
        for user_id, dingtalk_id, display_name, scope in result.all():
            for value in (
                user_id,
                dingtalk_id,
                (scope or {}).get("employee_id"),
                (scope or {}).get("erp_person_id"),
                (scope or {}).get("system_user_id"),
            ):
                if value:
                    names[str(value)] = str(display_name or "")
        return names

    async def _known_people_and_teams(self) -> tuple[set[str], set[str]]:
        result = await self.session.execute(
            select(Agent2IdentityBinding).where(
                Agent2IdentityBinding.tenant_id == self.tenant_id,
                Agent2IdentityBinding.active.is_(True),
            )
        )
        people: set[str] = set()
        teams: set[str] = set()
        for binding in result.scalars().all():
            people.add(binding.user_id)
            if binding.team_id:
                teams.add(binding.team_id)
        team_rows = await self.session.execute(
            select(Team.code).where(
                Team.active.is_(True),
                Team.department_name == "法务合约中心",
            )
        )
        teams.update(str(code) for code in team_rows.scalars() if code)
        return people, teams

    async def _known_nodes(self) -> set[str]:
        result = await self.session.execute(
            select(CaseLifecycleState.node)
            .where(CaseLifecycleState.tenant_id == self.tenant_id)
            .distinct()
        )
        return {
            str(value).strip() for value in result.scalars().all() if str(value).strip()
        }

    async def _require_new_case_identity(self, after: dict[str, Any]) -> None:
        source_system = str(after.get("source_system") or "")
        source_case_id = str(after.get("source_case_id") or "")
        source_map = await self.session.execute(
            text(
                """
                SELECT case_id
                FROM legal_ops_case_master_sources
                WHERE tenant_id = :tenant_id
                  AND source_system = :source_system
                  AND source_case_id = :source_case_id
                FOR UPDATE
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "source_system": source_system,
                "source_case_id": source_case_id,
            },
        )
        existing_case_id = source_map.scalar_one_or_none()
        if existing_case_id is None:
            existing_case_id = await self.session.scalar(
                select(Agent2Case.case_id)
                .where(
                    Agent2Case.tenant_id == self.tenant_id,
                    Agent2Case.source_type == source_system,
                    Agent2Case.external_case_id == source_case_id,
                )
                .with_for_update()
            )
        if existing_case_id is not None:
            raise DataIntakeError(
                "预览后已出现同一来源案件，请重新上传生成新预览",
                code="case_changed",
                status_code=409,
            )

    async def _lock_case_preview(self, before: dict[str, Any]) -> Agent2Case:
        case_id = self._uuid(str(before.get("case_id")), "案件")
        case = await self.session.scalar(
            select(Agent2Case)
            .where(
                Agent2Case.tenant_id == self.tenant_id,
                Agent2Case.case_id == case_id,
            )
            .with_for_update()
        )
        if case is None or _hash_json(_case_snapshot(case)) != _hash_json(before):
            raise DataIntakeError(
                "案件在预览后已经变化，请重新上传生成新预览",
                code="case_changed",
                status_code=409,
            )
        return case

    @staticmethod
    def _require_progress_preview_current(
        *,
        progress: CaseProgress,
        source_map: Any | None,
        before: dict[str, Any],
    ) -> None:
        if source_map is None:
            raise DataIntakeError(
                "文件来源进展在预览后已经变化，请重新上传",
                code="progress_changed",
                status_code=409,
            )
        expected_snapshot = {
            key: value
            for key, value in before.items()
            if key
            not in {
                "source_map_id",
                "progress_id",
                "version",
            }
        }
        if (
            int(source_map["source_version"]) != int(before.get("version") or 0)
            or int(progress.version) != int(before.get("version") or 0)
            or _hash_json(source_map["current_snapshot_json"])
            != _hash_json(expected_snapshot)
        ):
            raise DataIntakeError(
                "文件来源进展在预览后已经变化，请重新上传",
                code="progress_changed",
                status_code=409,
            )

    async def _case_index(self) -> CaseMasterIndex:
        result = await self.session.execute(
            select(Agent2Case).where(Agent2Case.tenant_id == self.tenant_id)
        )
        cases = []
        for case in result.scalars().all():
            cases.append(_case_snapshot(case))
        return CaseMasterIndex(cases)

    async def _current_case_source_hash(
        self,
        source_system: str,
        *,
        lock_rows: bool,
    ) -> str:
        lock_clause = "FOR SHARE" if lock_rows else ""
        result = await self.session.execute(
            text(
                f"""
                SELECT case_id::text AS case_id,
                       external_case_id AS source_case_id,
                       version
                FROM agent2_cases
                WHERE tenant_id = :tenant_id
                  AND UPPER(
                      COALESCE(
                          NULLIF(source_json ->> 'source_system', ''),
                          source_type
                      )
                  ) = UPPER(:source_system)
                ORDER BY case_id
                {lock_clause}
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "source_system": source_system,
            },
        )
        return _hash_json(
            [
                {
                    "case_id": str(row["case_id"]),
                    "source_case_id": str(row["source_case_id"]),
                    "version": int(row["version"]),
                }
                for row in result.mappings().all()
            ]
        )

    async def _progress_index(self) -> CaseProgressIndex:
        result = await self.session.execute(
            text(
                """
                SELECT source_map_id, source_system, external_progress_id,
                       fingerprint, progress_id, case_id, source_version,
                       current_snapshot_json
                FROM legal_ops_case_progress_sources
                WHERE tenant_id = :tenant_id
                """
            ),
            {"tenant_id": self.tenant_id},
        )
        values = []
        for row in result.mappings().all():
            item = dict(row["current_snapshot_json"] or {})
            item.update(
                {
                    "source_map_id": str(row["source_map_id"]),
                    "source_system": row["source_system"],
                    "external_progress_id": row["external_progress_id"],
                    "fingerprint": row["fingerprint"],
                    "progress_id": str(row["progress_id"]),
                    "case_id": str(row["case_id"]),
                    "content_origin": "imported_record",
                    "version": row["source_version"],
                }
            )
            values.append(item)
        return CaseProgressIndex(values)

    async def _upsert_case_source(
        self,
        *,
        batch: Any,
        row: Any,
        case_id: uuid.UUID,
        snapshot: dict[str, Any],
        missing: bool,
    ) -> None:
        source_system = str(
            snapshot.get("source_system")
            or snapshot.get("source_type")
            or (snapshot.get("source_json") or {}).get("source_system")
            or batch["data_source"]
        )
        source_case_id = str(
            snapshot.get("source_case_id") or snapshot.get("external_case_id") or ""
        )
        await self.session.execute(
            text(
                """
                INSERT INTO legal_ops_case_master_sources (
                    tenant_id, source_system, source_case_id, case_id,
                    source_batch_id, source_file_id, source_row_number,
                    source_version, source_snapshot_json, missing_from_latest_snapshot
                ) VALUES (
                    :tenant_id, :source_system, :source_case_id, :case_id,
                    :batch_id, :file_id, :row_number, 1, CAST(:snapshot AS jsonb), :missing
                )
                ON CONFLICT (tenant_id, source_system, source_case_id)
                DO UPDATE SET
                    case_id = EXCLUDED.case_id,
                    source_batch_id = EXCLUDED.source_batch_id,
                    source_file_id = EXCLUDED.source_file_id,
                    source_row_number = EXCLUDED.source_row_number,
                    source_version = legal_ops_case_master_sources.source_version + 1,
                    source_snapshot_json = EXCLUDED.source_snapshot_json,
                    missing_from_latest_snapshot = EXCLUDED.missing_from_latest_snapshot,
                    updated_at = now()
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "source_system": source_system,
                "source_case_id": source_case_id,
                "case_id": case_id,
                "batch_id": batch["batch_id"],
                "file_id": batch["file_id"],
                "row_number": row["source_row_number"],
                "snapshot": _json(snapshot),
                "missing": missing,
            },
        )

    async def _progress_source_row(self, before: dict[str, Any]) -> Any | None:
        external_id = str(before.get("external_progress_id") or "")
        if external_id:
            result = await self.session.execute(
                text(
                    """
                    SELECT * FROM legal_ops_case_progress_sources
                    WHERE tenant_id = :tenant_id AND source_system = :source_system
                      AND external_progress_id = :external_id
                    FOR UPDATE
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "source_system": before.get("source_system"),
                    "external_id": external_id,
                },
            )
        else:
            result = await self.session.execute(
                text(
                    """
                    SELECT * FROM legal_ops_case_progress_sources
                    WHERE tenant_id = :tenant_id AND fingerprint = :fingerprint
                    FOR UPDATE
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "fingerprint": before.get("fingerprint"),
                },
            )
        return result.mappings().first()

    async def _upsert_progress_source(
        self,
        *,
        batch: Any,
        row: Any,
        progress_id: uuid.UUID,
        snapshot: dict[str, Any],
    ) -> None:
        external_id = str(snapshot.get("external_progress_id") or "")
        existing = await self._progress_source_row(snapshot)
        if existing:
            await self.session.execute(
                text(
                    """
                    UPDATE legal_ops_case_progress_sources
                    SET progress_id = :progress_id, case_id = :case_id,
                        source_batch_id = :batch_id, source_file_id = :file_id,
                        source_row_number = :row_number,
                        source_version = source_version + 1,
                        current_snapshot_json = CAST(:snapshot AS jsonb),
                        fingerprint = :fingerprint, updated_at = now()
                    WHERE source_map_id = :source_map_id
                    """
                ),
                {
                    "progress_id": progress_id,
                    "case_id": self._uuid(str(snapshot["case_id"]), "案件"),
                    "batch_id": batch["batch_id"],
                    "file_id": batch["file_id"],
                    "row_number": row["source_row_number"],
                    "snapshot": _json(snapshot),
                    "fingerprint": snapshot["fingerprint"],
                    "source_map_id": existing["source_map_id"],
                },
            )
        else:
            await self.session.execute(
                text(
                    """
                    INSERT INTO legal_ops_case_progress_sources (
                        tenant_id, source_system, external_progress_id, fingerprint,
                        progress_id, case_id, source_batch_id, source_file_id,
                        source_row_number, source_version, current_snapshot_json
                    ) VALUES (
                        :tenant_id, :source_system, :external_id, :fingerprint,
                        :progress_id, :case_id, :batch_id, :file_id,
                        :row_number, 1, CAST(:snapshot AS jsonb)
                    )
                    """
                ),
                {
                    "tenant_id": self.tenant_id,
                    "source_system": snapshot["source_system"],
                    "external_id": external_id,
                    "fingerprint": snapshot["fingerprint"],
                    "progress_id": progress_id,
                    "case_id": self._uuid(str(snapshot["case_id"]), "案件"),
                    "batch_id": batch["batch_id"],
                    "file_id": batch["file_id"],
                    "row_number": row["source_row_number"],
                    "snapshot": _json(snapshot),
                },
            )

    async def _mark_batch_published(self, batch_id: uuid.UUID) -> None:
        await self.session.execute(
            text(
                """
                UPDATE legal_ops_intake_batches
                SET status = 'published', published_by = :actor,
                    published_at = now(), updated_at = now()
                WHERE tenant_id = :tenant_id AND batch_id = :batch_id
                """
            ),
            {
                "actor": self.actor_user_id,
                "tenant_id": self.tenant_id,
                "batch_id": batch_id,
            },
        )

    async def _invalidate_calculations(
        self,
        *,
        period_id: uuid.UUID | None,
        reason: str,
        rule_version_ids: tuple[uuid.UUID, ...] = (),
    ) -> None:
        condition = "AND period_id = :period_id" if period_id else ""
        rule_condition = (
            "AND rule_version_id = ANY(CAST(:rule_version_ids AS uuid[]))"
            if rule_version_ids
            else ""
        )
        await self.session.execute(
            text(
                f"""
                UPDATE legal_ops_performance_calculation_runs
                SET status = 'invalidated', invalidated_reason = :reason
                WHERE tenant_id = :tenant_id
                  AND status IN ('trial','published') {condition} {rule_condition}
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "period_id": period_id,
                "reason": reason,
                "rule_version_ids": list(rule_version_ids),
            },
        )

    async def _refresh_period_status(self, period_id: uuid.UUID) -> None:
        active_rules = await self._rules_for_period(period_id)
        if not active_rules:
            return
        packages_ready = []
        for active in active_rules:
            if not active.get("rule_spec_json"):
                packages_ready.append(False)
                continue
            versions = await self._current_source_versions(
                period_id,
                rule_version_id=active["rule_version_id"],
                rule_version=str(active["rule_version"]),
            )
            required = {
                str(item["key"])
                for item in active["rule_spec_json"].get("source_tables") or []
                if bool(item.get("required", True))
            }
            packages_ready.append(bool(required) and required <= set(versions))
        status_value = (
            "ready" if packages_ready and all(packages_ready) else "preparing"
        )
        await self.session.execute(
            text(
                """
                UPDATE legal_ops_performance_periods
                SET status = :status, updated_at = now()
                WHERE tenant_id = :tenant_id AND period_id = :period_id
                  AND status <> 'closed'
                """
            ),
            {
                "status": status_value,
                "tenant_id": self.tenant_id,
                "period_id": period_id,
            },
        )

    async def _audit(
        self,
        action: str,
        object_type: str,
        object_id: str,
        summary: str,
        details: dict[str, Any],
        *,
        batch_id: uuid.UUID | None = None,
    ) -> None:
        await self.session.execute(
            text(
                """
                INSERT INTO legal_ops_intake_audit_events (
                    tenant_id, actor_user_id, action, object_type,
                    object_id, batch_id, summary, details_json
                ) VALUES (
                    :tenant_id, :actor, :action, :object_type,
                    :object_id, :batch_id, :summary, CAST(:details AS jsonb)
                )
                """
            ),
            {
                "tenant_id": self.tenant_id,
                "actor": self.actor_user_id,
                "action": action,
                "object_type": object_type,
                "object_id": object_id,
                "batch_id": batch_id,
                "summary": summary,
                "details": _json(details),
            },
        )

    def _batch_payload(self, row: Any) -> dict[str, Any]:
        metadata = dict(row["metadata_json"] or {})
        return {
            "batch_no": row["batch_no"],
            "business_type": row["business_type"],
            "business_type_label": BUSINESS_LABELS[row["business_type"]],
            "data_source": row["data_source"],
            "file_name": row["file_name"],
            "file_hash": row["file_hash"],
            "uploaded_by": row["uploaded_by"],
            "uploaded_at": row["uploaded_at"].isoformat(),
            "import_mode": row["import_mode"],
            "status": row["status"],
            "status_label": STATUS_LABELS[row["status"]],
            "counts": self._counts_from_batch(row),
            "error_summary": row["error_summary"],
            "published_by": row["published_by"],
            "published_at": row["published_at"].isoformat()
            if row["published_at"]
            else "",
            "can_publish": row["status"] == "ready",
            "can_abandon": row["status"] not in {"published", "abandoned"},
            "source_profile": {
                "label": str(metadata.get("source_profile_label") or ""),
                "sheet": str(metadata.get("detected_sheet") or ""),
                "column_count": int(metadata.get("detected_columns") or 0),
                "profile_version": str(metadata.get("source_profile_version") or ""),
                "owner_link_policy": str(metadata.get("owner_link_policy") or ""),
                "team_link_policy": str(metadata.get("team_link_policy") or ""),
                "field_mapping": [
                    {
                        "business_field": str(item.get("business_field") or ""),
                        "source_header": str(item.get("source_header") or ""),
                    }
                    for item in (metadata.get("field_mapping") or [])
                    if isinstance(item, dict)
                ],
            },
        }

    @staticmethod
    def _counts_from_batch(row: Any) -> dict[str, int]:
        return {
            "原始行数": row["original_rows"],
            "新增": row["added_rows"],
            "更新": row["updated_rows"],
            "跳过": row["skipped_rows"],
            "冲突": row["conflict_rows"],
            "失败": row["failed_rows"],
            "提醒": row["warning_rows"],
        }

    @staticmethod
    def _rule_card(row: Any | None) -> dict[str, Any]:
        if not row:
            return {
                "configured": False,
                "status_label": "规则文件尚未配置，暂不可计算",
                "activation_ready": False,
            }
        understanding = row.get("understanding_json") or {}
        reference_documents = _rule_reference_documents(
            row.get("package_inventory_json")
        )
        return {
            "configured": bool(row.get("is_current")),
            "rule_ref": str(row["rule_version_id"]),
            "rule_version": row["rule_version"],
            "skill_name": row["skill_name"],
            "skill_version": row["skill_version"],
            "business_scope_name": row.get("business_scope_name") or "尚未填写",
            "draft_revision": int(row.get("draft_revision") or 1),
            "file_hash": row["source_file_hash"],
            "status": row["status"],
            "status_label": _rule_status_label(row["status"]),
            "understanding_hash": row["understanding_hash"],
            "rule_spec_hash": row["rule_spec_hash"],
            "activation_ready": bool(understanding.get("activation_ready", False)),
            "unresolved": understanding.get("unresolved", []),
            "issue_resolution_count": len(understanding.get("issue_resolutions", [])),
            "reference_count": len(reference_documents),
            "created_by": row["created_by"],
            "created_at": row["created_at"].isoformat(),
            "is_current": row["is_current"],
        }

    @staticmethod
    def _rule_amendment_payload(row: Any) -> dict[str, Any]:
        status_value = str(row["status"])
        validation = dict(row["validation_json"] or {})
        return {
            "amendment_ref": str(row["amendment_id"]),
            "amendment_no": int(row["amendment_no"]),
            "change_kind": str(row["change_kind"]),
            "change_kind_label": (
                "网页直接编辑"
                if row["change_kind"] == "manual_edit"
                else "中文沟通修改"
            ),
            "instruction": str(row["instruction"] or ""),
            "issue_resolutions": list(row["issue_resolutions_json"] or []),
            "base_understanding_hash": str(row["base_understanding_hash"]),
            "proposal_hash": str(row["proposed_understanding_hash"]),
            "changes": list(row["changes_json"] or []),
            "validation": validation,
            "status": status_value,
            "status_label": {
                "proposed": "等待查看差异",
                "applied": "已应用到草稿",
                "rejected": "已放弃",
                "stale": "草稿已变化，建议已失效",
            }.get(status_value, status_value),
            "can_apply": status_value == "proposed"
            and bool(validation.get("can_apply", True)),
            "created_by": str(row["created_by"]),
            "created_at": row["created_at"].isoformat(),
            "applied_by": str(row["applied_by"] or ""),
            "applied_at": (row["applied_at"].isoformat() if row["applied_at"] else ""),
        }

    @staticmethod
    def _rule_source_file_payload(row: Any) -> dict[str, Any]:
        inspection = dict(row["inspection_json"] or {})
        status_value = str(row["status"])
        return {
            "source_file_ref": str(row["source_file_id"]),
            "source_file_no": int(row["source_file_no"]),
            "business_label": str(row["business_label"] or row["safe_file_name"]),
            "file_name": str(row["safe_file_name"]),
            "file_hash": str(row["file_hash"]),
            "size_bytes": int(row["size_bytes"]),
            "inspection": inspection,
            "status": status_value,
            "status_label": (
                "用于当前规则核对" if status_value == "active" else "已移出当前规则核对"
            ),
            "can_abandon": status_value == "active",
            "uploaded_by": str(row["uploaded_by"]),
            "uploaded_at": row["uploaded_at"].isoformat(),
            "created_by": str(row["created_by"]),
            "created_at": row["created_at"].isoformat(),
            "abandoned_by": str(row["abandoned_by"] or ""),
            "abandoned_at": (
                row["abandoned_at"].isoformat() if row["abandoned_at"] else ""
            ),
        }

    @staticmethod
    def _uuid(value: str | None, label: str) -> uuid.UUID:
        try:
            return uuid.UUID(str(value or ""))
        except ValueError:
            raise DataIntakeError(
                f"{label}标识无效", code="invalid_reference"
            ) from None

    @staticmethod
    def _require_rule_interpretable(row: Any) -> None:
        if bool(row["is_current"]) or str(row["status"]) not in {
            "uploaded",
            "understanding_draft",
        }:
            raise DataIntakeError(
                "该规则已进入确认或启用流程，不能原地改写；请上传新的技能包版本",
                code="rule_immutable",
                status_code=409,
            )

    @staticmethod
    def _require_rule_draft_mutable(row: Any) -> None:
        if bool(row["is_current"]) or str(row["status"]) not in {
            "understanding_draft",
            "awaiting_confirmation",
        }:
            raise DataIntakeError(
                "请先读取并理解规则；已启用规则必须上传新版本修改",
                code="rule_draft_not_editable",
                status_code=409,
            )

    @staticmethod
    def _require_rule_source_files_mutable(row: Any) -> None:
        if bool(row["is_current"]) and str(row["status"]) == "active":
            return
        if bool(row["is_current"]) or str(row["status"]) not in {
            "uploaded",
            "understanding_draft",
            "awaiting_confirmation",
        }:
            raise DataIntakeError(
                "已启用规则不能改变核对底表；请上传新的技能版本",
                code="rule_source_files_immutable",
                status_code=409,
            )

    @staticmethod
    def _require_business(batch: Any, expected: str) -> None:
        if batch["business_type"] != expected:
            raise DataIntakeError(
                "批次类型与当前操作不一致", code="wrong_batch_type", status_code=409
            )

    @staticmethod
    def _require_batch_ready(batch: Any) -> None:
        if batch["status"] == "published":
            raise DataIntakeError(
                "该批次已经发布", code="already_published", status_code=409
            )
        if batch["status"] != "ready":
            raise DataIntakeError(
                "批次存在错误或已失效，不能发布",
                code="batch_not_ready",
                status_code=409,
            )


def _merge_parse_errors(preview: ImportPreview, parsed: ParsedTable) -> ImportPreview:
    parse_by_row = {
        row.row_number: tuple(row.errors) for row in parsed.rows if row.errors
    }
    if not parse_by_row:
        return preview
    output: list[PreviewRow] = []
    for row in preview.rows:
        extra = parse_by_row.get(row.row_number, ())
        if not extra:
            output.append(row)
            continue
        converted = tuple(_row_error_to_import(error) for error in extra)
        has_critical = any(error.critical for error in extra)
        output.append(
            PreviewRow(
                row.row_number,
                "失败" if has_critical else row.action,
                row.raw,
                row.normalized,
                row.before,
                None if has_critical else row.after,
                row.errors + converted,
            )
        )
    counts = dict(preview.counts)
    for original, changed in zip(preview.rows, output):
        if original.action != changed.action:
            counts[original.action] = max(0, counts.get(original.action, 0) - 1)
            counts["失败"] = counts.get("失败", 0) + 1
    return ImportPreview(tuple(output), counts, preview.deletes)


def _source_system(value: str) -> str:
    normalized = str(value or "").strip().upper()
    if (
        not normalized
        or len(normalized) > 64
        or any(ord(character) < 32 for character in normalized)
    ):
        raise DataIntakeError("数据源不能为空、不能超过64字且不能包含控制字符")
    return normalized


def _business_scope_key(value: str) -> str:
    normalized = str(value or "").strip() or "未命名板块"
    return (
        "scope-"
        + hashlib.md5(
            normalized.encode("utf-8"),
            usedforsecurity=False,
        ).hexdigest()
    )


def _batch_scope_key(value: str) -> str:
    normalized = str(value or "").strip()
    if not normalized:
        return ""
    return "scope-" + hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:32]


def _empty_period_package(message: str) -> dict[str, Any]:
    return {
        "rule_configured": False,
        "message": message,
        "required_table_count": 0,
        "uploaded_count": 0,
        "passed_count": 0,
        "error_count": 0,
        "missing_count": 0,
        "can_calculate": False,
        "tables": [],
    }


def _require_structural_assignment_chain(
    spec: dict[str, Any],
    understanding: dict[str, Any],
    skill_markdown: str,
    *,
    confirmed_additions: list[dict[str, Any]] | None = None,
) -> None:
    """Fail closed when Skill literals themselves form a two-hop ownership chain."""

    if str(spec.get("schema_version") or "") != "2":
        return
    catalog = extract_safe_rule_literals(skill_markdown)
    chain_candidates: list[tuple[dict[str, str], dict[str, str]]] = []
    mappings = list(catalog.mappings.values())
    for first in mappings:
        targets = {
            str(value).split("/", 1)[0].strip().casefold()
            for value in first.values()
            if str(value).strip()
        }
        for second in mappings:
            if first is second:
                continue
            second_keys = {
                str(value).strip().casefold() for value in second if str(value).strip()
            }
            overlap = targets.intersection(second_keys)
            if len(overlap) >= min(3, max(1, len(targets))):
                chain_candidates.append((first, second))
    if not chain_candidates:
        return

    subject = dict(spec.get("subject") or {})
    fields = dict(subject.get("fields") or {})
    lookups = dict(subject.get("lookups") or {})
    grouped_tables = {
        str(metric.get("source_table") or "")
        for metric in spec.get("metrics") or []
        if isinstance(metric, dict)
        and str(metric.get("subject_scope") or "grouped") == "grouped"
    }
    if fields.keys() & grouped_tables:
        raise DataIntakeError(
            "Skill 中存在可静态核验的两段归属映射，不能直接使用底表团队字段；"
            "请按“业务字段→负责人→负责人所属团队”形成固定映射链",
            code="assignment_chain_bypassed",
            status_code=409,
        )

    for table_key in sorted(grouped_tables):
        lookup = lookups.get(table_key)
        steps = list((lookup or {}).get("steps") or [])
        if len(steps) < 2:
            raise DataIntakeError(
                f"源表“{table_key}”尚未形成两段团队归属映射，不能启用规则",
                code="assignment_chain_incomplete",
                status_code=409,
            )
        matched = False
        for first, second in chain_candidates:
            confirmed_first = dict(first)
            confirmed_second = dict(second)
            for addition in confirmed_additions or []:
                if str(addition.get("table_key") or "") != table_key:
                    continue
                try:
                    step_index = int(addition.get("step_index"))
                except (TypeError, ValueError):
                    continue
                source_value = str(addition.get("source_value") or "").strip()
                target_value = str(addition.get("target_value") or "").strip()
                if step_index not in {0, 1} or not source_value or not target_value:
                    continue
                target_mapping = (
                    confirmed_first if step_index == 0 else confirmed_second
                )
                existing_value = target_mapping.get(source_value)
                if existing_value not in {None, target_value}:
                    raise DataIntakeError(
                        "人工确认的归属补充与 Skill 原有固定映射冲突",
                        code="confirmed_assignment_conflict",
                        status_code=409,
                    )
                target_mapping[source_value] = target_value
            if (
                dict(steps[0].get("mapping") or {}) == confirmed_first
                and dict(steps[1].get("mapping") or {}) == confirmed_second
            ):
                matched = True
                break
        if not matched:
            raise DataIntakeError(
                f"源表“{table_key}”的两段归属映射与 Skill 固定映射不一致",
                code="assignment_chain_mismatch",
                status_code=409,
            )

    understanding.setdefault("assignment_chain_enforced", True)


def _prefixed_employee_identifier(value: str) -> str:
    match = re.match(r"^\s*[\(（]\s*([A-Za-z]\d+)\s*[\)）]", value)
    return match.group(1) if match else ""


def _display_name_without_identifier(value: str) -> str:
    return re.sub(
        r"^\s*[\(（]\s*[A-Za-z]\d+\s*[\)）]\s*",
        "",
        value,
    ).strip()


def _unlinked_person_reference(source_system: str, source_value: str) -> str:
    digest = hashlib.sha256(
        f"{source_system.casefold()}|{source_value.strip().casefold()}".encode()
    ).hexdigest()
    return f"external-person:{digest[:40]}"


def _unlinked_team_reference(source_system: str, source_value: str) -> str:
    digest = hashlib.sha256(
        f"{source_system.casefold()}|team|{source_value.strip().casefold()}".encode()
    ).hexdigest()
    return f"external-team:{digest[:40]}"


def _file_error_preview(message: str) -> ImportPreview:
    return ImportPreview(
        (
            PreviewRow(
                1,
                "失败",
                {},
                {},
                None,
                None,
                (ImportErrorDetail("file_validation", message, "file"),),
            ),
        ),
        {
            "新增": 0,
            "更新": 0,
            "无变化": 0,
            "冲突": 0,
            "失败": 1,
            "待核验": 0,
            "重复跳过": 0,
        },
    )


def _row_error_to_import(error: RowError):
    return ImportErrorDetail(error.code, error.message, error.field, error.critical)


def _progress_details(after: dict[str, Any]) -> str:
    values = []
    if after.get("procedure_node"):
        values.append(f"程序节点：{after['procedure_node']}")
    if after.get("next_plan"):
        values.append(f"下一步计划：{after['next_plan']}")
    if after.get("plan_date"):
        values.append(f"计划日期：{after['plan_date']}")
    return "\n".join(values)


def _display_row(value: Any) -> Any:
    if not isinstance(value, dict):
        return value
    hidden = {
        "case_id",
        "progress_id",
        "source_map_id",
        "source_json",
        "fingerprint",
        "idempotency_key",
        "version",
        "external_case_id",
        "manual_case_id",
    }
    return {
        key: _display_row(item)
        for key, item in value.items()
        if key not in hidden and not key.startswith("__")
    }


def _case_snapshot(case: Agent2Case) -> dict[str, Any]:
    return {
        "case_id": str(case.case_id),
        "source_system": str(
            (case.source_json or {}).get("source_system") or case.source_type
        ),
        "source_case_id": case.external_case_id,
        "external_case_id": case.external_case_id,
        "case_number": case.case_number,
        "case_name": case.case_name,
        "case_type": case.case_type,
        "status": case.status,
        "owner_user_id": case.owner_user_id,
        "team_id": case.team_id,
        "company_id": case.company_id,
        "department_id": case.department_id,
        "source_type": case.source_type,
        "source_id": case.source_id,
        "source_json": dict(case.source_json or {}),
        "version": case.version,
    }


def _preview_case_source_hash(
    preview: ImportPreview,
    *,
    source_system: str,
) -> str:
    snapshots: dict[str, dict[str, Any]] = {}
    for row in preview.rows:
        before = row.before or {}
        before_source = str(
            before.get("source_system")
            or before.get("source_type")
            or (before.get("source_json") or {}).get("source_system")
            or ""
        )
        case_id = str(before.get("case_id") or "")
        if not case_id or before_source.casefold() != source_system.casefold():
            continue
        snapshots[case_id] = {
            "case_id": case_id,
            "source_case_id": str(
                before.get("source_case_id") or before.get("external_case_id") or ""
            ),
            "version": int(before.get("version") or 0),
        }
    return _hash_json([snapshots[key] for key in sorted(snapshots)])


def _lineage_summary(
    lineage: dict[str, Any],
    rule_spec: dict[str, Any],
    rule_version: str,
) -> list[dict[str, Any]]:
    table_labels: dict[str, str] = {}
    field_labels: dict[str, str] = {}
    for table in rule_spec.get("source_tables") or []:
        if not isinstance(table, dict):
            continue
        table_key = str(table.get("key") or "")
        table_name = str(table.get("name") or table_key or "源表")
        table_labels[table_key] = table_name
        for column in table.get("columns") or []:
            if not isinstance(column, dict):
                continue
            column_key = str(column.get("key") or "")
            field_labels[f"{table_key}.{column_key}"] = (
                f"{table_name}·{column.get('name') or column_key}"
            )
    output_labels = {
        str(output.get("key") or ""): str(
            output.get("name") or output.get("key") or "计算结果"
        )
        for output in rule_spec.get("outputs") or []
        if isinstance(output, dict)
    }

    def rows(items: Any) -> list[dict[str, Any]]:
        unique: dict[tuple[str, Any], dict[str, Any]] = {}
        for item in items or []:
            if not isinstance(item, dict):
                continue
            table_key = str(item.get("table") or "")
            row_number = item.get("row_number")
            unique[(table_key, row_number)] = {
                "table_name": table_labels.get(
                    table_key,
                    table_key or "源表",
                ),
                "row_number": row_number,
            }
        return list(unique.values())

    def ranges(items: Any) -> list[dict[str, Any]]:
        output: list[dict[str, Any]] = []
        for item in items or []:
            if not isinstance(item, dict):
                continue
            normalized_ranges = []
            for value in item.get("ranges") or []:
                if (
                    isinstance(value, list)
                    and len(value) == 2
                    and all(isinstance(number, int) for number in value)
                ):
                    normalized_ranges.append(value)
            if not normalized_ranges:
                continue
            table_key = str(item.get("table") or "")
            output.append(
                {
                    "table_name": table_labels.get(
                        table_key,
                        table_key or "源表",
                    ),
                    "ranges": normalized_ranges,
                    "row_count": int(item.get("row_count") or 0),
                }
            )
        return output

    summaries: list[dict[str, Any]] = []
    for output_key, detail in lineage.items():
        if output_key == "source_rows" or not isinstance(detail, dict):
            continue
        summaries.append(
            {
                "result_name": output_labels.get(str(output_key), str(output_key)),
                "rule_version": str(detail.get("rule_version") or rule_version),
                "fields": [
                    field_labels.get(str(field), str(field))
                    for field in detail.get("fields") or []
                ],
                "source_rows": rows(detail.get("source_rows")),
                "source_ranges": ranges(detail.get("source_row_ranges")),
                "source_row_count": int(detail.get("source_row_count") or 0),
                "metrics": [
                    {
                        "name": str(
                            item.get("metric_name")
                            or item.get("metric_key")
                            or "基础指标"
                        ),
                        "value": str(item.get("value") or ""),
                        "source_row_count": int(item.get("source_row_count") or 0),
                    }
                    for item in detail.get("metric_details") or []
                    if isinstance(item, dict)
                ],
            }
        )
    if not summaries and lineage.get("source_rows"):
        summaries.append(
            {
                "result_name": "计算异常的来源记录",
                "rule_version": rule_version,
                "fields": [],
                "source_rows": rows(lineage.get("source_rows")),
                "source_ranges": [],
                "source_row_count": len(rows(lineage.get("source_rows"))),
                "metrics": [],
            }
        )
    return summaries


def _lineage_source_text(item: dict[str, Any]) -> str:
    ranges = []
    for source in item.get("source_ranges") or []:
        parts = []
        for start, end in source.get("ranges") or []:
            parts.append(str(start) if start == end else f"{start}-{end}")
        if parts:
            ranges.append(
                f"{source.get('table_name') or '源表'}第"
                f"{'、'.join(parts)}行（共{source.get('row_count') or 0}条）"
            )
    if ranges:
        return "；".join(ranges)
    rows = [
        f"{row.get('table_name') or '源表'}第{row.get('row_number') or '未记录'}行"
        for row in item.get("source_rows") or []
    ]
    return "、".join(rows) or "规则常量"


def _matching_metric_target(
    targets: list[dict[str, Any]],
    *,
    key: str,
    name: str,
    subject_key: str = "",
    subject_name: str = "",
) -> dict[str, Any] | None:
    metric_candidates = {
        str(key).strip().casefold(),
        str(name).strip().casefold(),
    }
    subject_candidates = {
        str(subject_key).strip().casefold(),
        str(subject_name).strip().casefold(),
    }
    generic: dict[str, Any] | None = None
    for target in targets:
        metric = str(
            target.get("metric_name")
            or target.get("metric_key")
            or target.get("name")
            or ""
        ).strip()
        if metric.casefold() not in metric_candidates:
            continue
        scope = str(
            target.get("scope_name")
            or target.get("scope_key")
            or target.get("team_name")
            or ""
        ).strip()
        if scope and scope.casefold() in subject_candidates:
            return target
        if not scope and generic is None:
            generic = target
    return generic


def _metric_achievement(
    actual: Decimal,
    target: dict[str, Any] | None,
) -> dict[str, Any]:
    if not target:
        return {
            "status": "unknown",
            "status_label": "未配置目标",
            "target_value": "",
            "target_unit": "",
            "comparison": "",
            "comparison_label": "",
            "difference": "",
        }
    raw_target = str(target.get("target_value") or target.get("value") or "").strip()
    unit = str(target.get("unit") or "").strip()
    comparison = str(target.get("comparison") or "").strip()
    comparison_label = {
        "at_least": "不低于",
        "at_most": "不高于",
        "equal": "等于",
    }.get(comparison, "")
    numeric = raw_target.replace(",", "")
    if numeric.endswith("%"):
        numeric = numeric[:-1].strip()
        unit = unit or "%"
    try:
        target_value = Decimal(numeric)
    except (InvalidOperation, ValueError):
        target_value = None
    common = {
        "target_value": raw_target,
        "target_unit": unit,
        "comparison": comparison,
        "comparison_label": comparison_label,
        "target_scope": str(target.get("scope_name") or ""),
        "difference": (str(actual - target_value) if target_value is not None else ""),
    }
    if target_value is None:
        return {
            **common,
            "status": "unknown",
            "status_label": "目标值还不能比较",
        }
    if comparison == "at_least":
        achieved = actual >= target_value
    elif comparison == "at_most":
        achieved = actual <= target_value
    elif comparison == "equal":
        achieved = actual == target_value
    else:
        return {
            **common,
            "status": "unknown",
            "status_label": "请确认目标判断方向",
        }
    return {
        **common,
        "status": "achieved" if achieved else "not_achieved",
        "status_label": "已达到目标" if achieved else "尚未达到目标",
    }


def _metric_presentation(
    *,
    key: str,
    name: str,
    value: Decimal,
    unit: str,
) -> dict[str, str]:
    """Present signed change values in the business wording used by the Skill.

    The deterministic result remains signed: a negative value still participates
    in target comparisons as a decrease. Only the human-facing text converts
    ``-31.72%`` into ``下降31.72%``.
    """

    raw_value = str(value)
    normalized_unit = str(unit or "").strip()
    normalized_name = str(name or "").strip()
    is_decline_rate = normalized_unit == "%" and normalized_name.endswith("下降率")
    if not is_decline_rate:
        return {
            "display_value": raw_value,
            "display_prefix": "",
            "display_text": f"{raw_value}{normalized_unit}",
            "trend": "neutral",
        }
    if value < 0:
        display_value = str(abs(value))
        prefix = "下降"
        trend = "down"
    elif value > 0:
        display_value = raw_value
        prefix = "上升"
        trend = "up"
    else:
        display_value = raw_value
        prefix = "持平"
        trend = "flat"
    return {
        "display_value": display_value,
        "display_prefix": prefix,
        "display_text": f"{prefix}{display_value}{normalized_unit}",
        "trend": trend,
    }


def _period_type_label(value: str) -> str:
    return {
        "month": "月度",
        "quarter": "季度",
        "half_year": "半年度",
        "year": "年度",
        "custom": "自定义",
    }.get(value, value)


def _period_status_label(value: str) -> str:
    return {
        "preparing": "准备中",
        "ready": "可以试算",
        "calculated": "已试算",
        "published": "正式结果已发布",
        "closed": "已关闭",
    }.get(value, value)


def _rule_reference_documents(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, dict):
        documents = value.get("reference_documents") or []
        if isinstance(documents, list):
            output: list[dict[str, Any]] = []
            for item in documents:
                if not isinstance(item, dict):
                    continue
                path = str(item.get("path") or item.get("file_name") or "")
                output.append(
                    {
                        "file_name": str(
                            item.get("file_name")
                            or PurePosixPath(path).name
                            or "参考文档"
                        ),
                        "file_hash": str(item.get("file_hash") or ""),
                        "paragraph_count": int(item.get("paragraph_count") or 0),
                    }
                )
            return output
        return []
    if isinstance(value, list):
        return [
            {
                "file_name": PurePosixPath(str(path)).name,
                "file_hash": "",
                "paragraph_count": 0,
            }
            for path in value
            if str(path).lower().endswith(".docx")
        ]
    return []


def _rule_status_label(value: str) -> str:
    return {
        "uploaded": "已上传，待理解",
        "understanding_draft": "理解草稿仍有待确认项",
        "awaiting_confirmation": "等待发布管理员确认",
        "confirmed": "已确认",
        "active": "当前有效",
        "rejected": "已拒绝",
        "superseded": "已被新版本替换",
    }.get(value, value)


def _calculation_status_label(value: str) -> str:
    return {
        "trial": "试算结果，尚未正式发布",
        "published": "正式结果已发布",
        "invalidated": "数据已变化，需重新计算",
        "abandoned": "试算已放弃",
        "failed": "计算失败",
    }.get(value, value)


def _validation_status_label(value: str) -> str:
    return {
        "valid": "校验通过",
        "warning": "有提醒",
        "error": "有错误",
        "skipped": "无需写入",
    }.get(value, value)


def _column_type_label(value: str) -> str:
    return {
        "text": "文字",
        "date": "日期（如2026-07-01）",
        "decimal": "数字或金额",
        "percentage": "比例（如15%）",
        "integer": "整数",
    }.get(value, value)


def _rule_spec_review_items(spec: dict[str, Any]) -> list[str]:
    field_labels = {
        f"{table.get('key')}.{column.get('key')}": (
            f"{table.get('name') or table.get('key') or '源表'}·"
            f"{column.get('name') or column.get('key') or '字段'}"
        )
        for table in spec.get("source_tables") or []
        if isinstance(table, dict)
        for column in table.get("columns") or []
        if isinstance(column, dict)
    }
    metric_labels = {
        str(metric.get("key") or ""): str(
            metric.get("name") or metric.get("key") or "基础指标"
        )
        for metric in [
            *(spec.get("metrics") or []),
            *(spec.get("outputs") or []),
        ]
        if isinstance(metric, dict)
    }
    items: list[str] = []
    subject = spec.get("subject") or {}
    for lookup in (subject.get("lookups") or {}).values():
        if not isinstance(lookup, dict):
            continue
        for step in lookup.get("steps") or []:
            if not isinstance(step, dict):
                continue
            items.append(
                f"固定映射：{step.get('name') or '未命名映射'}"
                f"（{len(step.get('mapping') or {})}条；未匹配进入待处理）"
            )
    for metric in spec.get("metrics") or []:
        if not isinstance(metric, dict):
            continue
        items.append(
            f"基础指标 {metric.get('name') or metric.get('key') or '未命名'}："
            f"{_describe_aggregate_metric(metric, field_labels)}"
        )
    for output in spec.get("outputs") or []:
        if not isinstance(output, dict):
            continue
        items.append(
            f"结果 {output.get('name') or output.get('key') or '未命名'}："
            f"{_describe_expression(output.get('expression') or {}, field_labels, metric_labels)}；"
            f"{_rounding_label(str(output.get('rounding', 'half_up')))}，"
            f"保留{int(output.get('decimal_places', 2))}位；"
            f"{_divide_by_zero_label(str(output.get('on_divide_by_zero', 'error')))}"
        )
    return items or ["已形成并通过固定程序校验"]


def _rule_item_review(
    spec: dict[str, Any],
    item_type: str,
    item: dict[str, Any] | None,
) -> str:
    if item is None:
        return "尚未配置"
    field_labels = {
        f"{table.get('key')}.{column.get('key')}": (
            f"{table.get('name') or table.get('key') or '源表'}·"
            f"{column.get('name') or column.get('key') or '字段'}"
        )
        for table in spec.get("source_tables") or []
        if isinstance(table, dict)
        for column in table.get("columns") or []
        if isinstance(column, dict)
    }
    if item_type == "基础指标":
        return _describe_aggregate_metric(item, field_labels)
    metric_labels = {
        str(metric.get("key") or ""): str(
            metric.get("name") or metric.get("key") or "基础指标"
        )
        for metric in [
            *(spec.get("metrics") or []),
            *(spec.get("outputs") or []),
        ]
        if isinstance(metric, dict)
    }
    return (
        _describe_expression(
            item.get("expression") or {},
            field_labels,
            metric_labels,
        )
        + f"；{_rounding_label(str(item.get('rounding', 'half_up')))}，"
        f"保留{int(item.get('decimal_places', 2))}位；"
        f"{_divide_by_zero_label(str(item.get('on_divide_by_zero', 'error')))}"
    )


def _rule_lookup_details(spec: dict[str, Any]) -> list[dict[str, Any]]:
    table_labels = {
        str(table.get("key") or ""): str(
            table.get("name") or table.get("key") or "源表"
        )
        for table in spec.get("source_tables") or []
        if isinstance(table, dict)
    }
    details: list[dict[str, Any]] = []
    subject = spec.get("subject") or {}
    for table_key, lookup in (subject.get("lookups") or {}).items():
        if not isinstance(lookup, dict):
            continue
        for step in lookup.get("steps") or []:
            if not isinstance(step, dict):
                continue
            mapping = step.get("mapping") or {}
            details.append(
                {
                    "name": str(step.get("name") or "精确映射"),
                    "source_table": table_labels.get(
                        str(table_key),
                        str(table_key) or "源表",
                    ),
                    "entry_count": len(mapping),
                    "entries": [
                        {
                            "source": str(source),
                            "target": str(target),
                        }
                        for source, target in sorted(
                            mapping.items(),
                            key=lambda item: str(item[0]),
                        )
                    ],
                    "split_first": str(step.get("split_first") or ""),
                    "normalization_labels": [
                        {
                            "remove_whitespace": "忽略空格",
                            "normalize_parentheses": "统一中英文括号",
                            "casefold": "忽略英文字母大小写",
                        }.get(str(item), str(item))
                        for item in step.get("normalizers") or []
                    ],
                    "matching_labels": [
                        _lookup_match_strategy_label(str(item))
                        for item in step.get("match_strategies") or []
                    ],
                }
            )
        exclusions = [str(item) for item in lookup.get("exclude_source_values") or []]
        if exclusions:
            details.append(
                {
                    "name": "明确排除清单",
                    "source_table": table_labels.get(
                        str(table_key),
                        str(table_key) or "源表",
                    ),
                    "entry_count": len(exclusions),
                    "entries": [
                        {"source": item, "target": "不参与本规则统计"}
                        for item in sorted(exclusions)
                    ],
                    "split_first": "",
                    "normalization_labels": [],
                    "matching_labels": [],
                }
            )
    return details


def _rule_assignment_chains(spec: dict[str, Any]) -> list[dict[str, Any]]:
    table_definitions = {
        str(table.get("key") or ""): table
        for table in spec.get("source_tables") or []
        if isinstance(table, dict)
    }
    subject = spec.get("subject") or {}
    chains: list[dict[str, Any]] = []
    for table_key, lookup in (subject.get("lookups") or {}).items():
        if not isinstance(lookup, dict):
            continue
        table = table_definitions.get(str(table_key), {})
        table_name = str(table.get("name") or table.get("key") or table_key or "源表")
        source_field = str(lookup.get("source_field") or "")
        source_column = next(
            (
                column
                for column in table.get("columns") or []
                if str(column.get("key") or "") == source_field
            ),
            {},
        )
        entry_name = str(source_column.get("name") or source_field or "归属入口字段")
        steps = [
            {
                "name": str(step.get("name") or "精确映射"),
                "entry_count": len(step.get("mapping") or {}),
                "matching_labels": [
                    _lookup_match_strategy_label(str(item))
                    for item in step.get("match_strategies") or []
                ],
            }
            for step in lookup.get("steps") or []
            if isinstance(step, dict)
        ]
        path = [f"{table_name}·{entry_name}"]
        path.extend(
            f"{step['name']}（{step['entry_count']}条固定关系）" for step in steps
        )
        path.append(str(subject.get("name") or "核算对象"))
        chains.append(
            {
                "source_table": table_name,
                "entry_field": entry_name,
                "mode": "fixed_lookup_chain",
                "mode_label": "按固定映射链归属",
                "path": path,
                "steps": steps,
                "unmatched_policy": "无法唯一匹配的记录进入待处理，不计入任何团队",
                "direct_field_fallback": False,
            }
        )
    for table_key, column_key in (subject.get("fields") or {}).items():
        table = table_definitions.get(str(table_key), {})
        table_name = str(table.get("name") or table.get("key") or table_key or "源表")
        column = next(
            (
                item
                for item in table.get("columns") or []
                if str(item.get("key") or "") == str(column_key)
            ),
            {},
        )
        entry_name = str(column.get("name") or column_key or "分组字段")
        chains.append(
            {
                "source_table": table_name,
                "entry_field": entry_name,
                "mode": "direct_field",
                "mode_label": "直接按底表字段归属",
                "path": [
                    f"{table_name}·{entry_name}",
                    str(subject.get("name") or "核算对象"),
                ],
                "steps": [],
                "unmatched_policy": "空值记录不进入核算对象",
                "direct_field_fallback": True,
            }
        )
    return chains


def _lookup_match_strategy_label(value: str) -> str:
    return {
        "exact": "精确匹配",
        "expand_branch_short_form": "分公司简称转标准名称",
        "contains_unique": "包含关系仅限唯一结果",
        "strip_aftercare_suffix": "去除“善后”后缀后匹配",
        "strip_parenthetical": "去除括号说明后匹配",
    }.get(value, value)


def _describe_expression(
    expression: dict[str, Any],
    field_labels: dict[str, str],
    metric_labels: dict[str, str] | None = None,
) -> str:
    if "value" in expression:
        return str(expression["value"])
    if "field" in expression:
        field = str(expression["field"])
        return field_labels.get(field, field)
    if "metric" in expression:
        metric = str(expression["metric"])
        return (metric_labels or {}).get(metric, metric)
    if "result" in expression:
        result = str(expression["result"])
        return (metric_labels or {}).get(result, result)
    operator = {
        "add": " + ",
        "subtract": " − ",
        "multiply": " × ",
        "divide": " ÷ ",
        "min": "取较小值",
        "max": "取较大值",
        "abs": "取绝对值",
    }.get(str(expression.get("op") or ""), "未知运算")
    arguments = [
        _describe_expression(item, field_labels, metric_labels)
        for item in expression.get("args") or []
        if isinstance(item, dict)
    ]
    if operator == "取绝对值":
        return f"取绝对值（{arguments[0] if arguments else '待确认'}）"
    if operator in {"取较小值", "取较大值"}:
        return f"{operator}（{'，'.join(arguments)}）"
    return f"（{operator.join(arguments)}）"


def _describe_aggregate_metric(
    metric: dict[str, Any],
    field_labels: dict[str, str],
) -> str:
    aggregate = str(metric.get("aggregate") or "")
    where = metric.get("where")
    condition = (
        _describe_condition(where, field_labels)
        if isinstance(where, dict)
        else "全部有效记录"
    )
    if aggregate == "count":
        return f"满足“{condition}”的记录数量"
    value = metric.get("value")
    value_label = (
        _describe_expression(value, field_labels)
        if isinstance(value, dict)
        else "待确认字段"
    )
    null_policy = (
        "；参与加总的空金额按0处理" if metric.get("null_as_zero") is True else ""
    )
    return f"对满足“{condition}”的记录汇总“{value_label}”{null_policy}"


def _describe_condition(
    condition: dict[str, Any],
    field_labels: dict[str, str],
) -> str:
    if "all" in condition:
        return "，并且".join(
            _describe_condition(item, field_labels)
            for item in condition.get("all") or []
            if isinstance(item, dict)
        )
    if "any" in condition:
        return "，或者".join(
            _describe_condition(item, field_labels)
            for item in condition.get("any") or []
            if isinstance(item, dict)
        )
    operator = {
        "eq": "等于",
        "neq": "不等于",
        "lt": "早于或小于",
        "lte": "不晚于或不大于",
        "gt": "晚于或大于",
        "gte": "不早于或不小于",
        "in": "属于",
        "not_in": "不属于",
        "is_empty": "为空",
        "not_empty": "不为空",
    }.get(str(condition.get("op") or ""), "采用待确认比较")
    left = _describe_operand(condition.get("left"), field_labels)
    if str(condition.get("op") or "") in {"is_empty", "not_empty"}:
        return f"{left}{operator}"
    right = _describe_operand(condition.get("right"), field_labels)
    return f"{left}{operator}{right}"


def _describe_operand(
    operand: Any,
    field_labels: dict[str, str],
) -> str:
    if not isinstance(operand, dict):
        return "待确认值"
    if "field" in operand:
        field = str(operand["field"])
        return field_labels.get(field, field)
    if "value" in operand:
        return str(operand["value"])
    if "values" in operand:
        return "、".join(str(value) for value in operand["values"])
    boundary = {
        "period_start": "考核周期开始日",
        "period_end": "考核周期结束日",
        "month_start": "考核月第一天",
        "month_end": "考核月最后一天",
        "year_start": "考核年度第一天",
        "year_end": "考核年度最后一天",
    }.get(str(operand.get("boundary") or ""), "待确认日期")
    shift = operand.get("shift") or {}
    parts = []
    for key, label in (("years", "年"), ("months", "个月"), ("days", "天")):
        value = int(shift.get(key, 0) or 0)
        if value:
            parts.append(f"{value:+d}{label}")
    return boundary + (f"（偏移{'、'.join(parts)}）" if parts else "")


def _rounding_label(value: str) -> str:
    return {
        "half_up": "四舍五入",
        "half_even": "银行家舍入（五后取偶）",
    }.get(value, value)


def _divide_by_zero_label(value: str) -> str:
    return {
        "zero": "除数为0时按0计算（来自本版 Skill）",
        "error": "除数为0时标记为不可比较",
    }.get(value, value)


def _rule_explanation_items(values: Any) -> list[dict[str, str]]:
    output: list[dict[str, str]] = []
    if not isinstance(values, (list, tuple)):
        return output
    for index, item in enumerate(values, start=1):
        if not isinstance(item, dict):
            continue
        label = next(
            (
                str(item.get(key) or "").strip()
                for key in (
                    "result_name",
                    "result",
                    "metric",
                    "name",
                    "topic",
                )
                if str(item.get(key) or "").strip()
            ),
            f"计算口径{index}",
        )
        explanation = next(
            (
                str(item.get(key) or "").strip()
                for key in (
                    "explanation",
                    "formula",
                    "definition",
                    "logic",
                    "description",
                    "condition",
                )
                if str(item.get(key) or "").strip()
            ),
            "",
        )
        if not explanation:
            readable: list[str] = []
            for key, value in item.items():
                if key in {
                    "result_name",
                    "result",
                    "metric",
                    "name",
                    "topic",
                    "evidence",
                    "line_start",
                    "line_end",
                }:
                    continue
                if isinstance(value, (str, int, float, Decimal, bool)):
                    text_value = str(value).strip()
                    if text_value:
                        readable.append(text_value)
                elif isinstance(value, list):
                    text_value = "；".join(
                        str(part).strip()
                        for part in value
                        if isinstance(part, (str, int, float)) and str(part).strip()
                    )
                    if text_value:
                        readable.append(text_value)
            explanation = "；".join(readable)
        if explanation:
            output.append({"name": label, "explanation": explanation})
    return output
