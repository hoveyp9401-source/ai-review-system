from __future__ import annotations

import copy
import io
import re
from calendar import monthrange
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from typing import Any, Literal

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.legal_ops_data_intake.calculator import resolve_lookup_path
from app.legal_ops_data_intake.workbook import safe_excel_cell


class PerformanceReportError(ValueError):
    """The selected rule or source data cannot produce an auditable report."""


@dataclass(frozen=True)
class DefendantPerformanceReport:
    _payload: dict[str, Any]

    def as_dict(self) -> dict[str, Any]:
        return copy.deepcopy(self._payload)


@dataclass(frozen=True)
class _FieldRef:
    key: str
    source_name: str


@dataclass(frozen=True)
class _CaseRow:
    team_name: str
    lawyer_name: str
    branch_name: str
    case_name: str
    register_date: date
    is_closed: bool
    close_date: date | None
    source_row_number: int | None


@dataclass(frozen=True)
class _LossRow:
    is_closed: bool
    close_date: date | None
    opponent_type: str
    loss_reduction: Decimal | None
    case_reason: str
    subject_amount: Decimal | None
    interest: Decimal | None
    total_principal: Decimal | None
    total_appraisal: Decimal | None
    total_litigation: Decimal | None
    total_penalty: Decimal | None


@dataclass(frozen=True)
class _ReportPeriod:
    view: Literal["week", "month"]
    starts_on: date
    ends_on: date
    previous_cutoff: date
    last_year_cutoff: date
    label: str
    comparison_label: str

    def as_dict(self) -> dict[str, str]:
        return {
            "view": self.view,
            "starts_on": self.starts_on.isoformat(),
            "ends_on": self.ends_on.isoformat(),
            "cutoff_date": self.ends_on.isoformat(),
            "label": self.label,
            "comparison_label": self.comparison_label,
        }


def build_defendant_performance_report(
    rule_spec: dict[str, Any],
    sources: dict[str, list[dict[str, Any]]],
    *,
    view: Literal["week", "month"],
    anchor_date: date,
    targets: list[dict[str, Any]] | tuple[dict[str, Any], ...] = (),
) -> DefendantPerformanceReport:
    """Build one read-only weekly or monthly defendant performance report."""

    period = _report_period(view, anchor_date)
    table_key, fields = _case_table_definition(rule_spec)
    rows = list(sources.get(table_key) or ())
    lookup = _subject_lookup(rule_spec, table_key)
    excluded_teams = {
        str(value).strip()
        for value in (rule_spec.get("subject") or {}).get("exclude_values") or ()
        if str(value).strip()
    }
    cases: list[_CaseRow] = []
    assignment_errors: list[dict[str, Any]] = []
    data_quality_errors: list[dict[str, Any]] = []
    for row in rows:
        is_closed = _is_closed(_field_value(row, fields["is_closed"]))
        close_date = _optional_date(_field_value(row, fields["close_date"]))
        if is_closed and close_date is None:
            data_quality_errors.append(
                {
                    "source_row_number": _row_number(row),
                    "error_type": "closed_without_close_date",
                    "message": "已标记结案但结案日期为空，不计入存量",
                    "case_name": _case_name(row),
                }
            )
        resolution = resolve_lookup_path(lookup, row)
        if resolution.error:
            assignment_errors.append(
                {
                    "source_row_number": _row_number(row),
                    "step_name": str(resolution.error.get("step_name") or ""),
                    "source_value": str(resolution.error.get("value") or ""),
                    "message": _assignment_error_message(resolution.error),
                }
            )
            continue
        team_name = str(resolution.final_value or "").strip()
        if not team_name or team_name in excluded_teams:
            continue
        register_date = _optional_date(_field_value(row, fields["register_date"]))
        if register_date is None:
            continue
        lawyer_name = (
            _primary_lawyer_name(resolution.steps[0].mapped_value)
            if resolution.steps
            else ""
        )
        cases.append(
            _CaseRow(
                team_name=team_name,
                lawyer_name=lawyer_name,
                branch_name=str(
                    _field_value(row, fields["branch_name"]) or ""
                ).strip(),
                case_name=_case_name(row),
                register_date=register_date,
                is_closed=is_closed,
                close_date=close_date,
                source_row_number=_row_number(row),
            )
        )

    loss_metrics = (
        _loss_metrics(rows, fields, period)
        if period.view == "month"
        else _weekly_loss_not_applicable()
    )
    # A missing close date is still kept in the maintenance trace, but the
    # Skill already defines that row as excluded from stock. It must not hide
    # an otherwise deterministic target conclusion from business readers.
    data_incomplete = bool(assignment_errors)
    teams = sorted({item.team_name for item in cases})
    total_key = str((rule_spec.get("subject") or {}).get("total_key") or "__total__")
    total_label = str((rule_spec.get("subject") or {}).get("total_label") or "整体")
    scopes = [
        _scope_report(
            total_key,
            total_label,
            cases,
            period,
            targets,
            scope_type="overall",
            data_incomplete=data_incomplete,
            loss_metrics=loss_metrics,
        )
    ]
    scopes.extend(
        _scope_report(
            team,
            team,
            [item for item in cases if item.team_name == team],
            period,
            targets,
            scope_type="team",
            data_incomplete=data_incomplete,
        )
        for team in teams
    )
    lawyers = sorted({item.lawyer_name for item in cases if item.lawyer_name})
    personal_scopes = [
        _scope_report(
            f"person:{lawyer}",
            lawyer,
            [item for item in cases if item.lawyer_name == lawyer],
            period,
            targets,
            scope_type="person",
            data_incomplete=data_incomplete,
        )
        for lawyer in lawyers
    ]
    return DefendantPerformanceReport(
        {
            "period": period.as_dict(),
            "source_table_key": table_key,
            "source_row_count": len(rows),
            "mapped_row_count": len(cases),
            "assignment_error_count": len(assignment_errors),
            "data_quality_error_count": len(data_quality_errors),
            "target_conclusions_available": not data_incomplete,
            "loss_metrics": loss_metrics,
            "scopes": scopes,
            "personal_scopes": personal_scopes,
            "assignment_errors": assignment_errors,
            "data_quality_errors": data_quality_errors,
        }
    )


def export_defendant_performance_xlsx(
    report: DefendantPerformanceReport,
    *,
    scope_key: str,
) -> bytes:
    """Export a readable overview plus the auditable business detail sheets."""

    payload = report.as_dict()
    scope = _selected_scope(payload, scope_key)
    period = dict(payload["period"])
    weekly = str(period.get("view") or "") == "week"
    period_new_label = "本周新增" if weekly else "本月新增"
    period_closed_label = "本周结案" if weekly else "本月结案"
    workbook = Workbook()
    overview = workbook.active
    overview.title = "指标概览"
    stock = workbook.create_sheet("存量")
    additions = workbook.create_sheet("新增")
    combined = workbook.create_sheet("综合汇总")
    export_sheets = [stock, additions, combined]

    stock.append(
        ["分公司", "承办法务", "存量（件）", "去年存量", "存量同比", "存量环比"]
    )
    additions.append(
        [
            "分公司",
            "承办法务",
            "YTD新增（件）",
            "去年YTD新增",
            "新增同比",
            period_new_label,
        ]
    )
    combined.append(
        [
            "分公司",
            "承办法务",
            "存量",
            "存量同比",
            "存量环比",
            "YTD新增",
            "新增同比",
            period_new_label,
            period_closed_label,
        ]
    )
    for branch in scope["branches"]:
        stock.append(
            [
                safe_excel_cell(branch["branch_name"]),
                safe_excel_cell(branch["lawyer_name"]),
                int(branch["stock_count"]),
                int(branch["last_year_stock_count"]),
                _arrow_rate(branch["stock_yoy"]),
                _arrow_rate(branch["stock_period_change"]),
            ]
        )
        additions.append(
            [
                safe_excel_cell(branch["branch_name"]),
                safe_excel_cell(branch["lawyer_name"]),
                int(branch["year_to_date_new_count"]),
                int(branch["last_year_to_date_new_count"]),
                _arrow_rate(branch["new_yoy"]),
                int(branch["period_new_count"]),
            ]
        )
        combined.append(
            [
                safe_excel_cell(branch["branch_name"]),
                safe_excel_cell(branch["lawyer_name"]),
                int(branch["stock_count"]),
                _arrow_rate(branch["stock_yoy"]),
                _arrow_rate(branch["stock_period_change"]),
                int(branch["year_to_date_new_count"]),
                _arrow_rate(branch["new_yoy"]),
                int(branch["period_new_count"]),
                int(branch["period_closed_count"]),
            ]
        )
    stock.append(
        [
            "合计",
            None,
            int(scope["stock_count"]),
            int(scope["last_year_stock_count"]),
            _arrow_rate(scope["stock_yoy"]),
            _arrow_rate(scope["stock_period_change"]),
        ]
    )
    additions.append(
        [
            "合计",
            None,
            int(scope["year_to_date_new_count"]),
            int(scope["last_year_to_date_new_count"]),
            _arrow_rate(scope["new_yoy"]),
            int(scope["period_new_count"]),
        ]
    )
    combined.append(
        [
            "合计",
            None,
            int(scope["stock_count"]),
            _arrow_rate(scope["stock_yoy"]),
            _arrow_rate(scope["stock_period_change"]),
            int(scope["year_to_date_new_count"]),
            _arrow_rate(scope["new_yoy"]),
            int(scope["period_new_count"]),
            int(scope["period_closed_count"]),
        ]
    )
    if not weekly and scope.get("scope_type") == "overall":
        loss_metrics = (
            scope.get("loss_metrics")
            if isinstance(scope.get("loss_metrics"), dict)
            else {}
        )
        comprehensive = (
            loss_metrics.get("comprehensive_loss_rate")
            if isinstance(loss_metrics.get("comprehensive_loss_rate"), dict)
            else {}
        )
        substantial = (
            loss_metrics.get("substantial_loss_amount")
            if isinstance(loss_metrics.get("substantial_loss_amount"), dict)
            else {}
        )
        loss = workbook.create_sheet("减损")
        loss.append(
            [
                "指标",
                "结果",
                "纳入案件数",
                "标的额及利息合计（元）",
                "应付合计（元）",
            ]
        )
        loss.append(
            [
                "综合减损率",
                comprehensive.get("display")
                or comprehensive.get("status_label")
                or "暂不可计算",
                comprehensive.get("eligible_case_count"),
                _excel_amount(comprehensive.get("claim_amount")),
                _excel_amount(comprehensive.get("payable_amount")),
            ]
        )
        loss.append(
            [
                "实质减损金额",
                substantial.get("display")
                or substantial.get("status_label")
                or "暂不可计算",
                substantial.get("eligible_case_count"),
                None,
                None,
            ]
        )
        export_sheets.append(loss)

    new_case_details = workbook.create_sheet("新增案件明细")
    new_case_details.append(
        ["新增日期", "分公司", "对接法务", "案件名称", "原表行号"]
    )
    for item in scope.get("period_new_cases") or ():
        new_case_details.append(
            [
                _excel_date(item.get("register_date")),
                safe_excel_cell(item.get("branch_name")),
                safe_excel_cell(item.get("lawyer_name")),
                safe_excel_cell(item.get("case_name")),
                item.get("source_row_number"),
            ]
        )
    export_sheets.append(new_case_details)

    closed_case_details = workbook.create_sheet("结案案件明细")
    closed_case_details.append(
        ["结案日期", "分公司", "对接法务", "案件名称", "原表行号"]
    )
    for item in scope.get("period_closed_cases") or ():
        closed_case_details.append(
            [
                _excel_date(item.get("close_date")),
                safe_excel_cell(item.get("branch_name")),
                safe_excel_cell(item.get("lawyer_name")),
                safe_excel_cell(item.get("case_name")),
                item.get("source_row_number"),
            ]
        )
    export_sheets.append(closed_case_details)

    _populate_excel_overview(overview, scope=scope, period=period)
    _style_excel_overview(overview)
    for sheet in export_sheets:
        _style_export_sheet(sheet)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def export_defendant_performance_docx(
    report: DefendantPerformanceReport,
    *,
    scope_key: str,
) -> bytes:
    """Export the readable Word report used for weekly/monthly briefings."""

    payload = report.as_dict()
    scope = _selected_scope(payload, scope_key)
    period = dict(payload["period"])
    weekly = str(period.get("view") or "") == "week"
    period_name = "周" if weekly else "月"
    current_label = "本周" if weekly else "本月"
    scope_label = (
        "法务部门整体" if scope["scope_name"] == "整体" else str(scope["scope_name"])
    )
    starts_on = date.fromisoformat(str(period["starts_on"]))
    ends_on = date.fromisoformat(str(period["ends_on"]))
    document = Document()
    _configure_document_styles(document)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(f"{scope_label}被告案件{period_name}报")
    title_run.bold = True
    title_run.font.size = Pt(18)

    period_line = document.add_paragraph(
        "统计周期："
        f"{starts_on.year}年{starts_on.month}月{starts_on.day}日"
        f" — {ends_on.month}月{ends_on.day}日"
        f"（截止{ends_on:%m月%d日}）"
    )
    period_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    attribution = document.add_paragraph(
        "归属口径：底表分公司→对接人→法务团队映射链"
    )
    attribution.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in (*period_line.runs, *attribution.runs):
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(89, 89, 89)

    conclusion = document.add_paragraph()
    conclusion.paragraph_format.space_before = Pt(8)
    conclusion.paragraph_format.space_after = Pt(8)
    conclusion_run = conclusion.add_run("核心结论：")
    conclusion_run.bold = True
    conclusion.add_run(
        _target_summary_text(
            "存量同比",
            scope.get("stock_yoy"),
            scope.get("stock_target"),
        )
    )
    conclusion.add_run("；")
    conclusion.add_run(
        _target_summary_text(
            "年度累计新增同比",
            scope.get("new_yoy"),
            scope.get("new_target"),
        )
    )
    conclusion.add_run("。")

    _add_section_heading(document, "一、存量")
    document.add_paragraph(
        f"截至{ends_on:%m月%d日}，{scope_label}被告案件存量 "
        f"{scope['stock_count']} 件，同比{_sentence_rate(scope['stock_yoy'])}，"
        f"环比{_sentence_rate(scope['stock_period_change'])}。"
    )
    label = document.add_paragraph()
    label.add_run("存量明细（按分公司）：").bold = True
    stock_rows = [
        [
            branch["branch_name"],
            branch["lawyer_name"],
            branch["stock_count"],
            branch["last_year_stock_count"],
            _arrow_rate(branch["stock_yoy"]),
            _arrow_rate(branch["stock_period_change"]),
        ]
        for branch in scope["branches"]
    ]
    stock_rows.append(
        [
            "合计",
            "",
            scope["stock_count"],
            scope["last_year_stock_count"],
            _arrow_rate(scope["stock_yoy"]),
            _arrow_rate(scope["stock_period_change"]),
        ]
    )
    _add_business_table(
        document,
        ["分公司", "承办法务", "存量（件）", "去年同期", "同比", "环比"],
        stock_rows,
    )

    _add_section_heading(document, "二、新增")
    document.add_paragraph(
        f"本年度累计（1月1日-{ends_on:%m月%d日}）新增 "
        f"{scope['year_to_date_new_count']} 件，"
        f"同比{_sentence_rate(scope['new_yoy'])}。"
    )
    document.add_paragraph(
        f"{current_label}（{starts_on.month}/{starts_on.day}-"
        f"{ends_on.month}/{ends_on.day}）新增 "
        f"{scope['period_new_count']} 件。"
    )
    label = document.add_paragraph()
    label.add_run("新增明细（按分公司）：").bold = True
    new_rows = [
        [
            branch["branch_name"],
            branch["lawyer_name"],
            branch["year_to_date_new_count"],
            branch["last_year_to_date_new_count"],
            _arrow_rate(branch["new_yoy"]),
            branch["period_new_count"],
        ]
        for branch in scope["branches"]
    ]
    new_rows.append(
        [
            "合计",
            "",
            scope["year_to_date_new_count"],
            scope["last_year_to_date_new_count"],
            _arrow_rate(scope["new_yoy"]),
            scope["period_new_count"],
        ]
    )
    _add_business_table(
        document,
        [
            "分公司",
            "承办法务",
            "YTD新增（件）",
            "去年YTD",
            "同比",
            f"{current_label}新增",
        ],
        new_rows,
    )

    next_section = 3
    loss_metrics = (
        scope.get("loss_metrics")
        if isinstance(scope.get("loss_metrics"), dict)
        else {}
    )
    comprehensive = (
        loss_metrics.get("comprehensive_loss_rate")
        if isinstance(loss_metrics.get("comprehensive_loss_rate"), dict)
        else {}
    )
    substantial = (
        loss_metrics.get("substantial_loss_amount")
        if isinstance(loss_metrics.get("substantial_loss_amount"), dict)
        else {}
    )
    if (
        scope.get("scope_type") == "overall"
        and comprehensive.get("status") == "calculated"
        and substantial.get("status") == "calculated"
    ):
        _add_section_heading(document, "三、减损情况（部门整体口径）")
        document.add_paragraph(
            f"{current_label}被告案件综合减损率 "
            f"{comprehensive.get('display') or '暂不可计算'}；"
            "实质减损金额（本年度累计）"
            f"{substantial.get('display') or '暂不可计算'}。"
        )
        next_section = 4

    period_dates = f"{starts_on.month}/{starts_on.day}-{ends_on.month}/{ends_on.day}"
    _add_section_heading(
        document,
        f"{_chinese_section_number(next_section)}、"
        f"{current_label}结案明细（{period_dates}）",
    )
    closed_count = document.add_paragraph(
        f"{current_label}结案 {scope['period_closed_count']} 件。"
    )
    closed_count.paragraph_format.keep_with_next = bool(
        scope["period_closed_cases"]
    )
    if scope["period_closed_cases"]:
        _add_business_table(
            document,
            ["结案日期", "分公司", "对接法务", "案件名称"],
            [
                [
                    item["close_date"],
                    item["branch_name"],
                    item["lawyer_name"],
                    item["case_name"],
                ]
                for item in scope["period_closed_cases"]
            ],
        )
    else:
        document.add_paragraph(f"{current_label}无结案案件。")

    _add_section_heading(
        document,
        f"{_chinese_section_number(next_section + 1)}、"
        f"{current_label}新增明细（{period_dates}）",
    )
    new_count = document.add_paragraph(
        f"{current_label}新增 {scope['period_new_count']} 件。"
    )
    new_count.paragraph_format.keep_with_next = bool(scope["period_new_cases"])
    if scope["period_new_cases"]:
        _add_business_table(
            document,
            ["新增日期", "分公司", "对接法务", "案件名称"],
            [
                [
                    item["register_date"],
                    item["branch_name"],
                    item["lawyer_name"],
                    item["case_name"],
                ]
                for item in scope["period_new_cases"]
            ],
        )
    else:
        document.add_paragraph(f"{current_label}无新增案件。")

    document.core_properties.title = f"{scope_label}被告案件{period_name}报"
    document.core_properties.subject = "被告绩效指标及案件明细"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _export_skill_monthly_docx(
    payload: dict[str, Any],
    scope: dict[str, Any],
) -> bytes:
    """Generate the monthly briefing in the structure declared by SKILL.md."""

    period = dict(payload["period"])
    ends_on = date.fromisoformat(str(period["ends_on"]))
    overall = scope.get("scope_type") == "overall"
    scope_label = "被告案件" if overall else f"{scope['scope_name']}被告案件"
    stock_target = _target_magnitude(scope.get("stock_target"))
    new_target = _target_magnitude(scope.get("new_target"))

    document = Document()
    _configure_skill_monthly_document_styles(document)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _skill_run(
        title,
        "被告案件月度简报" if overall else f"{scope['scope_name']}被告案件月度简报",
        bold=True,
        size=18,
    )

    heading = document.add_paragraph()
    if stock_target is not None and new_target is not None:
        heading_text = (
            "一、案件存量下降"
            f"{_compact_decimal(stock_target)}%、新增案件数量下降"
            f"{_compact_decimal(new_target)}%"
        )
    else:
        heading_text = "一、案件存量与新增案件目标完成情况"
    _skill_run(heading, heading_text, bold=True)
    heading.paragraph_format.keep_with_next = True

    _add_skill_rate_line(
        document,
        prefix=f"截止{ends_on:%m月%d日}，{scope_label}新增",
        count=int(scope.get("year_to_date_new_count") or 0),
        rate=scope.get("new_yoy"),
        target=scope.get("new_target"),
    )
    _add_skill_rate_line(
        document,
        prefix=f"截止{ends_on:%m月%d日}，{scope_label}存量",
        count=int(scope.get("stock_count") or 0),
        rate=scope.get("stock_yoy"),
        target=scope.get("stock_target"),
    )

    target_line = document.add_paragraph()
    _skill_run(target_line, "■ ")
    if stock_target is not None and new_target is not None:
        effective_period = _target_period_label(
            scope.get("stock_target"),
            scope.get("new_target"),
        )
        _skill_run(target_line, f"计划{effective_period}，{scope_label}存量同比下降")
        _skill_number(target_line, _compact_decimal(stock_target))
        _skill_run(target_line, f"%，计划{effective_period}，新增案件数量同比下降")
        _skill_number(target_line, _compact_decimal(new_target))
        _skill_run(target_line, "%。")
    else:
        _skill_run(target_line, "当前目标值尚未在规则中配置。")

    team_heading = document.add_paragraph()
    _skill_run(team_heading, "各团队完成情况：", bold=True)
    team_heading.paragraph_format.keep_with_next = True
    team_scopes = [
        item
        for item in payload.get("scopes") or ()
        if isinstance(item, dict) and item.get("scope_type") == "team"
    ]
    if not overall:
        team_scopes = [scope]
    for team in team_scopes:
        _add_skill_team_line(document, team)

    section_two = document.add_paragraph()
    _skill_run(section_two, f"二、{ends_on:%m}月综合数据", bold=True)
    section_two.paragraph_format.keep_with_next = True
    _add_skill_rate_line(
        document,
        prefix=f"{scope_label}新增",
        count=int(scope.get("year_to_date_new_count") or 0),
        rate=scope.get("new_yoy"),
        target=scope.get("new_target"),
    )
    _add_skill_rate_line(
        document,
        prefix=f"{scope_label}存量",
        count=int(scope.get("stock_count") or 0),
        rate=scope.get("stock_yoy"),
        target=scope.get("stock_target"),
    )

    section_three = document.add_paragraph()
    _skill_run(section_three, "三、本月减损情况", bold=True)
    section_three.paragraph_format.keep_with_next = True
    loss_line = document.add_paragraph()
    _skill_run(loss_line, "■ ")
    loss_metrics = (
        scope.get("loss_metrics")
        if isinstance(scope.get("loss_metrics"), dict)
        else {}
    )
    comprehensive = (
        loss_metrics.get("comprehensive_loss_rate")
        if isinstance(loss_metrics.get("comprehensive_loss_rate"), dict)
        else {}
    )
    substantial = (
        loss_metrics.get("substantial_loss_amount")
        if isinstance(loss_metrics.get("substantial_loss_amount"), dict)
        else {}
    )
    if overall and comprehensive.get("status") == "calculated":
        _skill_run(loss_line, "本月被告案件综合减损率")
        _skill_number(
            loss_line,
            str(comprehensive.get("value") or "0.00"),
        )
        _skill_run(loss_line, "%，实质性减损金额累计完成")
        if substantial.get("status") == "calculated":
            _skill_number(
                loss_line,
                str(substantial.get("value_wan") or "0.00"),
            )
            _skill_run(loss_line, "万元；")
        else:
            _skill_run(
                loss_line,
                str(substantial.get("status_label") or "暂不可计算"),
            )
    elif overall:
        _skill_run(
            loss_line,
            str(
                comprehensive.get("status_label")
                or "本月减损指标暂不可计算"
            ),
        )
    else:
        _skill_run(loss_line, "减损指标按法务部门整体口径统计。")

    document.core_properties.title = (
        "被告案件月度简报"
        if overall
        else f"{scope['scope_name']}被告案件月度简报"
    )
    document.core_properties.subject = "按绩效 Skill 生成的被告案件月度报告"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _report_period(
    view: Literal["week", "month"],
    anchor_date: date,
) -> _ReportPeriod:
    if not isinstance(anchor_date, date):
        raise PerformanceReportError("统计截止日期无效")
    if view == "week":
        starts_on = anchor_date - timedelta(days=anchor_date.weekday())
        friday = starts_on + timedelta(days=4)
        ends_on = min(anchor_date, friday)
        first_day = starts_on.replace(day=1)
        week_number = ((starts_on.day + first_day.weekday() - 1) // 7) + 1
        return _ReportPeriod(
            view="week",
            starts_on=starts_on,
            ends_on=ends_on,
            previous_cutoff=starts_on - timedelta(days=3),
            last_year_cutoff=_shift_year(ends_on, -1),
            label=f"{starts_on.year}年{starts_on.month}月第{week_number}周",
            comparison_label="上周五",
        )
    if view == "month":
        starts_on = anchor_date.replace(day=1)
        return _ReportPeriod(
            view="month",
            starts_on=starts_on,
            ends_on=anchor_date,
            previous_cutoff=starts_on - timedelta(days=1),
            last_year_cutoff=_shift_year(anchor_date, -1),
            label=f"{anchor_date.year}年{anchor_date.month}月",
            comparison_label="上月末",
        )
    raise PerformanceReportError("只支持周维度或月维度")


def _case_table_definition(
    rule_spec: dict[str, Any],
) -> tuple[str, dict[str, _FieldRef]]:
    required_names = {
        "branch_name": "分公司名",
        "register_date": "系统登记日期",
        "is_closed": "是否结案",
        "close_date": "结案日期",
    }
    optional_names = {
        "opponent_type": "对方当事人单位/个人",
        "loss_reduction": "减损金额",
        "case_reason": "案情原因",
        "subject_amount": "标的额",
        "interest": "利息",
        "total_principal": "总应付本金",
        "total_appraisal": "总应付鉴定费",
        "total_litigation": "总应付诉讼费",
        "total_penalty": "总应付违约金",
    }
    for table in rule_spec.get("source_tables") or ():
        columns = {
            str(column.get("name") or "").strip(): _FieldRef(
                key=str(column.get("key") or "").strip(),
                source_name=str(column.get("name") or "").strip(),
            )
            for column in table.get("columns") or ()
            if isinstance(column, dict)
        }
        if all(name in columns for name in required_names.values()):
            fields = {
                key: columns[name] for key, name in required_names.items()
            }
            fields.update(
                {
                    key: columns[name]
                    for key, name in optional_names.items()
                    if name in columns
                }
            )
            return str(table.get("key") or ""), fields
    raise PerformanceReportError(
        "当前 Skill 未完整声明分公司、登记日期、结案状态和结案日期，暂不能生成被告绩效报告"
    )


def _subject_lookup(
    rule_spec: dict[str, Any],
    table_key: str,
) -> dict[str, Any]:
    lookup = (rule_spec.get("subject") or {}).get("lookups", {}).get(table_key)
    if not isinstance(lookup, dict) or not lookup.get("steps"):
        raise PerformanceReportError(
            "当前 Skill 未配置分公司到法务、再到团队的确认映射链"
        )
    return lookup


def _field_value(row: dict[str, Any], field: _FieldRef) -> Any:
    if field.key in row:
        return row.get(field.key)
    raw = row.get("__raw__")
    if isinstance(raw, dict):
        return raw.get(field.source_name)
    return None


def _loss_metrics(
    rows: list[dict[str, Any]],
    fields: dict[str, _FieldRef],
    period: _ReportPeriod,
) -> dict[str, Any]:
    required_fields = (
        "opponent_type",
        "loss_reduction",
        "case_reason",
        "subject_amount",
        "interest",
        "total_principal",
        "total_appraisal",
        "total_litigation",
        "total_penalty",
    )
    missing = [field for field in required_fields if field not in fields]
    if missing:
        unavailable = {
            "status": "unavailable",
            "status_label": "规则字段未完整配置",
            "display": "暂不可计算",
        }
        return {
            "comprehensive_loss_rate": dict(unavailable),
            "substantial_loss_amount": dict(unavailable),
        }

    loss_rows = [_loss_row(row, fields) for row in rows]
    period_closed = [
        row
        for row in loss_rows
        if row.is_closed
        and row.close_date is not None
        and period.starts_on <= row.close_date <= period.ends_on
    ]
    payment_rows = [
        row
        for row in period_closed
        if any(
            value is not None
            for value in (
                row.total_principal,
                row.total_appraisal,
                row.total_litigation,
                row.total_penalty,
            )
        )
    ]
    claim_amount = sum(
        (
            (row.subject_amount or Decimal(0))
            + (row.interest or Decimal(0))
            for row in payment_rows
        ),
        Decimal(0),
    )
    payable_amount = sum(
        (
            (row.total_principal or Decimal(0))
            + (row.total_appraisal or Decimal(0))
            + (row.total_litigation or Decimal(0))
            + (row.total_penalty or Decimal(0))
            for row in payment_rows
        ),
        Decimal(0),
    )
    comprehensive_rate = (
        (claim_amount - payable_amount) / claim_amount * Decimal(100)
        if claim_amount
        else Decimal(0)
    ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    year_start = period.ends_on.replace(month=1, day=1)
    substantial_rows = [
        row
        for row in loss_rows
        if row.is_closed
        and row.close_date is not None
        and year_start <= row.close_date <= period.ends_on
        and row.opponent_type in {"供应商", "班组"}
        and row.case_reason == "无争议-债权债务明确"
        and row.loss_reduction is not None
        and row.loss_reduction > 0
    ]
    substantial_yuan = sum(
        (row.loss_reduction or Decimal(0) for row in substantial_rows),
        Decimal(0),
    )
    substantial_wan = (substantial_yuan / Decimal(10000)).quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP,
    )
    return {
        "comprehensive_loss_rate": {
            "status": "calculated",
            "value": f"{comprehensive_rate:.2f}",
            "display": f"{comprehensive_rate:.2f}%",
            "eligible_case_count": len(payment_rows),
            "claim_amount": _decimal_text(claim_amount),
            "payable_amount": _decimal_text(payable_amount),
        },
        "substantial_loss_amount": {
            "status": "calculated",
            "value_yuan": _decimal_text(substantial_yuan),
            "value_wan": f"{substantial_wan:.2f}",
            "display": f"{substantial_wan:.2f}万元",
            "eligible_case_count": len(substantial_rows),
        },
    }


def _weekly_loss_not_applicable() -> dict[str, Any]:
    not_applicable = {
        "status": "not_applicable",
        "status_label": "当前 Skill 未配置周维度减损口径",
        "display": "仅支持月维度",
    }
    return {
        "comprehensive_loss_rate": dict(not_applicable),
        "substantial_loss_amount": dict(not_applicable),
    }


def _loss_row(
    row: dict[str, Any],
    fields: dict[str, _FieldRef],
) -> _LossRow:
    return _LossRow(
        is_closed=_is_closed(_field_value(row, fields["is_closed"])),
        close_date=_optional_date(_field_value(row, fields["close_date"])),
        opponent_type=str(
            _field_value(row, fields["opponent_type"]) or ""
        ).strip(),
        loss_reduction=_optional_decimal(
            _field_value(row, fields["loss_reduction"])
        ),
        case_reason=str(_field_value(row, fields["case_reason"]) or "").strip(),
        subject_amount=_optional_decimal(
            _field_value(row, fields["subject_amount"])
        ),
        interest=_optional_decimal(_field_value(row, fields["interest"])),
        total_principal=_optional_decimal(
            _field_value(row, fields["total_principal"])
        ),
        total_appraisal=_optional_decimal(
            _field_value(row, fields["total_appraisal"])
        ),
        total_litigation=_optional_decimal(
            _field_value(row, fields["total_litigation"])
        ),
        total_penalty=_optional_decimal(
            _field_value(row, fields["total_penalty"])
        ),
    )


def _scope_report(
    scope_key: str,
    scope_name: str,
    cases: list[_CaseRow],
    period: _ReportPeriod,
    targets: list[dict[str, Any]] | tuple[dict[str, Any], ...],
    *,
    scope_type: Literal["overall", "team", "person"],
    data_incomplete: bool,
    loss_metrics: dict[str, Any] | None = None,
) -> dict[str, Any]:
    current_stock = _stock_as_of(cases, period.ends_on)
    previous_stock = _stock_as_of(cases, period.previous_cutoff)
    last_year_stock = _stock_as_of(cases, period.last_year_cutoff)
    year_start = period.ends_on.replace(month=1, day=1)
    last_year_start = period.last_year_cutoff.replace(month=1, day=1)
    year_to_date_new = _new_between(cases, year_start, period.ends_on)
    last_year_to_date_new = _new_between(
        cases,
        last_year_start,
        period.last_year_cutoff,
    )
    period_new = sorted(
        _new_between(cases, period.starts_on, period.ends_on),
        key=lambda item: (
            item.register_date,
            item.branch_name,
            item.case_name,
            item.source_row_number or 0,
        ),
    )
    period_closed = sorted(
        _closed_between(cases, period.starts_on, period.ends_on),
        key=lambda item: (
            item.close_date or date.min,
            item.branch_name,
            item.case_name,
            item.source_row_number or 0,
        ),
    )
    stock_yoy_value = _change_rate(len(current_stock), len(last_year_stock))
    stock_period_value = _change_rate(len(current_stock), len(previous_stock))
    new_yoy_value = _change_rate(
        len(year_to_date_new),
        len(last_year_to_date_new),
    )
    branches = sorted({item.branch_name for item in cases if item.branch_name})
    branch_rows = [
        _branch_report(
            branch, [item for item in cases if item.branch_name == branch], period
        )
        for branch in branches
    ]
    branch_rows.sort(
        key=lambda item: (
            -int(item["stock_count"]),
            -int(item["year_to_date_new_count"]),
            str(item["branch_name"]),
        )
    )
    result = {
        "scope_key": scope_key,
        "scope_name": scope_name,
        "scope_type": scope_type,
        "stock_count": len(current_stock),
        "previous_stock_count": len(previous_stock),
        "last_year_stock_count": len(last_year_stock),
        "stock_yoy": _rate_result(stock_yoy_value),
        "stock_period_change": _rate_result(stock_period_value),
        "year_to_date_new_count": len(year_to_date_new),
        "last_year_to_date_new_count": len(last_year_to_date_new),
        "new_yoy": _rate_result(new_yoy_value),
        "period_new_count": len(period_new),
        "period_closed_count": len(period_closed),
        "stock_target": _target_result(
            "存量同比下降率",
            stock_yoy_value,
            targets,
            scope_key=scope_key,
            scope_name=scope_name,
            period=period,
            default_comparison="at_most",
            data_incomplete=data_incomplete,
        ),
        "new_target": _target_result(
            "新增同比下降率",
            new_yoy_value,
            targets,
            scope_key=scope_key,
            scope_name=scope_name,
            period=period,
            default_comparison="at_most",
            data_incomplete=data_incomplete,
        ),
        "branches": branch_rows,
        "period_new_cases": [_case_detail(item) for item in period_new],
        "period_closed_cases": [_case_detail(item) for item in period_closed],
    }
    if scope_type == "overall" and loss_metrics is not None:
        result["loss_metrics"] = copy.deepcopy(loss_metrics)
    return result


def _branch_report(
    branch_name: str,
    cases: list[_CaseRow],
    period: _ReportPeriod,
) -> dict[str, Any]:
    current_stock = _stock_as_of(cases, period.ends_on)
    previous_stock = _stock_as_of(cases, period.previous_cutoff)
    last_year_stock = _stock_as_of(cases, period.last_year_cutoff)
    year_to_date_new = _new_between(
        cases,
        period.ends_on.replace(month=1, day=1),
        period.ends_on,
    )
    last_year_to_date_new = _new_between(
        cases,
        period.last_year_cutoff.replace(month=1, day=1),
        period.last_year_cutoff,
    )
    return {
        "branch_name": branch_name,
        "lawyer_name": next(
            (item.lawyer_name for item in cases if item.lawyer_name),
            "",
        ),
        "stock_count": len(current_stock),
        "previous_stock_count": len(previous_stock),
        "last_year_stock_count": len(last_year_stock),
        "stock_yoy": _rate_result(
            _change_rate(len(current_stock), len(last_year_stock))
        ),
        "stock_period_change": _rate_result(
            _change_rate(len(current_stock), len(previous_stock))
        ),
        "year_to_date_new_count": len(year_to_date_new),
        "last_year_to_date_new_count": len(last_year_to_date_new),
        "new_yoy": _rate_result(
            _change_rate(len(year_to_date_new), len(last_year_to_date_new))
        ),
        "period_new_count": len(_new_between(cases, period.starts_on, period.ends_on)),
        "period_closed_count": len(
            _closed_between(cases, period.starts_on, period.ends_on)
        ),
    }


def _stock_as_of(cases: list[_CaseRow], cutoff: date) -> list[_CaseRow]:
    return [
        item
        for item in cases
        if item.register_date <= cutoff
        and (
            not item.is_closed
            or (item.close_date is not None and item.close_date > cutoff)
        )
    ]


def _new_between(
    cases: list[_CaseRow],
    starts_on: date,
    ends_on: date,
) -> list[_CaseRow]:
    return [item for item in cases if starts_on <= item.register_date <= ends_on]


def _closed_between(
    cases: list[_CaseRow],
    starts_on: date,
    ends_on: date,
) -> list[_CaseRow]:
    return [
        item
        for item in cases
        if item.is_closed
        and item.close_date is not None
        and starts_on <= item.close_date <= ends_on
    ]


def _case_detail(item: _CaseRow) -> dict[str, Any]:
    return {
        "branch_name": item.branch_name,
        "lawyer_name": item.lawyer_name,
        "case_name": item.case_name,
        "register_date": item.register_date.isoformat(),
        "close_date": item.close_date.isoformat() if item.close_date else None,
        "source_row_number": item.source_row_number,
    }


def _selected_scope(
    payload: dict[str, Any],
    scope_key: str,
) -> dict[str, Any]:
    for scope in payload.get("scopes") or ():
        if str(scope.get("scope_key") or "") == str(scope_key or ""):
            return dict(scope)
    raise PerformanceReportError("未找到要导出的团队或整体范围")


def _arrow_rate(rate: dict[str, Any]) -> str:
    try:
        value = Decimal(str(rate.get("value") or "0"))
    except (InvalidOperation, ValueError):
        value = Decimal(0)
    arrow = "↑" if value > 0 else "↓"
    return f"{arrow}{abs(value):.2f}%"


def _sentence_rate(rate: dict[str, Any]) -> str:
    display = str(rate.get("display") or "").strip()
    for label in ("下降", "增长", "持平"):
        if display.startswith(label):
            return f"{label} {display[len(label) :]}"
    return display or "暂不可比"


def _target_summary_text(
    metric_label: str,
    rate: Any,
    target: Any,
) -> str:
    actual_text = _sentence_rate(rate if isinstance(rate, dict) else {})
    if not isinstance(target, dict):
        return f"{metric_label}{actual_text}"
    target_display = str(target.get("target_display") or "").strip()
    target_text = (
        _sentence_rate({"display": target_display}) if target_display else ""
    )
    status = str(target.get("status") or "")
    gap = _target_gap_points(rate, target)
    if status == "achieved" and gap is not None:
        return (
            f"{metric_label}{actual_text}，达到目标"
            f"（目标{target_text}），超过目标 {gap:.2f} 个百分点"
        )
    if status == "not_achieved" and gap is not None:
        return (
            f"{metric_label}{actual_text}，未达到目标"
            f"（目标{target_text}），距离目标还差 {gap:.2f} 个百分点"
        )
    if target_text:
        return f"{metric_label}{actual_text}（目标{target_text}）"
    return f"{metric_label}{actual_text}"


def _target_gap_points(rate: Any, target: Any) -> Decimal | None:
    rate_value = _rate_decimal(rate)
    if rate_value is None or not isinstance(target, dict):
        return None
    try:
        target_value = Decimal(str(target.get("target_value")))
    except (InvalidOperation, TypeError, ValueError):
        return None
    return abs(rate_value - target_value).quantize(
        Decimal("0.01"),
        rounding=ROUND_HALF_UP,
    )


def _excel_target_gap_text(rate: Any, target: Any) -> str | None:
    if not isinstance(target, dict):
        return None
    gap = _target_gap_points(rate, target)
    if gap is None:
        return None
    status = str(target.get("status") or "")
    if status == "achieved":
        return f"超过目标{gap:.2f}个百分点"
    if status == "not_achieved":
        return f"距离目标还差{gap:.2f}个百分点"
    return None


def _excel_date(value: Any) -> date | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _populate_excel_overview(
    sheet: Any,
    *,
    scope: dict[str, Any],
    period: dict[str, Any],
) -> None:
    weekly = str(period.get("view") or "") == "week"
    period_name = "周报" if weekly else "月报"
    current_label = "本周" if weekly else "本月"
    scope_label = (
        "法务部门整体"
        if str(scope.get("scope_name") or "") == "整体"
        else str(scope.get("scope_name") or "")
    )
    starts_on = date.fromisoformat(str(period["starts_on"]))
    ends_on = date.fromisoformat(str(period["ends_on"]))

    sheet.merge_cells("A1:G1")
    sheet["A1"] = f"{scope_label}被告案件{period_name}"
    sheet.merge_cells("A2:B2")
    sheet["A2"] = "查看范围"
    sheet.merge_cells("C2:D2")
    sheet["C2"] = scope_label
    sheet.merge_cells("E2:F2")
    sheet["E2"] = "统计维度"
    sheet["G2"] = "周维度" if weekly else "月维度"
    sheet.merge_cells("A3:B3")
    sheet["A3"] = "统计期间"
    sheet.merge_cells("C3:G3")
    sheet["C3"] = (
        f"{starts_on.year}年{starts_on.month}月{starts_on.day}日"
        f" — {ends_on.month}月{ends_on.day}日"
    )

    sheet.merge_cells("A5:G5")
    sheet["A5"] = "核心指标"
    sheet.append(
        [
            "指标",
            "当前数量",
            "同比",
            "环比/本期",
            "目标",
            "完成情况",
            "差距说明",
        ]
    )
    stock_target = (
        scope.get("stock_target")
        if isinstance(scope.get("stock_target"), dict)
        else {}
    )
    new_target = (
        scope.get("new_target")
        if isinstance(scope.get("new_target"), dict)
        else {}
    )
    sheet.append(
        [
            "案件存量",
            int(scope.get("stock_count") or 0),
            str((scope.get("stock_yoy") or {}).get("display") or ""),
            str(
                (scope.get("stock_period_change") or {}).get("display") or ""
            ),
            str(stock_target.get("target_display") or ""),
            str(stock_target.get("status_label") or ""),
            _excel_target_gap_text(scope.get("stock_yoy"), stock_target),
        ]
    )
    sheet.append(
        [
            "年度累计新增",
            int(scope.get("year_to_date_new_count") or 0),
            str((scope.get("new_yoy") or {}).get("display") or ""),
            None,
            str(new_target.get("target_display") or ""),
            str(new_target.get("status_label") or ""),
            _excel_target_gap_text(scope.get("new_yoy"), new_target),
        ]
    )
    sheet.append(
        [
            f"{current_label}新增",
            int(scope.get("period_new_count") or 0),
            None,
            f"{starts_on:%m月%d日}—{ends_on:%m月%d日}",
            None,
            None,
            None,
        ]
    )
    sheet.append(
        [
            f"{current_label}结案",
            int(scope.get("period_closed_count") or 0),
            None,
            f"{starts_on:%m月%d日}—{ends_on:%m月%d日}",
            None,
            None,
            None,
        ]
    )

    next_row = 12
    loss_metrics = (
        scope.get("loss_metrics")
        if isinstance(scope.get("loss_metrics"), dict)
        else {}
    )
    if not weekly and scope.get("scope_type") == "overall" and loss_metrics:
        comprehensive = (
            loss_metrics.get("comprehensive_loss_rate")
            if isinstance(loss_metrics.get("comprehensive_loss_rate"), dict)
            else {}
        )
        substantial = (
            loss_metrics.get("substantial_loss_amount")
            if isinstance(loss_metrics.get("substantial_loss_amount"), dict)
            else {}
        )
        sheet.merge_cells(
            start_row=next_row,
            start_column=1,
            end_row=next_row,
            end_column=7,
        )
        sheet.cell(next_row, 1, "部门整体减损")
        sheet.cell(next_row + 1, 1, "综合减损率")
        sheet.cell(
            next_row + 1,
            2,
            comprehensive.get("display")
            or comprehensive.get("status_label")
            or "暂不可计算",
        )
        sheet.cell(
            next_row + 1,
            3,
            f"纳入{int(comprehensive.get('eligible_case_count') or 0)}件",
        )
        sheet.cell(next_row + 2, 1, "实质减损金额")
        sheet.cell(
            next_row + 2,
            2,
            substantial.get("display")
            or substantial.get("status_label")
            or "暂不可计算",
        )
        sheet.cell(
            next_row + 2,
            3,
            f"纳入{int(substantial.get('eligible_case_count') or 0)}件",
        )
        next_row += 5

    sheet.merge_cells(
        start_row=next_row,
        start_column=1,
        end_row=next_row,
        end_column=7,
    )
    sheet.cell(next_row, 1, "说明")
    sheet.merge_cells(
        start_row=next_row + 1,
        start_column=1,
        end_row=next_row + 1,
        end_column=7,
    )
    sheet.cell(
        next_row + 1,
        1,
        "归属按“底表分公司→法务对接人→团队”映射；"
        "详细计算结果及逐案记录请查看后续工作表。",
    )


def _style_excel_overview(sheet: Any) -> None:
    navy = "1F4E78"
    blue = "4477C2"
    pale_blue = "DCE6F1"
    border = Border(
        left=Side(style="thin", color="B8C4D1"),
        right=Side(style="thin", color="B8C4D1"),
        top=Side(style="thin", color="B8C4D1"),
        bottom=Side(style="thin", color="B8C4D1"),
    )
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet["A1"].fill = PatternFill("solid", fgColor=navy)
    sheet["A1"].font = Font(
        name="Microsoft YaHei",
        color="FFFFFF",
        bold=True,
        size=18,
    )
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34
    for row_number in (2, 3):
        for cell in sheet[row_number]:
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )
        label_columns = (1, 5) if row_number == 2 else (1,)
        for column in label_columns:
            cell = sheet.cell(row_number, column)
            cell.fill = PatternFill("solid", fgColor=pale_blue)
            cell.font = Font(name="Microsoft YaHei", bold=True, size=10)
    for row_number in range(1, sheet.max_row + 1):
        first_value = str(sheet.cell(row_number, 1).value or "")
        if first_value in {"核心指标", "部门整体减损", "说明"}:
            for cell in sheet[row_number]:
                cell.fill = PatternFill("solid", fgColor=pale_blue)
                cell.font = Font(
                    name="Microsoft YaHei",
                    color=navy,
                    bold=True,
                    size=11,
                )
            sheet.row_dimensions[row_number].height = 24
    for cell in sheet[6]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(
            name="Microsoft YaHei",
            color="FFFFFF",
            bold=True,
            size=10,
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for row_number in range(7, 11):
        for cell in sheet[row_number]:
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
    for row_number in range(1, sheet.max_row + 1):
        status_cell = sheet.cell(row_number, 6)
        status = str(status_cell.value or "")
        if status == "达到目标":
            status_cell.fill = PatternFill("solid", fgColor="E2F0D9")
            status_cell.font = Font(
                name="Microsoft YaHei",
                color="2E7D32",
                bold=True,
            )
        elif status == "未达到目标":
            status_cell.fill = PatternFill("solid", fgColor="FCE4D6")
            status_cell.font = Font(
                name="Microsoft YaHei",
                color="C62828",
                bold=True,
            )
    widths = {
        "A": 22,
        "B": 13,
        "C": 16,
        "D": 20,
        "E": 16,
        "F": 15,
        "G": 29,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row_number in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_number].height = max(
            sheet.row_dimensions[row_number].height or 0,
            23,
        )
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A7"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _style_export_sheet(sheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="4477C2")
    header_font = Font(color="FFFFFF", bold=True, size=12)
    total_font = Font(bold=True, size=11)
    border = Border(
        left=Side(style="thin", color="303030"),
        right=Side(style="thin", color="303030"),
        top=Side(style="thin", color="303030"),
        bottom=Side(style="thin", color="303030"),
    )
    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if str(sheet.cell(sheet.max_row, 1).value or "") == "合计":
        for cell in sheet[sheet.max_row]:
            cell.font = total_font
            cell.fill = PatternFill("solid", fgColor="DCE6F1")
    if sheet.title in {"新增案件明细", "结案案件明细"}:
        for row_number in range(2, sheet.max_row + 1):
            sheet.cell(row_number, 1).number_format = "yyyy-mm-dd"
            sheet.cell(row_number, 4).alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for index in range(1, sheet.max_column + 1):
        widest = max(
            len(str(sheet.cell(row=row, column=index).value or ""))
            for row in range(1, sheet.max_row + 1)
        )
        sheet.column_dimensions[
            sheet.cell(row=1, column=index).column_letter
        ].width = min(
            max(widest + 3, 13),
            36,
        )
    sheet.row_dimensions[1].height = 25
    sheet.page_setup.orientation = (
        "landscape" if sheet.max_column >= 7 else "portrait"
    )
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _configure_document_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for section in document.sections:
        section.top_margin = Pt(48)
        section.bottom_margin = Pt(48)
        section.left_margin = Pt(50)
        section.right_margin = Pt(50)


def _configure_skill_monthly_document_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(16)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.paragraph_format.line_spacing = Pt(23)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    for section in document.sections:
        section.top_margin = Pt(42)
        section.bottom_margin = Pt(42)
        section.left_margin = Pt(50)
        section.right_margin = Pt(50)


def _skill_run(
    paragraph: Any,
    text: Any,
    *,
    bold: bool = False,
    size: int = 16,
    color: RGBColor | None = None,
    underline: bool = False,
) -> Any:
    run = paragraph.add_run(str(text))
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    if color is not None:
        run.font.color.rgb = color
    return run


def _skill_number(paragraph: Any, value: Any) -> Any:
    return _skill_run(
        paragraph,
        value,
        bold=True,
        color=RGBColor(0, 0, 255),
        underline=True,
    )


def _target_magnitude(target: Any) -> Decimal | None:
    if not isinstance(target, dict):
        return None
    try:
        return abs(Decimal(str(target.get("target_value"))))
    except (InvalidOperation, TypeError, ValueError):
        return None


def _compact_decimal(value: Decimal) -> str:
    normalized = value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if normalized == normalized.to_integral():
        return str(int(normalized))
    return f"{normalized:.2f}"


def _target_period_label(*targets: Any) -> str:
    for target in targets:
        if not isinstance(target, dict):
            continue
        label = str(target.get("effective_period") or "").strip()
        if label:
            return label.removesuffix("目标").strip()
    return "当前考核期"


def _rate_decimal(rate: Any) -> Decimal | None:
    if not isinstance(rate, dict):
        return None
    try:
        return Decimal(str(rate.get("value")))
    except (InvalidOperation, TypeError, ValueError):
        return None


def _skill_rate_phrase(
    paragraph: Any,
    rate: Any,
    *,
    decline_word: str = "下降",
) -> None:
    value = _rate_decimal(rate)
    if value is None:
        _skill_run(paragraph, "暂不可比")
        return
    if value < 0:
        _skill_run(paragraph, decline_word, color=RGBColor(117, 189, 66))
    elif value > 0:
        _skill_run(paragraph, "增长", color=RGBColor(255, 0, 0))
    else:
        _skill_run(paragraph, "持平")
    _skill_run(paragraph, " ")
    _skill_number(paragraph, f"{abs(value):.2f}")
    _skill_run(paragraph, "%")


def _add_skill_rate_line(
    document: Document,
    *,
    prefix: str,
    count: int,
    rate: Any,
    target: Any,
) -> None:
    paragraph = document.add_paragraph()
    _skill_run(paragraph, f"■ {prefix}")
    _skill_number(paragraph, count)
    _skill_run(paragraph, "件，同比")
    _skill_rate_phrase(paragraph, rate, decline_word="降低")

    rate_value = _rate_decimal(rate)
    target_value = None
    if isinstance(target, dict):
        try:
            target_value = Decimal(str(target.get("target_value")))
        except (InvalidOperation, TypeError, ValueError):
            target_value = None
    status = str(target.get("status") or "") if isinstance(target, dict) else ""
    if (
        rate_value is not None
        and target_value is not None
        and status in {"achieved", "not_achieved"}
    ):
        _skill_run(paragraph, "，较目标值")
        achieved = status == "achieved"
        _skill_run(
            paragraph,
            "高" if achieved else "低",
            color=RGBColor(117, 189, 66) if achieved else RGBColor(255, 0, 0),
        )
        _skill_run(paragraph, " ")
        _skill_number(paragraph, f"{abs(rate_value - target_value):.2f}")
        _skill_run(paragraph, "个百分点；")
    else:
        _skill_run(paragraph, "；")


def _add_skill_team_line(document: Document, team: dict[str, Any]) -> None:
    paragraph = document.add_paragraph()
    _skill_run(paragraph, f"■ {team.get('scope_name') or '未命名团队'}：")
    _skill_run(paragraph, "存量")
    _skill_rate_phrase(paragraph, team.get("stock_yoy"))
    _skill_run(paragraph, "，新增")
    _skill_rate_phrase(paragraph, team.get("new_yoy"))


def _add_section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(12)


def _chinese_section_number(value: int) -> str:
    return {
        1: "一",
        2: "二",
        3: "三",
        4: "四",
        5: "五",
    }.get(value, str(value))


def _add_business_table(
    document: Document,
    headers: list[str],
    rows: list[list[Any]],
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    column_widths = _business_table_column_widths(headers)
    _set_docx_table_geometry(table, column_widths)
    _set_docx_row_repeat_header(table.rows[0])
    _set_docx_row_cant_split(table.rows[0])
    narrative_columns = {
        index for index, header in enumerate(headers) if header == "案件名称"
    }
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_docx_cell_margins(cell)
        _shade_docx_cell(cell, "3566B4")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].paragraph_format.keep_with_next = True
        cell.paragraphs[0].paragraph_format.space_before = Pt(0)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    for values in rows:
        cells = table.add_row().cells
        _set_docx_row_cant_split(table.rows[-1])
        for index, value in enumerate(values):
            cells[index].text = str(value if value is not None else "")
            cells[index].width = Inches(column_widths[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_docx_cell_margins(cells[index])
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if index in narrative_columns
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            for run in paragraph.runs:
                run.font.size = Pt(9)
        if values and str(values[0] or "").strip() == "合计":
            for cell in cells:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def _business_table_column_widths(headers: list[str]) -> list[float]:
    if len(headers) == 4 and headers[-1] == "案件名称":
        return [1.05, 1.40, 1.00, 3.55]
    if len(headers) == 6:
        return [1.50, 1.10, 1.10, 1.10, 1.10, 1.10]
    width = 7.0 / max(len(headers), 1)
    return [width for _ in headers]


def _set_docx_table_geometry(table: Any, widths_in: list[float]) -> None:
    widths = [round(width * 1440) for width in widths_in]
    total_width = sum(widths)
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(total_width))
    table_indent = properties.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        properties.append(table_indent)
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), "120")
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[index])
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                properties.append(cell_width)
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(widths[index]))


def _set_docx_cell_margins(cell: Any) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (
        ("top", 70),
        ("left", 80),
        ("bottom", 70),
        ("right", 80),
    ):
        margin = margins.find(qn(f"w:{edge}"))
        if margin is None:
            margin = OxmlElement(f"w:{edge}")
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def _set_docx_row_repeat_header(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _set_docx_row_cant_split(row: Any) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    properties.append(marker)


def _shade_docx_cell(cell: Any, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _change_rate(current: int, comparison: int) -> Decimal:
    if comparison == 0:
        return Decimal(0)
    return (
        (Decimal(current) - Decimal(comparison)) / Decimal(comparison) * Decimal(100)
    ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _rate_result(value: Decimal) -> dict[str, str]:
    if value < 0:
        direction = "decline"
        label = "下降"
    elif value > 0:
        direction = "growth"
        label = "增长"
    else:
        direction = "flat"
        label = "持平"
    magnitude = abs(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    return {
        "value": str(value),
        "direction": direction,
        "display": f"{label}{magnitude:.2f}%",
    }


def _target_result(
    metric_name: str,
    actual: Decimal,
    targets: list[dict[str, Any]] | tuple[dict[str, Any], ...],
    *,
    scope_key: str,
    scope_name: str,
    period: _ReportPeriod,
    default_comparison: Literal["at_least", "at_most", "equal"],
    data_incomplete: bool,
) -> dict[str, str]:
    target = _matching_target(
        metric_name,
        targets,
        scope_key=scope_key,
        scope_name=scope_name,
        period=period,
    )
    if target is None:
        return _unknown_target()
    try:
        target_value = Decimal(
            str(target.get("target_value") or "").replace("%", "").strip()
        )
    except (InvalidOperation, ValueError):
        return _unknown_target()
    comparison = str(target.get("comparison") or default_comparison).strip()
    if comparison == "at_least":
        achieved = actual >= target_value
    elif comparison == "at_most":
        achieved = actual <= target_value
    elif comparison == "equal":
        achieved = actual == target_value
    else:
        return {
            **_unknown_target(),
            "status_label": "目标判断方向待确认",
            "target_value": str(target_value),
            "target_display": _rate_result(target_value)["display"],
        }
    result = {
        "status": "achieved" if achieved else "not_achieved",
        "status_label": "达到目标" if achieved else "未达到目标",
        "target_value": str(target_value),
        "target_display": _rate_result(target_value)["display"],
        "comparison": comparison,
        "comparison_label": {
            "at_least": "不低于",
            "at_most": "不高于",
            "equal": "等于",
        }[comparison],
        "target_scope": str(
            target.get("scope_name")
            or target.get("scope_key")
            or target.get("team_name")
            or ""
        ),
        "effective_period": str(target.get("effective_period") or ""),
    }
    if data_incomplete:
        result["status"] = "data_incomplete"
        result["status_label"] = "数据不完整，目标待确认"
    return result


def _unknown_target() -> dict[str, str]:
    return {
        "status": "unknown",
        "status_label": "目标待确认",
        "target_value": "",
        "target_display": "",
        "comparison": "",
        "comparison_label": "",
        "target_scope": "",
        "effective_period": "",
    }


def _matching_target(
    metric_name: str,
    targets: list[dict[str, Any]] | tuple[dict[str, Any], ...],
    *,
    scope_key: str,
    scope_name: str,
    period: _ReportPeriod,
) -> dict[str, Any] | None:
    scope_candidates = {scope_key.casefold(), scope_name.casefold()}
    scoped: list[tuple[date, int, dict[str, Any]]] = []
    generic: list[tuple[date, int, dict[str, Any]]] = []
    for index, item in enumerate(targets):
        if str(item.get("metric_name") or "").strip() != metric_name:
            continue
        effective_date = _target_effective_date(item, period)
        if effective_date is None or effective_date > period.ends_on:
            continue
        target_scope = str(
            item.get("scope_name")
            or item.get("scope_key")
            or item.get("team_name")
            or ""
        ).strip()
        if target_scope and target_scope.casefold() in scope_candidates:
            scoped.append((effective_date, index, item))
        elif not target_scope:
            generic.append((effective_date, index, item))
    candidates = scoped or generic
    return max(candidates, key=lambda candidate: (candidate[0], candidate[1]))[2] if candidates else None


def _target_effective_date(
    target: dict[str, Any],
    period: _ReportPeriod,
) -> date | None:
    configured = str(target.get("effective_period") or "").strip()
    if not configured:
        return date.min
    compact = re.sub(r"\s+", "", configured)
    exact_day = re.fullmatch(
        r"(\d{4})(?:年|[-./])(\d{1,2})(?:月|[-./])(\d{1,2})(?:日)?(?:目标)?",
        compact,
    )
    if exact_day:
        return _safe_date(
            int(exact_day.group(1)),
            int(exact_day.group(2)),
            int(exact_day.group(3)),
        )
    month_end = re.fullmatch(
        r"(\d{4})年(\d{1,2})月(?:底|末)(?:目标)?",
        compact,
    )
    if month_end:
        year, month = int(month_end.group(1)), int(month_end.group(2))
        if not 1 <= month <= 12:
            return None
        return date(year, month, monthrange(year, month)[1])
    month_start = re.fullmatch(
        r"(\d{4})(?:年|[-./])(\d{1,2})(?:月)?(?:目标)?",
        compact,
    )
    if month_start:
        return _safe_date(
            int(month_start.group(1)),
            int(month_start.group(2)),
            1,
        )
    if compact == re.sub(r"\s+", "", period.label):
        return period.starts_on
    return None


def _safe_date(year: int, month: int, day: int) -> date | None:
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _primary_lawyer_name(value: Any) -> str:
    """Follow the Skill's explicit personal rule for slash-separated contacts."""

    return str(value or "").replace("／", "/").split("/", 1)[0].strip()


def _case_name(row: dict[str, Any]) -> str:
    raw = row.get("__raw__")
    if isinstance(raw, dict):
        for key in ("案件名称", "案件名", "项目名称"):
            value = str(raw.get(key) or "").strip()
            if value:
                return value
    row_number = _row_number(row)
    return f"原表第{row_number}行（案件名称未识别）" if row_number else "案件名称未识别"


def _optional_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _optional_decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value if value.is_finite() else None
    text = str(value).strip()
    if not text:
        return None
    normalized = (
        text.replace(",", "")
        .replace("，", "")
        .replace("￥", "")
        .replace("¥", "")
        .strip()
    )
    try:
        parsed = Decimal(normalized)
    except InvalidOperation:
        return None
    return parsed if parsed.is_finite() else None


def _excel_amount(value: Any) -> int | float | None:
    parsed = _optional_decimal(value)
    if parsed is None:
        return None
    if parsed == parsed.to_integral_value():
        return int(parsed)
    return float(parsed)


def _decimal_text(value: Decimal) -> str:
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _is_closed(value: Any) -> bool:
    return str(value or "").strip() == "是"


def _row_number(row: dict[str, Any]) -> int | None:
    try:
        value = int(row.get("__row_number__"))
    except (TypeError, ValueError):
        return None
    return value if value > 0 else None


def _assignment_error_message(error: dict[str, Any]) -> str:
    qualifier = "唯一" if bool(error.get("ambiguous")) else ""
    return (
        f"无法通过“{error.get('step_name') or '团队归属'}”"
        f"{qualifier}匹配：{error.get('value') or '空值'}"
    )


def _shift_year(value: date, years: int) -> date:
    try:
        return value.replace(year=value.year + years)
    except ValueError:
        return value.replace(year=value.year + years, day=28)
