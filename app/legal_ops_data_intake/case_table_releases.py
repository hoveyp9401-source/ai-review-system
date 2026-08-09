from __future__ import annotations

import hashlib
import json
import math
import os
import re
import shutil
import sqlite3
import tempfile
import time
import uuid
from collections import Counter
from collections.abc import Iterable
from contextlib import closing, contextmanager
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.agent2.case_table_rag import CaseTableDocument, write_case_table_index
from app.legal_ops_data_intake.storage import (
    FileStorageError,
    IntakeFileStore,
    safe_filename,
)

CASE_TABLE_TYPES = {
    "defendant": "defendant_case_table",
    "plaintiff": "plaintiff_case_table",
}
CASE_TABLE_LABELS = {
    "defendant_case_table": "被告案件底表",
    "plaintiff_case_table": "原告案件底表",
}
LIVE_INDEX_FILES = (
    "case_documents.jsonl",
    "case_index_meta.json",
    "case_index.sqlite",
)
SHANGHAI = ZoneInfo("Asia/Shanghai")


class CaseTableReleaseError(ValueError):
    def __init__(
        self,
        message: str,
        *,
        code: str = "case_table_release_error",
        status_code: int = 400,
    ):
        super().__init__(message)
        self.message = message
        self.code = code
        self.status_code = status_code


class CaseTableReleaseService:
    """Preview and atomically publish the case workbooks used by Agent2 queries."""

    def __init__(self, settings: Any, *, actor_user_id: str):
        self.settings = settings
        self.actor_user_id = str(actor_user_id or "").strip() or "unknown"
        storage_root = Path(settings.legal_ops_data_intake_storage_path).resolve()
        self.root = storage_root / "case_table_releases"
        self.pending_root = self.root / "pending"
        self.versions_root = self.root / "versions"
        self.current_path = self.root / "current.json"
        self.lock_path = self.root / ".publish.lock"
        self.raw_root = Path(
            getattr(settings, "legal_ops_case_table_raw_path", "data/rag_sources/raw")
        ).resolve()
        self.index_dir = Path(
            getattr(
                settings,
                "legal_ops_case_table_index_path",
                "data/rag_indexes/case_tables",
            )
        ).resolve()
        max_bytes = int(settings.legal_ops_data_intake_max_file_mb) * 1024 * 1024
        self.files = IntakeFileStore(self.root / "files", max_bytes=max_bytes)

    def status(self) -> dict[str, Any]:
        with _exclusive_lock(self.lock_path):
            current = self._ensure_baseline_locked()
            counts = self._index_counts(self.index_dir / "case_index.sqlite")
            sources = self._source_payloads(
                current.get("sources") or [], counts=counts
            )
        return {
            "current_version": str(current.get("version_id") or ""),
            "published_at": str(current.get("published_at") or ""),
            "published_by": str(current.get("published_by") or ""),
            "reason": str(current.get("reason") or ""),
            "document_count": sum(counts.values()),
            "tables": sources,
            "index_ready": (self.index_dir / "case_index.sqlite").is_file(),
        }

    def history(self, *, limit: int = 20) -> list[dict[str, Any]]:
        with _exclusive_lock(self.lock_path):
            current = self._ensure_baseline_locked()
        current_id = str(current.get("version_id") or "")
        versions: list[dict[str, Any]] = []
        if self.versions_root.is_dir():
            for path in self.versions_root.glob("*.json"):
                try:
                    item = _read_json(path)
                except (OSError, ValueError, json.JSONDecodeError):
                    continue
                item["is_current"] = str(item.get("version_id") or "") == current_id
                item["tables"] = self._source_payloads(
                    item.get("sources") or [],
                    counts=dict(item.get("table_counts") or {}),
                )
                item["can_restore"] = self._sources_available(item.get("sources") or [])
                versions.append(_public_version(item))
        versions.sort(
            key=lambda item: str(item.get("published_at") or ""), reverse=True
        )
        return versions[: max(1, min(int(limit or 20), 100))]

    def pending(self, *, limit: int = 20) -> list[dict[str, Any]]:
        previews: list[dict[str, Any]] = []
        with _exclusive_lock(self.lock_path):
            self._ensure_baseline_locked()
            if self.pending_root.is_dir():
                for manifest_path in self.pending_root.glob("*/manifest.json"):
                    try:
                        manifest = _read_json(manifest_path)
                    except (OSError, ValueError, json.JSONDecodeError):
                        continue
                    if manifest.get("status") != "ready" or not manifest.get(
                        "can_publish"
                    ):
                        continue
                    previews.append(_public_preview(manifest))
        previews.sort(
            key=lambda item: str(item.get("created_at") or ""), reverse=True
        )
        return previews[: max(1, min(int(limit or 20), 100))]

    def preview(
        self, *, content: bytes, filename: str, table_kind: str
    ) -> dict[str, Any]:
        table_type = self._table_type(table_kind)
        try:
            safe_name = safe_filename(filename)
        except FileStorageError as exc:
            raise CaseTableReleaseError(
                str(exc), code="invalid_file_name", status_code=422
            ) from exc
        if Path(safe_name).suffix.lower() != ".xlsx":
            raise CaseTableReleaseError(
                "案件底表只支持 .xlsx 文件", code="invalid_file_type", status_code=422
            )
        if table_kind == "defendant" and "原告" in safe_name:
            raise CaseTableReleaseError(
                "当前选择的是被告案件底表，但文件名显示为原告底表",
                code="table_type_mismatch",
                status_code=422,
            )
        if table_kind == "plaintiff" and "被告" in safe_name:
            raise CaseTableReleaseError(
                "当前选择的是原告案件底表，但文件名显示为被告底表",
                code="table_type_mismatch",
                status_code=422,
            )

        try:
            stored = self.files.store(content, safe_name)
        except FileStorageError as exc:
            raise CaseTableReleaseError(
                str(exc), code="invalid_upload", status_code=422
            ) from exc
        batch_no = _reference("CT")
        pending_dir = self.pending_root / batch_no
        candidate_dir = pending_dir / "candidate"
        pending_dir.mkdir(parents=True, exist_ok=False)
        input_path = pending_dir / safe_name
        _copy_atomic(stored.absolute_path, input_path)

        now = _now()
        archive_path = self._archive_target(safe_name, stored.file_hash, now[:10])
        current_index = self.index_dir / "case_index.sqlite"
        if not current_index.is_file():
            raise CaseTableReleaseError(
                "当前案件底表索引不存在，暂时不能生成更新预览",
                code="live_index_missing",
                status_code=409,
            )

        try:
            extracted = _extract_documents_from_workbook(
                input_path, refresh_date=now[:10]
            )
        except Exception as exc:
            raise CaseTableReleaseError(
                "无法读取该 Excel，请确认文件没有损坏且首行是表头",
                code="invalid_workbook",
                status_code=422,
            ) from exc
        documents = [
            replace(
                item,
                table_type=table_type,
                source_file=safe_name,
                updated_at=now[:10],
            )
            for item in extracted
        ]
        if not documents:
            raise CaseTableReleaseError(
                "文件中没有识别到案件明细", code="empty_workbook", status_code=422
            )

        # Read the live documents, source list, version and fingerprint as one
        # snapshot. A concurrent publish must not mix old rows with a new hash.
        with _exclusive_lock(self.lock_path):
            current_version = self._ensure_baseline_locked()
            current_documents = self._read_documents(current_index)
            current_sources = self._current_sources()
            current_hash = _sha256(current_index)
        old_documents = [
            item for item in current_documents if item.table_type == table_type
        ]
        other_documents = [
            item for item in current_documents if item.table_type != table_type
        ]
        checks = self._validate_documents(documents, table_type=table_type)
        comparison = self._compare_documents(old_documents, documents)
        same_as_current = (
            bool(old_documents)
            and comparison["added"] == 0
            and comparison["removed"] == 0
            and comparison["changed"] == 0
        )
        errors = list(checks["errors"])
        warnings = list(checks["warnings"])
        if same_as_current:
            warnings.append("上传文件与当前版本内容一致，无需再次发布")

        candidate_documents = [*other_documents, *documents]
        candidate_sources = [
            item
            for item in current_sources
            if str(item.get("table_type") or "") != table_type
        ]
        candidate_sources.append(
            {
                "table_type": table_type,
                "file": _portable_path(archive_path),
                "file_name": safe_name,
                "sha256": stored.file_hash.upper(),
                "size_bytes": stored.size_bytes,
                "document_count": len(documents),
            }
        )
        candidate_dir.mkdir(parents=True, exist_ok=False)
        write_case_table_index(
            candidate_documents,
            sqlite_path=candidate_dir / "case_index.sqlite",
            jsonl_path=candidate_dir / "case_documents.jsonl",
            metadata={
                "generated_at": now,
                "refresh_date": now[:10],
                "sources": candidate_sources,
                "preview_batch": batch_no,
            },
        )
        self._validate_index(
            candidate_dir / "case_index.sqlite", expected_count=len(candidate_documents)
        )

        manifest = {
            "batch_no": batch_no,
            "status": "validation_failed" if errors else "ready",
            "status_label": "存在问题" if errors else "待发布",
            "created_at": now,
            "created_by": self.actor_user_id,
            "table_kind": table_kind,
            "table_type": table_type,
            "table_label": CASE_TABLE_LABELS[table_type],
            "file_name": safe_name,
            "file_hash": stored.file_hash,
            "file_size": stored.size_bytes,
            "stored_file": _portable_path(stored.absolute_path),
            "archive_file": _portable_path(archive_path),
            "candidate_dir": _portable_path(candidate_dir),
            "candidate_hash": _sha256(candidate_dir / "case_index.sqlite"),
            "current_index_hash": current_hash,
            "current_version": str(current_version.get("version_id") or ""),
            "counts": {
                "current": len(old_documents),
                "uploaded": len(documents),
                **comparison,
                "blank_case_number": checks["blank_case_number"],
                "blank_case_name": checks["blank_case_name"],
                "duplicate_case_number_groups": checks["duplicate_case_number_groups"],
            },
            "errors": errors,
            "warnings": warnings,
            "sources": candidate_sources,
            "can_publish": not errors and not same_as_current,
        }
        _write_json_atomic(pending_dir / "manifest.json", manifest)
        return _public_preview(manifest)

    def publish(self, batch_no: str) -> dict[str, Any]:
        pending_dir = self._pending_dir(batch_no)
        manifest_path = pending_dir / "manifest.json"
        already_published = False
        with _exclusive_lock(self.lock_path):
            manifest = _read_json(manifest_path)
            if manifest.get("status") == "published":
                already_published = True
            elif manifest.get("status") != "ready" or not manifest.get("can_publish"):
                raise CaseTableReleaseError(
                    "该预览存在问题，不能发布",
                    code="preview_not_publishable",
                    status_code=409,
                )
            else:
                live_index = self.index_dir / "case_index.sqlite"
                live_hash = _sha256(live_index) if live_index.is_file() else ""
                if live_hash != str(manifest.get("current_index_hash") or ""):
                    raise CaseTableReleaseError(
                        "预览后底表已经被其他人更新，请重新上传并查看变化",
                        code="preview_stale",
                        status_code=409,
                    )

                stored_path = _resolve_existing_path(
                    str(manifest.get("stored_file") or "")
                )
                archive_path = self._safe_raw_path(
                    str(manifest.get("archive_file") or "")
                )
                _copy_atomic(stored_path, archive_path)
                if (
                    _sha256(archive_path).lower()
                    != str(manifest.get("file_hash") or "").lower()
                ):
                    raise CaseTableReleaseError(
                        "归档文件校验失败，未发布",
                        code="archive_hash_mismatch",
                        status_code=500,
                    )

                published_at = _now()
                version_id = _reference("V")
                candidate_dir = _resolve_existing_dir(
                    str(manifest.get("candidate_dir") or "")
                )
                version = {
                    "version_id": version_id,
                    "published_at": published_at,
                    "published_by": self.actor_user_id,
                    "reason": f"发布{manifest['table_label']}：{manifest['file_name']}",
                    "action": "publish",
                    "source_batch": batch_no,
                    "restored_from": "",
                    "sources": manifest.get("sources") or [],
                    "table_counts": self._index_counts(
                        candidate_dir / "case_index.sqlite"
                    ),
                    "index_sha256": str(manifest.get("candidate_hash") or ""),
                }
                manifest_before = dict(manifest)
                manifest.update(
                    {
                        "status": "published",
                        "status_label": "已发布",
                        "published_at": published_at,
                        "published_by": self.actor_user_id,
                        "version_id": version_id,
                        "can_publish": False,
                    }
                )
                self._activate_candidate_locked(
                    candidate_dir,
                    version=version,
                    pending_update=(manifest_path, manifest, manifest_before),
                )
        return {
            **_public_preview(manifest),
            "already_published": already_published,
            "current": self.status(),
        }

    def restore(self, version_id: str) -> dict[str, Any]:
        version_path = self._version_path(version_id)
        target = _read_json(version_path)
        sources = list(target.get("sources") or [])
        if not self._sources_available(sources):
            raise CaseTableReleaseError(
                "该历史版本的原始文件不完整，暂时不能恢复",
                code="source_file_missing",
                status_code=409,
            )

        with _exclusive_lock(self.lock_path):
            current = self._ensure_baseline_locked()
            if str(current.get("version_id") or "") == version_id:
                raise CaseTableReleaseError(
                    "该版本已经是当前版本", code="already_current", status_code=409
                )
            restore_ref = _reference("RESTORE")
            candidate_dir = self.root / "restore" / restore_ref / "candidate"
            candidate_dir.mkdir(parents=True, exist_ok=False)
            documents: list[CaseTableDocument] = []
            refreshed_sources: list[dict[str, Any]] = []
            for source in sources:
                source_path = self._safe_raw_path(str(source.get("file") or ""))
                table_type = str(source.get("table_type") or "")
                extracted = _extract_documents_from_workbook(
                    source_path, refresh_date=_now()[:10]
                )
                typed = [
                    replace(item, table_type=table_type, source_file=source_path.name)
                    for item in extracted
                ]
                documents.extend(typed)
                refreshed_sources.append(
                    {
                        **source,
                        "file": _portable_path(source_path),
                        "file_name": source_path.name,
                        "sha256": _sha256(source_path).upper(),
                        "size_bytes": source_path.stat().st_size,
                        "document_count": len(typed),
                    }
                )
            if not documents:
                raise CaseTableReleaseError(
                    "历史版本没有可恢复的案件数据",
                    code="empty_restore",
                    status_code=409,
                )
            now = _now()
            write_case_table_index(
                documents,
                sqlite_path=candidate_dir / "case_index.sqlite",
                jsonl_path=candidate_dir / "case_documents.jsonl",
                metadata={
                    "generated_at": now,
                    "refresh_date": now[:10],
                    "sources": refreshed_sources,
                    "restored_from": version_id,
                },
            )
            self._validate_index(
                candidate_dir / "case_index.sqlite", expected_count=len(documents)
            )
            new_version = {
                "version_id": _reference("V"),
                "published_at": now,
                "published_by": self.actor_user_id,
                "reason": f"恢复历史版本 {version_id}",
                "action": "restore",
                "source_batch": restore_ref,
                "restored_from": version_id,
                "sources": refreshed_sources,
                "table_counts": self._index_counts(candidate_dir / "case_index.sqlite"),
                "index_sha256": _sha256(candidate_dir / "case_index.sqlite"),
            }
            self._activate_candidate_locked(candidate_dir, version=new_version)
        return {
            "restored_from": version_id,
            "version": _public_version(new_version),
            "current": self.status(),
        }

    def _ensure_baseline(self) -> dict[str, Any]:
        with _exclusive_lock(self.lock_path):
            return self._ensure_baseline_locked()

    def _ensure_baseline_locked(self) -> dict[str, Any]:
        if self.current_path.is_file():
            return _read_json(self.current_path)
        live_index = self.index_dir / "case_index.sqlite"
        if not live_index.is_file():
            raise CaseTableReleaseError(
                "当前案件底表索引不存在", code="live_index_missing", status_code=409
            )
        meta = (
            _read_json(self.index_dir / "case_index_meta.json")
            if (self.index_dir / "case_index_meta.json").is_file()
            else {}
        )
        live_hash = _sha256(live_index)
        generated_at = str(meta.get("generated_at") or "")
        if not generated_at:
            generated_at = datetime.fromtimestamp(
                live_index.stat().st_mtime, SHANGHAI
            ).isoformat(timespec="seconds")
        baseline = {
            "version_id": f"BASE-{live_hash[:12].upper()}",
            "published_at": generated_at,
            "published_by": "系统已有版本",
            "reason": "启用自助更新前的底表版本",
            "action": "baseline",
            "source_batch": "",
            "restored_from": "",
            "sources": self._current_sources(meta=meta),
            "table_counts": self._index_counts(live_index),
            "index_sha256": live_hash,
        }
        self.versions_root.mkdir(parents=True, exist_ok=True)
        _write_json_atomic(self._version_path(baseline["version_id"]), baseline)
        _write_json_atomic(self.current_path, baseline)
        return baseline

    def _activate_candidate_locked(
        self,
        candidate_dir: Path,
        *,
        version: dict[str, Any],
        pending_update: tuple[Path, dict[str, Any], dict[str, Any]] | None = None,
    ) -> None:
        self._validate_index(
            candidate_dir / "case_index.sqlite",
            expected_count=sum(
                int(value or 0)
                for value in (version.get("table_counts") or {}).values()
            ),
        )
        self.index_dir.mkdir(parents=True, exist_ok=True)
        backup_dir = self.root / "swap_backups" / _reference("SWAP")
        backup_dir.mkdir(parents=True, exist_ok=False)
        old_current = (
            _read_json(self.current_path) if self.current_path.is_file() else None
        )
        installed: list[str] = []
        version_path = self._version_path(str(version["version_id"]))
        try:
            for name in LIVE_INDEX_FILES:
                source = candidate_dir / name
                if not source.is_file():
                    raise CaseTableReleaseError(
                        "候选索引文件不完整，未发布",
                        code="candidate_incomplete",
                        status_code=500,
                    )
                target = self.index_dir / name
                if target.is_file():
                    _link_or_copy(target, backup_dir / name)
            for name in LIVE_INDEX_FILES:
                _copy_atomic(candidate_dir / name, self.index_dir / name)
                installed.append(name)
            expected = sum(
                int(value or 0)
                for value in (version.get("table_counts") or {}).values()
            )
            self._validate_index(
                self.index_dir / "case_index.sqlite", expected_count=expected
            )
            version["index_sha256"] = _sha256(self.index_dir / "case_index.sqlite")
            self.versions_root.mkdir(parents=True, exist_ok=True)
            _write_json_atomic(version_path, version)
            _write_json_atomic(self.current_path, version)
            if pending_update:
                _write_json_atomic(pending_update[0], pending_update[1])
        except Exception:
            for name in reversed(installed):
                backup = backup_dir / name
                target = self.index_dir / name
                if backup.is_file():
                    _copy_atomic(backup, target)
                elif target.exists():
                    target.unlink()
            if old_current is not None:
                _write_json_atomic(self.current_path, old_current)
            elif self.current_path.exists():
                self.current_path.unlink()
            if version_path.exists():
                version_path.unlink()
            if pending_update:
                _write_json_atomic(pending_update[0], pending_update[2])
            raise

    def _validate_documents(
        self, documents: list[CaseTableDocument], *, table_type: str
    ) -> dict[str, Any]:
        case_numbers = [_case_number(item) for item in documents]
        duplicate_numbers = Counter(value for value in case_numbers if value)
        duplicate_groups = sum(count > 1 for count in duplicate_numbers.values())
        blank_numbers = sum(not value for value in case_numbers)
        blank_names = sum(not str(item.case_name or "").strip() for item in documents)
        errors: list[str] = []
        warnings: list[str] = []
        if table_type == "defendant_case_table":
            if blank_numbers:
                errors.append(f"有 {blank_numbers} 条案件缺少案件编号")
            if blank_names:
                errors.append(f"有 {blank_names} 条案件缺少案件名称")
            if duplicate_groups:
                errors.append(f"有 {duplicate_groups} 组案件编号重复")
        else:
            if blank_names:
                warnings.append(
                    f"有 {blank_names} 条原告案件未识别到案件名称，请重点核对变化明细"
                )
        return {
            "blank_case_number": blank_numbers,
            "blank_case_name": blank_names,
            "duplicate_case_number_groups": duplicate_groups,
            "errors": errors,
            "warnings": warnings,
        }

    def _compare_documents(
        self,
        old_documents: list[CaseTableDocument],
        new_documents: list[CaseTableDocument],
    ) -> dict[str, int]:
        old_map = _unique_document_map(old_documents)
        new_map = _unique_document_map(new_documents)
        old_keys = set(old_map)
        new_keys = set(new_map)
        common = old_keys & new_keys
        changed = sum(
            _document_signature(old_map[key]) != _document_signature(new_map[key])
            for key in common
        )
        return {
            "added": len(new_keys - old_keys),
            "removed": len(old_keys - new_keys),
            "changed": changed,
            "unchanged": len(common) - changed,
        }

    def _read_documents(self, index_path: Path) -> list[CaseTableDocument]:
        self._validate_index(index_path)
        with closing(sqlite3.connect(index_path)) as connection:
            connection.row_factory = sqlite3.Row
            rows = connection.execute("SELECT * FROM documents").fetchall()
        output: list[CaseTableDocument] = []
        for row in rows:
            try:
                facts = json.loads(str(row["facts_json"] or "{}"))
            except json.JSONDecodeError:
                facts = {}
            output.append(
                CaseTableDocument(
                    doc_id=str(row["doc_id"] or ""),
                    source_type=str(row["source_type"] or "case_table_rag"),
                    source_file=str(row["source_file"] or ""),
                    sheet_name=str(row["sheet_name"] or ""),
                    row_number=int(row["row_number"] or 0),
                    table_type=str(row["table_type"] or ""),
                    case_name=str(row["case_name"] or ""),
                    department=str(row["department"] or ""),
                    assignee_name=str(row["assignee_name"] or ""),
                    status=str(row["status"] or ""),
                    updated_at=str(row["updated_at"] or ""),
                    text=str(row["text"] or ""),
                    facts=facts,
                )
            )
        return output

    def _current_sources(
        self, *, meta: dict[str, Any] | None = None
    ) -> list[dict[str, Any]]:
        meta_value = meta
        if meta_value is None:
            meta_path = self.index_dir / "case_index_meta.json"
            meta_value = _read_json(meta_path) if meta_path.is_file() else {}
        source_types: dict[str, str] = {}
        index_path = self.index_dir / "case_index.sqlite"
        if index_path.is_file():
            with closing(sqlite3.connect(index_path)) as connection:
                for source_file, table_type in connection.execute(
                    "SELECT source_file, MIN(table_type) FROM documents GROUP BY source_file"
                ):
                    source_types[str(source_file or "")] = str(table_type or "")
        sources: list[dict[str, Any]] = []
        for source in meta_value.get("sources") or []:
            if not isinstance(source, dict):
                continue
            file_value = str(source.get("file") or "")
            file_name = Path(file_value).name
            table_type = str(
                source.get("table_type")
                or source_types.get(file_name)
                or _table_type_from_name(file_name)
            )
            sources.append(
                {
                    "table_type": table_type,
                    "file": file_value,
                    "file_name": file_name,
                    "sha256": str(source.get("sha256") or ""),
                    "size_bytes": int(source.get("size_bytes") or 0),
                    "document_count": int(source.get("document_count") or 0),
                }
            )
        return sources

    def _source_payloads(
        self, sources: Iterable[dict[str, Any]], *, counts: dict[str, int]
    ) -> list[dict[str, Any]]:
        by_type = {str(item.get("table_type") or ""): item for item in sources}
        output = []
        for kind, table_type in CASE_TABLE_TYPES.items():
            source = by_type.get(table_type, {})
            output.append(
                {
                    "table_kind": kind,
                    "table_type": table_type,
                    "label": CASE_TABLE_LABELS[table_type],
                    "file_name": str(
                        source.get("file_name")
                        or Path(str(source.get("file") or "")).name
                    ),
                    "file_hash": str(source.get("sha256") or ""),
                    "document_count": int(
                        counts.get(table_type) or source.get("document_count") or 0
                    ),
                }
            )
        return output

    def _sources_available(self, sources: Iterable[dict[str, Any]]) -> bool:
        values = list(sources)
        if not values:
            return False
        try:
            return all(
                self._safe_raw_path(str(item.get("file") or "")).is_file()
                for item in values
            )
        except CaseTableReleaseError:
            return False

    def _archive_target(self, filename: str, file_hash: str, date_label: str) -> Path:
        folder = (self.raw_root / date_label).resolve()
        if self.raw_root != folder and self.raw_root not in folder.parents:
            raise CaseTableReleaseError(
                "案件底表归档位置不安全", code="unsafe_archive_path", status_code=500
            )
        target = (folder / safe_filename(filename)).resolve()
        if folder not in target.parents:
            raise CaseTableReleaseError(
                "案件底表归档文件名不安全", code="unsafe_archive_path", status_code=422
            )
        if target.exists() and _sha256(target).lower() != file_hash.lower():
            target = folder / f"{target.stem}-{file_hash[:8]}{target.suffix}"
        return target

    def _safe_raw_path(self, value: str) -> Path:
        path = Path(value)
        target = (
            (Path.cwd() / path).resolve() if not path.is_absolute() else path.resolve()
        )
        if target != self.raw_root and self.raw_root not in target.parents:
            raise CaseTableReleaseError(
                "历史版本来源不在案件底表归档目录内",
                code="unsafe_source_path",
                status_code=409,
            )
        return target

    def _pending_dir(self, batch_no: str) -> Path:
        if not re.fullmatch(r"CT-[0-9]{14}-[A-F0-9]{8}", str(batch_no or "")):
            raise CaseTableReleaseError(
                "更新预览编号无效", code="invalid_batch_no", status_code=404
            )
        target = (self.pending_root / batch_no).resolve()
        if self.pending_root.resolve() not in target.parents or not target.is_dir():
            raise CaseTableReleaseError(
                "更新预览不存在", code="preview_not_found", status_code=404
            )
        return target

    def _version_path(self, version_id: str) -> Path:
        if not re.fullmatch(
            r"(?:BASE-[A-F0-9]{12}|V-[0-9]{14}-[A-F0-9]{8})", str(version_id or "")
        ):
            raise CaseTableReleaseError(
                "版本编号无效", code="invalid_version", status_code=404
            )
        return self.versions_root / f"{version_id}.json"

    @staticmethod
    def _validate_index(index_path: Path, *, expected_count: int | None = None) -> None:
        if not index_path.is_file():
            raise CaseTableReleaseError(
                "案件索引文件不存在", code="index_missing", status_code=500
            )
        try:
            with closing(sqlite3.connect(index_path)) as connection:
                integrity = str(
                    connection.execute("PRAGMA integrity_check").fetchone()[0]
                )
                document_count = int(
                    connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
                )
                fts_count = int(
                    connection.execute("SELECT COUNT(*) FROM documents_fts").fetchone()[
                        0
                    ]
                )
        except sqlite3.Error as exc:
            raise CaseTableReleaseError(
                "案件索引校验失败，未发布", code="index_invalid", status_code=500
            ) from exc
        if integrity != "ok" or document_count != fts_count:
            raise CaseTableReleaseError(
                "案件索引校验失败，未发布", code="index_invalid", status_code=500
            )
        if expected_count is not None and document_count != int(expected_count):
            raise CaseTableReleaseError(
                "案件索引数量与预览不一致，未发布",
                code="index_count_mismatch",
                status_code=500,
            )

    @staticmethod
    def _index_counts(index_path: Path) -> dict[str, int]:
        if not index_path.is_file():
            return {}
        with closing(sqlite3.connect(index_path)) as connection:
            return {
                str(table_type): int(count)
                for table_type, count in connection.execute(
                    "SELECT table_type, COUNT(*) FROM documents GROUP BY table_type"
                )
            }

    @staticmethod
    def _table_type(table_kind: str) -> str:
        try:
            return CASE_TABLE_TYPES[str(table_kind or "")]
        except KeyError as exc:
            raise CaseTableReleaseError(
                "案件底表类型无效", code="invalid_table_type", status_code=422
            ) from exc


def _extract_documents_from_workbook(
    source_path: Path,
    *,
    refresh_date: str = "",
) -> list[CaseTableDocument]:
    table_type = _table_type_from_name(source_path.name) or "case_table"
    documents: list[CaseTableDocument] = []
    # ERP exports sometimes declare an incorrect worksheet dimension (A1:A1).
    # Normal mode reads the real cell grid; read-only mode would silently see
    # just one cell and report an empty workbook.
    workbook = load_workbook(source_path, read_only=False, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            first_row = next(rows, None)
            if first_row is None:
                continue
            headers = [
                _clean_header(value, index) for index, value in enumerate(first_row)
            ]
            for row_number, values in enumerate(rows, start=2):
                facts = {
                    header: cleaned
                    for header, value in zip(headers, values, strict=False)
                    if (cleaned := _clean_cell(value))
                }
                if not facts:
                    continue
                case_name = _first_fact(facts, ("案件名称", "案名")) or _first_fact(
                    facts,
                    ("受理案由", "案由"),
                )
                text = _row_text(facts)
                if not case_name and "案" not in text:
                    continue
                doc_id = hashlib.sha256(
                    (
                        f"{source_path.name}|{sheet.title}|{row_number}|{text[:200]}"
                    ).encode()
                ).hexdigest()[:20]
                documents.append(
                    CaseTableDocument(
                        doc_id=doc_id,
                        source_type="case_table_rag",
                        source_file=source_path.name,
                        sheet_name=str(sheet.title),
                        row_number=row_number,
                        table_type=table_type,
                        case_name=case_name,
                        department=_first_fact(facts, ("法务部门", "部门")),
                        assignee_name=_first_fact(
                            facts,
                            ("负责人", "承办人", "承办", "经办"),
                        ),
                        status=_first_fact(
                            facts,
                            ("案件状态", "当前阶段", "阶段", "红黄绿", "状态"),
                        ),
                        updated_at=refresh_date,
                        text=text,
                        facts=facts,
                    )
                )
    finally:
        workbook.close()
    return documents


def _clean_header(value: Any, index: int) -> str:
    text = _clean_cell(value)
    if not text or text.lower().startswith("unnamed:"):
        return f"列{index + 1}"
    return text


def _clean_cell(value: Any) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return ""
    text = re.sub(r"\s+", " ", text)
    if text.endswith(".0") and re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def _first_fact(facts: dict[str, str], markers: tuple[str, ...]) -> str:
    for key, value in facts.items():
        if any(marker in key for marker in markers):
            return value
    return ""


def _row_text(facts: dict[str, str]) -> str:
    preferred: list[str] = []
    remaining: list[str] = []
    for key, value in facts.items():
        target = preferred if _preferred_case_field(key) else remaining
        target.append(f"{key}: {value}")
    return "；".join([*preferred, *remaining])


def _preferred_case_field(header: str) -> bool:
    return any(
        marker in header
        for marker in (
            "案件",
            "案号",
            "法务",
            "负责人",
            "承办",
            "公司",
            "分公司",
            "法院",
            "仲裁",
            "执行",
            "开庭",
            "阶段",
            "计划",
            "完成情况",
            "状态",
            "金额",
        )
    )


def _public_preview(manifest: dict[str, Any]) -> dict[str, Any]:
    return {
        key: manifest.get(key)
        for key in (
            "batch_no",
            "status",
            "status_label",
            "created_at",
            "created_by",
            "table_kind",
            "table_type",
            "table_label",
            "file_name",
            "file_hash",
            "file_size",
            "current_version",
            "counts",
            "errors",
            "warnings",
            "can_publish",
            "published_at",
            "published_by",
            "version_id",
        )
    }


def _public_version(version: dict[str, Any]) -> dict[str, Any]:
    return {
        key: version.get(key)
        for key in (
            "version_id",
            "published_at",
            "published_by",
            "reason",
            "action",
            "source_batch",
            "restored_from",
            "table_counts",
            "tables",
            "is_current",
            "can_restore",
        )
    }


def _case_number(document: CaseTableDocument) -> str:
    for key, value in document.facts.items():
        normalized = re.sub(r"[\s_\-（）()]+", "", str(key or ""))
        if "案件编号" in normalized or normalized in {"案件ID", "案件标识"}:
            return str(value or "").strip()
    return ""


def _document_key(document: CaseTableDocument) -> str:
    case_number = _case_number(document)
    if case_number:
        return f"number:{case_number}"
    case_name = str(document.case_name or "").strip()
    if case_name:
        return f"name:{case_name}"
    return f"row:{document.source_file}|{document.sheet_name}|{document.row_number}"


def _unique_document_map(
    documents: Iterable[CaseTableDocument],
) -> dict[str, CaseTableDocument]:
    grouped: dict[str, list[CaseTableDocument]] = {}
    for item in documents:
        grouped.setdefault(_document_key(item), []).append(item)
    return {key: values[0] for key, values in grouped.items() if len(values) == 1}


def _document_signature(document: CaseTableDocument) -> str:
    return hashlib.sha256(
        json.dumps(
            document.facts, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        ).encode("utf-8")
    ).hexdigest()


def _table_type_from_name(filename: str) -> str:
    if "被告" in filename:
        return "defendant_case_table"
    if "原告" in filename:
        return "plaintiff_case_table"
    return ""


def _reference(prefix: str) -> str:
    return f"{prefix}-{datetime.now(SHANGHAI).strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:8].upper()}"


def _now() -> str:
    return datetime.now(SHANGHAI).isoformat(timespec="seconds")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _portable_path(path: Path) -> str:
    target = path.resolve()
    try:
        return target.relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return str(target)


def _read_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise CaseTableReleaseError(
            "记录不存在", code="record_not_found", status_code=404
        )
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise TypeError("JSON object expected")
    return value


def _write_json_atomic(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2, sort_keys=True)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_name, path)
    finally:
        if os.path.exists(temporary_name):
            os.unlink(temporary_name)


def _copy_atomic(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{target.name}.", suffix=".tmp", dir=target.parent
    )
    os.close(descriptor)
    try:
        shutil.copyfile(source, temporary_name)
        with open(temporary_name, "r+b") as handle:
            os.fsync(handle.fileno())
        for attempt in range(8):
            try:
                os.replace(temporary_name, target)
                break
            except PermissionError:
                if attempt == 7:
                    raise
                time.sleep(0.05 * (attempt + 1))
    finally:
        if os.path.exists(temporary_name):
            os.unlink(temporary_name)


def _link_or_copy(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        os.link(source, target)
    except OSError:
        shutil.copy2(source, target)


def _resolve_existing_path(value: str) -> Path:
    path = Path(value)
    target = (Path.cwd() / path).resolve() if not path.is_absolute() else path.resolve()
    if not target.is_file():
        raise CaseTableReleaseError(
            "暂存的原始文件不存在", code="staged_file_missing", status_code=409
        )
    return target


def _resolve_existing_dir(value: str) -> Path:
    path = Path(value)
    target = (Path.cwd() / path).resolve() if not path.is_absolute() else path.resolve()
    if not target.is_dir():
        raise CaseTableReleaseError(
            "更新预览已失效，请重新上传", code="candidate_missing", status_code=409
        )
    return target


@contextmanager
def _exclusive_lock(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a+b") as handle:
        handle.seek(0, os.SEEK_END)
        if handle.tell() == 0:
            handle.write(b"0")
            handle.flush()
        handle.seek(0)
        if os.name == "nt":
            import msvcrt

            msvcrt.locking(handle.fileno(), msvcrt.LK_LOCK, 1)
            try:
                yield
            finally:
                handle.seek(0)
                msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl

            fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
            try:
                yield
            finally:
                fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
